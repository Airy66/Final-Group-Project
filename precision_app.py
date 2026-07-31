"""Role-based Precision Curator web application."""

import base64
import csv
from copy import copy
from collections import Counter
from datetime import date, datetime, timedelta, timezone
import hashlib
import hmac
import html
import io
import json
import logging
import os
from pathlib import Path
import secrets
import sys
from decimal import Decimal
from zoneinfo import ZoneInfo

from dotenv import load_dotenv

# Environment-backed services must not be imported before the application's
# .env file has been loaded. ``override=True`` also prevents a stale shell
# value from silently selecting a different database than the project config.
_test_environment = os.getenv("PRECISION_TESTING") == "true" and "pytest" in sys.modules
load_dotenv(override=True)
if _test_environment:
    os.environ.update(APP_ENV="testing", DEMO_MODE="true")
    os.environ.pop("MONGO_URI", None)

from bson import ObjectId
from flask import Flask, Response, abort, flash, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from flask.testing import FlaskClient
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pymongo.errors import DuplicateKeyError, PyMongoError
import requests
import re
from openpyxl import Workbook
from urllib.parse import unquote, urlparse, urlsplit
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from services.ai_search import AISearchError, ground_market_summary, predict_price_gemini, search_prices, summarize_market_gemini as summarize_market
from services.app_urls import application_email_url, canonical_app_base_url
from services.database import MongoRepository, utcnow
from services.mail_service import PasswordResetMailService
from services.price_alerts import VALID_DIRECTIONS, evaluate_price_alert, normalize_mail_error
from services.runtime_config import env_flag, is_production_environment, validate_production_configuration, web_security_configuration
from services.category_attributes import CATEGORY_PROFILES, PRODUCT_TYPE_LABELS, apply_category_filters, canonical_facet_value, category_scope, classify_query_intent, detect_product_category, enrich_listing_category, generate_available_facets, product_type_scope
from services.serpapi_search import build_serpapi_cache_key, call_serpapi_marketplace, engine_for_platform, normalize_query as normalize_serpapi_query, provider_sort_for_platform, serpapi_account, SerpApiError
from services.platforms import canonical_marketplace_platform
from src.ecommerce_price_monitor.collectors.walmart_collector import WalmartCollector
from src.ecommerce_price_monitor.utils.exceptions import CollectorError

validate_production_configuration()


class _ResetTokenLogFilter(logging.Filter):
    """Prevent reset secrets from appearing in development HTTP access logs."""

    _pattern = re.compile(r"/reset-password/[^\s?]+")

    def filter(self, record):
        if isinstance(record.msg, str):
            record.msg = self._pattern.sub("/reset-password/[redacted]", record.msg)
        if isinstance(record.args, tuple):
            record.args = tuple(self._pattern.sub("/reset-password/[redacted]", value) if isinstance(value, str) else value for value in record.args)
        return True


logging.getLogger("werkzeug").addFilter(_ResetTokenLogFilter())
app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True
_runtime_environment = (os.getenv("APP_ENV") or os.getenv("FLASK_ENV") or "development").strip().lower()


@app.errorhandler(PyMongoError)
def database_operation_unavailable(error):
    """Fail closed when an established MongoDB connection cannot serve a request."""
    app.logger.error(
        "MongoDB operation failed (%s)",
        error.__class__.__name__,
        exc_info=(type(error), error, error.__traceback__),
    )
    return Response(
        "Database temporarily unavailable. Please try again later.",
        status=503,
        mimetype="text/plain",
    )


def _env_int(name, default, minimum=None, maximum=None):
    try:
        value = int(os.getenv(name, str(default)))
    except (TypeError, ValueError):
        value = default
    if minimum is not None:
        value = max(minimum, value)
    if maximum is not None:
        value = min(maximum, value)
    return value


app.config.update(
    **web_security_configuration(),
    DEVELOPER_ROLE_PREVIEW=os.getenv("DEVELOPER_ROLE_PREVIEW", "false").lower() == "true",
    DEMO_TOOLS_ENABLED=os.getenv("DEMO_TOOLS_ENABLED", "false").lower() == "true",
    DEMO_MEMBERSHIP_UPGRADE_ENABLED=env_flag("DEMO_MEMBERSHIP_UPGRADE_ENABLED", False),
    REGISTRATION_WELCOME_EMAIL_ENABLED=env_flag("REGISTRATION_WELCOME_EMAIL_ENABLED", True),
    MONITOR_DAILY_REFRESH_ENABLED=env_flag("MONITOR_DAILY_REFRESH_ENABLED", False),
    MONITOR_AUTO_REFRESH_MAX_PER_USER=_env_int("MONITOR_AUTO_REFRESH_MAX_PER_USER", 1, 1, 1),
    MONITOR_SCHEDULED_RUN_LIMIT=_env_int("MONITOR_SCHEDULED_RUN_LIMIT", 5, 1),
    MONITOR_REFRESH_TIMEZONE=os.getenv("MONITOR_REFRESH_TIMEZONE", "Asia/Singapore"),
    MONITOR_DAILY_REFRESH_HOUR=_env_int("MONITOR_DAILY_REFRESH_HOUR", 8, 0, 23),
    MONITOR_TEST_PROVIDER_CALLS=False,
    PRICE_ALERTS_ENABLED=env_flag("PRICE_ALERTS_ENABLED", True),
    PRICE_ALERT_DEFAULT_THRESHOLD_PERCENT=os.getenv("PRICE_ALERT_DEFAULT_THRESHOLD_PERCENT", "5"),
    PRICE_ALERT_COOLDOWN_HOURS=_env_int("PRICE_ALERT_COOLDOWN_HOURS", 24, 1),
    RUNTIME_ENVIRONMENT=_runtime_environment,
    MAX_CONTENT_LENGTH=16 * 1024 * 1024,
    MAX_FORM_MEMORY_SIZE=16 * 1024 * 1024,
)


CSRF_SESSION_KEY = "_csrf_token"
CSRF_REJECTION_LOGGED_KEY = "_csrf_rejection_logged"
CSRF_HEADER_NAME = "X-CSRFToken"
CSRF_ERROR_MESSAGE = "Your session or form security token has expired. Refresh the page and try again."


def csrf_token():
    token = session.get(CSRF_SESSION_KEY)
    if not token:
        token = secrets.token_urlsafe(32)
        session[CSRF_SESSION_KEY] = token
    return token


def safe_internal_redirect_target(value, default=None, *, log_rejection=True):
    """Return a local absolute path, rejecting encoded and browser URL tricks."""
    fallback = default or url_for("dashboard_redirect")
    candidate = str(value or "").strip()
    if not candidate:
        return fallback
    decoded = candidate
    for _ in range(3):
        next_value = unquote(decoded)
        if next_value == decoded:
            break
        decoded = next_value
    parsed = urlsplit(decoded)
    safe = (
        decoded.startswith("/")
        and not decoded.startswith("//")
        and "\\" not in decoded
        and not parsed.scheme
        and not parsed.netloc
        and not any(ord(character) < 32 for character in decoded)
    )
    if safe:
        return decoded
    if log_rejection and session.get("user_id"):
        safe_call(lambda: repository.log_event(
            session.get("user_id"), "unsafe_redirect_rejected", session.get("role"), session.get("username"),
            {"endpoint": request.endpoint, "destination_kind": "external_or_malformed"},
        ), None)
    return fallback


class CSRFTestClient(FlaskClient):
    """Keep legacy tests practical while exercising enabled CSRF validation."""

    def open(self, *args, **kwargs):
        auto_csrf = kwargs.pop("auto_csrf", True)
        method = str(kwargs.get("method") or (args[1] if len(args) > 1 else "GET")).upper()
        if self.application.testing and auto_csrf and method not in {"GET", "HEAD", "OPTIONS", "TRACE"}:
            with self.session_transaction() as test_session:
                token = test_session.get(CSRF_SESSION_KEY) or secrets.token_urlsafe(32)
                test_session[CSRF_SESSION_KEY] = token
            headers = dict(kwargs.get("headers") or {})
            headers.setdefault(CSRF_HEADER_NAME, token)
            kwargs["headers"] = headers
        return super().open(*args, **kwargs)


app.test_client_class = CSRFTestClient


@app.before_request
def enforce_csrf_protection():
    if request.method in {"GET", "HEAD", "OPTIONS", "TRACE"}:
        return None
    expected = session.get(CSRF_SESSION_KEY)
    supplied = request.headers.get(CSRF_HEADER_NAME) or request.form.get("csrf_token")
    if expected and supplied and hmac.compare_digest(str(expected), str(supplied)):
        return None
    if session.get("user_id") and not session.get(CSRF_REJECTION_LOGGED_KEY):
        safe_call(lambda: repository.log_event(
            session.get("user_id"), "csrf_rejected", session.get("role"), session.get("username"),
            {"endpoint": request.endpoint, "method": request.method},
        ), None)
        session[CSRF_REJECTION_LOGGED_KEY] = True
    previous = safe_internal_redirect_target(request.referrer, url_for("dashboard_redirect"), log_rejection=False) if request.referrer else None
    return render_template("csrf_error.html", message=CSRF_ERROR_MESSAGE, previous_url=previous), 400


@app.context_processor
def inject_csrf_context():
    return {"csrf_token": csrf_token}


@app.after_request
def inject_csrf_form_fields(response):
    if response.direct_passthrough or not response.content_type or "text/html" not in response.content_type.lower():
        return response
    body = response.get_data(as_text=True)
    if "<form" not in body.lower():
        return response
    field = f'<input type="hidden" name="csrf_token" value="{html.escape(csrf_token(), quote=True)}">'
    pattern = re.compile(r'(<form\b[^>]*\bmethod\s*=\s*["\']?post["\']?[^>]*>)', re.IGNORECASE)
    response.set_data(pattern.sub(lambda match: match.group(1) + field, body))
    return response


@app.after_request
def password_reset_response_headers(response):
    if request.endpoint in {"forgot_password", "reset_password"}:
        response.headers["Cache-Control"] = "no-store"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["X-Robots-Tag"] = "noindex, nofollow"
    return response


repository = MongoRepository()
CLIENT_ID = os.getenv("EBAY_CLIENT_ID")
CLIENT_SECRET = os.getenv("EBAY_CLIENT_SECRET")
EBAY_MARKETPLACE_ID = os.getenv("EBAY_MARKETPLACE_ID", "EBAY_US")
CHART_DIR = Path("static/charts")
CHART_DIR.mkdir(parents=True, exist_ok=True)
AVATAR_UPLOAD_DIR = Path("static/uploads/avatars")
AVATAR_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
ALLOWED_AVATAR_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
MAX_AVATAR_SIZE_BYTES = 2 * 1024 * 1024
CURRENT_ANALYSIS_SCHEMA_VERSION = 3


@app.route("/favicon.ico")
def site_favicon():
    """Serve the site identity at the browser's conventional favicon URL."""
    return send_from_directory(app.static_folder, "favicon.svg", mimetype="image/svg+xml")


@app.get("/health")
def health():
    """Return process health without touching providers or persistence."""
    return jsonify({"status": "ok", "service": "precision-curator"})


class AnalysisScopeMismatch(Exception):
    def __init__(self, analysis_record, checks):
        self.analysis_record = analysis_record
        self.checks = checks

ROLES = {
    "consumer": "Consumer",
    "retailer": "Retailer / Reseller",
    "researcher": "Researcher",
    "administrator": "Administrator",
}
FEATURE_REQUIRED_TIERS = {
    "product_search": "basic",
    "compare_products": "basic",
    "save_evidence": "premium",
    "watchlist": "premium",
    "saved_research": "premium",
    "basic_analytics": "premium",
    "analytics_dashboard": "premium",
    "ai_summary": "premium",
    "advanced_analytics": "professional",
    "prediction": "premium",
    "standard_export": "premium",
    "prediction_validation": "professional",
    "source_audit": "professional",
    "logs": "professional",
    "export_report": "professional",
    "research_package": "professional",
}
MEMBERSHIP_PLANS = [
    {
        "id": "basic",
        "name": "Basic",
        "price": "Free",
        "description": "For simple product discovery and essential price comparison.",
        "features": [
            "Product search across available sources",
            "Simple product comparison",
            "View essential product details",
            "Limited search history",
            "Read-only preview of advanced modules",
        ],
    },
    {
        "id": "premium",
        "name": "Premium",
        "price": "$9.90 / month",
        "description": "For users who need saved workflows, watchlist monitoring, analytics, AI-assisted forecasting, and standard exports.",
        "features": [
            "Product search and comparison",
            "Save selected evidence",
            "Saved Research workspace",
            "Watchlist tracking and snapshot history",
            "Analytics dashboard and interactive charts",
            "AI-supported market summary",
            "AI price forecast in Watchlist",
            "Standard CSV, chart and report exports from unlocked pages",
        ],
    },
    {
        "id": "professional",
        "name": "Professional",
        "price": "$19.90 / month",
        "description": "For research and audit workflows requiring validation, provenance, activity records, and advanced research exports.",
        "features": [
            "All search, evidence, Watchlist, Analytics and AI forecasting capabilities",
            "Forecast validation workflow",
            "Source Audit",
            "Activity Logs",
            "Validation evidence summary",
            "Audit-log export",
            "Research package export",
            "Advanced reporting and provenance records",
        ],
    },
]
PLAN_BY_ID = {plan["id"]: plan for plan in MEMBERSHIP_PLANS}
LOCKED_FEATURE_VIEWS = {
    "saved_research": {
        "icon": "bookmarks",
        "eyebrow": "Evidence workspace",
        "headline": "Turn useful listings into reusable evidence.",
        "summary": "Save selected marketplace records with their source, observed price, confidence, and collection context intact.",
        "capabilities": [
            {"icon": "bookmark_added", "title": "Evidence library", "detail": "Keep deliberate observations separate from temporary search results."},
            {"icon": "inventory_2", "title": "Research packages", "detail": "Group evidence into focused, reusable research collections."},
            {"icon": "download", "title": "Portable records", "detail": "Export unlocked evidence workflows without losing provenance."},
        ],
    },
    "watchlist": {
        "icon": "monitoring",
        "eyebrow": "Price monitoring",
        "headline": "See how a market moves after the first search.",
        "summary": "Create product monitors, collect comparable price snapshots, and review changes against an established baseline.",
        "capabilities": [
            {"icon": "show_chart", "title": "Price history", "detail": "Follow average, low, and high observed prices over time."},
            {"icon": "notifications_active", "title": "Change signals", "detail": "Surface price movements and monitors that need attention."},
            {"icon": "storefront", "title": "Source comparison", "detail": "Track comparable listings across available marketplaces."},
        ],
    },
    "analytics_dashboard": {
        "icon": "analytics",
        "eyebrow": "Market analytics",
        "headline": "Move from listings to a comparable market view.",
        "summary": "Explore normalized price ranges, platform coverage, comparable records, and grounded analytical summaries.",
        "capabilities": [
            {"icon": "query_stats", "title": "Price distribution", "detail": "Understand range, average, outliers, and market position."},
            {"icon": "compare_arrows", "title": "Comparable sets", "detail": "Keep analytical conclusions tied to included records."},
            {"icon": "auto_awesome", "title": "Grounded summaries", "detail": "Use current normalized evidence as the basis for AI analysis."},
        ],
    },
    "source_audit": {
        "icon": "policy",
        "eyebrow": "Provenance and audit",
        "headline": "Verify where every research record came from.",
        "summary": "Review collection outcomes, source coverage, evidence events, and failures in one traceable audit workflow.",
        "capabilities": [
            {"icon": "fact_check", "title": "Source provenance", "detail": "Connect searches and saved evidence to their originating source."},
            {"icon": "warning", "title": "Failure visibility", "detail": "Distinguish no-result outcomes from provider and persistence errors."},
            {"icon": "description", "title": "Audit reports", "detail": "Generate exportable records for advanced review and documentation."},
        ],
    },
    "logs": {
        "icon": "receipt_long",
        "eyebrow": "Research activity",
        "headline": "Keep a chronological record of research actions.",
        "summary": "Review search, evidence, AI, and audit activity as an account-scoped timeline for reproducibility.",
        "capabilities": [
            {"icon": "timeline", "title": "Activity timeline", "detail": "Review important workflow events in chronological order."},
            {"icon": "smart_toy", "title": "AI traceability", "detail": "See model activity and the evidence scope used for analysis."},
            {"icon": "ios_share", "title": "Review exports", "detail": "Export activity records for professional research workflows."},
        ],
    },
    "export": {
        "icon": "download",
        "eyebrow": "Research export",
        "headline": "Take traceable research records outside the workspace.",
        "summary": "Export supported evidence, monitoring, analysis, or audit records from the page where they were created.",
        "capabilities": [
            {"icon": "table_view", "title": "Structured data", "detail": "Download supported records in portable tabular formats."},
            {"icon": "verified", "title": "Preserved context", "detail": "Keep source and workflow context attached to research outputs."},
            {"icon": "share", "title": "External review", "detail": "Prepare unlocked records for collaboration and reporting."},
        ],
    },
}
SOURCES = {"ebay": "eBay API", "ai": "AI Web Search", "demo": "Demo Data"}
ACCESSORY_KEYWORDS = {"case", "cover", "screen protector", "tempered glass", "film", "accessory", "charger", "cable", "adapter", "strap", "mount", "holder", "replacement", "repair", "housing", "frame", "parts", "lcd", "display", "digitizer", "lens"}
SERPAPI_MARKETPLACE_ENABLED = os.getenv("SERPAPI_MARKETPLACE_ENABLED", "true").lower() in {"1", "true", "on", "yes"}
SERPAPI_CACHE_TTL_HOURS = int(os.getenv("SERPAPI_CACHE_TTL_HOURS", "24"))
SERPAPI_DEMO_CACHE_TTL_HOURS = int(os.getenv("SERPAPI_DEMO_CACHE_TTL_HOURS", str(24 * 7)))
DEMO_SAFE_KEYWORDS = {"iphone", "iphone 14", "iphone 15", "iphone 17", "airpods", "airpods pro", "galaxy s24", "headphones"}

def current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    user = repository.get_user_by_id(user_id)
    if not user or user.get("account_status", "active") != "active":
        session.clear()
        return None
    session["username"] = user.get("display_name") or user.get("username") or "User"
    session["display_name"] = session["username"]
    session["email"] = user.get("email")
    roles = user.get("roles") or ([user.get("role")] if user.get("role") else ["consumer"])
    active_role = user.get("active_role") or session.get("active_role") or user.get("role") or roles[0]
    if active_role not in roles:
        active_role = roles[0]
    session["roles"] = roles
    session["active_role"] = active_role
    session["role"] = active_role
    return {
        "_id": user.get("_id"),
        "username": session["username"],
        "display_name": session["display_name"],
        "email": session.get("email"),
        "role": session["role"],
        "primary_role": user.get("primary_role") or user.get("role") or roles[0],
        "roles": roles,
        "active_role": active_role,
        "plan": normalize_membership_tier(user.get("plan")),
        "membership_tier": normalize_membership_tier(user.get("membership_tier") or user.get("plan")),
        "account_status": user.get("account_status"),
        "avatar_path": user.get("avatar_path") or "",
        "institution": user.get("institution") or "",
        "is_developer": bool(user.get("is_developer")),
    }


def user_initials(user):
    name = (user or {}).get("display_name") or (user or {}).get("username") or (user or {}).get("email") or "User"
    parts = re.findall(r"[A-Za-z0-9]+", name)
    if not parts:
        return "U"
    if len(parts) == 1:
        return parts[0][:2].upper()
    return (parts[0][0] + parts[-1][0]).upper()


def avatar_url(user):
    path = (user or {}).get("avatar_path")
    if not path:
        return None
    return url_for("static", filename=path)


def can_access(user, feature):
    if not user:
        return False
    if (user.get("active_role") or user.get("role")) == "administrator":
        return True
    tier = normalize_membership_tier(user.get("membership_tier") or user.get("plan"))
    return membership_rank(tier) >= membership_rank(required_tier(feature))


def membership_rank(tier):
    order = {"basic": 1, "premium": 2, "professional": 3}
    return order.get((tier or "basic").lower(), 0)


def required_tier(feature):
    return FEATURE_REQUIRED_TIERS.get((feature or "").strip(), "basic")


def tier_label(tier):
    return normalize_membership_tier(tier).capitalize()


def locked_message(feature):
    if feature == "watchlist":
        return "Premium unlocks evidence saving, watchlists, analytics dashboards, and AI-assisted price forecasting."
    if feature == "prediction":
        return "AI price forecasting is available on Premium and Professional plans."
    if feature == "source_audit":
        return "Source audit requires Professional membership."
    if feature == "logs":
        return "Activity logs are part of the Professional audit workflow because they support traceability and advanced review."
    if feature in {"saved_research", "save_evidence", "analytics_dashboard", "basic_analytics", "standard_export"}:
        return "Premium unlocks evidence saving, watchlists, analytics dashboards, and AI-assisted price forecasting."
    if feature in {"advanced_analytics", "prediction", "prediction_validation", "export_report", "research_package"}:
        return "Professional adds forecast validation, source audit, activity logs, and exportable research reports."
    return f"This feature requires {tier_label(required_tier(feature))} membership."


def locked_feature_response(feature, title=None, status_code=200):
    required = required_tier(feature)
    cta_label = f"Upgrade to {tier_label(required)}"
    locked_view = LOCKED_FEATURE_VIEWS.get(feature, {
        "icon": "lock",
        "eyebrow": "Workspace capability",
        "headline": f"Unlock {title or 'this feature'} when your workflow needs it.",
        "summary": locked_message(feature),
        "capabilities": [],
    })
    return render_template(
        "locked_feature.html",
        feature=feature,
        required_tier=required,
        required_label=tier_label(required),
        cta_label=cta_label,
        title=title or "Feature locked",
        message=locked_message(feature),
        locked_view=locked_view,
    ), status_code


def ensure_feature_access(feature, redirect_endpoint=None):
    if can_access(current_user(), feature):
        return None
    message = locked_message(feature)
    if redirect_endpoint:
        flash(message, "warning")
        return redirect(url_for(redirect_endpoint))
    return locked_feature_response(feature)


def normalize_membership_tier(value):
    tier = (value or "basic").strip().lower()
    if tier not in PLAN_BY_ID:
        return "basic"
    return tier


def is_valid_object_id(value):
    value = str(value or "").strip()
    if not value:
        return False
    return bool(re.fullmatch(r"[0-9a-fA-F]{24}|[0-9a-fA-F]{32}|[0-9a-fA-F-]{36}", value))


def is_safe_url(value):
    try:
        parsed = urlparse(str(value or "").strip())
    except Exception:
        return False
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def clean_search_query(value, max_length=120):
    text = re.sub(r"\s+", " ", str(value or "").strip())
    return text[:max_length]


SEARCH_SECURITY_PATTERNS = (
    "<script",
    "</script",
    "onerror=",
    "onload=",
    "javascript:",
    "' or '1'='1",
    " or 1=1",
    "$ne",
    "$gt",
    "$where",
)


def is_safe_search_query(value):
    text = str(value or "").strip()
    if not text or len(text) > 120:
        return False
    lowered = text.lower()
    suspicious_fragments = SEARCH_SECURITY_PATTERNS + ("$lt", "$regex", '{"$', "{'$")
    return not any(fragment in lowered for fragment in suspicious_fragments)


def is_blocked_search_query(value):
    text = clean_search_query(value)
    if not text or len(text) > 120:
        return True
    lowered = text.lower()
    suspicious_fragments = SEARCH_SECURITY_PATTERNS + ("$lt", "$regex", '{"$', "{'$")
    return any(fragment in lowered for fragment in suspicious_fragments)


def redact_search_query(value, limit=48):
    text = clean_search_query(value, max_length=limit)
    return re.sub(r"[\r\n\t]+", " ", text) if text else "[redacted]"


def record_blocked_search_attempt(user_id, query, reason):
    repository.log_event(user_id, "blocked_search_input", session.get("role"), session.get("username"), {
        "event_type": "blocked_search_input",
        "redacted_query": redact_search_query(query),
        "reason": reason,
        "action": "rejected",
    })


def display_search_source(value):
    source = str(value or "").strip().lower()
    mapping = {
        "ebay": "eBay",
        "walmart": "Walmart",
        "all": "Both sources",
        "mixed": "Both sources",
        "market": "Marketplace records",
        "market_records": "Marketplace records",
        "mongodb": "Marketplace records",
        "memory_fallback": "Marketplace records",
        "demo": "Marketplace records",
        "mock": "Marketplace records",
    }
    return mapping.get(source, "Marketplace records")


def display_search_status(record):
    status = str((record or {}).get("status") or "").strip().lower()
    result_count = int((record or {}).get("result_count") or 0)
    if status == "blocked":
        return "Blocked"
    if status == "failed":
        return "Failed"
    if status == "source_unavailable":
        return "Source unavailable"
    if status == "partial_success":
        return "Partial success"
    if result_count == 0:
        return "No results"
    if status in {"running", "completed", "complete"}:
        return "Completed"
    return "Completed" if result_count > 0 else "No results"


def empty_marketplace_search_notice(search_scope, diagnostics=None):
    diagnostics = diagnostics or {}
    source_status = diagnostics.get("source_status") or {}
    ebay_status = str(source_status.get("ebay") or diagnostics.get("ebay_status") or "").lower()
    walmart_status = str(source_status.get("walmart") or diagnostics.get("walmart_status") or "").lower()
    ebay_unavailable = ebay_status.startswith("unavailable")
    walmart_unavailable = walmart_status.startswith("unavailable")

    if search_scope == "ebay":
        return "eBay is temporarily unavailable. Try again later." if ebay_unavailable else "No comparable listings were found on eBay."
    if search_scope == "walmart":
        return "Walmart is temporarily unavailable. Try again later." if walmart_unavailable else "No comparable listings were found on Walmart."
    if ebay_unavailable and walmart_unavailable:
        return "No comparable listings could be collected because eBay and Walmart are currently unavailable. Try again later."
    if walmart_unavailable:
        return "No comparable listings were found. Walmart is currently unavailable, and no matching eBay listings were found."
    if ebay_unavailable:
        return "No comparable listings were found. eBay is currently unavailable, and no matching Walmart listings were found."
    if int(diagnostics.get("raw_ebay_count") or 0) or int(diagnostics.get("raw_walmart_count") or 0):
        return "Listings were returned, but none matched the requested product or model closely enough for comparison."
    return "No comparable listings were found on eBay or Walmart."


def cleanup_search_history(user_id=None):
    removed = 0
    keywords = []
    for row in safe_call(lambda: repository.list_searches(user_id, limit=0), []):
        keyword = str(row.get("keyword") or "")
        selected_source = str(row.get("selected_source") or "").lower()
        if is_blocked_search_query(keyword) or selected_source in {"mongodb", "demo", "mock", "memory_fallback"}:
            repository.delete_search(row["_id"], deleted_by=user_id, reason="security_cleanup")
            removed += 1
            keywords.append(keyword)
    return {"removed": removed, "keywords": keywords}


def safe_display_query(value, fallback="Search analytics"):
    if not isinstance(value, str):
        return fallback
    text = re.sub(r"\s+", " ", value).strip()
    if not text:
        return fallback
    lowered = text.lower()
    suspicious_fragments = ("$ne", "$gt", "$lt", "$regex", '{"$', "{'$", "<script", "onerror=", "javascript:")
    if any(fragment in lowered for fragment in suspicious_fragments):
        return "Filtered search result"
    return text[:80]


def analytics_display_query(record):
    raw_query = record.get("query")
    if isinstance(raw_query, str):
        display_query = safe_display_query(raw_query, fallback="Market analytics")
        if display_query == "Filtered search result":
            return display_query
        if display_query and display_query != "Search analytics":
            return display_query
    elif raw_query not in (None, "", [], (), {}):
        return "Filtered search result"

    fallback_query = record.get("keyword") or record.get("product_name")
    fallback_display = safe_display_query(fallback_query, fallback="Market analytics")
    return fallback_display or "Market analytics"


def make_json_safe(value):
    if isinstance(value, ObjectId):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, dict):
        return {str(key): make_json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [make_json_safe(item) for item in value]
    return value


def singapore_time_string(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime(value.year, value.month, value.day)
    else:
        return str(value)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S SGT")


def export_forbidden_response(message):
    required = "professional" if "Professional" in message else "premium"
    return render_template(
        "locked_feature.html",
        feature="export",
        required_tier=required,
        required_label=tier_label(required),
        cta_label=f"Upgrade to {tier_label(required)}",
        title="Export locked",
        message=message,
        locked_view=LOCKED_FEATURE_VIEWS["export"],
    ), 403


def export_allowed(member, required):
    return membership_rank(normalize_membership_tier(member.get("membership_tier") if member else "basic")) >= membership_rank(required)


def sanitize_csv_cell(value):
    """Neutralize spreadsheet formulas while preserving non-text cell types."""
    if not isinstance(value, str):
        return value
    first_visible = value.lstrip()
    if first_visible.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def csv_response(rows, filename):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    rows = list(rows or [])
    for row in rows:
        writer.writerow([sanitize_csv_cell(value) for value in row])
    content = buffer.getvalue()
    buffer.close()
    return Response(content, mimetype="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _current_membership_tier():
    user = current_user() or {}
    return normalize_membership_tier(user.get("membership_tier") or user.get("plan"))


@app.context_processor
def inject_product_context():
    return {"current_user": current_user(), "storage_status": repository.status(), "role_names": ROLES, "developer_role_preview": app.config["DEVELOPER_ROLE_PREVIEW"], "demo_tools_enabled": app.config["DEMO_TOOLS_ENABLED"], "demo_membership_upgrade_enabled": app.config["DEMO_MEMBERSHIP_UPGRADE_ENABLED"] and _is_non_production_runtime(), "can_access": can_access, "membership_plans": MEMBERSHIP_PLANS, "required_tier": required_tier, "normalize_membership_tier": normalize_membership_tier, "tier_label": tier_label, "membership_rank": membership_rank, "is_safe_url": is_safe_url, "display_search_source": display_search_source, "display_search_status": display_search_status, "user_initials": user_initials, "avatar_url": avatar_url, "is_test_account_row": is_test_account_row}


def login_required(view):
    from functools import wraps

    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user():
            flash("Please sign in to open the workspace.", "info")
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def role_required(*allowed):
    def decorator(view):
        from functools import wraps

        @wraps(view)
        @login_required
        def wrapped(*args, **kwargs):
            if session.get("active_role") not in allowed:
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


def _is_non_production_runtime():
    return not is_production_environment({"APP_ENV": str(app.config.get("RUNTIME_ENVIRONMENT") or "development")})


def demo_tools_required(view):
    """Allow demo mutations only for approved operators outside production."""
    from functools import wraps

    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        user = current_user() or {}
        approved_operator = (user.get("active_role") or user.get("role")) == "administrator" or user.get("is_developer")
        if not app.config.get("DEMO_TOOLS_ENABLED") or not _is_non_production_runtime() or not approved_operator:
            abort(404)
        return view(*args, **kwargs)
    return wrapped


def audit_scope_user_id(user=None):
    """Administrators may audit globally; every other role is owner-scoped."""
    user = user or current_user() or {}
    if (user.get("active_role") or user.get("role")) == "administrator":
        return None
    return user.get("_id")


def normalize_roles(role_list):
    if not role_list:
        return ["consumer"]
    ordered = []
    for role in role_list:
        if role and role not in ordered and role in ROLES:
            ordered.append(role)
    return ordered or ["consumer"]


def user_roles(user):
    return normalize_roles(user.get("roles") or ([user.get("role")] if user and user.get("role") else ["consumer"]))


def user_active_role(user):
    roles = user_roles(user)
    active_role = user.get("active_role") or user.get("role") or roles[0]
    return active_role if active_role in roles else roles[0]


def single_account_role(user):
    """Return the one canonical workspace role assigned to an account."""
    candidate = (user or {}).get("active_role") or (user or {}).get("role") or (user or {}).get("primary_role")
    return candidate if candidate in ROLES else "consumer"


def is_accessory(title):
    text = (title or "").lower()
    return any(keyword in text for keyword in ACCESSORY_KEYWORDS)


def filter_item_type(items, item_type):
    if item_type == "product_only":
        return [item for item in items if not is_accessory(item.get("title"))]
    if item_type == "accessories_only":
        return [item for item in items if is_accessory(item.get("title"))]
    return items


def get_service_health():
    mongo = repository.status()
    ebay_configured = bool(CLIENT_ID and CLIENT_SECRET)
    try:
        walmart_available = WalmartCollector is not None
    except Exception:
        walmart_available = False
    try:
        gemini_configured = bool(os.getenv("GEMINI_API_KEY"))
    except Exception:
        gemini_configured = False
    try:
        fallback_available = callable(summarize_market)
    except Exception:
        fallback_available = False
    return [
        {
            "name": "MongoDB Storage",
            "status": mongo["status"],
            "detail": mongo["detail"],
        },
        {
            "name": "eBay API",
            "status": "configured" if ebay_configured else "not configured",
            "detail": "Credentials configured; live search is checked during request." if ebay_configured else "Credentials are missing.",
        },
        {
            "name": "Walmart Source",
            "status": "available" if walmart_available else "unavailable",
            "detail": "Collected source; not presented as an official API.",
        },
        {
            "name": "Gemini AI Discover",
            "status": "configured" if gemini_configured else "not configured",
            "detail": "Analysis service, not a marketplace platform.",
        },
        {
            "name": "Rule-based fallback",
            "status": "available" if fallback_available else "unavailable",
            "detail": "Local summary used when Gemini is unavailable.",
        },
    ]


def safe_call(callable_fn, default):
    try:
        return callable_fn()
    except Exception:
        return default


SGT = timezone(timedelta(hours=8))


def format_sgt_datetime(value):
    if not value:
        return "—"
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return value
    if getattr(value, "tzinfo", None) is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(SGT).strftime("%d %b %Y, %H:%M SGT")


@app.context_processor
def inject_datetime_helpers():
    return {"format_sgt_datetime": format_sgt_datetime}


def display_sgt_datetime(value, fallback="—"):
    if not value:
        return fallback
    formatted = format_sgt_datetime(value)
    return formatted if formatted else fallback


def group_search_history(rows):
    """Collapse repeated searches into a compact current-user history without deleting records."""
    grouped = {}
    for row in rows or []:
        keyword = " ".join(str(row.get("keyword") or "").split())
        source = str(row.get("selected_source") or "all")
        key = (keyword.lower(), source.lower())
        entry = grouped.setdefault(key, {"row": row, "count": 0})
        entry["count"] += 1
    return [entry for entry in grouped.values()]


@app.context_processor
def inject_display_datetime_helper():
    return {"display_sgt_datetime": display_sgt_datetime, "group_search_history": group_search_history}


def _log_category(action_text, message_text="", source_text=""):
    action = re.sub(r"[^a-z0-9]+", "_", str(action_text or "").lower()).strip("_")
    context = f"{message_text} {source_text}".lower()
    if any(term in action for term in ("error", "failed", "warning")) or any(term in context for term in ("error", "failed", "warning", "fallback")):
        return "Errors / warnings"
    if action.startswith("save") or "evidence" in action or any(term in action for term in ("research_saved", "research_deleted")):
        return "Saved evidence"
    if any(term in action for term in ("ai_discover", "insight", "analysis", "gemini")):
        return "AI logs"
    if any(term in action for term in ("login", "logout", "register", "password", "archive", "role", "status", "user", "admin", "testing_records")):
        return "Admin actions"
    if any(term in action for term in ("search", "records_collected", "source_collected")):
        return "Search records"
    if any(term in action for term in ("compare", "comparison", "watchlist", "prediction", "forecast", "snapshot", "export")):
        return "Workflow events"
    return "Other events"


def _audit_source_display(value):
    normalized = str(value or "Internal").strip().lower()
    return {"both": "eBay and Walmart", "mixed": "eBay and Walmart", "ebay": "eBay", "walmart": "Walmart", "internal": "Internal"}.get(normalized, str(value or "Internal"))


def _audit_status_display(value):
    normalized = str(value or "recorded").strip().lower()
    return {
        "completed": "Completed", "complete": "Completed", "partial_success": "Partial",
        "failed": "Failed", "source_unavailable": "Failed", "no_results": "No results",
        "recorded": "Recorded", "running": "Recorded",
    }.get(normalized, "Recorded")


def _marketplace_status_display(value):
    raw = value.get("status") if isinstance(value, dict) else value
    normalized = str(raw or "not_requested").strip().lower().replace(" ", "_")
    return {
        "success": "Success", "cached": "Success", "live": "Success",
        "no_results": "No results", "no_result": "No results",
        "success_no_comparable_results": "No comparable results",
        "query_correction_suggested": "Query correction suggested",
        "timeout": "Timeout", "rate_limited": "Rate limited",
        "provider_error": "Provider error", "normalization_failure": "Provider error",
        "normalization_failed": "Normalization failed",
        "auth_error": "Provider error", "authentication_error": "Authentication failed", "unavailable": "Provider error",
        "not_requested": "Not requested",
    }.get(normalized, str(raw or "Not requested").replace("_", " ").title())


def _record_user_display(row):
    direct = row.get("display_name") or row.get("username") or row.get("user_email")
    if direct:
        return str(direct)
    user_id = row.get("user_id")
    user = safe_call(lambda: repository.get_user_by_id(user_id), None) if user_id else None
    return str((user or {}).get("display_name") or (user or {}).get("username") or "System")


def _workspace_display(row):
    value = row.get("workspace") or row.get("active_role") or row.get("role") or "Precision Curator"
    return str(value).replace("_", " ").title()


def _audit_details_text(value):
    if isinstance(value, dict):
        return "; ".join(f"{str(key).replace('_', ' ').title()}: {item}" for key, item in value.items() if item not in (None, "", [], {})) or "Recorded"
    return str(value or "Recorded")


def _audit_export_filters():
    return {
        "record_type": str(request.args.get("record_type") or request.args.get("tab") or "all").strip().lower(),
        "query": clean_search_query(request.args.get("q", "")),
        "date_from": str(request.args.get("date_from") or "").strip(),
        "date_to": str(request.args.get("date_to") or "").strip(),
    }


def _filter_audit_records(records, filters):
    record_type = filters.get("record_type") or "all"
    query = (filters.get("query") or "").lower()
    date_from, date_to = filters.get("date_from") or "", filters.get("date_to") or ""
    category_map = {"search": "search", "evidence": "evidence", "ai": "ai", "source": "source"}
    filtered = []
    for row in records:
        category = str(row.get("category") or "").lower()
        if record_type != "all" and category_map.get(record_type, record_type) not in category:
            continue
        haystack = " ".join(str(row.get(key) or "") for key in ("query", "user", "role", "source", "status", "detail")).lower()
        if query and query not in haystack:
            continue
        day = singapore_time_string(row.get("time"))[:10]
        if date_from and day and day < date_from:
            continue
        if date_to and day and day > date_to:
            continue
        filtered.append(row)
    return filtered


def _dated_export_filename(prefix):
    day = utcnow().astimezone(timezone(timedelta(hours=8))).strftime("%Y-%m-%d")
    return f"{prefix}-{day}.csv"


def _log_time_value(row):
    value = row.get("created_at") or row.get("timestamp")
    if hasattr(value, "isoformat"):
        return value
    return value


def _record_sort_key(row):
    value = row.get("raw_time")
    return value if isinstance(value, datetime) else datetime.min.replace(tzinfo=timezone.utc)


TEST_ACCOUNT_MARKERS = ("test", "demo", "sample", "mock", "fixture", "qa")


def is_test_account_row(user):
    if not user:
        return False
    if user.get("is_test_account"):
        return True
    text_parts = " ".join(str(user.get(field) or "") for field in ("display_name", "username", "email")).lower()
    if not text_parts.strip():
        return False
    if str(user.get("email") or "").lower().endswith(".test"):
        return True
    return any(marker in text_parts for marker in TEST_ACCOUNT_MARKERS)


def hide_test_accounts(users):
    return [user for user in users if not is_test_account_row(user)]


def _dashboard_price_value(row):
    for key in ("total_price", "normalized_price", "price"):
        value = row.get(key)
        if value is None:
            continue
        try:
            return float(value)
        except (TypeError, ValueError):
            continue
    return None


SHOW_DEMO_FALLBACK = os.getenv("SHOW_DEMO_FALLBACK", "false").strip().lower() in {"1", "true", "yes", "on"}
MOCK_SEARCH_MODE = os.getenv("MOCK_SEARCH_MODE", "false").strip().lower() in {"1", "true", "yes", "on"}
SHOW_SEARCH_DIAGNOSTICS = os.getenv("SHOW_SEARCH_DIAGNOSTICS", "false").strip().lower() in {"1", "true", "yes", "on"}
RELEVANCE_CONFIDENCE_THRESHOLD = float(os.getenv("SEARCH_CONFIDENCE_THRESHOLD", "0.50"))


def _normalize_search_text(value):
    return " ".join(re.findall(r"[a-z0-9]+", str(value or "").lower()))


def _compact_search_text(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _search_record_text(row):
    parts = [
        row.get("title"),
        row.get("product_name"),
        row.get("brand"),
        row.get("model"),
        row.get("normalized_title"),
        row.get("variant"),
        row.get("variation"),
        row.get("subtitle"),
        row.get("category"),
        row.get("seller"),
        row.get("availability"),
        row.get("condition"),
    ]
    return " ".join(part for part in (_normalize_search_text(part) for part in parts) if part)


def _search_terms(query):
    return [term for term in re.findall(r"[a-z0-9]+", str(query or "").lower()) if len(term) > 1]


def _has_token(text, token):
    text = f" {str(text or '').lower()} "
    token = str(token or "").lower().strip()
    if not token:
        return False
    return re.search(rf"(?<![a-z0-9]){re.escape(token)}(?![a-z0-9])", text) is not None


_ATTRIBUTE_TERMS = {
    "black", "white", "midnight", "purple", "blue", "red", "orange", "yellow", "green", "pink", "silver", "gold",
    "unlocked", "locked", "att", "verizon", "tmobile", "t-mobile", "sprint", "esim", "dual", "sim", "new", "used",
    "good", "very", "excellent", "refurbished", "condition", "open", "box", "boxed", "storage", "gb", "tb",
    "64gb", "128gb", "256gb", "512gb", "1tb", "2tb",
}
_ACCESSORY_TERMS = {"case", "cover", "charger", "cable", "screen protector", "lens protector", "accessories", "accessory", "strap", "adapter", "holder"}


def _is_accessory_query(query):
    normalized = _normalize_search_text(query)
    return any(keyword in normalized for keyword in _ACCESSORY_TERMS | ACCESSORY_KEYWORDS)


def _is_phone_like_query(query):
    compact = _compact_search_text(query)
    return any(term in compact for term in ("iphone", "galaxy", "pixel", "phone", "smartphone"))


def _is_product_like_query(query):
    terms = _search_terms(query)
    if len(terms) >= 2:
        return True
    compact = _compact_search_text(query)
    if re.search(r"(iphone|galaxy|airpods|pixel|ipad|macbook|watch)\d+", compact):
        return True
    return compact not in {"phone", "smartphone", "laptop", "tablet", "headphones", "earbuds", "case", "charger", "cable"}


def parse_canonical_product_model(value):
    """Parse query/title text once for every marketplace adapter.

    Edition checks deliberately use token boundaries and longest-first order so
    that ``Pro Max`` can never be mistaken for ``Pro``.
    """
    normalized = _normalize_search_text(value)
    compact = _compact_search_text(value)
    brand = family = generation = None
    if re.search(r"\b(?:apple\s+)?iphone\b", normalized):
        brand, family = "Apple", "iPhone"
        match = re.search(r"\biphone\s*(\d{1,2}|se)\b", normalized)
        generation = match.group(1).upper() if match else None
    elif re.search(r"\b(?:samsung\s+)?galaxy\b", normalized):
        brand, family = "Samsung", "Galaxy"
        match = re.search(r"\bgalaxy\s*(s\d{1,2}|z\s*(?:fold|flip)\s*\d*)\b", normalized)
        generation = re.sub(r"\s+", " ", match.group(1)).upper() if match else None
    elif re.search(r"\b(?:google\s+)?pixel\b", normalized):
        brand, family = "Google", "Pixel"
        match = re.search(r"\bpixel\s*(\d{1,2})\b", normalized)
        generation = match.group(1) if match else None

    if re.search(r"\bpro\s+max\b", normalized) or "promax" in compact:
        edition = "Pro Max"
    elif re.search(r"\bpro\b", normalized):
        edition = "Pro"
    elif re.search(r"\bair\b", normalized):
        edition = "Air"
    elif re.search(r"\b(?:se|e)\b", normalized) and family == "iPhone":
        edition = "SE/e"
    elif re.search(r"\bplus\b", normalized):
        edition = "Plus"
    else:
        edition = "Base"

    storage = None
    storage_match = re.search(r"(?<![a-z0-9])(\d+(?:\.\d+)?)\s*(tb|gb)(?![a-z0-9])", normalized)
    if storage_match:
        amount = storage_match.group(1).rstrip("0").rstrip(".") if "." in storage_match.group(1) else storage_match.group(1)
        storage = f"{amount}{storage_match.group(2).upper()}"
    carrier = None
    for pattern, label in (
        (r"\bunlocked\b", "Unlocked"), (r"\bverizon\b", "Verizon"),
        (r"\b(?:at\s*&\s*t|att)\b", "AT&T"), (r"\bt\s*mobile\b", "T-Mobile"),
        (r"\bsprint\b", "Sprint"),
    ):
        if re.search(pattern, normalized):
            carrier = label
            break

    if any(re.search(rf"\b{re.escape(term)}\b", normalized) for term in ("case", "cover", "charger", "cable", "adapter", "protector", "strap", "holder")):
        product_type = "accessory"
    elif re.search(r"\b(?:plan|contract|installment|monthly)\b", normalized):
        product_type = "plan"
    elif re.search(r"\b(?:replacement|replacement part|spare part|screen only|housing only)\b", normalized):
        product_type = "part"
    elif family:
        product_type = "device"
    else:
        product_type = "unknown"
    return {
        "brand": brand, "product_family": family, "generation": generation,
        "edition": edition, "storage": storage, "carrier": carrier,
        "product_type": product_type, "normalized": normalized, "compact": compact,
        # Compatibility aliases used by older display code/tests.
        "family": family.lower() if family else None,
        "pro": edition in {"Pro", "Pro Max"}, "plus": edition == "Plus",
        "max": edition == "Pro Max", "air": edition == "Air",
    }


def _query_model_parts(query):
    return parse_canonical_product_model(query)


def _record_model_parts(row):
    return parse_canonical_product_model(_search_record_text(row))


def _classify_search_match(row, query, score):
    flags = []
    query_parts = _query_model_parts(query)
    record_parts = _record_model_parts(row)
    text = record_parts["normalized"]
    terms = _search_terms(query)
    compact_query = query_parts["compact"]
    compact_text = record_parts["compact"]
    query_tokens = set(query_parts["normalized"].split())

    if not text or not any(term in text or term in compact_text for term in terms if term not in _ATTRIBUTE_TERMS):
        flags.append("keyword_mismatch")
    row["parsed_model"] = {key: record_parts.get(key) for key in ("brand", "product_family", "generation", "edition", "storage", "carrier", "product_type")}
    if _is_phone_like_query(query) and not _is_accessory_query(query) and record_parts["product_type"] == "accessory":
        flags.append("accessory")
    elif query_parts["family"] and record_parts["product_type"] in {"plan", "part"}:
        flags.append(f"product_type_{record_parts['product_type']}")
    if row.get("analytics_eligible") is False:
        flags.append("installment_price" if row.get("price_type") in NON_COMPARABLE_PRICE_TYPES else "invalid_price")
    family = query_parts["family"]
    if family:
        if record_parts["family"] != family or not record_parts["generation"]:
            flags.append("missing_model")
        elif query_parts["generation"] and query_parts["generation"] != record_parts["generation"]:
            flags.append("missing_model")
        elif query_parts["edition"] != record_parts["edition"]:
            suffix = {"Pro Max": "pro_max", "Air": "air", "Base": "base", "Pro": "pro", "Plus": "plus", "SE/e": "se"}.get(record_parts["edition"], "variant")
            flags.append(f"edition_conflict_{suffix}")

    if score < RELEVANCE_CONFIDENCE_THRESHOLD:
        flags.append("low_confidence")
    if row.get("demo_mode") or row.get("data_mode") == "synthetic":
        flags.append("demo_data")

    edition_conflict = any(flag.startswith("edition_conflict_") for flag in flags)
    if "accessory" in flags:
        match_type = "accessory_risk"
        label = "Accessory risk"
    elif "missing_model" in flags or "low_confidence" in flags or "keyword_mismatch" in flags or "invalid_price" in flags or "installment_price" in flags or any(flag.startswith("product_type_") for flag in flags):
        match_type = "low_confidence"
        label = "Low confidence"
    elif edition_conflict:
        match_type = "related_variant"
        label = "Related variant"
    elif compact_query and compact_query in compact_text:
        match_type = "exact_match"
        label = "Exact match"
    elif query_parts["family"] and record_parts["family"] == query_parts["family"] and query_parts["generation"] == record_parts["generation"]:
        match_type = "exact_match"
        label = "Exact match"
    elif not family and score >= 0.78:
        match_type = "exact_match"
        label = "Exact match"
    else:
        match_type = "low_confidence"
        label = "Low confidence"
        if "low_confidence" not in flags:
            flags.append("low_confidence")
    return match_type, label, flags


def _match_rank(row):
    return {"exact_match": 0, "related_variant": 1, "accessory_risk": 2, "low_confidence": 3}.get(row.get("match_type"), 4)


def _score_search_match(row, query):
    terms = _search_terms(query)
    if not terms:
        return 0.0
    text = _search_record_text(row)
    if not text:
        return 0.0
    normalized_query = _normalize_search_text(query)
    compact_query = _compact_search_text(query)
    compact_text = _compact_search_text(text)
    exact = 1.0 if normalized_query and normalized_query in text else 0.0
    compact_exact = 1.0 if compact_query and compact_query in compact_text else 0.0
    matches = sum(1 for term in terms if term in text or term in compact_text)
    coverage = matches / len(terms)
    confidence = row.get("confidence")
    try:
        confidence = float(confidence) if confidence is not None else None
    except (TypeError, ValueError):
        confidence = None
    confidence_score = confidence if confidence is not None else (0.8 if str(row.get("confidence_level") or "").lower() in {"high", "medium"} else 0.0)
    return max(exact, compact_exact, coverage * 0.7 + confidence_score * 0.3)


def _filter_relevant_search_items(items, query, *, confidence_threshold=RELEVANCE_CONFIDENCE_THRESHOLD):
    terms = _search_terms(query)
    if not terms:
        return items
    filtered = []
    for row in items:
        score = _score_search_match(row, query)
        row["match_confidence"] = round(score, 2)
        match_type, label, flags = _classify_search_match(row, query, score)
        row["warning_flags"] = flags
        row["match_type"] = match_type
        row["match_label"] = label
        row["product_fit"] = {"exact_match": "Exact model", "related_variant": "Related variant", "accessory_risk": "Possible accessory", "low_confidence": "Review needed"}.get(match_type, "Review needed")
        if match_type == "exact_match":
            filtered.append(row)
    return filtered


def _filter_relevant_search_items_with_rejections(items, query, *, confidence_threshold=RELEVANCE_CONFIDENCE_THRESHOLD):
    terms = _search_terms(query)
    accepted, rejected = [], []
    if not terms:
        return items, rejected
    for row in items:
        score = _score_search_match(row, query)
        row["match_confidence"] = round(score, 2)
        match_type, label, flags = _classify_search_match(row, query, score)
        row["warning_flags"] = flags
        row["match_type"] = match_type
        row["match_label"] = label
        row["product_fit"] = {"exact_match": "Exact model", "related_variant": "Related variant", "accessory_risk": "Possible accessory", "low_confidence": "Review needed"}.get(match_type, "Review needed")
        include_related = bool(row.get("_include_related_variants"))
        if match_type == "exact_match" or (include_related and match_type == "related_variant"):
            accepted.append(row)
        else:
            rejected_row = dict(row)
            rejected_row["rejection_reason"] = ", ".join(flags or ["low_confidence"])
            rejected.append(rejected_row)
    return accepted, rejected


def _apply_search_relevance(rows, query):
    return _filter_relevant_search_items(normalize_price_items(rows), query)


def _normalize_search_scope(value):
    mapping = {"all": "both", "both platforms": "both", "mongodb": "both", "market_records": "both", "demo": "both"}
    value = str(value or "both").strip().lower()
    return mapping.get(value, value)


def _diagnostic_sample(rows):
    sample = []
    for row in list(rows or [])[:5]:
        sample.append({
            "title": row.get("title") or row.get("product_name"),
            "platform": row.get("platform"),
            "match_type": row.get("match_type"),
            "match_reason": row.get("match_label") or row.get("rejection_reason") or row.get("warning_flags"),
            "price": row.get("price"),
            "match_confidence": row.get("match_confidence"),
            "warning_flags": row.get("warning_flags"),
            "rejection_reason": row.get("rejection_reason"),
        })
    return sample


def _diagnostic_sample_n(rows, limit=3):
    return _diagnostic_sample(list(rows or [])[:limit])


def _new_search_diagnostics(query, search_scope):
    return {
        "query": query,
        "normalized_query": normalize_serpapi_query(query),
        "search_scope": search_scope,
        "force_refresh": False,
        "ebay_cache_hit": False,
        "walmart_cache_hit": False,
        "serpapi_status": "not_requested",
        "serpapi_error": None,
        "raw_ebay_count": 0,
        "raw_walmart_count": 0,
        "parsed_walmart_count": 0,
        "submitted_source_queries": {"ebay": query, "walmart": normalize_serpapi_query(query)},
        "api_response_status": {"ebay": None, "walmart": None},
        "normalized_ebay_count": 0,
        "normalized_walmart_count": 0,
        "walmart_valid_price_count": 0,
        "excluded_walmart_count": 0,
        "comparable_walmart_count": 0,
        "displayed_walmart_count": 0,
        "retained_result_ids": [],
        "comparable_result_ids": [],
        "filtered_result_ids": [],
        "visible_result_ids": [],
        "search_run_mode": "not_started",
        "page_load_mode": "direct",
        "result_id_source": "none",
        "walmart_provider_response_mode": "not_requested",
        "product_type_eligible_ebay_count": 0,
        "product_type_eligible_walmart_count": 0,
        "exact_ebay_count": 0,
        "exact_walmart_count": 0,
        "related_ebay_count": 0,
        "related_walmart_count": 0,
        "accessory_ebay_count": 0,
        "accessory_walmart_count": 0,
        "low_confidence_ebay_count": 0,
        "low_confidence_walmart_count": 0,
        "category_filtered_count": 0,
        "price_filtered_count": 0,
        "final_display_count": 0,
        "final_ebay_count": 0,
        "final_walmart_count": 0,
        "rejected_count_by_reason": {},
        "walmart_parser_rejection_counts": {},
        "walmart_rejected_count_by_reason": {},
        "first_raw_walmart_records": [],
        "first_normalized_ebay_records": [],
        "first_normalized_walmart_records": [],
        "first_rejected_records": [],
        "first_rejected_walmart_records": [],
        "first_records": [],
        "source_status": {"ebay": "not_requested", "walmart": "not_requested"},
        "ebay_status": "not_requested",
        "walmart_status": "not_requested",
        "walmart_error": None,
        "walmart_response_type": "not_requested",
        "raw_walmart_status": "not_requested",
    }


def _record_rejections(diagnostics, rejected):
    walmart_parser_counts = Counter(diagnostics.get("walmart_parser_rejection_counts") or {})
    counter = Counter(walmart_parser_counts)
    walmart_counter = Counter(walmart_parser_counts)
    for row in rejected:
        row_platform = canonical_marketplace_platform(row.get("platform"), row.get("source_type"))
        if row_platform == "Walmart":
            reasons = [row.get("rejection_reason_code") or _canonical_walmart_rejection_reason(row)]
        else:
            reasons = row.get("warning_flags") or [row.get("rejection_reason") or "unknown"]
        for reason in reasons:
            counter[reason] += 1
            if row_platform == "Walmart":
                walmart_counter[reason] += 1
    diagnostics["rejected_count_by_reason"] = dict(counter)
    diagnostics["walmart_rejected_count_by_reason"] = dict(walmart_counter)
    diagnostics["excluded_walmart_count"] = sum(walmart_counter.values())
    diagnostics["first_rejected_records"] = _diagnostic_sample_n(rejected)
    diagnostics["first_rejected_walmart_records"] = _diagnostic_sample_n([
        row for row in rejected
        if canonical_marketplace_platform(row.get("platform"), row.get("source_type")) == "Walmart"
    ])


WALMART_REJECTION_LABELS = {
    "accessory": "Accessory rather than the requested product",
    "installment_or_plan": "Installment or service-plan offer",
    "incompatible_model": "Different product model or variant",
    "category_mismatch": "Different product category",
    "duplicate": "Duplicate marketplace listing",
    "missing_title": "Missing product title",
    "missing_price": "Missing comparable price",
    "invalid_price": "Invalid comparable price",
    "missing_source_url": "Missing or invalid source link",
    "unsupported_record_shape": "Unsupported marketplace record",
}


def _canonical_walmart_rejection_reason(row):
    flags = {str(value or "").strip().lower() for value in row.get("warning_flags") or []}
    price_type = str(row.get("price_type") or "").strip().lower()
    if "accessory" in flags or row.get("match_type") == "accessory_risk":
        return "accessory"
    if price_type == "installment" or any("plan" in flag or "installment" in flag for flag in flags):
        return "installment_or_plan"
    if any(flag.startswith(("edition_conflict", "generation_conflict", "brand_conflict", "family_conflict")) for flag in flags):
        return "incompatible_model"
    if any("category" in flag for flag in flags):
        return "category_mismatch"
    return "incompatible_model"


def _prepare_walmart_rejection(row, reason=None):
    rejected = dict(row)
    reason = reason or _canonical_walmart_rejection_reason(rejected)
    rejected["rejection_reason_code"] = reason
    rejected["rejection_reason"] = reason
    rejected["exclusion_reason"] = WALMART_REJECTION_LABELS.get(reason, "Not comparable to the requested product")
    rejected["platform"] = "Walmart"
    return rejected


def _prepare_scope_exclusion(row, selected_category_key=None, selected_product_type=None):
    excluded = dict(row)
    role = str(excluded.get("product_role") or "")
    category_key = str(excluded.get("category_key") or "generic")
    product_type = str((excluded.get("attributes") or {}).get("product_type") or "")
    if role == "accessory":
        code, label = "accessory", "Accessory rather than the requested product"
    elif role in {"service_or_plan", "installment"}:
        code, label = "installment_or_plan", "Installment or service-plan offer"
    elif role == "replacement_part":
        code, label = "replacement_part", "Replacement part rather than a complete product"
    elif role != "complete_product":
        code, label = "product_role_mismatch", "Not verified as a complete product"
    elif selected_category_key and category_key != selected_category_key:
        code, label = "category_mismatch", "Different product category"
    elif selected_product_type and product_type != selected_product_type:
        code, label = "product_type_mismatch", "Different product type"
    else:
        code, label = "comparison_scope_mismatch", "Outside the selected comparison scope"
    excluded["rejection_reason_code"] = code
    excluded["rejection_reason"] = code
    excluded["exclusion_reason"] = label
    return excluded


def _excluded_walmart_snapshot(row):
    source_url = row.get("source_url") or row.get("item_url")
    return {
        "platform": "Walmart",
        "title": str(row.get("title") or row.get("product_name") or "Walmart listing")[:300],
        "product_name": str(row.get("product_name") or row.get("title") or "Walmart listing")[:300],
        "price": row.get("price"),
        "raw_price_text": str(row.get("raw_price_text") or "")[:80],
        "currency": str(row.get("currency") or "USD")[:8],
        "source_url": source_url if is_safe_url(source_url) else "",
        "item_url": source_url if is_safe_url(source_url) else "",
        "match_type": row.get("match_type"),
        "price_type": row.get("price_type"),
        "rejection_reason_code": row.get("rejection_reason_code") or _canonical_walmart_rejection_reason(row),
        "exclusion_reason": row.get("exclusion_reason") or WALMART_REJECTION_LABELS.get(
            row.get("rejection_reason_code") or _canonical_walmart_rejection_reason(row),
            "Not comparable to the requested product",
        ),
    }


def _deduplicate_search_items(items, rejected=None):
    kept, seen = [], set()
    for item in items:
        source_url = str(item.get("source_url") or "").strip().casefold()
        key = ("url", source_url) if source_url else (
            "listing",
            str(item.get("platform") or "").casefold(),
            re.sub(r"\W+", " ", str(item.get("title") or item.get("product_name") or "").casefold()).strip(),
            item.get("total_price") if item.get("total_price") is not None else item.get("price"),
        )
        if key in seen:
            if rejected is not None and canonical_marketplace_platform(item.get("platform"), item.get("source_type")) == "Walmart":
                rejected.append(_prepare_walmart_rejection(item, "duplicate"))
            continue
        seen.add(key)
        kept.append(item)
    return kept


def _canonical_source_statuses(query, search_scope, diagnostics, items):
    completed_at = utcnow()
    statuses = {}
    for source, platform_name in (("ebay", "eBay"), ("walmart", "Walmart")):
        requested = search_scope in {source, "both"}
        legacy = str((diagnostics.get("source_status") or {}).get(source) or "not_requested").lower().replace(" ", "_")
        spelling = diagnostics.get("provider_spelling_suggestion") if source == "walmart" else None
        raw_count = int(diagnostics.get(f"raw_{source}_count") or 0)
        if source == "walmart" and spelling:
            raw_count = int(diagnostics.get("corrected_walmart_raw_count") or raw_count)
        normalized_count = int(diagnostics.get(f"normalized_{source}_count") or 0)
        retained_count = sum(
            1 for item in items
            if canonical_marketplace_platform(item.get("platform"), item.get("source_type")) == platform_name
        )
        error_text = str(diagnostics.get(f"{source}_error") or diagnostics.get("serpapi_error") or "")
        safe_message = str(diagnostics.get(f"{source}_safe_message") or "")
        error_code = None
        if source == "walmart" and query and requested:
            lowered = error_text.casefold()
            if spelling:
                state = "query_correction_suggested"
            elif legacy.startswith("unavailable") or error_text:
                if "normalization_failed" in lowered:
                    state, error_code = "normalization_failed", "normalization_failed"
                elif "provider_connection_error" in lowered:
                    state, error_code = "provider_connection_error", "provider_connection_error"
                elif "timeout" in lowered:
                    state, error_code = "timeout", "timeout"
                elif "429" in lowered or "rate" in lowered:
                    state, error_code = "rate_limited", "rate_limited"
                elif any(term in lowered for term in ("api_key", "api key", "auth", "credential", "invalid key", "401", "403")):
                    state, error_code = "authentication_error", "authentication_error"
                else:
                    state, error_code = "provider_error", "provider_error"
            elif retained_count > 0:
                state = "success"
            elif raw_count == 0:
                state = "no_results"
            else:
                state = "success_no_comparable_results"
        elif not query or not requested:
            state = "not_requested"
            error_code = "not_requested"
        elif spelling:
            state = "query_correction_suggested"
        elif legacy in {"live", "cached", "matched", "matched_stored_evidence", "filtered_out"}:
            state = "normalization_error" if raw_count and not normalized_count else "success"
        elif legacy in {"no_results", "no_results_", "no_results_found"} or (not raw_count and legacy not in {"unavailable", "unavailable_bot_check"}):
            state = "no_results"
        else:
            lowered = error_text.casefold()
            if "provider_connection_error" in lowered:
                state, error_code = "provider_connection_error", "provider_connection_error"
            elif "timeout" in lowered:
                state, error_code = "timeout", "timeout"
            elif "429" in lowered or "rate" in lowered:
                state, error_code = "rate_limited", "rate_limited"
            elif any(term in lowered for term in ("api_key", "auth", "credential", "401", "403")):
                state, error_code = "authentication_error", "authentication_error"
            else:
                state, error_code = "provider_error", "provider_error"
        statuses[source] = {
            "status": state,
            "status_code": state,
            "submitted_query": query,
            "effective_query": query,
            "raw_count": raw_count,
            "parsed_count": int(diagnostics.get(f"parsed_{source}_count") or normalized_count),
            "normalized_count": normalized_count,
            "excluded_count": int(diagnostics.get(f"excluded_{source}_count") or 0),
            "comparable_count": int(diagnostics.get(f"comparable_{source}_count") or retained_count),
            "displayed_count": int(diagnostics.get(f"displayed_{source}_count") or retained_count),
            "retained_count": retained_count,
            "error_code": error_code,
            "safe_message": safe_message,
            "spelling_suggestion": spelling,
            "completed_at": completed_at,
        }
    return statuses


def _source_status_messages(source_statuses):
    walmart = (source_statuses or {}).get("walmart") or {}
    state = walmart.get("status")
    if state == "not_requested" or walmart.get("error_code") == "not_requested":
        return []
    if state == "success":
        return []
    if state == "no_results":
        return ["No Walmart listings were returned for this search."]
    if state == "success_no_comparable_results":
        return ["Walmart returned listings, but none met the current comparable-product criteria."]
    if state == "query_correction_suggested":
        return [f"Walmart suggested ‘{walmart.get('spelling_suggestion')}’. Search using the suggested spelling to view those results."]
    if state in {"timeout", "rate_limited", "authentication_error", "provider_connection_error", "provider_error", "normalization_error", "normalization_failed"}:
        return ["Walmart is temporarily unavailable. Results shown are limited to other available sources."]
    return []


def _edit_distance(left, right):
    left, right = left.casefold(), right.casefold()
    previous = list(range(len(right) + 1))
    for index, left_char in enumerate(left, 1):
        current = [index]
        for other_index, right_char in enumerate(right, 1):
            current.append(min(current[-1] + 1, previous[other_index] + 1, previous[other_index - 1] + (left_char != right_char)))
        previous = current
    return previous[-1]


def _search_spelling_suggestion(query, items, source_statuses):
    provider = ((source_statuses or {}).get("walmart") or {}).get("spelling_suggestion")
    exact_quality = sum(bool(re.search(rf"\b{re.escape(query)}\b", str(item.get("title") or ""), re.I)) for item in items)
    if provider and provider.casefold() != query.casefold():
        return {"query": provider, "kind": "related_brand" if exact_quality >= 2 else "provider", "subtle": exact_quality >= 2}
    if exact_quality >= 2 or len(query) < 5:
        return None
    curated = ("Apple", "Samsung", "Google Pixel", "Chanel", "Gucci", "iPhone", "smartphone", "laptop", "headphones", "camera")
    matches = [(candidate, _edit_distance(query, candidate)) for candidate in curated]
    candidate, distance = min(matches, key=lambda pair: pair[1])
    return {"query": candidate, "kind": "curated", "subtle": False} if distance <= 1 else None


def _serpapi_walmart_enabled():
    return SERPAPI_MARKETPLACE_ENABLED and bool(os.getenv("SERPAPI_API_KEY"))


def _legacy_walmart_scraper_enabled():
    return os.getenv("WALMART_LEGACY_SCRAPER", "false").lower() in {"1", "true", "on", "yes"} and (app.debug or _search_diagnostics_enabled())


def _serpapi_demo_safe_ttl(query):
    normalized = normalize_serpapi_query(query)
    return SERPAPI_DEMO_CACHE_TTL_HOURS if normalized in DEMO_SAFE_KEYWORDS else SERPAPI_CACHE_TTL_HOURS


def _serpapi_cache_params(query, platform, *, category="", min_price=None, max_price=None, sort="", location="", country="", limit=20):
    normalized_query = normalize_serpapi_query(query)
    engine = engine_for_platform(platform)
    params = {
        "provider": "serpapi",
        "engine": engine,
        "platform": platform,
        "q": normalized_query,
        "category": category or "",
        "min_price": "" if min_price is None else min_price,
        "max_price": "" if max_price is None else max_price,
        "sort": provider_sort_for_platform(platform, sort),
        "location": location or os.getenv("SERPAPI_LOCATION", ""),
        "country": country or os.getenv("SERPAPI_COUNTRY", "us"),
        "currency": os.getenv("SERPAPI_CURRENCY", "USD"),
        "store_id": os.getenv("SERPAPI_WALMART_STORE_ID", "") if engine == "walmart" else "",
        "limit": limit,
    }
    params["cache_key"] = build_serpapi_cache_key(params)
    return params


def _provider_cache_key(provider, platform, query, *, category="", min_price=None, max_price=None, sort="", limit=20):
    payload = {
        "provider": str(provider or "").lower(),
        "platform": str(platform or "").lower(),
        "q": normalize_serpapi_query(query),
        "category": str(category or "").strip().lower(),
        "min_price": "" if min_price is None else str(min_price),
        "max_price": "" if max_price is None else str(max_price),
        "sort": str(sort or "").strip().lower(),
        "limit": str(limit or ""),
    }
    return f"{payload['provider']}:" + hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


def search_ebay_cached(query, *, category="", min_price=None, max_price=None, sort="", limit=20, refresh_live=False):
    cache_key = _provider_cache_key("ebay", "ebay", query, category=category, min_price=min_price, max_price=max_price, sort=sort, limit=limit)
    now = utcnow()
    cache = repository.get_search_cache(cache_key)
    cache_expires_at = _normalized_utc_datetime(cache.get("expires_at")) if cache else None
    if cache and cache_expires_at and cache_expires_at > now and not refresh_live:
        records = normalize_price_items(cache.get("normalized_records") or [])
        for row in records:
            row["source_type"] = "local_cache"
            row["data_source_label"] = "local_cache"
            row["cached_from_source_type"] = "live_ebay"
            row["cache_key"] = cache_key
            row["cache_hit"] = True
            row["collected_at"] = row.get("collected_at") or cache.get("collected_at")
        return records, {"cache_key": cache_key, "cache_status": "hit", "source_type": "local_cache", "collected_at": cache.get("collected_at"), "expires_at": cache_expires_at, "record_count": len(records)}
    raw_rows = search_ebay_items(query, limit=limit)
    records = normalize_price_items(raw_rows)
    for row in records:
        row["source_type"] = "live_ebay"
        row["data_source_label"] = "live_ebay"
        row["cache_key"] = cache_key
        row["cache_hit"] = False
    expires_at = now + timedelta(hours=SERPAPI_CACHE_TTL_HOURS)
    repository.save_search_cache(cache_key, {
        "provider": "ebay",
        "platform": "ebay",
        "normalized_query": normalize_serpapi_query(query),
        "params": {"category": category or "", "min_price": min_price, "max_price": max_price, "sort": sort or "", "limit": limit},
        "raw_response": raw_rows,
        "normalized_records": records,
        "live_source_type": "live_ebay",
        "collected_at": now,
        "expires_at": expires_at,
        "ttl_hours": SERPAPI_CACHE_TTL_HOURS,
    })
    return records, {"cache_key": cache_key, "cache_status": "refresh_live" if refresh_live else "miss", "source_type": "live_ebay", "collected_at": now, "expires_at": expires_at, "record_count": len(records)}


def search_serpapi_cached(query, platform, *, category="", min_price=None, max_price=None, sort="", limit=20, refresh_live=False):
    params = _serpapi_cache_params(query, platform, category=category, min_price=min_price, max_price=max_price, sort=sort, limit=limit)
    cache_key = params["cache_key"]
    now = utcnow()
    live_source_type = "serpapi_walmart" if platform == "walmart" else "serpapi_live"
    cache = repository.get_search_cache(cache_key)
    cache_expires_at = _normalized_utc_datetime(cache.get("expires_at")) if cache else None
    if cache and cache_expires_at and cache_expires_at > now and not refresh_live:
        records = normalize_price_items(cache.get("normalized_records") or [])
        parser_diagnostics = cache.get("parser_diagnostics") or {}
        for row in records:
            row["source_type"] = "local_cache"
            row["data_source_label"] = "local_cache"
            row["cache_key"] = cache_key
            row["cache_hit"] = True
            row["cached_from_source_type"] = cache.get("live_source_type") or live_source_type
            row["collected_at"] = row.get("collected_at") or cache.get("collected_at")
        return records, {
            "cache_key": cache_key,
            "cache_status": "hit",
            "source_type": "local_cache",
            "collected_at": cache.get("collected_at"),
            "expires_at": cache_expires_at,
            "record_count": len(records),
            "serpapi_params": (cache.get("params") or {}),
            "spelling_suggestion": cache.get("spelling_suggestion"),
            "parser_diagnostics": parser_diagnostics,
        }
    live = call_serpapi_marketplace(
        query,
        platform=platform,
        category=category,
        min_price=min_price,
        max_price=max_price,
        sort=params["sort"],
        location=params["location"],
        country=params["country"],
        store_id=params["store_id"],
        limit=limit,
        no_cache=refresh_live,
    )
    records = normalize_price_items(live.get("normalized_records") or [])
    for row in records:
        row["source_type"] = live_source_type
        row["data_source_label"] = live_source_type
        row["cache_key"] = cache_key
        row["cache_hit"] = False
        row["collected_at"] = row.get("collected_at") or live.get("collected_at")
    expires_at = now + timedelta(hours=_serpapi_demo_safe_ttl(query))
    repository.save_search_cache(cache_key, {
        "provider": "serpapi",
        "engine": params["engine"],
        "platform": platform,
        "normalized_query": params["q"],
        "params": {k: v for k, v in params.items() if k != "cache_key"},
        "raw_response": live.get("raw_response"),
        "normalized_records": records,
        "live_source_type": live_source_type,
        "collected_at": live.get("collected_at") or now,
        "expires_at": expires_at,
        "ttl_hours": _serpapi_demo_safe_ttl(query),
        "spelling_suggestion": live.get("spelling_suggestion"),
        "parser_diagnostics": live.get("parser_diagnostics") or {},
    })
    return records, {
        "cache_key": cache_key,
        "cache_status": "refresh_live" if refresh_live else "miss",
        "source_type": live_source_type,
        "collected_at": live.get("collected_at"),
        "expires_at": expires_at,
        "record_count": len(records),
        "serpapi_params": live.get("request_params") or {},
        "spelling_suggestion": live.get("spelling_suggestion"),
        "parser_diagnostics": live.get("parser_diagnostics") or {},
    }


def _search_diagnostics_enabled():
    return SHOW_SEARCH_DIAGNOSTICS or MOCK_SEARCH_MODE or app.debug or os.getenv("DEBUG_SEARCH", "").lower() == "true"


def _search_ui_diagnostics_enabled():
    """Expose provider stages only for explicitly enabled development diagnostics."""
    runtime = str(app.config.get("RUNTIME_ENVIRONMENT") or _runtime_environment or "").strip().lower()
    return runtime == "development" and SHOW_SEARCH_DIAGNOSTICS


def _should_show_demo_fallback():
    return SHOW_DEMO_FALLBACK or MOCK_SEARCH_MODE


def _has_stored_walmart_evidence(query):
    evidence = safe_call(lambda: repository.list_evidence(session["user_id"], limit=200), [])
    query_norm = _compact_search_text(query)
    count = 0
    for row in evidence:
        source_type = str(row.get("source_type") or "").lower()
        platform = str(row.get("platform") or "").lower()
        if "walmart" not in source_type and platform != "walmart":
            continue
        if query_norm and query_norm not in _compact_search_text(row.get("query") or row.get("product_name") or row.get("title") or ""):
            continue
        count += 1
    return count


def _stored_walmart_evidence(query):
    evidence = safe_call(lambda: repository.list_evidence(session["user_id"], limit=200), [])
    query_norm = _compact_search_text(query)
    rows = []
    for row in evidence:
        source_type = str(row.get("source_type") or "").lower()
        platform = str(row.get("platform") or "").lower()
        if "walmart" not in source_type and platform != "walmart":
            continue
        if query_norm and query_norm not in _compact_search_text(row.get("query") or row.get("product_name") or row.get("title") or ""):
            continue
        item = dict(row)
        item["platform"] = "Walmart"
        item["source_type"] = "stored_walmart_evidence"
        item["data_mode"] = "stored"
        item["demo_mode"] = False
        item["source_label"] = "Stored Walmart evidence"
        item["data_source_label"] = "stored_walmart_evidence"
        item["price_numeric"] = item.get("normalized_price", item.get("price"))
        item["normalized_price"] = item.get("normalized_price", item.get("price"))
        rows.append(item)
    return rows


def _normalize_export_source_type(item):
    source_type = str(item.get("source_type") or "").lower()
    if source_type in {"live_ebay", "live_walmart", "serpapi_walmart", "serpapi_live", "local_cache", "stored_walmart_evidence", "demo_sample"}:
        return source_type
    platform = str(item.get("platform") or "").lower()
    data_mode = str(item.get("data_mode") or "").lower()
    if "demo" in source_type or data_mode == "synthetic":
        return "demo_sample"
    if platform == "ebay":
        return "live_ebay"
    if platform == "walmart":
        return "live_walmart"
    return source_type or "unknown"


def _record_source_label(item):
    source_type = _normalize_export_source_type(item)
    labels = {
        "live_ebay": "eBay Browse API",
        "live_walmart": "Walmart marketplace evidence",
        "serpapi_walmart": "Walmart via SerpAPI",
        "serpapi_live": "Marketplace evidence via SerpAPI",
        "stored_walmart_evidence": "Stored marketplace evidence",
        "local_cache": "Stored marketplace evidence",
    }
    return labels.get(source_type, "Stored marketplace evidence")


def _is_demo_market_record(row):
    source_type = str(row.get("source_type") or "").lower()
    data_mode = str(row.get("data_mode") or "").lower()
    platform = str(row.get("platform") or "").lower()
    return bool(row.get("demo_mode")) or data_mode == "synthetic" or "demo" in source_type or "mongodb" in platform


def _label_record_source(row, source_kind):
    if source_kind == "demo":
        row["source_type"] = row.get("source_type") or "Demo records"
    elif source_kind == "fallback":
        row["source_type"] = row.get("source_type") or "Fallback records"
    else:
        row["source_type"] = row.get("source_type") or "Collected records"
    row["data_source_label"] = "demo" if source_kind == "demo" else ("fallback" if source_kind == "fallback" else "live")
    row["demo_mode"] = bool(row.get("demo_mode")) or source_kind == "demo" or _is_demo_market_record(row)
    return row


def _market_records_for_identity(user_id, search_record=None, allow_fallback=True):
    saved_evidence = safe_call(lambda: repository.list_evidence(user_id, limit=100), [])
    if saved_evidence:
        return normalize_price_items(saved_evidence), "saved"

    search_record = search_record or {}
    latest_search = search_record if search_record.get("_id") else None
    if latest_search is None:
        latest_search = safe_call(lambda: repository.list_searches(user_id, limit=1), [])
        latest_search = latest_search[0] if latest_search else None
    if latest_search:
        if latest_search.get("result_tokens"):
            records = normalize_price_items(resolve_result_tokens(latest_search["result_tokens"]))
        elif latest_search.get("market_record_ids"):
            records = normalize_price_items(repository.get_market_records(latest_search["market_record_ids"]))
        else:
            records = normalize_price_items(repository.get_products(latest_search["_id"]))
        if records:
            return records, "search"
    if allow_fallback and _should_show_demo_fallback():
        return normalize_price_items(demo_search_items("", limit=40, platform="demo_all")), "demo"
    return [], "fallback"


def build_retailer_dashboard_view(user_id, monitor_id=None):
    def retailer_summary(records, distribution, prices):
        total = sum(distribution.values()) or 0
        distribution_rows = [
            {"platform": platform, "count": count, "percent": round((count / total) * 100, 1) if total else 0}
            for platform, count in distribution.most_common()
        ]
        price_range = round(max(prices) - min(prices), 2) if prices else None
        top = distribution_rows[0] if distribution_rows else None
        if not records:
            insight = "No active monitor snapshots are available yet. Create a monitor from a product search, then refresh it to collect current listings."
        elif top and len(distribution_rows) > 1:
            insight = f"{top['platform']} currently contributes most monitored listings, while the other source adds comparison coverage. Review the observed price range before selecting sourcing candidates."
        elif top:
            insight = f"{top['platform']} is the only visible source in this retailer view. Add another source before making a cross-platform sourcing decision."
        else:
            insight = "Platform coverage is limited. Collect or save more marketplace records before making sourcing decisions."
        return distribution_rows, price_range, insight

    def monitor_metrics(monitor):
        snapshots = safe_call(lambda: repository.list_price_snapshots(monitor["_id"], limit=90), [])
        normalized_snapshots = []
        for snapshot in snapshots:
            row = dict(snapshot)
            listing_rows = normalize_price_items(row.get("listing_records") or [])
            listing_prices = [_dashboard_price_value(item) for item in listing_rows]
            listing_prices = [value for value in listing_prices if value is not None]
            if _snapshot_float(row.get("average_price")) is None and listing_prices:
                row["average_price"] = round(sum(listing_prices) / len(listing_prices), 2)
                row["lowest_price"] = round(min(listing_prices), 2)
                row["highest_price"] = round(max(listing_prices), 2)
                row["record_count"] = len(listing_prices)
            if _snapshot_float(row.get("average_price")) is not None:
                normalized_snapshots.append(row)
        normalized_snapshots.sort(
            key=lambda row: _normalized_utc_datetime(row.get("collected_at") or row.get("created_at"))
            or datetime.min.replace(tzinfo=timezone.utc)
        )
        latest = normalized_snapshots[-1] if normalized_snapshots else None
        previous = normalized_snapshots[-2] if len(normalized_snapshots) > 1 else None
        current_average = _snapshot_float((latest or {}).get("average_price"))
        previous_average = _snapshot_float((previous or {}).get("average_price"))
        current_low = _snapshot_float((latest or {}).get("lowest_price"))
        current_high = _snapshot_float((latest or {}).get("highest_price"))
        change_amount = round(current_average - previous_average, 2) if current_average is not None and previous_average is not None else None
        change_percentage = round(change_amount / previous_average * 100, 2) if change_amount is not None and previous_average else None
        latest_listings = normalize_price_items((latest or {}).get("listing_records") or [])
        priced_listings = [
            (item, _dashboard_price_value(item))
            for item in latest_listings
            if item.get("analytics_eligible")
        ]
        priced_listings = [(item, value) for item, value in priced_listings if value is not None]
        best_listing, best_price = min(priced_listings, key=lambda pair: pair[1]) if priced_listings else (None, current_low)
        best_platform = (best_listing or {}).get("platform") or (latest or {}).get("best_platform")
        below_average_percent = (
            round((current_average - best_price) / current_average * 100, 1)
            if current_average and best_price is not None
            else None
        )
        if not normalized_snapshots:
            trend_state, trend_label = "empty", "Waiting for baseline"
        elif len(normalized_snapshots) == 1:
            trend_state, trend_label = "baseline", "Baseline collected"
        elif change_amount is not None and change_amount < 0:
            trend_state, trend_label = "falling", "Price falling"
        elif change_amount is not None and change_amount > 0:
            trend_state, trend_label = "rising", "Price rising"
        else:
            trend_state, trend_label = "stable", "Price stable"
        average_position = 50.0
        if current_low is not None and current_high is not None and current_high > current_low and current_average is not None:
            average_position = round((current_average - current_low) / (current_high - current_low) * 100, 1)
        return {
            "snapshots": normalized_snapshots,
            "snapshot_count": len(normalized_snapshots),
            "trend_chart": _watchlist_trend_chart_data(normalized_snapshots),
            "sparkline": [_snapshot_float(row.get("average_price")) for row in normalized_snapshots[-14:]],
            "current_low": current_low,
            "current_average": current_average,
            "current_high": current_high,
            "change_amount": change_amount,
            "change_percentage": change_percentage,
            "trend_state": trend_state,
            "trend_label": trend_label,
            "best_platform": best_platform,
            "best_price": best_price,
            "below_average_percent": below_average_percent,
            "average_position": average_position,
            "latest_listing_count": int((latest or {}).get("listing_count") or (latest or {}).get("record_count") or len(latest_listings)),
            "latest_collected_at": (latest or {}).get("collected_at") or (latest or {}).get("created_at"),
            "data_quality": (latest or {}).get("data_quality"),
        }

    try:
        monitored = _watchlist_items_for_user(user_id, include_archived=False)
        for monitor in monitored:
            monitor["dashboard"] = monitor_metrics(monitor)
        selected_monitor = next((row for row in monitored if str(row.get("_id")) == str(monitor_id)), None) if monitor_id else None
        scoped_monitors = [selected_monitor] if selected_monitor else monitored
        records = []
        for monitor in scoped_monitors:
            snapshot = monitor.get("latest_snapshot") or {}
            snapshot_records = snapshot.get("listing_records") or monitor.get("selected_records_snapshot") or []
            for row in normalize_price_items(snapshot_records):
                row["watchlist_id"] = str(monitor.get("_id"))
                row["snapshot_collected_at"] = snapshot.get("collected_at") or snapshot.get("created_at")
                records.append(row)
        # A listing can occur in multiple monitors; count it once using its stable URL/id and platform.
        deduped, seen = [], set()
        for row in records:
            listing_identity = row.get("listing_id") or row.get("item_id") or row.get("external_id")
            if not listing_identity:
                listing_identity = str(row.get("source_url") or row.get("item_url") or row.get("_id") or row.get("title") or "").strip().lower().rstrip("/")
            key = (str(row.get("platform") or "").lower(), str(listing_identity))
            if key not in seen:
                seen.add(key)
                deduped.append(row)
        visible_records = [row for row in deduped if row.get("analytics_eligible") and not _is_demo_market_record(row)]
        source_kind = "watchlist" if visible_records else "empty"
        market_message = "Market overview based on the latest snapshot for each active monitor." if visible_records else "No active Watchlist monitor snapshots are available yet."
        market_badge = "Latest monitor snapshots" if visible_records else "No monitor data"
        active_monitor_count = len(scoped_monitors)
        latest_snapshot_time = max((monitor.get("latest_snapshot", {}).get("collected_at") or monitor.get("latest_snapshot", {}).get("created_at") for monitor in scoped_monitors if monitor.get("latest_snapshot")), default=None)
        distribution = Counter(row.get("platform") or "Unknown" for row in visible_records if row.get("platform") and "mongodb" not in str(row.get("platform")).lower())
        prices = [_dashboard_price_value(row) for row in visible_records]
        prices = [price for price in prices if price is not None]
        categories = sorted({row.get("category") or "Other" for row in visible_records if row.get("category")})
        platforms = sorted({row.get("platform") for row in visible_records if row.get("platform") and "mongodb" not in str(row.get("platform")).lower()})
        summary = {
            "listings": len(visible_records),
            "lowest": round(min(prices), 2) if prices else None,
            "average": round(sum(prices) / len(prices), 2) if prices else None,
            "highest": round(max(prices), 2) if prices else None,
        }
        distribution_rows, price_range, sourcing_insight = retailer_summary(visible_records, distribution, prices)
        summary["range"] = price_range
        # Portfolio scope deliberately avoids price arithmetic across unrelated
        # products. Price KPIs exist only for one explicitly selected monitor.
        if selected_monitor is None:
            summary.update(lowest=None, average=None, highest=None, range=None)
        selected_metrics = selected_monitor.get("dashboard") if selected_monitor else None
        if selected_metrics:
            summary.update(
                lowest=selected_metrics.get("current_low"),
                average=selected_metrics.get("current_average"),
                highest=selected_metrics.get("current_high"),
                range=(
                    round(selected_metrics["current_high"] - selected_metrics["current_low"], 2)
                    if selected_metrics.get("current_high") is not None and selected_metrics.get("current_low") is not None
                    else None
                ),
            )
            if selected_metrics.get("best_platform") and selected_metrics.get("best_price") is not None:
                gap = selected_metrics.get("below_average_percent")
                sourcing_insight = (
                    f"{selected_metrics['best_platform']} currently has the lowest comparable listing at "
                    f"USD {selected_metrics['best_price']:.2f}"
                    + (f", {gap:.1f}% below the observed market average." if gap is not None else ".")
                )
        portfolio_metrics = {
            "active_monitors": len(monitored),
            "price_drops": sum((monitor.get("dashboard") or {}).get("trend_state") == "falling" for monitor in monitored),
            "needs_attention": sum(
                not monitor.get("latest_snapshot")
                or monitor.get("last_refresh_status") in {"failed", "partial"}
                for monitor in monitored
            ),
            "baselines": sum((monitor.get("dashboard") or {}).get("trend_state") == "baseline" for monitor in monitored),
        }
        return {
            "records": visible_records,
            "distribution": distribution,
            "distribution_rows": distribution_rows,
            "prices": summary,
            "platforms": platforms,
            "categories": categories,
            "market_message": market_message,
            "market_badge": market_badge,
            "source_kind": source_kind,
            "demo_records": 0,
            "active_monitor_count": active_monitor_count,
            "monitors": monitored,
            "selected_monitor_id": str(selected_monitor.get("_id")) if selected_monitor else "",
            "portfolio_scope": selected_monitor is None,
            "selected_monitor": selected_monitor,
            "selected_metrics": selected_metrics,
            "trend_chart": (selected_metrics or {}).get("trend_chart") or _watchlist_trend_chart_data([]),
            "portfolio_metrics": portfolio_metrics,
            "latest_snapshot_time": latest_snapshot_time,
            "scope_label": ((f"Based on {len(monitored)} active monitor{'s' if len(monitored) != 1 else ''} · {len(visible_records)} unique listings") if selected_monitor is None else f"{selected_monitor.get('product_label') or selected_monitor.get('keyword') or 'Selected monitor'} · {len(visible_records)} listings") if active_monitor_count else "No active monitoring data yet.",
            "sourcing_insight": sourcing_insight,
        }
    except Exception as exc:
        app.logger.warning("Retailer dashboard view failed (%s)", exc.__class__.__name__)
        fallback_records, fallback_distribution, prices = [], Counter(), []
        distribution_rows, price_range, sourcing_insight = retailer_summary([], fallback_distribution, prices)
        return {
            "records": fallback_records,
            "distribution": fallback_distribution,
            "distribution_rows": distribution_rows,
            "prices": {
                "listings": 0,
                "lowest": round(min(prices), 2) if prices else None,
                "average": round(sum(prices) / len(prices), 2) if prices else None,
                "highest": round(max(prices), 2) if prices else None,
                "range": price_range,
            },
            "platforms": [],
            "categories": [],
            "market_message": "No active Watchlist monitor snapshots are available yet.",
            "market_badge": "No monitor data",
            "source_kind": "empty",
            "demo_records": 0,
            "active_monitor_count": 0,
            "monitors": [],
            "selected_monitor_id": "",
            "portfolio_scope": True,
            "selected_monitor": None,
            "selected_metrics": None,
            "trend_chart": _watchlist_trend_chart_data([]),
            "portfolio_metrics": {"active_monitors": 0, "price_drops": 0, "needs_attention": 0, "baselines": 0},
            "latest_snapshot_time": None,
            "scope_label": "No active monitoring data yet.",
            "sourcing_insight": sourcing_insight,
        }


def _analytics_records_for_search(search_record_id, user_scope):
    record = repository.get_search(search_record_id, user_scope)
    if not record:
        return None, [], "fallback"
    evidence = safe_call(lambda: repository.list_evidence(record.get("user_id"), limit=200), [])
    evidence = [row for row in evidence if str(row.get("search_record_id")) == str(record.get("_id"))]
    if evidence:
        return record, normalize_price_items(evidence), "live"
    if record.get("result_tokens"):
        items = normalize_price_items(resolve_result_tokens(record["result_tokens"]))
    elif record.get("market_record_ids"):
        items = normalize_price_items(repository.get_market_records(record["market_record_ids"]))
    else:
        items = normalize_price_items(repository.get_products(search_record_id))
    if items:
        return record, items, "live"
    return record, normalize_price_items(demo_search_items(record.get("keyword", ""), limit=20, platform="demo_all")), "demo"


def _evidence_save_plan(selected_products, user_id, search_id):
    planned = []
    seen = set()
    excluded = []
    for product in selected_products:
        title = str(product.get("title") or product.get("product_name") or "").strip().lower()
        platform = str(product.get("platform") or "").strip().lower()
        dedupe_key = (title, platform)
        reason = None
        if product.get("is_accessory") or product.get("relevance_warning"):
            reason = "accessory/invalid"
        elif dedupe_key in seen:
            reason = "duplicate"
        if reason:
            excluded.append({"product": product, "reason": reason})
            continue
        seen.add(dedupe_key)
        planned.append(product)
    saved = 0
    for product in planned:
        _evidence_id, created = repository.save_evidence(user_id, search_id, product)
        saved += int(created)
        if not created:
            excluded.append({"product": product, "reason": "duplicate"})
    return saved, excluded


GENERIC_MONITOR_CATEGORY_VALUES = {"", "all", "marketplace", "general products", "generic", "not available"}


def _meaningful_monitor_category(value):
    return str(value or "").strip().lower() not in GENERIC_MONITOR_CATEGORY_VALUES


def _watchlist_category_metadata(record=None, items=None, monitor=None):
    """Resolve one authoritative category for creation, refresh and exports."""
    record = record or {}
    monitor = monitor or {}
    items = list(items or [])
    frozen_scope = monitor.get("frozen_scope") or {}

    key_candidates = [
        monitor.get("category_key"),
        frozen_scope.get("category_key"),
        record.get("selected_category_key"),
        record.get("detected_category_key"),
    ]
    key_candidates.extend(item.get("category_key") for item in items)
    category_key = next((str(value).strip() for value in key_candidates if _meaningful_monitor_category(value)), None)

    display_candidates = [
        monitor.get("category_display"),
        frozen_scope.get("category_display"),
        record.get("detected_category_display"),
    ]
    if category_key and category_key in CATEGORY_PROFILES:
        display_candidates.append(CATEGORY_PROFILES[category_key]["display_name"])
    display_candidates.extend(item.get("category_display") for item in items)
    display_candidates.extend([monitor.get("category"), frozen_scope.get("category")])
    category_display = next(
        (str(value).strip() for value in display_candidates if _meaningful_monitor_category(value)),
        None,
    )
    if not category_display and category_key:
        category_display = category_key.replace("_", " ").title()
    return {
        "category_key": category_key,
        "category_display": category_display or "General products",
        "category_query": category_display or category_key or "",
    }


def _monitor_with_category_metadata(monitor):
    monitor_view = dict(monitor or {})
    search_record = None
    if monitor_view.get("search_record_id") and monitor_view.get("user_id"):
        search_record = repository.get_search(monitor_view["search_record_id"], monitor_view["user_id"])
    category = _watchlist_category_metadata(
        record=search_record,
        items=monitor_view.get("selected_records_snapshot") or monitor_view.get("initial_records_snapshot") or [],
        monitor=monitor_view,
    )
    monitor_view.update(category)
    return monitor_view


def _watchlist_scope_from_products(products, record=None, tracking_mode="search_scope"):
    products = normalize_price_items(products)
    priced = [item for item in products if item.get("analytics_eligible") and item.get("total_price") not in (None, 0)]
    if not priced:
        return None
    keyword = (record.get("keyword") if record else None) or priced[0].get("keyword") or priced[0].get("title") or "Tracked products"
    platform_scope = sorted({item.get("platform") for item in priced if item.get("platform")}) or ["all"]
    category = _watchlist_category_metadata(record=record, items=priced)
    best_platform = Counter(item.get("platform") for item in priced if item.get("platform")).most_common(1)
    return {
        "tracking_mode": tracking_mode,
        "keyword": keyword,
        "product_label": keyword,
        "platform_scope": ", ".join(platform_scope),
        "category": category["category_query"] or "All",
        "category_key": category["category_key"],
        "category_display": category["category_display"],
        "condition_scope": Counter(item.get("availability") or item.get("condition") or "Not provided" for item in priced).most_common(1)[0][0],
        "source_scope": record.get("selected_source") if record else "search",
        "record_scope": "selected_comparison_records" if tracking_mode == "selected_records" else "all_current_search_results",
        "selected_record_ids": [str(item.get("_id")) for item in priced if item.get("_id")] if tracking_mode == "selected_records" else [],
        "selected_records_snapshot": [dict(item) for item in priced] if tracking_mode == "selected_records" else [],
        "data_source_label": "fallback" if any(item.get("data_mode") == "synthetic" for item in priced) else ("demo" if any(item.get("demo_mode") for item in priced) else "live"),
        "best_platform": best_platform[0][0] if best_platform else None,
        "record_count": len(priced),
        "items": priced,
    }


def _snapshot_payload_from_items(watchlist_item, items):
    items = [item for item in normalize_price_items(items) if item.get("analytics_eligible") and item.get("total_price") not in (None, 0)]
    if not items:
        return None
    prices = [float(item["total_price"]) for item in items]
    best_platform = Counter(item.get("platform") or "Unknown" for item in items).most_common(1)[0][0]
    data_source_label = "fallback" if any(item.get("data_mode") == "synthetic" for item in items) else ("demo" if any(item.get("demo_mode") for item in items) else "live")
    average_price = round(sum(prices) / len(prices), 2)
    return {
        "watchlist_id": watchlist_item["_id"],
        "monitor_id": watchlist_item.get("monitor_id") or watchlist_item["_id"],
        "comparison_set_id": watchlist_item.get("comparison_set_id") or watchlist_item.get("comparison_group_id"),
        "snapshot_id": secrets.token_hex(12),
        "selected_result_ids": list(watchlist_item.get("selected_record_ids") or []),
        "comparable_result_ids": [item.get("_selection_token") or item.get("_id") for item in items],
        "tracking_mode": watchlist_item.get("tracking_mode") or "search_scope",
        "keyword": watchlist_item["keyword"],
        "product_label": watchlist_item["product_label"],
        "platform_scope": watchlist_item.get("platform_scope"),
        "source_scope": watchlist_item.get("source_scope"),
        "record_count": len(items),
        "listing_count": len(items),
        "lowest_price": round(min(prices), 2),
        "highest_price": round(max(prices), 2),
        "average_price": average_price,
        "price_range": round(max(prices) - min(prices), 2),
        "best_platform": best_platform,
        "collected_at": utcnow(),
        "source_label": data_source_label,
        "data_quality": "good" if len(items) >= 3 else "limited",
        "listing_records": [
            {key: value for key, value in item.items() if key not in {"_id", "user_id", "search_record_id"}}
            for item in items
        ],
        "notes": None,
    }


def _watchlist_items_for_user(user_id, include_archived=False):
    items = safe_call(lambda: repository.list_watchlist_items(user_id, include_archived=include_archived, limit=100), [])
    for item in items:
        snapshots = safe_call(lambda item_id=item["_id"]: repository.list_price_snapshots(item_id, limit=50), [])
        item["snapshot_count"] = safe_call(lambda item_id=item["_id"]: repository.count_price_snapshots(item_id), len(snapshots))
        item["latest_snapshot"] = snapshots[0] if snapshots else None
    return items


def _watchlist_selected_item(items, selected_id=None):
    if not items:
        return None
    if selected_id:
        for item in items:
            if str(item.get("_id")) == str(selected_id):
                return item
    return max(items, key=lambda row: row.get("updated_at") or row.get("created_at") or datetime.min.replace(tzinfo=timezone.utc))


def _watchlist_redirect(item_id):
    return redirect(url_for("watchlist", item_id=item_id))


def _snapshot_float(value):
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _forecast_context_presentation(snapshot_count):
    """Describe the forecast method supported by eligible snapshot history."""
    try:
        count = max(int(snapshot_count or 0), 0)
    except (TypeError, ValueError):
        count = 0
    noun = "snapshot" if count == 1 else "snapshots"
    if count >= 3:
        return {
            "key": "trend",
            "label": f"Trend-ready · {count} {noun}",
            "basis": "Three or more eligible snapshots support the linear-trend benchmark.",
        }
    if count == 2:
        return {
            "key": "moving_average",
            "label": "Moving average · 2 snapshots",
            "basis": "Two eligible snapshots support the moving-average benchmark.",
        }
    if count == 1:
        return {
            "key": "persistence",
            "label": "Initial baseline · 1 snapshot",
            "basis": "One eligible snapshot supports only the persistence benchmark.",
        }
    return {
        "key": "ineligible",
        "label": "Not forecast eligible",
        "basis": "This snapshot does not contain a valid observed average.",
    }


def _snapshots_with_forecast_context(snapshots):
    rows = [dict(snapshot) for snapshot in snapshots]
    chronological = sorted(
        enumerate(rows),
        key=lambda indexed: indexed[1].get("collected_at")
        or indexed[1].get("created_at")
        or datetime.min.replace(tzinfo=timezone.utc),
    )
    eligible_count = 0
    for _, row in chronological:
        eligible = (
            _snapshot_float(row.get("average_price")) is not None
            and row.get("data_quality") not in {"failed", "empty", "invalid"}
        )
        if eligible:
            eligible_count += 1
            row["forecast_context"] = _forecast_context_presentation(eligible_count)
        else:
            row["forecast_context"] = _forecast_context_presentation(0)
    return rows


def _normalized_utc_datetime(value):
    if not isinstance(value, datetime):
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def is_snapshot_after_prediction(snapshot, prediction):
    # Collection time is the business event time.  ``created_at`` is only the
    # time the repository persisted the document and must not make an old
    # imported snapshot look like a later validation observation.
    snapshot_time = _normalized_utc_datetime(snapshot.get("collected_at") or snapshot.get("created_at"))
    prediction_time = _normalized_utc_datetime(prediction.get("prediction_created_at"))
    if not snapshot_time or not prediction_time:
        return False
    return snapshot_time > prediction_time


def _sorted_watchlist_snapshots(watchlist_id, user_id):
    snapshots = safe_call(lambda: repository.list_price_snapshots(watchlist_id, limit=50), [])
    snapshots = [row for row in snapshots if str(row.get("user_id")) == str(user_id)]
    snapshots.sort(key=lambda row: row.get("collected_at") or row.get("created_at") or datetime.min.replace(tzinfo=timezone.utc))
    return snapshots


def _monitor_snapshot_history(monitor, user_id):
    """Return only the selected user's snapshots for this exact monitor."""
    monitor_id = str(monitor.get("monitor_id") or monitor.get("_id"))
    watchlist_id = str(monitor.get("_id"))
    snapshots = _sorted_watchlist_snapshots(monitor["_id"], user_id)
    return [
        row for row in snapshots
        if str(row.get("user_id")) == str(user_id)
        and str(row.get("watchlist_id")) == watchlist_id
        # Older documents used the Watchlist document id as ``monitor_id``.
        # It remains safe to read them because the watchlist id and user scope
        # above are both exact; new writes use the stable monitor id.
        and str(row.get("monitor_id") or watchlist_id) in {monitor_id, watchlist_id}
    ]


def _display_price(value):
    try:
        value = float(value)
    except (TypeError, ValueError):
        value = None
    return f"{value:.2f}" if value is not None else ""


def _display_percentage(value):
    try:
        value = float(value)
    except (TypeError, ValueError):
        value = None
    return f"{value:.2f}%" if value is not None else ""


def _valid_prediction_validation(prediction, monitor, snapshots):
    """Validate forecast lineage without changing historical source records."""
    watchlist_id = str(monitor.get("_id"))
    monitor_id = str(monitor.get("monitor_id") or monitor.get("_id"))
    if str(prediction.get("watchlist_id")) != watchlist_id:
        return False, None
    if str(prediction.get("monitor_id") or watchlist_id) != monitor_id:
        return False, None
    if prediction.get("validation_monitor_id") and str(prediction.get("validation_monitor_id")) != monitor_id:
        return False, None
    if prediction.get("validation_prediction_id") and str(prediction.get("validation_prediction_id")) != str(prediction.get("prediction_id") or prediction.get("_id")):
        return False, None
    if prediction.get("status") != "evaluated":
        return False, None
    baseline_id = prediction.get("baseline_snapshot_id")
    validation_id = prediction.get("actual_snapshot_id") or prediction.get("validation_snapshot_id")
    if not baseline_id or not validation_id or str(baseline_id) == str(validation_id):
        return False, None
    snapshot_by_id = {str(row.get("_id")): row for row in snapshots}
    baseline = snapshot_by_id.get(str(baseline_id))
    validation = snapshot_by_id.get(str(validation_id))
    if not baseline or not validation:
        return False, None
    if not is_snapshot_after_prediction(validation, prediction):
        return False, None
    baseline_time = _normalized_utc_datetime(baseline.get("collected_at") or baseline.get("created_at"))
    validation_time = _normalized_utc_datetime(validation.get("collected_at") or validation.get("created_at"))
    if not baseline_time or not validation_time or validation_time <= baseline_time:
        return False, None
    return True, validation


def _baseline_prediction(values):
    if not values:
        return None, None
    if len(values) == 1:
        return round(values[-1], 2), "last_value"
    if len(values) == 2:
        return round(sum(values) / len(values), 2), "moving_average"
    if len(values) >= 3:
        x = list(range(len(values)))
        x_mean = sum(x) / len(x)
        y_mean = sum(values) / len(values)
        denom = sum((i - x_mean) ** 2 for i in x) or 1
        slope = sum((i - x_mean) * (y - y_mean) for i, y in zip(x, values)) / denom
        trend = round(y_mean + slope * len(values), 2)
        moving = round(sum(values) / len(values), 2)
        return trend if trend > 0 else moving, "linear_trend"
    return None, None


def _prediction_metrics(actual, predicted):
    if actual is None or predicted is None:
        return None, None
    abs_error = round(abs(actual - predicted), 2)
    pct_error = round(abs_error / actual * 100, 2) if actual else None
    return abs_error, pct_error


def _benchmark_method_presentation(method):
    presentations = {
        "last_value": (
            "Persistence benchmark",
            "Uses the baseline Snapshot observed average unchanged as the benchmark prediction.",
        ),
        "moving_average": (
            "Historical moving-average benchmark",
            "Uses the mean of the observed averages in the Forecast Cycle's training history.",
        ),
        "linear_trend": (
            "Linear-trend benchmark",
            "Extrapolates the observed-average trend in the Forecast Cycle's training history.",
        ),
    }
    return presentations.get(method, (
        str(method).replace("_", " ").title() if method else "Benchmark method unavailable",
        "The benchmark method was not recorded for this legacy Forecast Cycle.",
    ))


def _prediction_training_payload(item, snapshots):
    values = []
    for snap in snapshots:
        avg = _snapshot_float(snap.get("average_price"))
        if avg is not None:
            values.append(avg)
    return values


def _valid_forecast_snapshots(snapshots):
    return [row for row in snapshots if _snapshot_float(row.get("average_price")) is not None and row.get("data_quality") not in {"failed", "empty", "invalid"}]


def _build_prediction_record(item, snapshots):
    values = _prediction_training_payload(item, snapshots)
    baseline_predicted_average_price, baseline_method = _baseline_prediction(values)
    if baseline_predicted_average_price is None:
        return {
            "status": "insufficient_history",
            "source_label": item.get("source_label") or item.get("data_source_label") or "live",
            "data_quality": "insufficient",
            "baseline_predicted_average_price": None,
            "baseline_method": None,
            "ai_provider": "gemini",
            "ai_predicted_average_price": None,
            "ai_predicted_direction": None,
            "ai_confidence_level": None,
            "ai_status": "not_attempted",
            "ai_failure_reason": "insufficient_history",
            "ai_reason": "No snapshots available for prediction.",
            "explanation_text": "More snapshots are needed before a reliable price prediction can be produced.",
            "snapshot_count": len(values),
            "training_snapshot_ids": [row.get("_id") for row in snapshots if row.get("_id")],
            "prediction_created_at": utcnow(),
            "created_at": utcnow(),
        }
    ai_payload, ai_status, ai_reason = predict_price_gemini(
        item.get("product_label") or item.get("keyword") or "Tracked item",
        item.get("tracking_mode") or "search_scope",
        item.get("platform_scope") or "all",
        item.get("source_label") or item.get("data_source_label") or "live",
        snapshots,
    )
    baseline_snapshot = snapshots[-1] if snapshots else None
    baseline_snapshot_at = (baseline_snapshot or {}).get("collected_at") or (baseline_snapshot or {}).get("created_at")
    baseline_average_price = _snapshot_float((baseline_snapshot or {}).get("average_price"))
    created_at = utcnow()
    forecast_id = secrets.token_hex(12)
    predicted_average_price = ai_payload.get("predicted_average_price") if ai_payload.get("predicted_average_price") is not None else baseline_predicted_average_price
    prediction = {
        "prediction_id": forecast_id,
        "forecast_id": forecast_id,
        "monitor_id": item.get("monitor_id") or item.get("_id"),
        "comparison_set_id": item.get("comparison_set_id") or item.get("comparison_group_id"),
        "baseline_snapshot_id": baseline_snapshot.get("_id") if baseline_snapshot else None,
        "baseline_snapshot_at": baseline_snapshot_at,
        "baseline_average_price": baseline_average_price,
        "baseline_snapshot_observed_price": baseline_average_price,
        "tracking_mode": item.get("tracking_mode"),
        "product_label": item.get("product_label"),
        "keyword": item.get("keyword"),
        "platform_scope": item.get("platform_scope"),
        "source_scope": item.get("source_scope"),
        "snapshot_count": len(values),
        "training_snapshot_ids": [row.get("_id") for row in snapshots if row.get("_id")],
        "baseline_method": baseline_method,
        "baseline_predicted_average_price": baseline_predicted_average_price,
        "benchmark_method": baseline_method,
        "benchmark_predicted_price": baseline_predicted_average_price,
        "ai_provider": "gemini",
        "ai_predicted_average_price": ai_payload.get("predicted_average_price"),
        "ai_method": "gemini" if ai_payload.get("predicted_average_price") is not None else None,
        "ai_predicted_price": ai_payload.get("predicted_average_price"),
        "ai_predicted_direction": ai_payload.get("predicted_direction"),
        "ai_confidence_level": ai_payload.get("confidence_level"),
        "ai_status": ai_status,
        "ai_failure_reason": ai_reason,
        "ai_reason": ai_payload.get("reason") or ai_reason or "AI prediction unavailable.",
        "predicted_average_price": predicted_average_price,
        "forecast_method": "Gemini" if ai_payload.get("predicted_average_price") is not None else baseline_method,
        "forecast_provider": "gemini" if ai_payload.get("predicted_average_price") is not None else "deterministic_baseline",
        "cycle_status": "pending",
        "status": "pending_actual" if ai_payload.get("predicted_average_price") is not None else "ai_unavailable",
        "prediction_created_at": created_at,
        "created_at": created_at,
        "forecast_horizon": "next_eligible_snapshot",
        "forecast_target_description": "Next eligible Monitor snapshot",
        "forecast_schema_version": 1,
        "data_quality": "good" if len(values) >= 3 else "limited",
        "source_label": item.get("data_source_label") or item.get("source_label") or "live",
        "explanation_text": "The baseline prediction is deterministic. Gemini predicts the direction and average price from the same watchlist snapshot history." if ai_payload.get("predicted_average_price") is not None else "Gemini prediction is unavailable, so only the deterministic baseline is shown.",
    }
    return prediction


def _evaluate_prediction_record(prediction, actual_snapshot):
    actual = _snapshot_float(actual_snapshot.get("average_price"))
    baseline = _snapshot_float(prediction.get("benchmark_predicted_price"))
    if baseline is None:
        baseline = _snapshot_float(prediction.get("baseline_predicted_average_price"))
    ai_value = _snapshot_float(prediction.get("ai_predicted_price"))
    if ai_value is None:
        ai_value = _snapshot_float(prediction.get("ai_predicted_average_price"))
    baseline_abs, baseline_pct = _prediction_metrics(actual, baseline)
    ai_abs, ai_pct = _prediction_metrics(actual, ai_value)
    validation_at = actual_snapshot.get("collected_at") or actual_snapshot.get("created_at")
    evaluated_at = utcnow()
    prediction.update({
        "status": "evaluated",
        "cycle_status": "validated",
        "actual_snapshot_id": actual_snapshot.get("_id"),
        "actual_average_price": round(actual, 2) if actual is not None else None,
        "validation_snapshot_id": actual_snapshot.get("_id"),
        "validation_snapshot_at": validation_at,
        "observed_average_price": round(actual, 2) if actual is not None else None,
        "validation_observed_price": round(actual, 2) if actual is not None else None,
        "baseline_absolute_error": baseline_abs,
        "baseline_percentage_error": baseline_pct,
        "ai_absolute_error": ai_abs,
        "ai_percentage_error": ai_pct,
        "benchmark_absolute_error": baseline_abs,
        "benchmark_error_percent": baseline_pct,
        "ai_error_percent": ai_pct,
        "baseline_mae": baseline_abs,
        "baseline_mape": baseline_pct,
        "ai_mae": ai_abs,
        "ai_mape": ai_pct,
        "absolute_error": ai_abs if ai_value is not None else baseline_abs,
        "percentage_error": ai_pct if ai_value is not None else baseline_pct,
        "evaluated_at": evaluated_at,
        "validated_at": evaluated_at,
    })
    return prediction


def _forecast_cycle_view(prediction, snapshots):
    """Return one legacy-safe, owner-scoped Forecast Cycle presentation model."""
    if not prediction:
        return None
    snapshots_by_id = {str(row.get("_id")): row for row in snapshots}
    baseline = snapshots_by_id.get(str(prediction.get("baseline_snapshot_id")))
    validation_id = prediction.get("validation_snapshot_id") or prediction.get("actual_snapshot_id")
    validation = snapshots_by_id.get(str(validation_id))
    ai_value = _snapshot_float(prediction.get("ai_predicted_price"))
    if ai_value is None:
        ai_value = _snapshot_float(prediction.get("ai_predicted_average_price"))
    benchmark_prediction = _snapshot_float(prediction.get("benchmark_predicted_price"))
    if benchmark_prediction is None:
        benchmark_prediction = _snapshot_float(prediction.get("baseline_predicted_average_price"))
    predicted = ai_value if ai_value is not None else benchmark_prediction
    if predicted is None:
        predicted = _snapshot_float(prediction.get("predicted_average_price"))
    status = str(prediction.get("status") or "failed")
    cycle_status = prediction.get("cycle_status")
    visible_status = cycle_status if cycle_status in {"pending", "validated", "failed", "superseded"} else ("validated" if status == "evaluated" else ("pending" if status in {"pending_actual", "ai_unavailable", "waiting_for_validation"} else status))
    created_at = prediction.get("prediction_created_at") or prediction.get("created_at")
    baseline_at = prediction.get("baseline_snapshot_at") or ((baseline or {}).get("collected_at") or (baseline or {}).get("created_at"))
    validation_at = prediction.get("validation_snapshot_at") or ((validation or {}).get("collected_at") or (validation or {}).get("created_at"))
    observed = _snapshot_float(prediction.get("observed_average_price"))
    if observed is None:
        observed = _snapshot_float(prediction.get("actual_average_price"))
    if observed is None:
        observed = _snapshot_float((validation or {}).get("average_price"))
    source_observed = _snapshot_float(prediction.get("baseline_snapshot_observed_price"))
    if source_observed is None:
        source_observed = _snapshot_float(prediction.get("baseline_average_price"))
    if source_observed is None:
        source_observed = _snapshot_float((baseline or {}).get("average_price"))
    benchmark_method = prediction.get("benchmark_method") or prediction.get("baseline_method")
    benchmark_method_label, benchmark_explanation = _benchmark_method_presentation(benchmark_method)
    derived_benchmark_absolute_error, derived_benchmark_error_percent = _prediction_metrics(observed, benchmark_prediction)
    derived_ai_absolute_error, derived_ai_error_percent = _prediction_metrics(observed, ai_value)
    benchmark_absolute_error = derived_benchmark_absolute_error
    if benchmark_absolute_error is None:
        benchmark_absolute_error = prediction.get("benchmark_absolute_error")
    if benchmark_absolute_error is None:
        benchmark_absolute_error = prediction.get("baseline_absolute_error")
    benchmark_error_percent = derived_benchmark_error_percent
    if benchmark_error_percent is None:
        benchmark_error_percent = prediction.get("benchmark_error_percent")
    if benchmark_error_percent is None:
        benchmark_error_percent = prediction.get("baseline_percentage_error")
    ai_absolute_error = derived_ai_absolute_error
    if ai_absolute_error is None:
        ai_absolute_error = prediction.get("ai_absolute_error")
    ai_error_percent = derived_ai_error_percent
    if ai_error_percent is None:
        ai_error_percent = prediction.get("ai_error_percent")
    if ai_error_percent is None:
        ai_error_percent = prediction.get("ai_percentage_error")
    cycle = dict(prediction)
    for legacy_field in (
        "ai_predicted_average_price", "baseline_predicted_average_price", "ai_predicted_direction",
        "ai_confidence_level", "source_label", "baseline_absolute_error", "baseline_percentage_error",
        "ai_absolute_error", "ai_percentage_error", "baseline_forecast_value", "gemini_forecast_value",
    ):
        cycle.setdefault(legacy_field, None)
    cycle.update({
        "forecast_id": prediction.get("forecast_id") or prediction.get("prediction_id") or prediction.get("_id"),
        "owner_user_id": prediction.get("owner_user_id") or prediction.get("user_id"),
        "generated_at": created_at,
        "forecast_created_at": created_at,
        "baseline_snapshot_id": prediction.get("baseline_snapshot_id"),
        "baseline_snapshot_at": baseline_at,
        "baseline_snapshot_at_display": baseline_at,
        "baseline_snapshot_observed_price": source_observed,
        "baseline_average_price_display": source_observed,
        "benchmark_method": benchmark_method,
        "benchmark_method_label": benchmark_method_label,
        "benchmark_explanation": benchmark_explanation,
        "benchmark_predicted_price": benchmark_prediction,
        "ai_method": prediction.get("ai_method") or prediction.get("ai_provider") or ("gemini" if ai_value is not None else None),
        "ai_predicted_price": ai_value,
        "predicted_average_price": predicted,
        "forecast_method": prediction.get("forecast_method") or ("Gemini" if ai_value is not None else (prediction.get("baseline_method") or "Not available in this legacy forecast record")),
        "forecast_provider": prediction.get("forecast_provider") or (prediction.get("ai_provider") if ai_value is not None else "Deterministic baseline"),
        "visible_status": visible_status,
        "status": visible_status,
        "validation_snapshot_id": validation_id,
        "validation_snapshot_at": validation_at,
        "validation_snapshot_at_display": validation_at,
        "validation_observed_price": observed,
        "observed_average_price": observed,
        "benchmark_absolute_error": benchmark_absolute_error,
        "benchmark_error_percent": benchmark_error_percent,
        "ai_absolute_error": ai_absolute_error,
        "ai_error_percent": ai_error_percent,
        "validated_at_display": prediction.get("validated_at") or prediction.get("evaluated_at"),
        "absolute_error": prediction.get("absolute_error") if prediction.get("absolute_error") is not None else (prediction.get("ai_absolute_error") if ai_value is not None else prediction.get("baseline_absolute_error")),
        "percentage_error": prediction.get("percentage_error") if prediction.get("percentage_error") is not None else (prediction.get("ai_percentage_error") if ai_value is not None else prediction.get("baseline_percentage_error")),
        "short_id": str(prediction.get("forecast_id") or prediction.get("prediction_id") or prediction.get("_id") or "")[:8],
        "is_legacy": not bool(prediction.get("forecast_schema_version")),
    })
    return cycle


def _forecast_cycle_chart_data(cycle):
    if not cycle:
        return {"available": False, "pending": False, "labels": [], "values": [], "tooltip_rows": [], "status_text": "No forecast has been generated for the latest snapshot.", "reason": "no_forecast"}
    labels = ["Benchmark prediction", "AI-assisted forecast"]
    values = [cycle.get("benchmark_predicted_price"), cycle.get("ai_predicted_price")]
    rows = [{"label": label, "value": value, "abs_error": None, "pct_error": None} for label, value in zip(labels, values)]
    validated = cycle.get("visible_status") == "validated" and cycle.get("observed_average_price") is not None
    if validated:
        labels.append("Observed market average")
        values.append(cycle.get("validation_observed_price"))
        rows.append({"label": "Observed market average", "value": cycle.get("validation_observed_price"), "abs_error": None, "pct_error": None})
    available = all(value is not None for value in values)
    return {
        "available": available and validated,
        "pending": cycle.get("visible_status") == "pending",
        "labels": labels,
        "values": values,
        "tooltip_rows": rows,
        "status_text": "Benchmark prediction, AI-assisted forecast and observed market average for the selected Forecast Cycle." if validated else "Observation pending until a later eligible snapshot is collected.",
        "error_level": "pending" if not validated else "validated",
        "reason": None if available else "legacy_metadata_missing",
        "cycle_id": cycle.get("short_id"),
    }


def _prediction_has_later_actual(prediction, snapshots):
    return _prediction_actual_snapshot(prediction, snapshots) is not None


def _prediction_actual_snapshot(prediction, snapshots):
    if not prediction or prediction.get("status") not in {"evaluated", "validated"}:
        return None
    prediction_created_at = prediction.get("prediction_created_at")
    actual_snapshot_id = prediction.get("actual_snapshot_id")
    actual_snapshot = None
    if actual_snapshot_id:
        actual_snapshot = next((row for row in snapshots if str(row.get("_id")) == str(actual_snapshot_id)), None)
    if not actual_snapshot:
        return None
    return actual_snapshot if is_snapshot_after_prediction(actual_snapshot, prediction) else None


def _prediction_chart_filename(prediction, actual_snapshot=None):
    if not prediction:
        return None
    values = [
        ("Benchmark prediction", _snapshot_float(prediction.get("benchmark_predicted_price")) if prediction.get("benchmark_predicted_price") is not None else _snapshot_float(prediction.get("baseline_predicted_average_price"))),
        ("AI-assisted forecast", _snapshot_float(prediction.get("ai_predicted_price")) if prediction.get("ai_predicted_price") is not None else _snapshot_float(prediction.get("ai_predicted_average_price"))),
    ]
    if actual_snapshot:
        values.append(("Observed market average", _snapshot_float(actual_snapshot.get("average_price"))))
    data = [(label, value) for label, value in values if value is not None]
    if not data:
        return None
    filename = f"prediction_{prediction.get('_id') or prediction.get('prediction_id')}.png"
    labels = [label for label, _ in data]
    heights = [value for _, value in data]
    plt.figure(figsize=(6.5, 4))
    bars = plt.bar(labels, heights, color=["#4f46e5", "#0ea5e9", "#10b981"][:len(labels)])
    plt.title("Predicted vs actual price")
    plt.ylabel("Average price")
    for bar, value in zip(bars, heights):
        plt.text(bar.get_x() + bar.get_width() / 2, value, f"{value:.2f}", ha="center", va="bottom", fontsize=9)
    plt.tight_layout()
    plt.savefig(CHART_DIR / filename)
    plt.close()
    return filename


def _watchlist_trend_chart_data(snapshots):
    snapshots = _snapshots_with_forecast_context(snapshots)
    valid_snapshots = []
    for snap in snapshots:
        average_price = _snapshot_float(snap.get("average_price"))
        if average_price is None:
            continue
        valid_snapshots.append((snap, average_price))
    valid_snapshots.sort(key=lambda row: row[0].get("created_at") or row[0].get("collected_at") or datetime.min.replace(tzinfo=timezone.utc))
    labels = []
    average = []
    low = []
    high = []
    counts = []
    source_labels = []
    forecast_context_labels = []
    for snap, average_price in valid_snapshots:
        low_price = _snapshot_float(snap.get("lowest_price")) or average_price
        high_price = _snapshot_float(snap.get("highest_price")) or average_price
        label = format_sgt_datetime(snap.get("created_at") or snap.get("collected_at"))
        labels.append(label)
        average.append(float(average_price))
        low.append(float(low_price))
        high.append(float(high_price))
        count = int(snap.get("record_count") or 0)
        source_label = snap.get("source_label") or "live"
        forecast_context_label = snap["forecast_context"]["label"]
        counts.append(count)
        source_labels.append(source_label)
        forecast_context_labels.append(forecast_context_label)
    tooltip_rows = [
        {
            "label": label,
            "average": average_value,
            "low": low_value,
            "high": high_value,
            "count": count,
            "source": source_label,
            "forecast_context": forecast_context_label,
        }
        for label, average_value, low_value, high_value, count, source_label, forecast_context_label in zip(labels, average, low, high, counts, source_labels, forecast_context_labels)
    ]
    available = len(average) >= 2
    stable = available and min(average) == max(average)
    trend_summary = None
    if available:
        first_average = average[0]
        latest_average = average[-1]
        change_amount = round(latest_average - first_average, 2)
        change_percentage = round(change_amount / first_average * 100, 2) if first_average else None
        if stable:
            movement = "stable"
            summary_text = "Stable trend: average price stayed unchanged across recent snapshots."
        elif latest_average > first_average:
            movement = "rising"
            summary_text = "Rising trend: latest average is higher than the first recorded snapshot."
        elif latest_average < first_average:
            movement = "falling"
            summary_text = "Falling trend: latest average is lower than the first recorded snapshot."
        else:
            movement = "mixed"
            summary_text = "Mixed movement: prices changed across the tracked period."
        trend_summary = {
            "movement": movement,
            "first_average": first_average,
            "latest_average": latest_average,
            "change_amount": change_amount,
            "change_percentage": change_percentage,
            "snapshot_count": len(average),
            "summary_text": summary_text,
        }
    y_min = None
    y_max = None
    if available:
        min_value = min(low + average)
        max_value = max(high + average)
        if min_value == max_value:
            padding = max(abs(min_value) * 0.05, 1.0)
            y_min = round(min_value - padding, 2)
            y_max = round(max_value + padding, 2)
        else:
            padding = max((max_value - min_value) * 0.1, 1.0)
            y_min = round(max(min_value - padding, 0), 2)
            y_max = round(max_value + padding, 2)
    return {
        "available": available,
        "stable": stable,
        "snapshot_count": len(average),
        "y_min": y_min,
        "y_max": y_max,
        "labels": labels,
        "average": average,
        "low": low,
        "high": high,
        "counts": counts,
        "source_labels": source_labels,
        "forecast_context_labels": forecast_context_labels,
        "tooltip_rows": tooltip_rows,
        "trend_summary": trend_summary,
    }


def build_analytics_payload(records, *, query="", source="", analysis_scope="", mode="full"):
    rows = list(records or [])
    payload = {
        "summary": {
            "total_records": 0,
            "platform_count": 0,
            "category_count": 0,
            "average_price": 0.0,
            "median_price": 0.0,
            "minimum_price": 0.0,
            "maximum_price": 0.0,
            "price_range": 0.0,
            "best_platform": "",
        },
        "price_distribution": {"bins": [], "counts": [], "percentages": []},
        "platform_distribution": {"labels": [], "counts": [], "percentages": []},
        "condition_distribution": {"labels": [], "counts": [], "percentages": []},
        "category_distribution": {"labels": [], "counts": [], "percentages": []},
        "platform_price_comparison": {"labels": [], "average": [], "median": [], "minimum": [], "maximum": []},
        "comparable_price_curve": {
            "available": False, "labels": [], "prices": [], "record_ids": [],
            "median": None, "x_axis": "Listing rank", "y_axis": "Price (USD)",
        },
        "time_series": {"labels": [], "average": [], "low": [], "high": []},
        "platform_condition_breakdown": {"platforms": [], "conditions": [], "series": []},
        "records": [],
        "filters": {
            "platforms": ["All"],
            "metric_modes": [
                {"value": "count", "label": "Count"},
                {"value": "percentage", "label": "Percentage"},
            ],
            "sort_options": [
                {"value": "value_desc", "label": "Value: high to low"},
                {"value": "value_asc", "label": "Value: low to high"},
                {"value": "label_asc", "label": "Label: A to Z"},
            ],
        },
        "charts": {
            "priceHistogram": {"labels": [], "counts": []},
            "platformDistribution": [],
            "conditionDistribution": [],
            "availabilityDistribution": [],
            "platformComparison": [],
        },
        "analysisScope": analysis_scope or "all current search results",
        "singlePlatformMessage": None,
        "analyticsLabel": "Saved evidence analytics" if mode == "saved" else ("Selected record analytics" if mode == "selected" else "Search analytics"),
        "platformChartTitle": "Platform distribution",
    }
    if not rows:
        return payload

    normalized_rows = []
    metric_rows = []
    numeric_prices = []
    platform_values = []
    category_values = []
    condition_values = []
    storage_values = []
    for item in rows:
        parsed_model = item.get("parsed_model") or parse_canonical_product_model(item.get("title"))
        attributes = item.get("attributes") or {}
        item_storage = item.get("storage") or attributes.get("storage") or parsed_model.get("storage")
        price_value = item.get("normalized_price", item.get("price"))
        try:
            numeric_price = float(price_value)
        except (TypeError, ValueError):
            numeric_price = None
        eligible = bool(item.get("analytics_eligible"))
        if eligible and numeric_price is not None and not pd.isna(numeric_price):
            numeric_prices.append(numeric_price)
        platform = str(item.get("platform") or "Unknown")
        category = str(item.get("category_display") or item.get("category") or "Not available")
        condition = str(item.get("condition_display") or "Not specified by source")
        if eligible:
            platform_values.append(platform)
            category_values.append(category)
            condition_values.append(condition)
            metric_rows.append(item)
            if item_storage:
                storage_values.append(item_storage)
        normalized_rows.append({
            "record_id": str(item.get("_selection_token") or item.get("_id") or ""),
            "title": item.get("title") or item.get("product_name") or "Untitled",
            "platform": platform,
            "price": round(float(numeric_price), 2) if numeric_price is not None else None,
            "currency": item.get("currency") or "USD",
            "seller": item.get("seller") or "",
            "condition": condition,
            "condition_display": condition,
            "category": category,
            "category_display": category,
            "normalized_price": round(float(numeric_price), 2) if numeric_price is not None else None,
            "price_quality": item.get("price_quality_status") or ("Comparable full price" if eligible else "Excluded"),
            "exclusion_reason": item.get("exclusion_reason") or "",
            "analytics_eligible": eligible,
            "condition_source": item.get("condition_source") or "not_specified_by_source",
            "storage": item_storage,
            "collected_at": display_sgt_datetime(item.get("collected_at") or item.get("created_at")),
            "source_url": item.get("source_url") or item.get("item_url") or "",
            "record_source": _record_source_label(item),
        })
    payload["records"] = normalized_rows
    if not numeric_prices:
        platform_counts = Counter(platform_values)
        condition_counts = Counter(condition_values)
        category_counts = Counter(category_values)
        payload["summary"] = {
            "total_records": 0,
            "platform_count": int(len(platform_counts)),
            "category_count": int(len(category_counts)),
            "average_price": 0.0,
            "median_price": 0.0,
            "minimum_price": 0.0,
            "maximum_price": 0.0,
            "price_range": 0.0,
            "best_platform": "",
        }
        payload["platform_distribution"] = {
            "labels": [str(label) for label in platform_counts.keys()],
            "counts": [int(value) for value in platform_counts.values()],
            "percentages": [0.0 for _ in platform_counts.values()],
        }
        payload["condition_distribution"] = {
            "labels": [str(label) for label in condition_counts.keys()],
            "counts": [int(value) for value in condition_counts.values()],
            "percentages": [0.0 for _ in condition_counts.values()],
        }
        payload["platform_price_comparison"] = {"labels": [], "average": [], "median": [], "minimum": [], "maximum": []}
        payload["category_distribution"] = {
            "labels": [str(label) for label in category_counts.keys()],
            "counts": [int(value) for value in category_counts.values()],
            "percentages": [0.0 for _ in category_counts.values()],
        }
        payload["filters"]["platforms"] = ["All"] + [str(label) for label in platform_counts.keys()]
        payload["charts"] = {
            "priceHistogram": {"labels": [], "counts": []},
            "platformDistribution": [{"name": str(label), "count": int(value), "percentage": 0.0} for label, value in platform_counts.items()],
            "conditionDistribution": [{"name": str(label), "count": int(value), "percentage": 0.0} for label, value in condition_counts.items()],
            "availabilityDistribution": [{"name": str(label), "count": int(value), "percentage": 0.0} for label, value in condition_counts.items()],
            "platformComparison": [],
            "timeSeries": {"labels": [], "average": [], "low": [], "high": []},
            "platformConditionBreakdown": {"platforms": [], "conditions": [], "series": []},
        }
        payload["singlePlatformMessage"] = "This analysis contains one platform. Price and listing insights remain valid; cross-platform comparisons are not shown." if len(platform_counts) == 1 else None
        payload["platformChartTitle"] = "Platform distribution" if len(platform_counts) > 1 else "Single-platform distribution"
        return payload

    series = pd.Series(numeric_prices, dtype="float64")
    curve_rows = sorted(
        (
            row for row in normalized_rows
            if row.get("analytics_eligible") and row.get("normalized_price") is not None
        ),
        key=lambda row: (float(row["normalized_price"]), row.get("record_id") or ""),
    )
    payload["comparable_price_curve"] = {
        "available": len(curve_rows) >= 2,
        "labels": list(range(1, len(curve_rows) + 1)),
        "prices": [round(float(row["normalized_price"]), 2) for row in curve_rows],
        "record_ids": [row["record_id"] for row in curve_rows],
        "median": round(float(series.median()), 2),
        "x_axis": "Listing rank",
        "y_axis": "Price (USD)",
    }
    platform_counts = Counter(platform_values)
    category_counts = Counter(category_values)
    condition_counts = Counter(condition_values)
    min_price = float(series.min())
    max_price = float(series.max())
    median_price = float(series.median())
    average_price = float(series.mean())
    bin_count = min(12, max(1, int(round(len(series) ** 0.5))))
    if len(set(series.tolist())) == 1:
        bin_count = 1
    cut = pd.cut(series, bins=bin_count, include_lowest=True, duplicates="drop")
    histogram = series.groupby(cut, observed=True).size().reset_index(name="count")
    histogram.columns = ["bucket", "count"]
    histogram_labels = [str(bucket).replace("(", "").replace("]", "").replace(",", " -") for bucket in histogram["bucket"].tolist()]
    histogram_counts = [int(value) for value in histogram["count"].tolist()]
    histogram_total = sum(histogram_counts) or 1
    storage_variants = sorted(set(storage_values))
    mixed_storage = len(storage_variants) > 1
    mixed_conditions = len(condition_counts) > 1
    # Storage changes the product variant being compared. Condition is a
    # separate price-quality dimension and must not silently turn a
    # single-capacity comparison into a mixed-variant analysis.
    mixed_variants = mixed_storage
    lowest_platform = str(min(metric_rows, key=lambda item: float(item.get("normalized_price", item.get("price")))).get("platform")) if metric_rows else ""
    payload["summary"] = {
        "total_records": int(len(series)),
        "platform_count": int(len(platform_counts)),
        "category_count": int(len(category_counts)),
        "average_price": round(average_price, 2),
        "median_price": round(median_price, 2),
        "minimum_price": round(min_price, 2),
        "maximum_price": round(max_price, 2),
        "price_range": round(max_price - min_price, 2),
        "best_platform": "" if mixed_variants else lowest_platform,
        "storage_variants": storage_variants,
        "mixed_storage": mixed_storage,
        "mixed_variants": mixed_variants,
        "mixed_conditions": mixed_conditions,
        "scope_label": "Aggregate across mixed storage variants" if mixed_storage else "Comparable single-storage scope",
    }
    payload["price_distribution"] = {
        "bins": histogram_labels,
        "counts": histogram_counts,
        "percentages": [round((count / histogram_total) * 100, 1) for count in histogram_counts],
    }
    platform_labels = [str(label) for label in platform_counts.keys()]
    platform_values_counts = [int(value) for value in platform_counts.values()]
    platform_total = sum(platform_values_counts) or 1
    payload["platform_distribution"] = {
        "labels": platform_labels,
        "counts": platform_values_counts,
        "percentages": [round((count / platform_total) * 100, 1) for count in platform_values_counts],
    }
    condition_sorted = sorted(condition_counts.items(), key=lambda item: item[1], reverse=True)
    condition_total = sum(value for _, value in condition_sorted) or 1
    payload["condition_distribution"] = {
        "labels": [str(label) for label, _ in condition_sorted],
        "counts": [int(value) for _, value in condition_sorted],
        "percentages": [round((int(value) / condition_total) * 100, 1) for _, value in condition_sorted],
    }
    comparison_labels = []
    comparison_average = []
    comparison_median = []
    comparison_minimum = []
    comparison_maximum = []
    for platform in platform_labels:
        platform_prices = [
            float(item.get("normalized_price", item.get("price")))
            for item in metric_rows
            if str(item.get("platform") or "Unknown") == platform and item.get("normalized_price", item.get("price")) is not None and not pd.isna(item.get("normalized_price", item.get("price")))
        ]
        if not platform_prices:
            continue
        values = pd.Series(platform_prices, dtype="float64")
        comparison_labels.append(platform)
        comparison_average.append(round(float(values.mean()), 2))
        comparison_median.append(round(float(values.median()), 2))
        comparison_minimum.append(round(float(values.min()), 2))
        comparison_maximum.append(round(float(values.max()), 2))
    payload["platform_price_comparison"] = {
        "labels": comparison_labels,
        "average": comparison_average,
        "median": comparison_median,
        "minimum": comparison_minimum,
        "maximum": comparison_maximum,
    }
    category_total = sum(category_counts.values()) or 1
    payload["category_distribution"] = {
        "labels": [str(label) for label in category_counts.keys()],
        "counts": [int(value) for value in category_counts.values()],
        "percentages": [round((int(value) / category_total) * 100, 1) for value in category_counts.values()],
    }
    timeline_rows = []
    for item in metric_rows:
        label = str(item.get("collected_at") or item.get("created_at") or "")
        try:
            price = float(item.get("normalized_price", item.get("price")))
        except (TypeError, ValueError):
            continue
        if pd.isna(price):
            continue
        timeline_rows.append((label, price))
    timeline_rows = [row for row in timeline_rows if row[0]]
    timeline_rows.sort(key=lambda row: row[0])
    payload["time_series"] = {
        "labels": [label for label, _ in timeline_rows],
        "average": [round(value, 2) for _, value in timeline_rows],
        "low": [round(value * 0.95, 2) for _, value in timeline_rows],
        "high": [round(value * 1.05, 2) for _, value in timeline_rows],
    }
    platform_condition_matrix = {}
    condition_order = []
    for item in metric_rows:
        platform = str(item.get("platform") or "Unknown")
        condition = str(item.get("condition_display") or "Not specified by source")
        platform_condition_matrix.setdefault(platform, Counter())
        platform_condition_matrix[platform][condition] += 1
        if condition not in condition_order:
            condition_order.append(condition)
    platforms_for_stack = list(platform_condition_matrix.keys())
    stacked_series = []
    for condition in condition_order:
        stacked_series.append({
            "name": condition,
            "type": "bar",
            "stack": "condition",
            "data": [int(platform_condition_matrix[platform].get(condition, 0)) for platform in platforms_for_stack],
        })
    payload["platform_condition_breakdown"] = {
        "platforms": platforms_for_stack,
        "conditions": condition_order,
        "series": stacked_series,
    }
    payload["filters"]["platforms"] = ["All"] + comparison_labels
    payload["charts"] = {
        "priceHistogram": {"labels": histogram_labels, "counts": histogram_counts},
        "platformDistribution": [
            {"name": label, "count": count, "percentage": pct}
            for label, count, pct in zip(platform_labels, platform_values_counts, payload["platform_distribution"]["percentages"])
        ],
        "conditionDistribution": [
            {"name": label, "count": count, "percentage": pct}
            for label, count, pct in zip(payload["condition_distribution"]["labels"], payload["condition_distribution"]["counts"], payload["condition_distribution"]["percentages"])
        ],
        "availabilityDistribution": [
            {"name": label, "count": count, "percentage": pct}
            for label, count, pct in zip(payload["condition_distribution"]["labels"], payload["condition_distribution"]["counts"], payload["condition_distribution"]["percentages"])
        ],
        "platformComparison": [
            {"name": label, "average": avg, "median": med, "minimum": low, "maximum": high}
            for label, avg, med, low, high in zip(comparison_labels, comparison_average, comparison_median, comparison_minimum, comparison_maximum)
        ],
        "timeSeries": payload["time_series"],
        "platformConditionBreakdown": payload["platform_condition_breakdown"],
    }
    payload["singlePlatformMessage"] = "This analysis contains one platform. Price and listing insights remain valid; cross-platform comparisons are not shown." if len(platform_labels) == 1 else None
    payload["platformChartTitle"] = "Platform distribution" if len(platform_labels) > 1 else "Single-platform distribution"
    return payload


def _watchlist_trend_svg_data(chart):
    if not chart or not chart.get("available") or not chart.get("average"):
        return None
    values = [float(value) for value in chart.get("average", []) if value is not None]
    if len(values) < 2:
        return None
    width = 720
    height = 260
    padding_x = 36
    padding_y = 30
    min_value = min(values)
    max_value = max(values)
    if min_value == max_value:
        min_value -= max(abs(min_value) * 0.05, 1.0)
        max_value += max(abs(max_value) * 0.05, 1.0)
    span = max(max_value - min_value, 1.0)
    step_x = (width - padding_x * 2) / max(len(values) - 1, 1)
    points = []
    for index, value in enumerate(values):
        x = padding_x + step_x * index
        y = height - padding_y - ((value - min_value) / span) * (height - padding_y * 2)
        points.append({
            "x": round(x, 2),
            "y": round(y, 2),
            "label": chart.get("labels", [""] * len(values))[index],
            "average": value,
            "low": chart.get("low", [None] * len(values))[index],
            "high": chart.get("high", [None] * len(values))[index],
            "count": chart.get("counts", [0] * len(values))[index],
            "source": chart.get("source_labels", ["live"] * len(values))[index],
            "forecast_context": chart.get("forecast_context_labels", ["Not forecast eligible"] * len(values))[index],
        })
    return {
        "width": width,
        "height": height,
        "points": points,
        "labels": chart.get("labels", []),
        "average": values,
        "y_min": round(min_value, 2),
        "y_max": round(max_value, 2),
        "stable": bool(chart.get("stable")),
    }


def _watchlist_trend_summary(chart):
    return chart.get("trend_summary") if chart else None


def _prediction_validation_chart_data(prediction):
    if not prediction:
        return {"available": False, "pending": False, "labels": [], "values": [], "tooltip_rows": [], "status_text": "No prediction has been generated yet.", "error_level": "pending", "reason": "no_prediction"}
    labels = ["Baseline forecast"]
    values = [_snapshot_float(prediction.get("baseline_predicted_average_price"))]
    if prediction.get("ai_predicted_average_price") is not None:
        labels.append("Gemini forecast")
        values.append(_snapshot_float(prediction.get("ai_predicted_average_price")))
    actual = _snapshot_float(prediction.get("actual_average_price"))
    if actual is not None and values[0] is not None:
        labels.append("Actual observed average")
        values.append(actual)
        baseline_error = _snapshot_float(prediction.get("baseline_percentage_error"))
        ai_error = _snapshot_float(prediction.get("ai_percentage_error"))
        if (prediction.get("baseline_absolute_error") or 0) == 0 or (prediction.get("ai_absolute_error") or 0) == 0 or (baseline_error or 0) == 0 or (ai_error or 0) == 0:
            error_level = "matched"
        elif baseline_error is not None and baseline_error <= 5:
            error_level = "low"
        elif baseline_error is not None and baseline_error <= 15:
            error_level = "moderate"
        elif baseline_error is not None:
            error_level = "high"
        else:
            error_level = "matched" if actual == values[0] else "pending"
        tooltip_rows = [
            {
                "label": "Baseline forecast",
                "value": values[0],
                "abs_error": _snapshot_float(prediction.get("baseline_absolute_error")),
                "pct_error": baseline_error,
            }
        ]
        if prediction.get("ai_predicted_average_price") is not None:
            tooltip_rows.append({
                "label": "Gemini forecast",
                "value": _snapshot_float(prediction.get("ai_predicted_average_price")),
                "abs_error": _snapshot_float(prediction.get("ai_absolute_error")),
                "pct_error": ai_error,
            })
        tooltip_rows.append({
            "label": "Actual observed average",
            "value": actual,
            "abs_error": None,
            "pct_error": None,
        })
        return {
            "available": True,
            "pending": False,
            "labels": labels,
            "values": values,
            "tooltip_rows": tooltip_rows,
            "status_text": "Predicted vs actual average price.",
            "error_level": error_level,
            "reason": None,
        }
    error_level = "pending"
    tooltip_rows = [
        {
            "label": "Baseline forecast",
            "value": values[0],
            "abs_error": None,
            "pct_error": None,
        }
    ]
    if prediction.get("ai_predicted_average_price") is not None:
        tooltip_rows.append({
            "label": "Gemini forecast",
            "value": values[1],
            "abs_error": None,
            "pct_error": None,
        })
    return {
        "available": False,
        "pending": True,
        "labels": labels,
        "values": values,
        "tooltip_rows": tooltip_rows,
        "status_text": "Actual validation is pending.",
        "error_level": error_level,
        "reason": "actual_missing" if actual is None else "baseline_missing",
    }


def _normalize_gemini_status_reason(reason):
    text = (reason or "").lower()
    if "missing_key" in text:
        return "missing_key"
    if "invalid_key" in text or "api_key_invalid" in text or "api key not valid" in text:
        return "invalid_key"
    if "permission_denied" in text or "permission" in text:
        return "permission_denied"
    if "quota" in text or "resource_exhausted" in text:
        return "quota_exceeded"
    if "invalid_response" in text:
        return "invalid_response"
    if "timeout" in text or "deadline" in text:
        return "api_call_failed"
    if "service_unavailable" in text or "unavailable" in text:
        return "api_call_failed"
    if "sdk_unavailable" in text:
        return "api_call_failed"
    if "api_error" in text:
        return "api_call_failed"
    if "gemini prediction is unavailable" in text:
        return "api_call_failed"
    return "unknown"


def seed_demo_prediction_validation(user_id=None):
    if user_id is None:
        return None
    user = repository.get_user_by_id(user_id)
    if not user:
        return None
    demo_item = None
    for item in repository.list_watchlist_items(user_id=user_id, include_archived=False, limit=20):
        if (item.get("source_label") or item.get("data_source_label")) == "demo":
            demo_item = item
            break
    if not demo_item:
        item_id, _created = repository.create_watchlist_item(
            user_id,
            {
                "keyword": "demo validation",
                "product_label": "Demo validation watchlist",
                "tracking_mode": "search_scope",
                "platform_scope": "demo_all",
                "source_label": "demo",
                "data_source_label": "demo",
            },
        )
        demo_item = repository.get_watchlist_item(item_id, user_id)
    if not demo_item:
        return None
    base_time = datetime(2026, 7, 1, 2, 0, tzinfo=timezone.utc)
    initial_snapshots = [
        {"average_price": 135.14, "lowest_price": 130.10, "highest_price": 140.20, "record_count": 4, "source_label": "demo", "data_quality": "demo", "demo": True, "collected_at": base_time},
        {"average_price": 138.60, "lowest_price": 133.10, "highest_price": 143.40, "record_count": 5, "source_label": "demo", "data_quality": "limited", "demo": True, "collected_at": base_time + timedelta(days=1)},
        {"average_price": 141.50, "lowest_price": 136.20, "highest_price": 145.60, "record_count": 5, "source_label": "demo", "data_quality": "limited", "demo": True, "collected_at": base_time + timedelta(days=2)},
    ]
    for payload in initial_snapshots:
        repository.save_price_snapshot(demo_item["_id"], demo_item["user_id"], payload)
    snapshots = _sorted_watchlist_snapshots(demo_item["_id"], user_id)
    prediction = _build_prediction_record(demo_item, snapshots)
    if prediction.get("status") == "insufficient_history":
        return {"watchlist_id": demo_item["_id"], "prediction_id": None, "snapshot_id": None}
    prediction_id = repository.save_prediction(user_id, demo_item["_id"], prediction)
    actual_snapshot = {
        "average_price": 143.72,
        "lowest_price": 139.80,
        "highest_price": 146.90,
        "record_count": 6,
        "source_label": "demo",
        "data_quality": "demo",
        "demo": True,
        "collected_at": prediction["prediction_created_at"] + timedelta(hours=6),
        "notes": "Demo validation snapshot collected after prediction.",
    }
    actual_snapshot_id = repository.save_price_snapshot(demo_item["_id"], demo_item["user_id"], actual_snapshot)
    actual_record = repository.latest_price_snapshot(demo_item["_id"])
    evaluated = _evaluate_prediction_record(prediction, actual_record)
    repository.update_prediction(prediction_id, {
        "status": evaluated.get("status"),
        "cycle_status": evaluated.get("cycle_status"),
        "actual_snapshot_id": evaluated.get("actual_snapshot_id"),
        "actual_average_price": evaluated.get("actual_average_price"),
        "validation_snapshot_at": evaluated.get("validation_snapshot_at"),
        "observed_average_price": evaluated.get("observed_average_price"),
        "validation_observed_price": evaluated.get("validation_observed_price"),
        "baseline_absolute_error": evaluated.get("baseline_absolute_error"),
        "baseline_percentage_error": evaluated.get("baseline_percentage_error"),
        "benchmark_absolute_error": evaluated.get("benchmark_absolute_error"),
        "benchmark_error_percent": evaluated.get("benchmark_error_percent"),
        "ai_absolute_error": evaluated.get("ai_absolute_error"),
        "ai_percentage_error": evaluated.get("ai_percentage_error"),
        "ai_error_percent": evaluated.get("ai_error_percent"),
        "baseline_mae": evaluated.get("baseline_mae"),
        "baseline_mape": evaluated.get("baseline_mape"),
        "ai_mae": evaluated.get("ai_mae"),
        "ai_mape": evaluated.get("ai_mape"),
        "absolute_error": evaluated.get("absolute_error"),
        "percentage_error": evaluated.get("percentage_error"),
        "evaluated_at": evaluated.get("evaluated_at"),
        "validated_at": evaluated.get("validated_at"),
    })
    return {"watchlist_id": demo_item["_id"], "prediction_id": prediction_id, "snapshot_id": actual_snapshot_id}


def seed_demo_snapshots(watchlist_ids=None, user_id=None):
    demo_rows = [
        {"average_price": 135.14, "lowest_price": 120.00, "highest_price": 150.00},
        {"average_price": 138.20, "lowest_price": 121.00, "highest_price": 154.00},
        {"average_price": 132.90, "lowest_price": 118.00, "highest_price": 149.00},
        {"average_price": 141.50, "lowest_price": 123.00, "highest_price": 158.00},
        {"average_price": 139.80, "lowest_price": 122.00, "highest_price": 156.00},
    ]
    if watchlist_ids is None:
        watchlist_ids = [
            row["_id"]
            for row in repository.list_watchlist_items(user_id=user_id, include_archived=False, limit=20)
            if (row.get("source_label") or row.get("data_source_label")) == "demo"
        ]
    created = []
    for watchlist_id in watchlist_ids or []:
        item = repository.get_watchlist_item(watchlist_id, user_id) if user_id is not None else repository.get_watchlist_item(watchlist_id)
        if not item:
            continue
        base_time = datetime(2026, 7, 2, 1, 0, tzinfo=timezone.utc)
        for index, row in enumerate(demo_rows[:5]):
            payload = {
                **row,
                "record_count": 5,
                "source_label": "demo",
                "data_quality": "demo" if index == 0 else "limited",
                "demo": True,
                "collected_at": base_time + timedelta(days=index),
                "notes": "Demo snapshot history for presentation.",
            }
            created.append(repository.save_price_snapshot(item["_id"], item["user_id"], payload))
    return created


def build_activity_records(user_id):
    records = []
    for row in safe_call(lambda: repository.list_audit_logs(user_id=user_id, limit=50), []):
        action = row.get("event_type") or row.get("action") or "audit"
        details = row.get("details") if isinstance(row.get("details"), dict) else {}
        action_key = str(action).lower()
        action_label = {
            "login": "Login", "logout": "Logout", "role_switch": "Role switch",
            "search_started": "Search submitted", "search_query": "Search completed",
            "records_collected": "Records collected", "ebay_api_call": "eBay source call",
            "walmart_source_call": "Walmart source call", "api_call_failed": "Source call failed",
            "evidence_save": "Evidence saved", "save_evidence": "Evidence saved",
            "research_saved": "Research saved", "ai_discover": "AI price insight generated",
            "analysis_deleted": "Analysis deleted", "prediction_created": "Price forecast generated",
            "password_reset_requested": "Password reset requested",
            "password_reset_email_printed": "Reset email printed",
            "password_reset_email_queued": "Reset email queued",
            "password_reset_completed": "Password reset completed",
            "password_reset_invalid_attempt": "Invalid reset link attempted",
            "password_reset_throttled": "Password reset throttled",
            "password_reset_mail_failed": "Reset email failed",
        }.get(action_key, str(action).replace("_", " ").title())
        source = details.get("search_scope") or details.get("source") or details.get("actual_source") or details.get("provider")
        category = _log_category(action, row.get("message", ""), source or "")
        if action_key in {"ebay_api_call", "walmart_source_call", "api_call_failed"}:
            category = "Errors / warnings" if action_key == "api_call_failed" else "Search records"
        query = details.get("query")
        count = details.get("record_count")
        scope = _audit_source_display(details.get("search_scope") or details.get("requested_source") or details.get("actual_source") or details.get("source"))
        if action_key == "login":
            detail = "Signed in to the workspace."
        elif action_key == "logout":
            detail = "Signed out of the workspace."
        elif action_key == "role_switch":
            detail = f"Active workspace changed to {str(details.get('new_active_role') or details.get('role') or 'assigned role').replace('_', ' ').title()}."
        elif action_key == "search_started":
            detail = f"Submitted search for “{query or 'product'}” using {scope}."
        elif action_key == "records_collected":
            detail = f"Retained {count or 0} record(s) for “{query or 'product'}” from {scope}."
        elif action_key == "ebay_api_call":
            detail = f"eBay returned {count or 0} record(s) for “{query or 'product'}”; cache: {details.get('cache_status') or 'not stated'}."
        elif action_key == "walmart_source_call":
            detail = f"Walmart retrieval returned {count or 0} record(s) for “{query or 'product'}”; provider: {details.get('provider') or 'not stated'}."
        elif action_key == "api_call_failed":
            detail = f"{details.get('source') or 'Marketplace source'} failed for “{query or 'product'}”: {details.get('error_type') or 'provider error'}."
        elif action_key in {"evidence_save", "save_evidence", "research_saved"}:
            detail = f"Saved {details.get('saved_count') or details.get('record_count') or 1} evidence record(s) for “{query or details.get('title') or 'selected product'}”."
        elif action_key == "ai_discover":
            detail = f"Generated a {details.get('generation_mode') or 'grounded'} insight for “{query or 'analysis'}” from {count or 0} record(s)."
        elif action_key == "password_reset_requested":
            detail = "A password reset was requested for this account."
        elif action_key in {"password_reset_email_printed", "password_reset_email_queued"}:
            detail = f"Password reset delivery was prepared using {details.get('delivery_mode') or 'configured mail'} mode."
        elif action_key == "password_reset_completed":
            detail = "The account password was updated using a valid single-use reset link."
        elif action_key == "password_reset_invalid_attempt":
            detail = "An invalid, expired, revoked, or already-used reset link was rejected."
        elif action_key == "password_reset_throttled":
            detail = "A repeated password reset request was safely throttled."
        elif action_key == "password_reset_mail_failed":
            detail = "Password reset delivery failed; the unused reset token was revoked."
        else:
            detail = _audit_details_text(details or row.get("message") or action_label)
        raw_status = str(details.get("status") or "").lower().replace(" ", "_")
        if any(term in action_key for term in ("failed", "error", "blocked", "invalid_attempt")) or raw_status in {"failed", "error", "provider_error"}:
            status = "Failed"
        elif raw_status in {"success", "completed", "complete"} or action_key in {"records_collected", "search_query", "evidence_save", "save_evidence", "research_saved", "ai_discover", "prediction_created", "password_reset_email_printed", "password_reset_email_queued", "password_reset_completed"}:
            status = "Completed"
        elif raw_status in {"no_results", "no_result"}:
            status = "No results"
        else:
            status = "Recorded"
        records.append({
            "category": category,
            "action": action_label,
            "user": _record_user_display(row),
            "role": row.get("role"),
            "workspace": _workspace_display(row),
            "detail": detail,
            "status": status,
            "time": _log_time_value(row),
            "raw_time": row.get("created_at") or row.get("timestamp"),
            "record_id": str(row.get("_id") or ""),
            "record_collection": "audit_logs",
        })
    records.sort(key=_record_sort_key, reverse=True)
    return records[:50]


def build_audit_records(user_id):
    records = []
    for row in safe_call(lambda: repository.list_searches(user_id=user_id, limit=50), []):
        source_scope = _audit_source_display(row.get("selected_source"))
        statuses = row.get("source_statuses") or {}
        overall_status = _audit_status_display(row.get("status"))
        result_count = int(row.get("result_count") or 0)
        if statuses:
            ebay_status = _marketplace_status_display(statuses.get("ebay"))
            walmart_status = _marketplace_status_display(statuses.get("walmart"))
        else:
            requested = str(row.get("selected_source") or "").lower()
            inferred = "No results" if result_count == 0 else ("Provider error" if overall_status == "Failed" else "Success")
            ebay_status = inferred if requested in {"ebay", "both", "mixed"} else "Not requested"
            walmart_status = inferred if requested in {"walmart", "both", "mixed"} else "Not requested"
        result_categories = row.get("result_categories") or {}
        complete_count = sum(1 for item in result_categories.values() if isinstance(item, dict) and item.get("product_role") == "complete_product")
        comparable_count = complete_count if result_categories else (result_count if row.get("comparison_enabled") is not False else 0)
        source_diagnostics = row.get("source_diagnostics") or {}
        walmart_rejections = source_diagnostics.get("walmart_rejected_count_by_reason") or {}
        walmart_trace = (
            f" Walmart pipeline — raw: {int(source_diagnostics.get('raw_walmart_count') or 0)},"
            f" parsed: {int(source_diagnostics.get('parsed_walmart_count') or 0)},"
            f" normalized: {int(source_diagnostics.get('normalized_walmart_count') or 0)},"
            f" excluded: {int(source_diagnostics.get('excluded_walmart_count') or 0)},"
            f" comparable: {int(source_diagnostics.get('comparable_walmart_count') or 0)},"
            f" displayed: {int(source_diagnostics.get('displayed_walmart_count') or 0)}."
            + (f" Rejection reasons: {_audit_details_text(walmart_rejections)}." if walmart_rejections else "")
        ) if str(row.get("selected_source") or "").lower() in {"walmart", "both", "mixed"} else ""
        records.append({
            "category": "Search records",
            "type_label": "Search",
            "audit_id": str(row.get("search_run_id") or row.get("_id") or ""),
            "query": row.get("keyword") or "—",
            "role": row.get("role"),
            "user": _record_user_display(row),
            "workspace": _workspace_display(row),
            "source": source_scope,
            "source_scope": source_scope,
            "ebay_status": ebay_status,
            "walmart_status": walmart_status,
            "matched_records": result_count,
            "comparable_records": comparable_count,
            "action": "Search submitted",
            "time": _log_time_value(row),
            "status": overall_status,
            "detail": f"Listings returned and retained for this search: {result_count}; comparable complete products: {comparable_count}.{walmart_trace}",
            "raw_time": row.get("created_at") or row.get("timestamp"),
            "record_id": str(row.get("_id") or ""),
            "record_collection": "search_records",
        })
    for row in safe_call(lambda: repository.list_evidence(user_id=user_id, limit=50), []):
        records.append({
            "category": "Saved evidence",
            "type_label": "Saved evidence",
            "audit_id": str(row.get("_id") or ""),
            "query": row.get("title") or "—",
            "role": row.get("role") or "saved evidence",
            "user": _record_user_display(row),
            "workspace": _workspace_display(row),
            "source": row.get("platform") or row.get("source_type") or "Evidence",
            "source_scope": row.get("platform") or row.get("source_type") or "Evidence",
            "ebay_status": "Not requested",
            "walmart_status": "Not requested",
            "matched_records": 1,
            "comparable_records": 1 if row.get("analytics_eligible", True) else 0,
            "action": "Evidence saved",
            "time": _log_time_value(row),
            "status": "Recorded",
            "detail": f"Evidence status: {row.get('evidence_status') or 'saved'}; confidence: {row.get('confidence_level') or 'not stated'}.",
            "raw_time": row.get("created_at") or row.get("saved_at") or row.get("timestamp"),
            "record_id": str(row.get("_id") or ""),
            "record_collection": "evidence_records",
        })
    for row in safe_call(lambda: repository.list_ai_logs(user_id=user_id, limit=50), []):
        records.append({
            "category": "AI analysis",
            "type_label": "AI analysis",
            "audit_id": str(row.get("_id") or ""),
            "query": row.get("keyword") or "—",
            "role": row.get("role") or "researcher",
            "user": _record_user_display(row),
            "workspace": _workspace_display(row),
            "source": row.get("provider") or row.get("model") or "Gemini",
            "source_scope": "AI analysis",
            "ebay_status": "Not requested",
            "walmart_status": "Not requested",
            "matched_records": row.get("record_count") or "",
            "comparable_records": row.get("record_count") or "",
            "action": "AI analysis generated",
            "time": _log_time_value(row),
            "status": "Recorded",
            "detail": row.get("raw_response_summary") or row.get("prompt_summary") or "Read only",
            "raw_time": row.get("created_at") or row.get("timestamp"),
            "record_id": str(row.get("_id") or ""),
            "record_collection": "ai_search_logs",
        })
    for row in safe_call(lambda: repository.list_audit_logs(user_id=user_id, limit=50), []):
        action = row.get("event_type") or row.get("action") or "audit"
        if not any(term in action.lower() for term in ("source", "search", "records_collected", "api_call", "save_evidence", "research_saved")):
            continue
        details = row.get("details") if isinstance(row.get("details"), dict) else {}
        source_value = details.get("search_scope") or details.get("requested_source") or details.get("source") or details.get("actual_source") or "Internal"
        event_status = "Failed" if any(term in action.lower() for term in ("failed", "error")) else ("Completed" if action in {"records_collected", "source_collected"} else "Recorded")
        records.append({
            "category": "Source events",
            "type_label": "Source event",
            "audit_id": str(row.get("_id") or ""),
            "query": details.get("query") or row.get("message") or row.get("label") or "—",
            "role": row.get("role"),
            "user": _record_user_display(row),
            "workspace": _workspace_display(row),
            "source": _audit_source_display(source_value),
            "source_scope": _audit_source_display(source_value),
            "ebay_status": "Not requested",
            "walmart_status": "Not requested",
            "matched_records": details.get("record_count") or "",
            "comparable_records": details.get("comparable_result_count") or "",
            "action": str(action).replace("_", " ").title(),
            "time": _log_time_value(row),
            "status": event_status,
            "detail": _audit_details_text(details or row.get("message") or row.get("label")),
            "raw_time": row.get("created_at") or row.get("timestamp"),
            "record_id": str(row.get("_id") or ""),
            "record_collection": "audit_logs",
        })
    records.sort(key=_record_sort_key, reverse=True)
    return records[:200]


def build_ai_activity_export_records(user_id):
    records = []
    correlated_operations = []
    for row in safe_call(lambda: repository.list_activity_logs(user_id=user_id, limit=200), []):
        action = str(row.get("action") or row.get("event_type") or "").lower()
        if not any(term in action for term in ("ai", "insight", "analysis", "prediction", "forecast")):
            continue
        operation = "Price insight generated"
        if "validation" in action or "evaluat" in action:
            operation = "Forecast validation generated"
        elif "prediction" in action or "forecast" in action:
            operation = "Price forecast generated"
        elif "summary" in action or "analysis" in action:
            operation = "Analysis summary generated"
        created = row.get("started_at") or row.get("created_at") or row.get("timestamp")
        completed = row.get("completed_at") or row.get("created_at") or row.get("timestamp")
        fallback_reason = row.get("fallback_reason")
        correlated_operations.append((str(row.get("query") or "").strip().lower(), str(row.get("model") or "").strip().lower(), _normalized_utc_datetime(created)))
        records.append({
            "event_id": str(row.get("_id") or ""), "operation": operation,
            "user": _record_user_display(row), "role": row.get("role") or "researcher",
            "workspace": _workspace_display(row), "query": row.get("query") or row.get("analysis") or "",
            "provider": row.get("provider") or "Gemini", "model": row.get("model") or "",
            "input_record_count": row.get("record_count") or 0, "scope": row.get("scope") or "Current search results",
            "fallback_used": "Yes" if fallback_reason else "No", "status": "Failed" if row.get("error") else "Completed",
            "started_at": created, "completed_at": completed, "duration_ms": row.get("duration_ms") or "",
            "error_summary": row.get("error_summary") or row.get("error") or fallback_reason or "", "raw_time": created,
        })
    for row in safe_call(lambda: repository.list_ai_logs(user_id=user_id, limit=200), []):
        query_key = str(row.get("keyword") or "").strip().lower()
        model_key = str(row.get("model") or "").strip().lower()
        created_time = _normalized_utc_datetime(row.get("created_at"))
        if created_time and any(query_key == query and model_key == model and correlated_time and abs((created_time - correlated_time).total_seconds()) <= 30 for query, model, correlated_time in correlated_operations):
            continue
        records.append({
            "event_id": str(row.get("_id") or ""), "operation": "Analysis summary generated",
            "user": _record_user_display(row), "role": row.get("role") or "researcher",
            "workspace": _workspace_display(row), "query": row.get("keyword") or "",
            "provider": row.get("provider") or "Gemini", "model": row.get("model") or "",
            "input_record_count": row.get("record_count") or 0, "scope": "Current search results",
            "fallback_used": "No", "status": "Completed", "started_at": row.get("created_at"),
            "completed_at": row.get("completed_at") or row.get("created_at"), "duration_ms": row.get("duration_ms") or "",
            "error_summary": "", "raw_time": row.get("created_at"),
        })
    records.sort(key=_record_sort_key, reverse=True)
    return records[:200]


def get_access_token():
    if not CLIENT_ID or not CLIENT_SECRET:
        raise RuntimeError("eBay search is unavailable because its API credentials are not configured.")
    credentials = base64.b64encode(f"{CLIENT_ID}:{CLIENT_SECRET}".encode()).decode()
    response = backend_http_session().post(
        "https://api.ebay.com/identity/v1/oauth2/token",
        headers={"Authorization": f"Basic {credentials}", "Content-Type": "application/x-www-form-urlencoded"},
        data={"grant_type": "client_credentials", "scope": "https://api.ebay.com/oauth/api_scope"},
        timeout=20,
    )
    response.raise_for_status()
    return response.json()["access_token"]


def search_ebay_items(query, limit=20, item_type="all", region="US"):
    marketplace = EBAY_MARKETPLACE_ID or {"US": "EBAY_US", "GB": "EBAY_GB", "AU": "EBAY_AU", "SG": "EBAY_SG"}.get(region, "EBAY_US")
    response = backend_http_session().get(
        "https://api.ebay.com/buy/browse/v1/item_summary/search",
        headers={"Authorization": f"Bearer {get_access_token()}", "X-EBAY-C-MARKETPLACE-ID": marketplace},
        params={"q": query, "limit": limit}, timeout=25,
    )
    response.raise_for_status()
    now = datetime.now(timezone.utc).isoformat()
    items = []
    for item in response.json().get("itemSummaries", []):
        try:
            price = float(item.get("price", {}).get("value"))
        except (TypeError, ValueError):
            price = None
        item_url = item.get("itemWebUrl")
        shipping_options = item.get("shippingOptions") or []
        shipping_value = (shipping_options[0].get("shippingCost") or {}).get("value") if shipping_options else None
        try:
            shipping = float(shipping_value) if shipping_value is not None else None
        except (TypeError, ValueError):
            shipping = None
        items.append({
            "platform": "eBay", "title": item.get("title"), "product_name": item.get("title"), "price": price, "price_numeric": price, "normalized_price": price,
            "currency": item.get("price", {}).get("currency", "USD"), "image_url": item.get("image", {}).get("imageUrl"),
            "item_url": item_url, "seller": item.get("seller", {}).get("username"), "condition": item.get("condition") or "Not stated",
            "source_url": item_url, "source_type": "live_ebay", "collected_at": now, "confidence_level": "high", "confidence": .95,
            "category": "Marketplace", "brand": item.get("brand") or "Not stated", "availability": "Not provided", "evidence_status": "not_saved", "data_mode": "api", "demo_mode": False,
            "shipping": shipping, "region": region,
        })
    return filter_item_type(items, item_type)


def search_walmart_items(query, limit=12, item_type="all", diagnostics=None):
    collector = WalmartCollector()
    try:
        products = collector.search_products(query, max_results=limit)
        if diagnostics is not None:
            diagnostics.update(_copy_walmart_diagnostics(collector))
    finally:
        collector.close()
    now = datetime.now(timezone.utc).isoformat()
    items = []
    for product in products:
        price = float(product.price) if product.price is not None else None
        items.append({
            "platform": "Walmart", "title": product.name, "product_name": product.name,
            "price": price, "price_numeric": price, "normalized_price": price,
            "currency": product.currency or "USD", "image_url": product.image_url, "item_url": product.url,
            "seller": product.seller or product.brand or "Walmart", "condition": product.availability or "Not stated",
            "brand": product.brand or "Not stated", "category": "Marketplace", "availability": product.availability or "Not provided",
            "source_url": product.url, "source_type": "live_walmart", "collected_at": now,
            "confidence_level": "medium", "confidence": .72, "demo_mode": False,
            "evidence_status": "not_saved", "data_mode": "collected",
        })
    return filter_item_type(items, item_type)


def _copy_walmart_diagnostics(collector):
    diagnostics = getattr(collector, "last_search_diagnostics", {}) or {}
    return json.loads(json.dumps(diagnostics, default=str))


def walmart_demo_items(query, limit=12, category=None, min_price=None, max_price=None):
    """Return clearly labeled Walmart demo records when collection is unavailable."""
    if not _should_show_demo_fallback():
        return []
    rows = repository.search_market_records(query, "Walmart Demo", category, min_price, max_price, limit=limit)
    if not rows:
        rows = repository.search_market_records(query, "Walmart", category, min_price, max_price, limit=limit)
    if not rows:
        rows = demo_search_items(query, limit=limit)
    now = datetime.now(timezone.utc).isoformat()
    demo_rows = []
    for index, item in enumerate(normalize_price_items(rows[:limit])):
        row = dict(item)
        row.pop("_id", None)
        row.update(
            platform="Walmart Demo",
            source_type="Walmart Demo Data",
            source_url=row.get("source_url") or f"https://example.com/walmart-demo/{index + 1}",
            item_url=row.get("source_url") or f"https://example.com/walmart-demo/{index + 1}",
            collected_at=row.get("collected_at") or now,
            confidence=row.get("confidence") or .70,
            confidence_level=row.get("confidence_level") or "sample",
            demo_mode=True,
            data_mode="synthetic",
            evidence_status="not_saved",
        )
        demo_rows.append(row)
    demo_rows = apply_result_filters(demo_rows, category=category, min_price=min_price, max_price=max_price)
    return _filter_relevant_search_items(demo_rows, query)


def load_demo_data():
    path = Path("docs/sample_exports/demo_sample.json")
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8") as stream:
        return json.load(stream)


def backend_http_session():
    """Match the tested eBay client and never inherit broken workstation proxies."""
    client = requests.Session()
    client.trust_env = False
    return client


def demo_search_items(query, limit=12, item_type="all", platform=None):
    if all(term in re.sub(r"[^a-z0-9]", "", query.lower()) for term in ("iphone17pro", "256gb")):
        now = datetime.now(timezone.utc).isoformat()
        rows = [
            ("eBay", "Apple iPhone 17 Pro 256GB - New Unlocked", 1099.00, 12.99, "New", "high", .96, "https://example.com/test/ebay-iphone17"),
            ("Amazon", "Apple iPhone 17 Pro 256GB Unlocked", 1129.00, 0.00, "New", "high", .95, "https://example.com/test/amazon-iphone17"),
            ("Best Buy", "Apple iPhone 17 Pro 256GB", 1149.99, 0.00, "New", "high", .94, "https://example.com/test/bestbuy-iphone17"),
            ("Walmart", "Apple iPhone 17 Pro 256GB Smartphone", 1089.00, 24.95, "New", "medium", .87, "https://example.com/test/walmart-iphone17"),
            ("Shopee", "iPhone 17 Pro 256GB Global Model", 1039.00, 38.00, "New", "medium", .82, "https://example.com/test/shopee-iphone17"),
            ("ReMarket", "Apple flagship phone 256GB", 899.00, 15.00, "Refurbished", "low", .63, "https://example.com/test/remarket-iphone17"),
            ("eBay", "iPhone 17 Pro 256GB Protective Case Bundle", 79.00, 5.99, "New", "medium", .78, "https://example.com/test/ebay-accessory"),
        ]
        items = [{"platform": platform_name, "title": title, "price": price, "price_numeric": price, "shipping": shipping,
                  "currency": "USD", "condition": condition, "availability": "In Stock", "category": "Electronics",
                  "brand": "Apple", "seller": platform_name, "source_url": url, "item_url": url, "product_link": url,
                  "source_type": "synthetic_test_data", "collected_at": now, "confidence_level": confidence,
                  "confidence_score": score, "evidence_status": "not saved", "data_mode": "synthetic", "demo_mode": True,
                  "region": "US"} for platform_name, title, price, shipping, condition, confidence, score, url in rows]
        return filter_item_type(items[:limit], item_type)
    rows = load_demo_data()
    platform_names = {"demo_amazon": "Amazon", "demo_douyin": "Douyin", "demo_xiaohongshu": "Xiaohongshu"}
    if platform in platform_names:
        rows = [row for row in rows if row.get("platform") == platform_names[platform]]
    shopee_mode = platform == "demo_shopee"
    if shopee_mode:
        rows = [row for row in rows if row.get("platform") == "Taobao"] or rows
    words = {word.lower() for word in query.split()}
    matched = [row for row in rows if not words or words & set(str(row.get("name", "")).lower().split())] or rows
    now = datetime.now(timezone.utc).isoformat()
    items = [{
        "platform": "Shopee" if shopee_mode else row.get("platform", "Demo merchant"), "title": row.get("name"), "price": row.get("price"),
        "price_numeric": row.get("price"), "currency": row.get("currency", "USD"), "image_url": None, "item_url": None,
        "seller": row.get("brand", "Generated seller"), "condition": row.get("availability", "Generated"), "source_url": None,
        "source_type": "synthetic_demo", "collected_at": row.get("timestamp") or now, "confidence_level": "sample",
        "category": row.get("category", "Other"), "brand": row.get("brand", "Generated brand"),
        "availability": row.get("availability", "Generated"), "evidence_status": "not saved", "data_mode": "synthetic", "demo_mode": True,
        "shipping": None, "region": "Global",
    } for row in matched[:limit]]
    return filter_item_type(items, item_type)


def sort_items(items, sort_option):
    if sort_option in {"price_asc", "price_desc"}:
        return sorted(items, key=lambda row: row.get("total_price") if row.get("total_price") is not None else float("inf"), reverse=sort_option == "price_desc")
    if sort_option == "platform":
        return sorted(items, key=lambda row: str(row.get("platform", "")).lower())
    if sort_option == "confidence":
        rank = {"high": 0, "medium": 1, "low": 2, "sample": 3}
        return sorted(items, key=lambda row: rank.get(row.get("confidence_level"), 4))
    return items


def calculate_summary(items):
    eligible = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    storage_variants = sorted({(item.get("parsed_model") or parse_canonical_product_model(item.get("title"))).get("storage") for item in eligible if (item.get("parsed_model") or parse_canonical_product_model(item.get("title"))).get("storage")})
    mixed_storage = len(storage_variants) > 1
    totals = [float(item["total_price"]) for item in eligible]
    best = min(eligible, key=lambda item: item["total_price"]) if eligible else None
    return {"total_results": len(items), "comparable_results": len(eligible), "excluded_results": len(items) - len(eligible), "lowest_price": round(min(totals), 2) if totals else None, "highest_price": round(max(totals), 2) if totals else None, "average_price": round(sum(totals) / len(totals), 2) if totals else None, "potential_saving": round(max(totals)-min(totals), 2) if totals else None, "best_platform": None if mixed_storage else (best.get("platform") if best else None), "best_source": None if mixed_storage else (best.get("source_type") if best else None), "storage_variants": storage_variants, "mixed_storage": mixed_storage, "scope_label": "Mixed storage variants" if mixed_storage else "Comparable model variant", "accessories_excluded": sum(bool(item.get("is_accessory")) for item in items)}


NON_COMPARABLE_PRICE_TYPES = {"installment", "subscription", "deposit", "down_payment", "contract", "starting_price", "unknown_suspicious"}


def _price_quality(item):
    explicit_type = str(item.get("price_type") or "").strip().lower().replace("-", "_")
    raw_text = str(item.get("raw_price_text") or item.get("price_text") or item.get("display_price") or "").strip()
    searchable = " ".join([raw_text, str(item.get("title") or ""), str(item.get("offer_type") or "")]).lower()
    if explicit_type in NON_COMPARABLE_PRICE_TYPES:
        price_type = explicit_type
    elif re.search(r"(?:/\s*(?:mo|month)|per\s+month|monthly|installments?|instalments?)", searchable):
        price_type = "installment"
    elif re.search(r"subscription|membership\s+price", searchable):
        price_type = "subscription"
    elif re.search(r"deposit", searchable):
        price_type = "deposit"
    elif re.search(r"down\s*payment", searchable):
        price_type = "down_payment"
    elif re.search(r"with\s+(?:a\s+)?(?:plan|contract)|contract\s+price", searchable):
        price_type = "contract"
    elif re.search(r"starting\s+(?:at|from)|from\s+\$", searchable):
        price_type = "starting_price"
    else:
        return "full_price", None, None
    reasons = {
        "installment": "Monthly instalment price, not a full product price",
        "subscription": "Subscription price, not a full product price",
        "deposit": "Deposit amount, not a full product price",
        "down_payment": "Down-payment amount, not a full product price",
            "contract": "Contract-dependent price, not a standalone full price",
            "starting_price": "Starting price, not a confirmed full product price",
            "unknown_suspicious": "Price structure could not be verified as a full product price",
    }
    billing_period = item.get("billing_period") or ("monthly" if price_type in {"installment", "subscription"} else None)
    return price_type, billing_period, reasons[price_type]


def _normalized_condition(item):
    """Prefer a source condition, then use an explicit title cue as a transparent fallback."""
    value = str(item.get("condition") or "").strip()
    source = "source" if value and value.lower() not in {"unknown", "not stated", "not provided", "n/a"} else None
    if not value or source is None:
        title = str(item.get("title") or item.get("product_name") or "").lower()
        if re.search(r"pre[- ]?owned", title):
            return "Pre-owned", "title_inferred"
        if "refurbished" in title or "renewed" in title:
            return "Refurbished", "title_inferred"
        if "restored" in title:
            return "Restored", "title_inferred"
        if re.search(r"\bused\b", title):
            return "Used", "title_inferred"
        if re.search(r"\b(new|brand new|factory sealed|sealed)\b", title):
            return "New", "title_inferred"
        if re.search(r"\bvery good\b", title):
            return "Very good", "title_inferred"
        if re.search(r"\bexcellent\b", title):
            return "Excellent", "title_inferred"
        return "Not specified by source", "not_specified_by_source"
    normalized = value.lower().replace("_", " ")
    if "pre" in normalized and "own" in normalized:
        return "Pre-owned", source
    if "refurb" in normalized or "renew" in normalized:
        return "Refurbished", source
    if "restor" in normalized:
        return "Restored", source
    if "used" in normalized:
        return "Used", source
    if "new" in normalized:
        return "New", source
    if "very good" in normalized:
        return "Very good", source
    if "excellent" in normalized:
        return "Excellent", source
    return value[:1].upper() + value[1:], source


def display_source_label(item):
    source = str(item.get("source_type") or item.get("source") or item.get("data_source") or "").lower()
    platform = str(item.get("platform") or "")
    if "serpapi" in source or platform.lower() == "walmart":
        return "Walmart via SerpAPI"
    if "ebay" in source or platform.lower() == "ebay":
        return "eBay Browse API"
    return item.get("source_label") or item.get("source_type") or platform or "Source record"


def normalize_price_items(items):
    normalized = []
    for item in items:
        row = dict(item)
        row["platform"] = canonical_marketplace_platform(
            row.get("platform"),
            row.get("source_type") or row.get("source") or row.get("data_source"),
        )
        row["title"] = row.get("title") or row.get("product_name") or "Untitled product"
        row["is_accessory"] = is_accessory(row["title"])
        row["relevance_label"] = "Likely accessory" if row["is_accessory"] else ""
        row["product_name"] = row.get("product_name") or row["title"]
        row["parsed_model"] = parse_canonical_product_model(row["title"])
        row["storage"] = row.get("storage") or row["parsed_model"].get("storage")
        # Canonical condition fields consumed by every workspace/template.
        row["condition_normalized"], row["condition_source"] = _normalized_condition(row)
        row["condition_display"] = row["condition_normalized"]
        row["condition"] = row["condition_display"]  # legacy/export compatibility
        for url_key in ("source_url", "item_url", "product_link", "image_url"):
            if row.get(url_key) and not is_safe_url(row.get(url_key)):
                row[url_key] = None
        confidence = row.get("confidence")
        if not row.get("confidence_level") and confidence is not None:
            row["confidence_level"] = "high" if confidence >= .9 else ("medium" if confidence >= .75 else "low")
        try:
            price = float(row.get("price"))
        except (TypeError, ValueError):
            price = None
        raw_price_text = row.get("raw_price_text") or row.get("price_text") or row.get("display_price")
        if price is not None and (
            not raw_price_text
            or re.fullmatch(r"\s*[-+]?\d+(?:\.\d+)?\s*", str(raw_price_text))
        ):
            raw_price_text = f"{row.get('currency') or 'USD'} {price:.2f}"
        row["raw_price_text"] = raw_price_text or "Price unavailable"
        price_type, billing_period, exclusion_reason = _price_quality(row)
        try:
            comparison_price = float(row.get("normalized_price")) if row.get("normalized_price") is not None else price
        except (TypeError, ValueError):
            comparison_price = price
        if exclusion_reason:
            comparison_price = None
        try:
            shipping = float(row.get("shipping")) if row.get("shipping") is not None else None
        except (TypeError, ValueError):
            shipping = None
        eligible = bool(comparison_price is not None and not exclusion_reason)
        row.update(price=price, price_numeric=comparison_price, normalized_price=comparison_price, price_type=price_type, billing_period=billing_period, analytics_eligible=eligible, exclusion_reason=exclusion_reason, price_quality_status="Comparable full price" if eligible else "Non-comparable", shipping=shipping, shipping_known=shipping is not None,
                   total_price=round(comparison_price + (shipping or 0), 2) if comparison_price is not None else None,
                   total_currency="USD" if row.get("normalized_price") is not None else row.get("currency", "USD"),
                   product_link=row.get("source_url") or row.get("item_url"), display_source=display_source_label(row))
        enrich_listing_category(row, row.get("query") or "")
        normalized.append(row)
    return normalized


def resolve_result_tokens(tokens):
    market_ids, external_ids = [], []
    for token in tokens:
        kind, separator, value = str(token).partition(":")
        if not separator or not value:
            continue
        if kind == "market":
            market_ids.append(value)
        elif kind == "external":
            external_ids.append(value)
    rows = []
    for row in repository.get_market_records(market_ids):
        row["_selection_token"] = f"market:{row['_id']}"
        rows.append(row)
    for row in repository.get_product_results_by_ids(external_ids):
        row["_selection_token"] = f"external:{row['_id']}"
        rows.append(row)
    order = {token: index for index, token in enumerate(tokens)}
    rows.sort(key=lambda row: order.get(row.get("_selection_token"), len(order)))
    return rows


def _comparison_debug(event, **details):
    if app.debug or os.getenv("DEBUG_SEARCH", "false").lower() == "true":
        app.logger.info("comparison workflow %s: %s", event, details)


# There is no product-level reason to cap a user-selected comparison set. The
# browser keeps the selection practical while the server accepts any selected
# result tokens owned by the current Search Run.
COMPARISON_MAX_RECORDS = None


def comparison_metrics(items):
    comparable = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    prices = [float(item["total_price"]) for item in comparable]
    excluded = [item for item in items if item not in comparable]
    return {
        "selected_count": len(items),
        "comparable_count": len(comparable),
        "excluded_count": len(excluded),
        "lowest": round(min(prices), 2) if prices else None,
        "highest": round(max(prices), 2) if prices else None,
        "average": round(sum(prices) / len(prices), 2) if prices else None,
        "spread": round(max(prices) - min(prices), 2) if prices else None,
        "excluded": excluded,
        "comparable": comparable,
    }


def _report_forecast_chart_data(cycle):
    values = [
        ("Benchmark prediction", _snapshot_float(cycle.get("benchmark_predicted_price")) if cycle else None),
        ("AI-assisted forecast", _snapshot_float(cycle.get("ai_predicted_price")) if cycle else None),
        ("Observed market average", _snapshot_float(cycle.get("validation_observed_price")) if cycle else None),
    ]
    values = [(label, value) for label, value in values if value is not None]
    maximum = max((value for _, value in values), default=0) or 1
    return [{"label": label, "value": value, "width": round(value / maximum * 100, 1)} for label, value in values]


def comparison_conclusion(items, metrics):
    comparable = metrics["comparable"]
    if not comparable:
        return "No selected records contain a comparable full product price. Review the excluded rows before continuing."
    lowest_item = min(comparable, key=lambda item: float(item["total_price"]))
    scope_note = "among the selected comparable full-price records"
    conclusion = f"{lowest_item.get('platform') or 'The selected sources'} currently provides the lowest comparable full-price offer {scope_note}."
    if metrics["excluded_count"]:
        types = Counter(item.get("price_type") or "non-comparable" for item in metrics["excluded"])
        labels = {"installment": "monthly-payment", "subscription": "subscription", "deposit": "deposit", "contract": "contract", "starting_price": "starting-price"}
        summary = ", ".join(f"{count} {labels.get(kind, 'non-comparable')} listing{'s' if count != 1 else ''}" for kind, count in types.items())
        conclusion += f" {summary.capitalize()} {'were' if metrics['excluded_count'] != 1 else 'was'} excluded from the summary."
    if metrics["comparable_count"] < 3:
        conclusion += " This is a small selected set and should not be treated as a broad market conclusion."
    return conclusion


def _submitted_selection_tokens(source):
    tokens = []
    for name in ("selected_record_ids", "selected_records", "result_token"):
        tokens.extend(source.getlist(name))
    seen = set()
    ordered = []
    for token in tokens:
        token = str(token or "").strip()
        if token and token not in seen:
            seen.add(token)
            ordered.append(token)
    return ordered


def _tokens_allowed_for_search(record, tokens):
    allowed = set(record.get("result_tokens", [])) if record else set()
    return [token for token in tokens if token in allowed]


def _comparison_group_records(group):
    return normalize_price_items(group.get("selected_records") or [])


def annotate_comparison(items, query):
    comparable_prices = [float(item["total_price"]) for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    if len(comparable_prices) >= 3:
        median = float(pd.Series(comparable_prices, dtype="float64").median())
        for item in items:
            value = item.get("total_price")
            if item.get("analytics_eligible") and value is not None and median > 0 and float(value) < median * 0.12:
                item.update(analytics_eligible=False, normalized_price=None, total_price=None,
                            price_type="unknown_suspicious", price_quality_status="Excluded - suspected non-comparable",
                            exclusion_reason="Suspiciously low price relative to comparable listings")
    totals = [item["total_price"] for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    average = sum(totals) / len(totals) if totals else 0
    stopwords = {"the", "and", "for", "with", "new", "model"}
    terms = [term for term in re.findall(r"[a-z0-9]+", query.lower()) if len(term) > 1 and term not in stopwords]
    for item in items:
        item["is_outlier"] = bool(item.get("price_type") == "unknown_suspicious" or (average and item.get("total_price") is not None and item["total_price"] < average * .40))
        item["outlier_flag"] = "Potential outlier / possible accessory" if item["is_outlier"] else "—"
        title = str(item.get("title", "")).lower()
        matched = sum(term in title for term in terms)
        item["relevance_warning"] = "Title may not match key query terms" if terms and matched < max(1, (len(terms)+1)//2) else ""
        if item.get("is_accessory") or (item.get("relevance_warning") and item.get("data_mode") != "synthetic"):
            item["analytics_eligible"] = False
            item["exclusion_reason"] = item.get("exclusion_reason") or ("Accessory listing, not a comparable product" if item.get("is_accessory") else item["relevance_warning"])
        item["price_quality_status"] = "Comparable full price" if item.get("analytics_eligible") else (item.get("price_quality_status") or "Excluded")
        if item.get("analytics_eligible") and average and item.get("total_price") is not None:
            item["price_gap"] = round(item["total_price"] - average, 2)
            item["market_position"] = "Within selected range"
        else:
            item["price_gap"], item["market_position"] = None, "Excluded"
    comparable = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    if comparable:
        lowest = min(float(item["total_price"]) for item in comparable)
        highest = max(float(item["total_price"]) for item in comparable)
        for item in comparable:
            value = float(item["total_price"])
            item["market_position"] = "Lowest comparable" if value == lowest else ("Highest comparable" if value == highest else "Within selected range")
    return items


def build_price_insight(items, summary):
    priced = [item for item in items if item.get("total_price") is not None]
    if not priced:
        return "No priced listings are available for comparison. Broaden the filters or choose another source."
    cheapest = min(priced, key=lambda item: item["total_price"])
    eligible = [item for item in priced if not item.get("is_accessory") and not item.get("relevance_warning")]
    recommended = min(eligible, key=lambda item: item["total_price"]) if eligible else cheapest
    credible_saving = max(item["total_price"] for item in priced) - recommended["total_price"]
    unknown_shipping = sum(not item.get("shipping_known") for item in priced)
    low_confidence = sum(item.get("confidence_level") in {"low", "sample"} for item in priced)
    shipping_note = ("Shipping is included where reported; " + (f"{unknown_shipping} listing(s) do not disclose shipping, so their totals may be understated." if unknown_shipping else "all compared listings report shipping."))
    limitation = f"Review condition, seller reputation and source confidence before deciding; {low_confidence} listing(s) use low or sample confidence."
    outlier_note = f"The absolute cheapest listing from {cheapest.get('platform')} is flagged as a potential outlier and is not the recommended baseline. " if cheapest.get("is_outlier") else ""
    return f"{outlier_note}{recommended.get('platform')} has the best comparable total at {recommended.get('currency')} {recommended['total_price']:.2f}. Choosing it instead of the highest-priced option could save {recommended.get('currency')} {credible_saving:.2f}. {shipping_note} {limitation}"


def perform_search(query, source, item_type, region="US"):
    if source == "ebay":
        return search_ebay_items(query, item_type=item_type, region=region)
    if source == "ai":
        items, model, raw = search_prices(query)
        repository.log_ai_search(session["user_id"], query, model, "Retail price collection through backend AI web search.", raw)
        return filter_item_type(items, item_type)
    if source == "walmart":
        return search_walmart_items(query, item_type=item_type)
    if source == "demo" or source.startswith("demo_"):
        return demo_search_items(query, item_type=item_type, platform="demo_all")
    raise ValueError("Unsupported search source.")


def apply_result_filters(items, platform="", category="", min_price=None, max_price=None, currency="", condition=""):
    def accepted(item):
        if platform and platform.lower() not in str(item.get("platform", "")).lower():
            return False
        if category and category.lower() != str(item.get("category", "Other")).lower():
            return False
        if currency and currency != str(item.get("currency", "")):
            return False
        if condition and condition.lower() not in str(item.get("condition", "")).lower():
            return False
        try:
            price = float(item.get("total_price", item.get("price")))
        except (TypeError, ValueError):
            return min_price is None and max_price is None
        return (min_price is None or price >= min_price) and (max_price is None or price <= max_price)
    return [item for item in items if accepted(item)]


def optional_float(value):
    try:
        return float(value) if value not in (None, "") else None
    except ValueError:
        return None


def present_audit_logs(rows):
    labels = {"login": "Login", "logout": "Logout", "register": "Register", "password_reset_requested": "Password reset requested", "password_reset": "Password reset",
              "admin_user_update": "Admin user update", "search_started": "Search started", "records_collected": "Records collected",
              "comparison_generated": "Comparison generated", "insight_generated": "Insight generated",
              "compare": "Compare selected", "ai_discover": "AI discover", "save_evidence": "Save evidence",
              "research_saved": "Research saved", "api_call_failed": "API call failed", "mongodb_save_failed": "MongoDB save failed",
              "evidence_deleted": "Evidence deleted", "research_deleted": "Research deleted", "search_deleted": "Search deleted",
              "audit_log_archived": "Audit log archived", "activity_log_archived": "Activity log archived", "testing_records_deleted": "Test data deleted", "testing_records_reset": "Test data reset"}
    presented = []
    for row in rows:
        details, event = row.get("details") or {}, row.get("event_type") or row.get("action")
        if event == "login": message = "Signed in to the workspace."
        elif event == "register": message = f"Created a new {details.get('role', 'user')} account."
        elif event == "password_reset_requested": message = "Generated a prototype reset link."
        elif event == "password_reset": message = "Completed a password reset."
        elif event == "admin_user_update": message = f"Updated {details.get('target_user', 'a user')}."
        elif event == "logout": message = "Signed out of the workspace."
        elif event == "search_started": message = f"Started search for “{details.get('query', '')}” using {details.get('source', 'selected source')}."
        elif event == "records_collected": message = f"Collected {details.get('record_count', 0)} normalized record(s) from {details.get('actual_source', 'source')}."
        elif event == "comparison_generated": message = f"Compared {details.get('record_count', details.get('product_count', 0))} record(s) using total price."
        elif event == "compare": message = f"Compared {details.get('product_count', 0)} selected MongoDB market record(s)."
        elif event == "ai_discover": message = f"Generated a {details.get('generation_mode', 'rule-based')} market insight from {details.get('record_count', 0)} record(s)."
        elif event == "save_evidence": message = f"Saved {details.get('saved_count', details.get('record_count', 0))} selected record(s) to evidence_records."
        elif event == "insight_generated": message = f"Generated {details.get('mode') or 'price insight'}."
        elif event == "research_saved": message = f"Saved research for “{details.get('query', '')}” with {details.get('record_count', 0)} record(s)."
        elif event == "api_call_failed": message = f"{details.get('source', 'API')} failed ({details.get('error_type', 'connection error')}); fallback applied."
        elif event == "mongodb_save_failed": message = "MongoDB was unavailable; the research package remains in memory fallback storage."
        elif event == "evidence_deleted": message = "Saved evidence was removed from the active view."
        elif event == "research_deleted": message = "Saved research was archived from the active view."
        elif event == "search_deleted": message = "Search history was cleared from the active view."
        elif event == "audit_log_archived": message = "An audit log entry was archived."
        elif event == "activity_log_archived": message = "An activity log entry was archived."
        elif event == "testing_records_deleted": message = "Synthetic test data was deleted."
        elif event == "testing_records_reset": message = "Synthetic test data was reset."
        else: message = "Workspace activity recorded."
        presented.append({**row, "event_type": event, "created_at": row.get("created_at") or row.get("timestamp"), "label": labels.get(event, event.replace("_", " ").title() if event else "Event"), "message": message})
    return presented


def user_can_manage_record(record_user_id):
    return session.get("role") == "administrator" or str(record_user_id) == str(session.get("user_id"))


def results_dataframe(items):
    rows = []
    for item in items:
        try:
            price = float(item.get("normalized_price") if item.get("normalized_price") is not None else item.get("total_price", item.get("price")))
        except (TypeError, ValueError):
            continue
        rows.append({"price": price, "condition": item.get("condition") or "Unknown", "seller": item.get("seller") or "Unknown", "platform": item.get("platform") or "Unknown"})
    return pd.DataFrame(rows, columns=["price", "condition", "seller", "platform"])


def create_chart(df, search_id, chart_type):
    filename = f"{chart_type}_{search_id}.png"
    plt.figure(figsize=(8, 4.8))
    if chart_type == "price_distribution":
        plt.hist(df["price"], bins=min(10, max(3, len(df))), color="#60a5fa", edgecolor="white")
        plt.xlabel("Price")
        plt.ylabel("Listings")
        plt.title("Price distribution")
    else:
        column = "platform" if chart_type == "platforms" else "condition"
        df[column].value_counts().head(8).plot(kind="bar", color="#2563eb")
        plt.xlabel(column.title())
        plt.ylabel("Listings")
        plt.title(f"{column.title()} distribution")
        plt.xticks(rotation=25, ha="right")
    plt.tight_layout()
    plt.savefig(CHART_DIR / filename)
    plt.close()
    return filename


@app.route("/")
def home():
    return render_template("landing.html", membership_plans=MEMBERSHIP_PLANS)


@app.route("/membership")
@login_required
def membership():
    user = current_user()
    if (user.get("active_role") or user.get("role")) == "administrator":
        flash("Administrator accounts use system access and do not manage a membership tier.", "info")
        return redirect(url_for("administrator_dashboard"))
    current_tier = normalize_membership_tier(user.get("membership_tier") or user.get("plan"))
    current_role = user.get("active_role") or user.get("role")
    return render_template(
        "membership.html",
        current_tier=current_tier,
        current_role=current_role,
        membership_plans=MEMBERSHIP_PLANS,
    )


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    user = current_user()
    current_tier = normalize_membership_tier(user.get("membership_tier") or user.get("plan"))
    if request.method == "POST":
        display_name = request.form.get("display_name", "").strip()
        institution = request.form.get("institution", "").strip()
        if not display_name:
            flash("Display name is required.", "error")
        elif len(display_name) > 80:
            flash("Display name must be 80 characters or fewer.", "error")
        elif len(institution) > 120:
            flash("Institution or company must be 120 characters or fewer.", "error")
        else:
            existing = repository.get_user_by_display_name(display_name)
            if existing and str(existing.get("_id")) != str(user["_id"]):
                flash("That display name is already in use.", "error")
            else:
                updated = repository.update_user(user["_id"], {"display_name": display_name, "username": display_name, "institution": institution})
                session["username"] = updated.get("display_name") or updated.get("username") or "User"
                session["display_name"] = session["username"]
                repository.log_event(user["_id"], "profile_updated", user.get("active_role") or user.get("role"), session["username"], {"fields": ["display_name", "institution"]})
                flash("Profile updated.", "success")
                return redirect(url_for("profile"))
    return render_template("profile.html", current_tier=current_tier, current_plan=PLAN_BY_ID.get(current_tier, PLAN_BY_ID["basic"]))


@app.post("/profile/avatar")
@login_required
def upload_avatar():
    user = current_user()
    upload = request.files.get("avatar")
    if not upload or not upload.filename:
        flash("Choose an image to upload.", "error")
        return redirect(url_for("profile"))
    filename = secure_filename(upload.filename)
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ALLOWED_AVATAR_EXTENSIONS:
        flash("Avatar must be a PNG, JPG, JPEG, or WEBP image.", "error")
        return redirect(url_for("profile"))
    upload.stream.seek(0, os.SEEK_END)
    size = upload.stream.tell()
    upload.stream.seek(0)
    if size > MAX_AVATAR_SIZE_BYTES:
        flash("Avatar image must be 2MB or smaller.", "error")
        return redirect(url_for("profile"))
    avatar_name = f"{user['_id']}-{secrets.token_hex(8)}.{extension}"
    avatar_path = AVATAR_UPLOAD_DIR / avatar_name
    upload.save(avatar_path)
    previous = user.get("avatar_path")
    repository.update_user(user["_id"], {"avatar_path": f"uploads/avatars/{avatar_name}"})
    if previous:
        previous_path = Path("static") / previous
        if previous_path.exists() and previous_path.parent == AVATAR_UPLOAD_DIR:
            previous_path.unlink(missing_ok=True)
    repository.log_event(user["_id"], "avatar_uploaded", user.get("active_role") or user.get("role"), user.get("username"), {"filename": avatar_name})
    flash("Profile photo updated.", "success")
    return redirect(url_for("profile"))


@app.post("/profile/avatar/remove")
@login_required
def remove_avatar():
    user = current_user()
    previous = user.get("avatar_path")
    repository.update_user(user["_id"], {"avatar_path": ""})
    if previous:
        previous_path = Path("static") / previous
        if previous_path.exists() and previous_path.parent == AVATAR_UPLOAD_DIR:
            previous_path.unlink(missing_ok=True)
    repository.log_event(user["_id"], "avatar_removed", user.get("active_role") or user.get("role"), user.get("username"), {})
    flash("Profile photo removed. Initials will be shown instead.", "success")
    return redirect(url_for("profile"))


@app.post("/membership/upgrade")
@login_required
def upgrade_membership():
    user = current_user()
    if (user.get("active_role") or user.get("role")) == "administrator":
        flash("Administrator accounts use system access and do not manage a membership tier.", "info")
        return redirect(url_for("administrator_dashboard"))
    if not app.config.get("DEMO_MEMBERSHIP_UPGRADE_ENABLED") or not _is_non_production_runtime():
        abort(403)
    target_tier = normalize_membership_tier(request.form.get("membership_tier"))
    if target_tier not in {"premium", "professional"}:
        abort(400)
    current_tier = normalize_membership_tier(user.get("membership_tier") or user.get("plan"))
    if target_tier == current_tier:
        flash(f"You are already on the {tier_label(current_tier)} plan.", "info")
        return redirect(url_for("membership"))
    if membership_rank(target_tier) < membership_rank(current_tier):
        flash("Plan downgrades are managed by the administrator in this prototype.", "warning")
        return redirect(url_for("membership"))
    repository.update_user(user["_id"], {"membership_tier": target_tier, "plan": target_tier})
    repository.log_event(user["_id"], "membership_upgrade", user.get("active_role") or user.get("role"), user.get("username"), {"from": current_tier, "to": target_tier, "billing": "prototype"})
    flash(f"Prototype upgrade complete. Your workspace now uses the {tier_label(target_tier)} plan.", "success")
    return redirect(url_for("membership"))


@app.route("/sample-analysis")
def sample_analysis():
    """Public, read-only example: no search, comparison, or persistence actions."""
    sample_records = [
        {"platform": "eBay", "title": "iPhone sample 256GB - eBay listing", "price": 699.00, "normalized_price": 699.00, "currency": "USD", "availability": "In stock", "condition": "New", "source_type": "Marketplace"},
        {"platform": "eBay", "title": "iPhone sample 256GB - renewed listing", "price": 729.00, "normalized_price": 729.00, "currency": "USD", "availability": "In stock", "condition": "Refurbished", "source_type": "Marketplace"},
        {"platform": "Walmart", "title": "iPhone sample 256GB - store listing", "price": 679.00, "normalized_price": 679.00, "currency": "USD", "availability": "Limited stock", "condition": "New", "source_type": "Marketplace"},
        {"platform": "Walmart", "title": "iPhone sample 256GB - online listing", "price": 689.00, "normalized_price": 689.00, "currency": "USD", "availability": "In stock", "condition": "New", "source_type": "Marketplace"},
        {"platform": "Demo Source", "title": "iPhone sample 256GB - demo record A", "price": 689.00, "normalized_price": 689.00, "currency": "USD", "availability": "In stock", "condition": "New", "source_type": "Demo Source"},
        {"platform": "Demo Source", "title": "iPhone sample 256GB - demo record B", "price": 709.00, "normalized_price": 709.00, "currency": "USD", "availability": "In stock", "condition": "New", "source_type": "Demo Source"},
    ]
    items = normalize_price_items(sample_records)
    available_items = [item for item in items if str(item.get("availability", "")).lower() not in {"out of stock", "unavailable", "not available"}]
    best = min(available_items, key=lambda item: item["normalized_price"]) if available_items else None
    prices = [item["normalized_price"] for item in items]
    platform_stats = []
    by_platform = {}
    for item in items:
        bucket = by_platform.setdefault(item["platform"], {"count": 0, "total": 0.0, "available": 0})
        bucket["count"] += 1
        bucket["total"] += float(item["normalized_price"])
        if str(item.get("availability", "")).lower() not in {"out of stock", "unavailable", "not available"}:
            bucket["available"] += 1
    total_records = len(items)
    for platform, values in by_platform.items():
        average_price = round(values["total"] / values["count"], 2)
        share = round((values["count"] / total_records) * 100, 1) if total_records else 0
        platform_stats.append({
            "platform": platform,
            "count": values["count"],
            "average_price": average_price,
            "share": share,
            "available": values["available"],
            "is_best": best and best["platform"] == platform,
        })
    platform_stats.sort(key=lambda row: row["average_price"])
    lowest_price = round(min(prices), 2)
    highest_price = round(max(prices), 2)
    price_span = max(highest_price - lowest_price, 1)
    summary = {
        "lowest_price": lowest_price,
        "average_price": round(sum(prices) / len(prices), 2),
        "highest_price": highest_price,
        "best_platform": best["platform"] if best else "Not available",
        "total_records": total_records,
        "platform_count": len(by_platform),
    }
    workflow_steps = [
        {"title": "Search", "description": "Collect sample price records from available sources."},
        {"title": "Compare", "description": "Review price differences and platform coverage."},
        {"title": "Save Evidence", "description": "Preserve selected records for later review."},
        {"title": "Analyze", "description": "Summarise price range, average price, and best visible option."},
        {"title": "Review Results", "description": "Inspect sample records and workflow output."},
    ]
    sample_insight = (
        f"The sample records show a price range from {summary['lowest_price']:.2f} to {summary['highest_price']:.2f}, "
        f"with {summary['best_platform']} currently the strongest visible option among the displayed records."
    )
    return render_template(
        "public_sample.html",
        items=items,
        summary=summary,
        platform_stats=platform_stats,
        workflow_steps=workflow_steps,
        sample_insight=sample_insight,
        demo_keyword="iPhone sample",
        coverage_note=f"{summary['total_records']} demonstration records across {summary['platform_count']} platforms",
        price_span=price_span,
    )


AUTH_REGISTER_ROLES = {key: label for key, label in ROLES.items() if key != "administrator"}


def dashboard_for_role(role):
    return url_for(f"{role}_dashboard") if role in ROLES else url_for("login")


def reset_token_hash(token):
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def is_future_datetime(value):
    if not value:
        return False
    if getattr(value, "tzinfo", None) is None:
        value = value.replace(tzinfo=timezone.utc)
    return value > utcnow()


PASSWORD_RESET_GENERIC_MESSAGE = "If an account exists for this email, a password reset link has been sent."


def password_reset_token_ttl_minutes():
    try:
        return max(1, min(int(os.getenv("PASSWORD_RESET_TOKEN_TTL_MINUTES", "30")), 1440))
    except (TypeError, ValueError):
        return 30


def password_reset_base_url():
    return canonical_app_base_url()


def password_reset_url(token):
    return application_email_url("password_reset", token=token)


def registration_login_url():
    return application_email_url("login")


def price_alert_monitor_url(monitor_id):
    # Scheduled refreshes run without a Flask request/application context.
    return application_email_url("watchlist_monitor", monitor_id=monitor_id)


def valid_password_reset_record(token):
    record = repository.get_password_reset_token(reset_token_hash(token))
    if not record or record.get("used_at") is not None or record.get("revoked_at") is not None or not is_future_datetime(record.get("expires_at")):
        return None, None
    user = repository.get_user_by_id(record.get("user_id"))
    if not user or user.get("account_status") != "active":
        return None, None
    return record, user


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if current_user():
            return redirect(url_for("dashboard_redirect"))
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = repository.get_user_by_email(email) if email and password else None
        if not user or user.get("account_status") != "active" or not check_password_hash(user.get("password_hash") or "", password):
            if user:
                repository.record_login_failure(user["_id"])
            flash("Invalid email or password.", "error")
            return render_template("login.html"), 200
        assigned_role = single_account_role(user)
        if user_roles(user) != [assigned_role] or user.get("role") != assigned_role or user.get("active_role") != assigned_role:
            repository.update_user(user["_id"], {
                "roles": [assigned_role],
                "primary_role": assigned_role,
                "active_role": assigned_role,
                "role": assigned_role,
            })
            user = repository.get_user_by_id(user["_id"])
        repository.record_login_success(user["_id"])
        session.clear()
        session.update(
            user_id=user["_id"],
            username=user.get("display_name") or user.get("username"),
            display_name=user.get("display_name") or user.get("username"),
            email=user.get("email"),
            roles=[assigned_role],
            active_role=assigned_role,
            role=assigned_role,
        )
        session.permanent = True
        repository.log_event(user["_id"], "login", session["active_role"], user.get("display_name") or user.get("username"), {"login_method": "password"})
        return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("dashboard_redirect")))
    if current_user():
        return redirect(url_for("dashboard_redirect"))
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        display_name = request.form.get("display_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")
        role = request.form.get("role", "consumer")
        membership_tier = normalize_membership_tier(request.form.get("membership_tier") or "basic")
        if not display_name or not email or not password or not confirm_password:
            flash("All register fields are required.", "error")
        elif role not in AUTH_REGISTER_ROLES:
            flash("Normal users cannot register as Administrator.", "error")
        elif len(password) < 8:
            flash("Password must be at least 8 characters long.", "error")
        elif password != confirm_password:
            flash("Passwords do not match.", "error")
        elif repository.get_user_by_display_name(display_name):
            flash("That display name is already in use.", "error")
        elif repository.get_user_by_email(email):
            flash("An account with that email already exists.", "error")
        else:
            user = repository.create_user(display_name, email, generate_password_hash(password), role, account_status="active", roles=[role], active_role=role, plan=membership_tier, membership_tier=membership_tier)
            repository.log_event(user["_id"], "register", role, user.get("display_name"), {"email": user.get("email"), "primary_role": user.get("primary_role"), "assigned_roles": user.get("roles"), "active_role": user.get("active_role"), "plan": user.get("plan"), "membership_tier": user.get("membership_tier")})
            welcome_failed = False
            if app.config.get("REGISTRATION_WELCOME_EMAIL_ENABLED", True):
                try:
                    delivery = PasswordResetMailService().send_registration_welcome(
                        user.get("email"),
                        user.get("display_name"),
                        registration_login_url(),
                        user.get("membership_tier") or user.get("plan") or "basic",
                    )
                    repository.log_event(
                        user["_id"], "registration_welcome_email_delivered", role, user.get("display_name"),
                        {"delivery_mode": os.getenv("EMAIL_MODE", "console").lower(), "delivery_status": delivery},
                    )
                except Exception as exc:
                    welcome_failed = True
                    repository.log_event(
                        user["_id"], "registration_welcome_email_failed", role, user.get("display_name"),
                        {"delivery_mode": os.getenv("EMAIL_MODE", "console").lower(), "error_type": exc.__class__.__name__},
                    )
                    app.logger.warning("Registration welcome email delivery failed (%s)", exc.__class__.__name__)
            if welcome_failed:
                flash("Account created. The welcome email could not be delivered, but you can sign in normally.", "info")
            else:
                flash("Account created. Please log in.", "success")
            return redirect(url_for("login"))
    return render_template("register.html", roles=AUTH_REGISTER_ROLES, membership_plans=MEMBERSHIP_PLANS, selected_plan=normalize_membership_tier(request.args.get("plan") or "basic"))


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        if not email:
            flash("Email address is required.", "error")
            return render_template("forgot_password.html"), 200
        user = repository.get_user_by_email(email)
        if user and user.get("account_status") == "active":
            now = utcnow()
            latest = repository.latest_password_reset_token(user["_id"])
            latest_created = latest.get("created_at") if latest else None
            if latest_created and getattr(latest_created, "tzinfo", None) is None:
                latest_created = latest_created.replace(tzinfo=timezone.utc)
            throttled = bool(latest_created and latest_created > now - timedelta(seconds=60))
            if throttled:
                repository.log_event(user["_id"], "password_reset_throttled", user.get("role"), user.get("display_name"), {"reason": "request_interval"})
            else:
                token = secrets.token_urlsafe(32)
                ttl_minutes = password_reset_token_ttl_minutes()
                repository.create_password_reset_token(user["_id"], reset_token_hash(token), now + timedelta(minutes=ttl_minutes))
                repository.log_event(user["_id"], "password_reset_requested", user.get("role"), user.get("display_name"), {"delivery_mode": os.getenv("EMAIL_MODE", "console").lower()})
                try:
                    delivery = PasswordResetMailService().send_password_reset(user.get("email"), password_reset_url(token), ttl_minutes)
                    repository.log_event(user["_id"], f"password_reset_email_{delivery}", user.get("role"), user.get("display_name"), {"delivery_mode": os.getenv("EMAIL_MODE", "console").lower()})
                except Exception as exc:
                    repository.revoke_unused_password_reset_tokens(user["_id"])
                    repository.log_event(user["_id"], "password_reset_mail_failed", user.get("role"), user.get("display_name"), {"error_type": exc.__class__.__name__})
                    app.logger.warning("Password reset delivery failed (%s)", exc.__class__.__name__)
        flash(PASSWORD_RESET_GENERIC_MESSAGE, "success")
    return render_template("forgot_password.html")


@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    user = current_user()
    account = repository.get_user_by_id(user["_id"])
    if request.method == "POST":
        current_password = request.form.get("current_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")
        if not account or not check_password_hash(account.get("password_hash") or "", current_password):
            flash("Current password is incorrect.", "error")
        elif len(new_password) < 8:
            flash("New password must be at least 8 characters long.", "error")
        elif new_password != confirm_password:
            flash("New password and confirmation do not match.", "error")
        else:
            repository.update_user(user["_id"], {"password_hash": generate_password_hash(new_password)})
            repository.log_event(user["_id"], "password_changed", user.get("active_role") or user.get("role"), user.get("display_name"), {})
            flash("Password changed successfully.", "success")
            return redirect(url_for("profile"))
    return render_template("change_password.html")


@app.route("/reset-password/<token>", methods=["GET", "POST"])
def reset_password(token):
    reset_record, user = valid_password_reset_record(token)
    token_valid = bool(reset_record and user)
    if request.method == "POST":
        new_password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")
        if not token_valid:
            linked_record = repository.get_password_reset_token(reset_token_hash(token))
            linked_user = repository.get_user_by_id(linked_record.get("user_id")) if linked_record else None
            if linked_user:
                repository.log_event(linked_user["_id"], "password_reset_invalid_attempt", linked_user.get("role"), linked_user.get("display_name"), {"reason": "invalid_or_expired"})
            flash("This password reset link is invalid or has expired.", "error")
        elif not new_password or not confirm_password:
            flash("Both password fields are required.", "error")
        elif len(new_password) < 8:
            flash("Password must be at least 8 characters long.", "error")
        elif new_password != confirm_password:
            flash("Passwords do not match.", "error")
        else:
            completed_at = utcnow()
            repository.update_user(user["_id"], {"password_hash": generate_password_hash(new_password), "password_updated_at": completed_at})
            repository.mark_password_reset_token_used(reset_record["_id"], completed_at)
            repository.revoke_unused_password_reset_tokens(user["_id"], exclude_token_id=reset_record["_id"], revoked_at=completed_at)
            repository.log_event(user["_id"], "password_reset_completed", user.get("role"), user.get("display_name"), {})
            session.clear()
            flash("Password updated successfully. Sign in with your new password.", "success")
            return redirect(url_for("login"))
    return render_template("reset_password.html", token_valid=token_valid)


@app.post("/role-switch")
@login_required
def role_switch():
    flash("Each account has one assigned workspace role. Contact an administrator if your role needs to change.", "info")
    return redirect(url_for("dashboard_redirect"))


@app.post("/logout")
def logout():
    if current_user():
        repository.log_event(session["user_id"], "logout", session.get("role"), session.get("username"), {})
    session.clear()
    return redirect(url_for("home"))


@app.route("/dashboard")
@login_required
def dashboard_redirect():
    user = current_user()
    return redirect(dashboard_for_role(user["active_role"]))


@app.route("/dashboard/<role>")
def invalid_dashboard_role(role):
    """Catch unknown dashboard roles without exposing a server error."""
    if role not in ROLES:
        flash("Choose a valid workspace role.", "error")
        return redirect(url_for("login"))
    user = current_user()
    if not user:
        return redirect(url_for("login"))
    return redirect(dashboard_for_role(user["active_role"]))


@app.route("/dashboard/consumer")
@role_required("consumer")
def consumer_dashboard():
    records, source_kind = _market_records_for_identity(session["user_id"])
    records = [_label_record_source(row, source_kind) for row in records]
    visible_records = [row for row in records if not _is_demo_market_record(row)] or records
    searches = repository.list_searches(session["user_id"], limit=8)
    evidence = repository.list_evidence(session["user_id"], limit=4)
    watchlist_items = _watchlist_items_for_user(session["user_id"], include_archived=False)
    return render_template(
        "dashboard_consumer_role.html",
        searches=searches,
        evidence=evidence,
        watchlist_items=watchlist_items,
        evidence_count=len(evidence),
        watchlist_count=len(watchlist_items),
        latest_search=searches[0] if searches else None,
        latest_evidence=evidence[0] if evidence else None,
        latest_watchlist=watchlist_items[0] if watchlist_items else None,
        platforms=sorted({row.get("platform") for row in visible_records if row.get("platform") and "mongodb" not in str(row.get("platform")).lower()}),
        categories=sorted({row.get("category") for row in visible_records if row.get("category")}),
        market_message="Market overview based on saved and collected records." if source_kind != "demo" else "No saved market records found. Demo records are shown for preview.",
        source_kind=source_kind,
    )


@app.route("/dashboard/retailer")
@role_required("retailer")
def retailer_dashboard():
    requested_monitor_id = request.args.get("monitor_id")
    monitor_id = requested_monitor_id or None
    dashboard = build_retailer_dashboard_view(session["user_id"], monitor_id)
    return render_template(
        "dashboard_retailer_saas.html",
        market_message=dashboard["market_message"],
        market_badge=dashboard["market_badge"],
        active_monitor_count=dashboard.get("active_monitor_count", 0),
        latest_snapshot_time=dashboard.get("latest_snapshot_time"),
        scope_label=dashboard.get("scope_label", "No active monitoring data yet."),
        source_kind=dashboard["source_kind"],
        demo_records=dashboard["demo_records"],
        real_records=dashboard["records"],
        searches=repository.list_searches(session["user_id"], limit=10),
        distribution=dashboard["distribution"],
        distribution_rows=dashboard["distribution_rows"],
        prices=dashboard["prices"],
        sourcing_insight=dashboard["sourcing_insight"],
        products=dashboard["records"],
        platforms=dashboard["platforms"],
        categories=dashboard["categories"],
        monitors=dashboard.get("monitors", []),
        selected_monitor_id=dashboard.get("selected_monitor_id", ""),
        selected_monitor=dashboard.get("selected_monitor"),
        selected_metrics=dashboard.get("selected_metrics"),
        trend_chart=dashboard.get("trend_chart", {}),
        portfolio_metrics=dashboard.get("portfolio_metrics", {}),
        portfolio_scope=dashboard.get("portfolio_scope", True),
        synthetic=dashboard["source_kind"] == "demo",
    )


def build_researcher_dashboard_view(user_id):
    """Build an owner-scoped research workflow view; never expose global records."""
    health = get_service_health()
    searches = safe_call(lambda: repository.list_searches(user_id, limit=20), [])
    products = safe_call(lambda: repository.list_products(user_id, limit=20), [])
    evidence = safe_call(lambda: repository.list_evidence(user_id, limit=20), [])
    research_packages = safe_call(lambda: repository.list_research(user_id, limit=20), [])
    analyses = safe_call(lambda: repository.list_analysis_records(user_id, limit=20), [])
    ai_logs = safe_call(lambda: repository.list_ai_logs(user_id, limit=12), [])
    audit_logs = present_audit_logs(safe_call(lambda: repository.list_audit_logs(user_id, limit=20), []))
    source_warning_count = sum(
        row.get("event_type") in {"api_call_failed", "mongodb_save_failed"}
        for row in audit_logs
    )
    snapshot = {
        "search_records": len(safe_call(lambda: repository.list_searches(user_id, limit=0), [])),
        "product_results": len(safe_call(lambda: repository.list_products(user_id, limit=0), [])),
        "evidence_records": len(safe_call(lambda: repository.list_evidence(user_id, limit=0), [])),
        "research_records": len(safe_call(lambda: repository.list_research(user_id, limit=0), [])),
        "ai_search_logs": len(safe_call(lambda: repository.list_ai_logs(user_id, limit=0), [])),
        "analytics_reports": len(safe_call(lambda: repository.list_analysis_records(user_id, limit=0), [])),
    }
    sources = [
        {"name": "Atlas evidence store", "status": health[0]["status"], "type": "Persistent research and evidence storage", "detail": health[0]["detail"]},
        {"name": "eBay marketplace source", "status": health[1]["status"], "type": "Comparable marketplace listing collection", "detail": health[1]["detail"]},
        {"name": "Walmart marketplace source", "status": health[2]["status"], "type": "Comparable marketplace listing collection", "detail": health[2]["detail"]},
        {"name": "Gemini analysis service", "status": health[3]["status"], "type": "Grounded analysis support, not a marketplace source", "detail": health[3]["detail"]},
    ]
    return {
        "searches": searches,
        "products": products,
        "evidence": normalize_price_items(evidence),
        "research_packages": research_packages,
        "analyses": analyses,
        "ai_logs": ai_logs,
        "audit_logs": audit_logs,
        "source_warning_count": source_warning_count,
        "snapshot": snapshot,
        "pipeline": [
            {"label": "Search runs", "count": snapshot["search_records"], "endpoint": "search"},
            {"label": "Saved evidence", "count": snapshot["evidence_records"], "endpoint": "saved"},
            {"label": "Research packages", "count": snapshot["research_records"], "endpoint": "saved"},
            {"label": "Analysis reports", "count": snapshot["analytics_reports"], "endpoint": "analytics_compatibility"},
        ],
        "sources": sources,
    }


@app.route("/dashboard/researcher")
@role_required("researcher")
def researcher_dashboard():
    dashboard = build_researcher_dashboard_view(session["user_id"])
    return render_template(
        "dashboard_researcher_role.html",
        **dashboard,
    )


@app.route("/dashboard/administrator")
@role_required("administrator")
def administrator_dashboard():
    health = get_service_health()
    users = safe_call(lambda: repository.list_users(limit=50), [])
    test_account_count = sum(1 for user in users if is_test_account_row(user))
    search_counts = Counter(str(row.get("user_id")) for row in safe_call(lambda: repository.list_searches(limit=0), []))
    evidence_counts = Counter(str(row.get("user_id")) for row in safe_call(lambda: repository.list_evidence(limit=0), []))
    last_activity = {}
    for row in safe_call(lambda: repository.list_audit_logs(limit=0), []):
        user_id = str(row.get("user_id"))
        if user_id not in last_activity:
            last_activity[user_id] = row.get("created_at") or row.get("timestamp")
    for user in users:
        user_id = str(user.get("_id"))
        user["search_count"] = search_counts[user_id]
        user["evidence_count"] = evidence_counts[user_id]
        user["last_activity"] = last_activity.get(user_id) or user.get("created_at")
        user["account_status"] = user.get("account_status") or user.get("status") or "active"
        user["roles"] = user_roles(user)
        user["active_role"] = user_active_role(user)
        user["plan"] = user.get("plan") or "basic"
    sources = health
    return render_template(
        "dashboard_administrator.html",
        users=users,
        test_account_count=test_account_count,
        sources=sources,
        snapshot=safe_call(repository.admin_snapshot, {}),
        evidence=safe_call(lambda: repository.list_evidence(limit=20), []),
        audit_logs=present_audit_logs(safe_call(lambda: repository.list_audit_logs(limit=30), [])),
    )


@app.post("/dashboard/administrator/users/delete")
@role_required("administrator")
def administrator_delete_selected_users():
    selected_ids = [value for value in request.form.getlist("user_id") if value]
    users = {str(user.get("_id")): user for user in safe_call(lambda: repository.list_users(limit=0), [])}
    changed = 0
    skipped = 0
    for user_id in selected_ids:
        target = users.get(str(user_id))
        if not target or str(target.get("_id")) == str(session.get("user_id")):
            skipped += 1
            continue
        if (target.get("role") or "") == "administrator" or "administrator" in user_roles(target):
            skipped += 1
            continue
        if repository.delete_user(user_id):
            changed += 1
    if changed:
        repository.log_event(
            session["user_id"],
            "admin_selected_users_deleted",
            session.get("role"),
            session.get("username"),
            {"changed_count": changed, "skipped_count": skipped, "selected_ids": selected_ids, "action": "delete_selected_users"},
        )
        flash(f"{changed} selected account(s) were permanently deleted.", "success")
    else:
        flash("No selected accounts could be deleted. Administrator accounts and your own account cannot be deleted.", "info")
    return redirect(url_for("administrator_dashboard"))


@app.post("/dashboard/administrator/users/<user_id>")
@role_required("administrator")
def administrator_update_user(user_id):
    target = repository.get_user_by_id(user_id)
    if not target:
        flash("That user could not be found.", "error")
        return redirect(url_for("administrator_dashboard") + "#user-management")
    if str(target.get("_id")) == str(session.get("user_id")):
        flash("Your own administrator account cannot be changed from this screen.", "error")
        return redirect(url_for("administrator_dashboard") + "#user-management")
    target_roles = user_roles(target)
    target_is_administrator = "administrator" in target_roles or target.get("active_role") == "administrator" or target.get("role") == "administrator"
    legacy_roles = normalize_roles(request.form.getlist("roles")) if request.form.getlist("roles") else []
    requested_role = (
        request.form.get("workspace_role")
        or request.form.get("active_role")
        or (legacy_roles[0] if len(legacy_roles) == 1 else None)
        or single_account_role(target)
    )
    if target_is_administrator:
        requested_role = "administrator"
    submitted_tier = request.form.get("membership_tier") or request.form.get("plan") or target.get("membership_tier") or target.get("plan") or "basic"
    plan = normalize_membership_tier(submitted_tier)
    membership_tier = plan
    account_status = request.form.get("account_status", target.get("account_status", "active"))
    if requested_role not in ROLES or (requested_role == "administrator" and not target_is_administrator):
        flash("Please choose a valid single workspace role.", "error")
        return redirect(url_for("administrator_dashboard") + "#user-management")
    if account_status not in {"active", "inactive"}:
        flash("Please choose a valid account status.", "error")
        return redirect(url_for("administrator_dashboard") + "#user-management")
    updates = {}
    if target_roles != [requested_role] or single_account_role(target) != requested_role:
        updates.update(
            roles=[requested_role],
            primary_role=requested_role,
            active_role=requested_role,
            role=requested_role,
        )
    if not target_is_administrator:
        if plan != normalize_membership_tier(target.get("plan")):
            updates["plan"] = plan
        if membership_tier != normalize_membership_tier(target.get("membership_tier") or target.get("plan")):
            updates["membership_tier"] = membership_tier
    if account_status != target.get("account_status"):
        updates["account_status"] = account_status
    if updates:
        repository.update_user(user_id, updates)
        repository.log_event(
            session["user_id"],
            "admin_user_update",
            session.get("role"),
            session.get("username"),
            {"target_user": target.get("display_name") or target.get("username") or target.get("email"), "updates": updates},
        )
        flash("User updated.", "success")
    else:
        flash("No changes were made.", "info")
    return redirect(url_for("administrator_dashboard") + "#user-management")


@app.route("/saved")
@login_required
def saved():
    locked = not can_access(current_user(), "saved_research")
    if locked:
        return locked_feature_response("saved_research", title="Saved Research")
    scope = None if session.get("role") in {"researcher", "administrator"} else session["user_id"]
    return render_template("saved_packages.html", evidence=normalize_price_items(repository.list_evidence(scope, limit=100)), research=repository.list_research(scope, limit=50), comparison_sets=repository.list_comparison_groups(scope, saved_only=True, limit=50), synthetic_notice=True, search_record_id=session.get("last_search_record_id"))


@app.post("/saved/comparison/<comparison_id>")
@login_required
def save_comparison_set(comparison_id):
    blocked = ensure_feature_access("saved_research", redirect_endpoint="search")
    if blocked:
        return blocked
    group = repository.get_comparison_group(comparison_id, session["user_id"])
    if not group:
        abort(404)
    items = annotate_comparison(_comparison_group_records(group), group.get("query", ""))
    metrics = comparison_metrics(items)
    repository.update_comparison_group(comparison_id, {
        "document_type": "comparison_set", "saved_status": "saved", "status": "saved", "saved_at": utcnow(),
        "frozen_summary": {key: value for key, value in metrics.items() if key not in {"excluded", "comparable"}},
        "monitoring_enabled": False,
    })
    repository.log_event(session["user_id"], "comparison_set_saved", session.get("role"), session.get("username"), {"comparison_id": comparison_id, "record_count": len(items)})
    flash("Comparison set saved to Saved Research. Monitoring was not enabled.", "success")
    return redirect(url_for("saved"))


@app.post("/saved/comparison/<comparison_id>/delete")
@login_required
def delete_saved_comparison_set(comparison_id):
    group = repository.get_comparison_group(comparison_id, session["user_id"])
    if not group:
        abort(404)
    repository.update_comparison_group(comparison_id, {"is_deleted": True, "deleted_at": utcnow(), "saved_status": "deleted"})
    flash("Comparison set deleted.", "success")
    return redirect(url_for("saved"))


@app.post("/saved")
@login_required
def save_evidence():
    blocked = ensure_feature_access("save_evidence", redirect_endpoint="search")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    try:
        index = int(request.form.get("product_index", "-1"))
    except ValueError:
        index = -1
    record = repository.get_search(search_id, session["user_id"])
    products = normalize_price_items((repository.get_market_records(record.get("market_record_ids", [])) if record and record.get("market_record_ids") else repository.get_products(search_id)) if record else [])
    if index < 0 or index >= len(products):
        flash("That result could not be found.", "error")
    else:
        product = products[index]
        product.update(user_role=session["role"], display_name=session["username"], query=record["keyword"], product_name=product.get("title"), timestamp=product.get("collected_at"), confidence=product.get("confidence_level"), demo_mode=product.get("data_mode") == "synthetic")
        _evidence_id, created = repository.save_evidence(session["user_id"], search_id, product)
        if created:
            repository.log_event(session["user_id"], "evidence_save", session["role"], session["username"], {"query": record["keyword"], "title": product.get("title"), "platform": product.get("platform")})
        flash("Evidence saved to your research." if created else "This evidence is already saved.", "info")
    return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("saved")))


@app.post("/saved/bulk")
@login_required
def save_selected_evidence():
    blocked = ensure_feature_access("save_evidence", redirect_endpoint="search")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"])
    products = normalize_price_items((repository.get_market_records(record.get("market_record_ids", [])) if record and record.get("market_record_ids") else repository.get_products(search_id)) if record else [])
    try:
        indices = {int(value) for value in request.form.getlist("product")}
    except ValueError:
        indices = set()
    selected = []
    for index in sorted(indices):
        if 0 <= index < len(products):
            product = products[index]
            product.update(user_role=session["role"], display_name=session["username"], query=record["keyword"], product_name=product.get("title"), timestamp=product.get("collected_at"), confidence=product.get("confidence_level"), demo_mode=product.get("data_mode") == "synthetic")
            selected.append(product)
    if not selected:
        flash("Select at least one result to save.", "error")
        return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("saved")))
    saved_count, excluded = _evidence_save_plan(selected, session["user_id"], search_id)
    for product in selected:
        repository.log_event(session["user_id"], "evidence_save", session["role"], session["username"], {"query": record["keyword"], "title": product.get("title"), "platform": product.get("platform")})
    excluded_count = len(excluded)
    if excluded_count:
        flash(f"{len(selected)} selected, {saved_count} saved, {excluded_count} excluded because they were accessory/invalid/duplicate records.", "info")
    else:
        flash(f"Saved {saved_count} selected evidence record(s).", "info")
    destination = safe_internal_redirect_target(request.form.get("next"), url_for("saved"))
    return redirect(destination + "#saved-evidence-records")


@app.post("/saved/evidence/<evidence_id>/delete")
@login_required
def delete_saved_evidence(evidence_id):
    evidence = repository.get_evidence(evidence_id)
    if not evidence or not user_can_manage_record(evidence.get("user_id")):
        flash("You do not have permission to delete this record.", "error")
        return redirect(url_for("saved"))
    repository.delete_evidence(evidence_id, deleted_by=session.get("user_id"), reason="saved_evidence_deleted")
    repository.log_event(session["user_id"], "evidence_deleted", session.get("role"), session.get("username"), {"record_id": evidence_id, "target_user": evidence.get("display_name") or evidence.get("user_id")})
    flash("Saved evidence deleted.", "success")
    return redirect(url_for("saved") + "#saved-evidence-records")


@app.post("/saved/evidence/delete")
@login_required
def delete_saved_evidence_selected():
    try:
        evidence_ids = [value for value in request.form.getlist("evidence_id") if value]
    except ValueError:
        evidence_ids = []
    if not evidence_ids:
        flash("Select at least one saved evidence record.", "error")
        return redirect(url_for("saved"))
    deleted = 0
    for evidence_id in evidence_ids:
        evidence = repository.get_evidence(evidence_id)
        if not evidence or not user_can_manage_record(evidence.get("user_id")):
            continue
        repository.delete_evidence(evidence_id, deleted_by=session.get("user_id"), reason="bulk_saved_evidence_deleted")
        deleted += 1
    if deleted:
        repository.log_event(session["user_id"], "evidence_deleted", session.get("role"), session.get("username"), {"deleted_count": deleted, "delete_mode": "bulk"})
        flash(f"{deleted} records deleted.", "success")
    else:
        flash("You do not have permission to delete the selected records.", "error")
    return redirect(url_for("saved") + "#saved-evidence-records")


@app.post("/saved/market")
@login_required
def save_market_evidence():
    blocked = ensure_feature_access("save_evidence", redirect_endpoint="search")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"])
    products = repository.get_market_records(request.form.getlist("record_id"))
    products = annotate_comparison(normalize_price_items(products), record.get("keyword", "") if record else "")
    if not record or not products:
        flash("Select at least one MongoDB market record to save.", "error")
        return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("search")))
    created_count = 0
    for product in products:
        product.update(market_record_id=product.get("_id"), user_role=session["role"], display_name=session["username"], query=record["keyword"], product_name=product.get("title"), timestamp=product.get("collected_at"), confidence=product.get("confidence"), evidence_status="saved", demo_mode=product.get("demo_mode", False))
        _evidence_id, created = repository.save_evidence(session["user_id"], search_id, product)
        created_count += int(created)
    repository.log_event(session["user_id"], "save_evidence", session["role"], session["username"], {"query": record["keyword"], "record_count": len(products), "saved_count": created_count, "collection": "evidence_records"})
    flash(f"Saved {created_count} selected MongoDB record(s) to evidence_records.", "info")
    return redirect(url_for("saved"))


@app.post("/saved/results")
@login_required
def save_result_evidence():
    blocked = ensure_feature_access("save_evidence", redirect_endpoint="search")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    comparison_id = request.form.get("comparison_id", "")
    group = repository.get_comparison_group(comparison_id, session["user_id"]) if comparison_id else None
    record = repository.get_search(search_id, session["user_id"]) if search_id else None
    if group:
        search_id = group.get("search_record_id") or search_id
        record = repository.get_search(search_id, session["user_id"]) if search_id else record
        products = annotate_comparison(_comparison_group_records(group), group.get("query") or (record.get("keyword", "") if record else ""))
        tokens = [str(token) for token in group.get("selected_record_ids", [])]
    else:
        requested_tokens = _submitted_selection_tokens(request.form)
        tokens = _tokens_allowed_for_search(record, requested_tokens)
        products = annotate_comparison(normalize_price_items(resolve_result_tokens(tokens)), record.get("keyword", "") if record else "")
    _comparison_debug("save_evidence", comparison_id=comparison_id, search_record_id=search_id, selected_record_ids=tokens, record_count=len(products))
    if not record or not products:
        flash("Select at least one collected result to save.", "error")
        return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("search")))
    for product in products:
        product.update(result_token=product.get("_selection_token"), market_record_id=product.get("_id") if product.get("_selection_token", "").startswith("market:") else None, user_role=session["role"], display_name=session["username"], query=record["keyword"], product_name=product.get("title"), timestamp=product.get("collected_at"), evidence_status="saved")
    saved_count, excluded = _evidence_save_plan(products, session["user_id"], search_id)
    repository.log_event(session["user_id"], "save_evidence", session["role"], session["username"], {"query": record["keyword"], "record_count": len(products), "saved_count": saved_count, "excluded_count": len(excluded), "source_types": sorted({item.get("source_type") for item in products})})
    if excluded:
        flash(f"{len(products)} selected, {saved_count} saved, {len(excluded)} excluded because they were accessory/invalid/duplicate records.", "info")
    else:
        flash(f"Saved {saved_count} selected result(s) to evidence_records.", "info")
    return redirect(url_for("saved") + "#saved-evidence-records")


@app.post("/saved/research")
@login_required
def save_research():
    blocked = ensure_feature_access("saved_research", redirect_endpoint="search")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"])
    products = (repository.get_market_records(record.get("market_record_ids", [])) if record and record.get("market_record_ids") else repository.get_products(search_id)) if record else []
    try:
        indices = {int(value) for value in request.form.getlist("product")}
    except ValueError:
        indices = set()
    if not record or not products:
        flash("Run a product search before saving research.", "error")
        return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("search")))
    normalized = annotate_comparison(normalize_price_items(products), record["keyword"])
    selected_indices = sorted(index for index in indices if 0 <= index < len(normalized))
    selected_records = [normalized[index] for index in selected_indices] if selected_indices else normalized
    summary = calculate_summary(selected_records)
    insight, insight_mode = build_price_insight(selected_records, summary), "Rule-based insight"
    research_id = repository.save_research(session["user_id"], {
        "search_record_id": repository._id(search_id), "display_name": session["username"], "role": session["role"],
        "query": record["keyword"], "selected_source": record["selected_source"],
        "selected_sources": sorted({item.get("source_type") or item.get("platform") for item in selected_records}),
        "selected_indices": selected_indices, "selected_records": selected_records, "normalized_price_records": selected_records,
        "analysis_summary": summary, "ai_price_insight": insight, "insight_mode": insight_mode,
        "demo_mode": record.get("data_mode") == "synthetic", "source_evidence_status": "synthetic" if record.get("data_mode") == "synthetic" else "API-sourced",
    })
    for product in selected_records:
        product.update(user_role=session["role"], display_name=session["username"], query=record["keyword"], product_name=product.get("title"), timestamp=product.get("collected_at"), confidence=product.get("confidence_level"), demo_mode=record.get("data_mode") == "synthetic")
    saved_count, excluded = _evidence_save_plan(selected_records, session["user_id"], search_id)
    repository.log_event(session["user_id"], "research_saved", session["role"], session["username"], {"query": record["keyword"], "research_id": research_id, "record_count": len(selected_records), "saved_count": saved_count, "excluded_count": len(excluded), "storage_mode": repository.mode})
    if repository.mode != "mongodb":
        repository.log_event(session["user_id"], "mongodb_save_failed", session["role"], session["username"], {"query": record["keyword"], "fallback": "memory", "reason": repository.status()["detail"]})
    if excluded:
        flash(f"{len(selected_records)} selected, {saved_count} saved, {len(excluded)} excluded because they were accessory/invalid/duplicate records.", "info")
    else:
        flash("Research package saved with normalized listings, analysis and insight.", "info")
    return redirect(url_for("saved") + "#saved-evidence-records")


@app.post("/saved/research/<research_id>/delete")
@login_required
def delete_saved_research(research_id):
    research = repository.get_research(research_id)
    if not research or not user_can_manage_record(research.get("user_id")):
        flash("You do not have permission to delete this record.", "error")
        return redirect(url_for("saved"))
    repository.delete_research(research_id, deleted_by=session.get("user_id"), reason="saved_research_deleted")
    repository.log_event(session["user_id"], "research_deleted", session.get("role"), session.get("username"), {"research_id": research_id, "target_user": research.get("display_name") or research.get("user_id")})
    flash("Saved research archived.", "success")
    return redirect(url_for("saved"))


@app.post("/saved/research/delete")
@login_required
def delete_saved_research_selected():
    research_ids = [value for value in request.form.getlist("research_id") if value]
    if not research_ids:
        flash("Select at least one saved research package.", "error")
        return redirect(url_for("saved"))
    deleted = 0
    for research_id in research_ids:
        research = repository.get_research(research_id)
        if not research or not user_can_manage_record(research.get("user_id")):
            continue
        repository.delete_research(research_id, deleted_by=session.get("user_id"), reason="bulk_saved_research_deleted")
        deleted += 1
    if deleted:
        repository.log_event(session["user_id"], "research_deleted", session.get("role"), session.get("username"), {"deleted_count": deleted, "delete_mode": "bulk"})
        flash(f"Archived {deleted} saved research package(s).", "success")
    else:
        flash("You do not have permission to delete the selected packages.", "error")
    return redirect(url_for("saved"))


@app.route("/saved/research/<research_id>")
@login_required
def research_detail(research_id):
    scope = None if session.get("role") in {"researcher", "administrator"} else session["user_id"]
    research = repository.get_research(research_id, scope)
    if not research:
        abort(404)
    return render_template("research_detail.html", research=research)


@app.route("/saved/export/evidence.csv")
@login_required
def saved_export_evidence_csv():
    if not can_access(current_user(), "save_evidence"):
        return export_forbidden_response("Export is available on Premium and Professional plans.")
    evidence = safe_call(lambda: repository.list_evidence(session["user_id"], limit=200), [])
    selected_ids = {str(value) for value in request.args.getlist("evidence_id") if value}
    if selected_ids:
        evidence = [item for item in evidence if str(item.get("_id")) in selected_ids]
    rows = [["Query", "Platform", "Product", "Price", "Condition", "Source type", "Collected", "Source link"]]
    for item in evidence:
        source_type = _normalize_export_source_type(item)
        rows.append([
            item.get("query", ""),
            item.get("platform", ""),
            item.get("product_name") or item.get("title") or "",
            item.get("normalized_price") if item.get("normalized_price") is not None else item.get("price", ""),
            item.get("condition", ""),
            source_type,
            singapore_time_string(item.get("timestamp") or item.get("collected_at")),
            item.get("source_url") or "",
        ])
    return csv_response(rows, "saved-evidence.csv")


@app.route("/saved/report")
@login_required
def saved_export_report():
    if not can_access(current_user(), "standard_export"):
        return export_forbidden_response("Saved evidence report is available on Premium and Professional plans.")
    evidence = safe_call(lambda: repository.list_evidence(session["user_id"], limit=200), [])
    report_rows = []
    for item in evidence:
        normalized_price = item.get("normalized_price")
        observed_price = normalized_price if normalized_price is not None else item.get("price")
        try:
            price_text = f"{'USD' if normalized_price is not None else item.get('currency') or 'USD'} {float(observed_price):.2f}"
        except (TypeError, ValueError):
            price_text = "Not available"
        report_rows.append({
            "query": item.get("query") or "Not available",
            "platform": item.get("platform") or "Not available",
            "title": item.get("product_name") or item.get("title") or "Untitled listing",
            "price": price_text,
            "condition": item.get("condition_display") or item.get("condition") or "Not specified by source",
            "source": _record_source_label(item),
            "collected_at": display_sgt_datetime(item.get("timestamp") or item.get("collected_at") or item.get("created_at")),
            "source_url": item.get("source_url") if is_safe_url(item.get("source_url")) else "",
        })
    summary = {
        "record_count": len(report_rows),
        "platform_count": len({row["platform"] for row in report_rows if row["platform"] != "Not available"}),
        "query_count": len({row["query"] for row in report_rows if row["query"] != "Not available"}),
    }
    return render_template(
        "saved_evidence_report.html",
        rows=report_rows,
        summary=summary,
        generated_at=utcnow(),
    )


@app.post("/search-history/<search_id>/delete")
@login_required
def delete_search_history(search_id):
    destination = safe_internal_redirect_target(request.form.get("next"), url_for("dashboard_redirect"))
    record = repository.get_search(search_id)
    if not record or not user_can_manage_record(record.get("user_id")):
        flash("You do not have permission to delete this record.", "error")
        return redirect(destination)
    repository.delete_search(search_id, deleted_by=session.get("user_id"), reason="search_history_deleted")
    repository.log_event(session["user_id"], "search_deleted", session.get("role"), session.get("username"), {"search_id": search_id, "query": record.get("keyword")})
    flash("Search history deleted.", "success")
    return redirect(destination)


@app.post("/search-history/clear")
@login_required
def clear_search_history():
    destination = safe_internal_redirect_target(request.form.get("next"), url_for("dashboard_redirect"))
    search_ids = [row["_id"] for row in repository.list_searches(session["user_id"], limit=0)]
    if not search_ids:
        flash("No search history found to clear.", "info")
        return redirect(destination)
    deleted = repository.delete_many_searches(search_ids, deleted_by=session.get("user_id"), reason="search_history_cleared")
    repository.log_event(session["user_id"], "search_deleted", session.get("role"), session.get("username"), {"deleted_count": deleted, "delete_mode": "clear_history"})
    flash("Search history cleared.", "success")
    return redirect(destination)


@app.post("/compare")
@login_required
def compare():
    params = request.form if request.form else request.args
    search_id = params.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"])
    requested_tokens = _submitted_selection_tokens(params)
    tokens = _tokens_allowed_for_search(record, requested_tokens)
    _comparison_debug("compare_submit", search_record_id=search_id, submitted=requested_tokens, accepted=tokens)
    if not record:
        flash("Run a search before comparing selected records.", "error")
        return redirect(url_for("search"))
    if record.get("comparison_enabled") is False:
        flash("Choose one product category or product type before comparing listings.", "error")
        return redirect(safe_internal_redirect_target(params.get("next"), url_for("search_results", search_run_id=search_id)))
    if len(tokens) < 2:
        flash("Select at least two results to compare.", "error")
        return redirect(safe_internal_redirect_target(params.get("next"), url_for("search", search_record_id=search_id)))
    products = annotate_comparison(normalize_price_items(resolve_result_tokens(tokens)), record.get("keyword", ""))
    products = [enrich_listing_category(item, record.get("keyword", "")) for item in products]
    products = [item for item in products if item.get("product_role") == "complete_product"]
    if record.get("selected_category_key"):
        products = [item for item in products if item.get("category_key") == record["selected_category_key"]]
    if record.get("selected_product_type"):
        products = [item for item in products if (item.get("attributes") or {}).get("product_type") == record["selected_product_type"]]
    tokens = [item.get("_selection_token") for item in products if item.get("_selection_token")]
    if len(products) < 2:
        flash("Selected records could not be resolved. Please select records from the current search again.", "error")
        return redirect(safe_internal_redirect_target(params.get("next"), url_for("search", search_record_id=search_id)))
    comparison_id = repository.create_comparison_group(session["user_id"], {
        "name": f"{record.get('keyword', 'Selected')} comparison",
        "query": record.get("keyword", ""),
        "workspace": "comparison",
        "search_record_id": record.get("_id"),
        "selected_record_ids": tokens,
        "selected_records": products,
        "source_types": [item.get("source_type") or item.get("platform") for item in products],
        "document_type": "comparison_set",
        "saved_status": "draft",
        "status": "draft",
    })
    session["selected_result_ids"] = tokens
    session["temporary_comparison_state"] = comparison_id
    repository.log_event(session["user_id"], "comparison_generated", session["role"], session["username"], {"query": record["keyword"], "comparison_id": comparison_id, "record_count": len(products), "selected_record_ids": tokens})
    return compare_group(comparison_id)


@app.route("/compare/<comparison_id>")
@login_required
def compare_group(comparison_id):
    group = repository.get_comparison_group(comparison_id, session["user_id"])
    if not group:
        flash("Selected comparison records were not found. Please select records again.", "error")
        return redirect(url_for("search"))
    search_id = group.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"]) if search_id else None
    products = annotate_comparison(_comparison_group_records(group), group.get("query") or (record.get("keyword", "") if record else ""))
    selected_record_ids = [str(token) for token in group.get("selected_record_ids", [])]
    _comparison_debug("compare_display", comparison_id=comparison_id, selected_record_ids=selected_record_ids, displayed_count=len(products))
    products = annotate_comparison(normalize_price_items(products), record.get("keyword", "") if record else "")
    if len(products) < 2:
        flash("Select at least two results to compare.", "error")
    stats = comparison_metrics(products)
    conclusion = comparison_conclusion(products, stats)
    platforms = sorted({item.get("platform") for item in products if item.get("platform")})
    source_summary = f"{len(products)} selected records" + (f" · {' and '.join(platforms)}" if platforms else "")
    storage_variants = sorted({item.get("storage") or (item.get("parsed_model") or {}).get("storage") for item in products if item.get("storage") or (item.get("parsed_model") or {}).get("storage")})
    mixed_storage = len(storage_variants) > 1
    preview_payload = _analysis_record_payload("selected_comparison", record or {"_id": search_id, "keyword": group.get("query", "")}, products, selected_record_ids, comparison_set_id=comparison_id)
    current_analysis = repository.find_analysis_by_signature(session["user_id"], preview_payload["analysis_signature"])
    if current_analysis:
        analysis_action, analysis_label = "open", "Open analysis"
    else:
        analysis_action, analysis_label = "create", "Analyze comparison"
    return render_template("compare_complete.html", products=products, record=record, stats=stats, conclusion=conclusion, source_summary=source_summary, comparison_id=comparison_id, selected_record_ids=selected_record_ids, comparison_created_at=group.get("created_at"), storage_variants=storage_variants, mixed_storage=mixed_storage, current_analysis=current_analysis, analysis_action=analysis_action, analysis_label=analysis_label, comparison_saved=group.get("saved_status") == "saved")


def _watchlist_item_payload_from_search(record, items, tracking_mode="search_scope"):
    scope = _watchlist_scope_from_products(items, record, tracking_mode=tracking_mode)
    if not scope:
        return None
    scope["tracking_mode"] = tracking_mode
    scope["record_scope"] = "selected_comparison_records" if tracking_mode == "selected_records" else "all_current_search_results"
    frozen_records = [dict(item) for item in normalize_price_items(items) if item.get("analytics_eligible")]
    return {
        "tracking_mode": scope["tracking_mode"],
        "keyword": scope["keyword"],
        "product_label": scope["product_label"],
        "platform_scope": scope["platform_scope"],
        "category": scope["category"],
        "category_key": scope["category_key"],
        "category_display": scope["category_display"],
        "condition_scope": scope["condition_scope"],
        "source_scope": "search",
        "data_source_label": scope["data_source_label"],
        "source_label": scope["data_source_label"],
        "record_scope": scope["record_scope"],
        "selected_record_ids": scope["selected_record_ids"],
        "selected_records_snapshot": scope["selected_records_snapshot"],
        "search_record_id": record.get("_id"),
        "comparison_record_ids": [item.get("_id") for item in items if item.get("_id")],
        "initial_records_snapshot": frozen_records,
        "frozen_scope": {
            "keyword": scope["keyword"],
            "category": scope["category"],
            "category_key": scope["category_key"],
            "category_display": scope["category_display"],
            "platform_scope": scope["platform_scope"],
            "condition_scope": scope["condition_scope"],
            "tracking_mode": scope["tracking_mode"],
            "record_scope": scope["record_scope"],
            "selected_record_ids": list(scope["selected_record_ids"]),
        },
    }


@app.post("/watchlist/from-search")
@login_required
def watchlist_from_search():
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"]) if search_id else None
    if not record:
        flash("No tracked products yet. Add a product or comparison from Product Search to begin price monitoring.", "info")
        return redirect(url_for("watchlist"))
    tracking_mode = request.form.get("tracking_mode") or "search_scope"
    if tracking_mode == "search_scope" and record.get("comparison_enabled") is False:
        flash("Choose one product category or product type before tracking the full result set.", "error")
        return redirect(url_for("search_results", search_run_id=search_id))
    product_indices = request.form.getlist("product")
    tokens = _submitted_selection_tokens(request.form)
    if tracking_mode == "search_scope":
        if record.get("result_tokens"):
            products = resolve_result_tokens(record["result_tokens"])
        elif record.get("market_record_ids"):
            products = repository.get_market_records(record["market_record_ids"])
        else:
            products = repository.get_products(search_id)
    elif product_indices and record.get("result_tokens"):
        allowed = list(record.get("result_tokens", []))
        products = resolve_result_tokens([allowed[int(index)] for index in product_indices if index.isdigit() and int(index) < len(allowed)])
    elif product_indices and record.get("market_record_ids"):
        allowed = list(record.get("market_record_ids", []))
        products = repository.get_market_records([allowed[int(index)] for index in product_indices if index.isdigit() and int(index) < len(allowed)])
    elif tokens and record.get("result_tokens"):
        allowed = set(record.get("result_tokens", []))
        products = resolve_result_tokens([token for token in tokens if token in allowed])
    elif tokens:
        products = repository.get_market_records(tokens)
    else:
        flash("Select at least one collected result to track.", "error")
        return redirect(safe_internal_redirect_target(request.form.get("next"), url_for("search", search_record_id=search_id)))
    products = [enrich_listing_category(item, record.get("keyword", "")) for item in normalize_price_items(products)]
    products = [item for item in products if item.get("product_role") == "complete_product"]
    if record.get("selected_category_key"):
        products = [item for item in products if item.get("category_key") == record["selected_category_key"]]
    if record.get("selected_product_type"):
        products = [item for item in products if (item.get("attributes") or {}).get("product_type") == record["selected_product_type"]]
    _comparison_debug("watchlist_from_search", search_record_id=search_id, tracking_mode=tracking_mode, selected_record_ids=tokens, record_count=len(products))
    payload = _watchlist_item_payload_from_search(record, products, tracking_mode=tracking_mode)
    if not payload:
        flash("No valid comparable records available for Watchlist tracking.", "error")
        return redirect(url_for("search", search_record_id=search_id))
    watchlist_id, created = repository.create_watchlist_item(session["user_id"], payload)
    watchlist_item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    snapshot_result = _collect_watchlist_snapshot(watchlist_item) if watchlist_item else None
    if snapshot_result:
        flash("Watchlist item created with initial snapshot.", "success")
    elif created:
        flash("No valid records available for Watchlist snapshot.", "error")
    else:
        flash("Watchlist item updated.", "success")
    repository.log_event(session["user_id"], "watchlist_created", session["role"], session["username"], {"watchlist_id": watchlist_id, "search_record_id": search_id})
    return redirect(url_for("watchlist"))


@app.post("/watchlist/from-compare")
@login_required
def watchlist_from_compare():
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    search_id = request.form.get("search_record_id", "")
    comparison_id = request.form.get("comparison_id", "")
    group = repository.get_comparison_group(comparison_id, session["user_id"]) if comparison_id else None
    record = repository.get_search(search_id, session["user_id"]) if search_id else None
    if group:
        search_id = group.get("search_record_id") or search_id
        record = repository.get_search(search_id, session["user_id"]) if search_id else record
        tokens = [str(token) for token in group.get("selected_record_ids", [])]
        products = _comparison_group_records(group)
    else:
        tokens = _submitted_selection_tokens(request.form)
        products = []
    if not group and record and tokens and record.get("result_tokens"):
        allowed = set(record.get("result_tokens", []))
        products = resolve_result_tokens([token for token in tokens if token in allowed])
    _comparison_debug("watchlist_from_compare", comparison_id=comparison_id, search_record_id=search_id, selected_record_ids=tokens, record_count=len(products))
    payload = _watchlist_item_payload_from_search(record or {}, products, tracking_mode="selected_records")
    if not payload:
        flash("No valid comparable records available for Watchlist tracking.", "error")
        return redirect(url_for("compare_group", comparison_id=comparison_id) if comparison_id else url_for("search", search_record_id=search_id))
    payload["comparison_group_id"] = comparison_id
    payload["comparison_set_id"] = comparison_id
    watchlist_id, created = repository.create_watchlist_item(session["user_id"], payload)
    watchlist_item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    snapshot_result = _collect_watchlist_snapshot(watchlist_item) if watchlist_item else None
    if snapshot_result:
        flash(f"{len(products)} selected records added to watchlist.", "success")
    elif created:
        flash("No valid records available for Watchlist snapshot.", "error")
    else:
        flash("Watchlist item updated.", "success")
    repository.log_event(session["user_id"], "watchlist_created", session["role"], session["username"], {"watchlist_id": watchlist_id, "search_record_id": search_id, "source": "compare"})
    return _watchlist_redirect(watchlist_id)


def _collect_watchlist_snapshot(watchlist_item):
    record = repository.get_search(watchlist_item.get("search_record_id"), watchlist_item.get("user_id")) if watchlist_item.get("search_record_id") else None
    if watchlist_item.get("tracking_mode") == "selected_records":
        items = watchlist_item.get("selected_records_snapshot") or []
        if not items and watchlist_item.get("selected_record_ids"):
            items = repository.get_market_records(watchlist_item.get("selected_record_ids", []))
    elif record and record.get("result_tokens"):
        items = resolve_result_tokens(record["result_tokens"])
    elif record and record.get("market_record_ids"):
        items = repository.get_market_records(record["market_record_ids"])
    elif record:
        items = repository.get_products(record["_id"])
    else:
        items = []
    if not items and watchlist_item.get("tracking_mode") != "selected_records":
        items = demo_search_items(watchlist_item.get("keyword", ""), limit=20, platform="demo_all")
    payload = _snapshot_payload_from_items(watchlist_item, items)
    if not payload or not payload.get("average_price"):
        return None
    latest = repository.latest_price_snapshot(watchlist_item["_id"])
    if latest and latest.get("created_at"):
        created_date = latest["created_at"].date() if hasattr(latest["created_at"], "date") else None
        if created_date == utcnow().date():
            payload["notes"] = "Updated same-day snapshot"
    snapshot_id = repository.save_price_snapshot(watchlist_item["_id"], watchlist_item["user_id"], payload)
    return snapshot_id, payload


def monitor_refresh_timezone():
    name = app.config.get("MONITOR_REFRESH_TIMEZONE") or "Asia/Singapore"
    try:
        return ZoneInfo(name)
    except Exception:
        return ZoneInfo("Asia/Singapore")


def next_monitor_refresh_at(now=None):
    now_utc = _normalized_utc_datetime(now) or utcnow()
    local_now = now_utc.astimezone(monitor_refresh_timezone())
    hour = int(app.config.get("MONITOR_DAILY_REFRESH_HOUR", 8))
    candidate = local_now.replace(hour=hour, minute=0, second=0, microsecond=0)
    if candidate <= local_now:
        candidate += timedelta(days=1)
    return candidate.astimezone(timezone.utc)


def monitor_scheduled_date(now=None):
    now_utc = _normalized_utc_datetime(now) or utcnow()
    return now_utc.astimezone(monitor_refresh_timezone()).date().isoformat()


def _monitor_platforms(monitor):
    scope = str((monitor.get("frozen_scope") or {}).get("platform_scope") or monitor.get("platform_scope") or "all").lower()
    platforms = set()
    if "ebay" in scope or scope == "all":
        platforms.add("ebay")
    if "walmart" in scope or scope == "all":
        platforms.add("walmart")
    return platforms or {"ebay", "walmart"}


def _record_identity_values(record):
    values = set()
    platform = str(record.get("platform") or "").strip().lower()
    for key in ("item_id", "product_id", "source_id", "listing_id", "_selection_token", "_id", "url", "product_url"):
        value = str(record.get(key) or "").strip().lower()
        if value:
            values.add(f"{platform}:{value}")
    title = " ".join(str(record.get("title") or record.get("product_name") or "").lower().split())
    if title:
        values.add(f"{platform}:title:{title}")
    return values


def _filter_monitor_identity_scope(monitor, records):
    if monitor.get("tracking_mode") != "selected_records":
        return records
    frozen = monitor.get("selected_records_snapshot") or monitor.get("initial_records_snapshot") or []
    identities = set().union(*(_record_identity_values(row) for row in frozen)) if frozen else set()
    if not identities:
        return []
    return [row for row in records if identities.intersection(_record_identity_values(row))]


def _retrieve_monitor_provider_records(monitor):
    """Retrieve the Monitor's frozen scope through the canonical providers."""
    monitor = _monitor_with_category_metadata(monitor)
    frozen_scope = monitor.get("frozen_scope") or {}
    query = frozen_scope.get("keyword") or monitor.get("keyword") or monitor.get("product_label")
    category = monitor.get("category_query") or ""
    requested = _monitor_platforms(monitor)
    records = []
    source_statuses = {}
    if "ebay" in requested:
        try:
            rows, meta = search_ebay_cached(query, category=category, limit=20, refresh_live=True)
            records.extend(rows)
            source_statuses["ebay"] = {"status": "success", "record_count": len(rows), "source_type": meta.get("source_type")}
        except Exception as exc:
            source_statuses["ebay"] = {"status": "failed", "error_code": f"ebay_{exc.__class__.__name__.lower()}"}
    if "walmart" in requested:
        try:
            rows, meta = search_serpapi_cached(query, "walmart", category=category, limit=20, refresh_live=True)
            records.extend(rows)
            source_statuses["walmart"] = {"status": "success", "record_count": len(rows), "source_type": meta.get("source_type")}
        except Exception as exc:
            source_statuses["walmart"] = {"status": "failed", "error_code": f"walmart_{exc.__class__.__name__.lower()}"}
    return _filter_monitor_identity_scope(monitor, records), source_statuses


def _offline_monitor_records(monitor):
    records = monitor.get("selected_records_snapshot") or monitor.get("initial_records_snapshot") or []
    platforms = _monitor_platforms(monitor)
    statuses = {platform: {"status": "success", "record_count": sum(str(row.get("platform") or "").lower() == platform for row in records), "source_type": "offline_frozen_scope"} for platform in platforms}
    return [dict(row) for row in records], statuses


def _evaluate_pending_monitor_forecast(monitor, snapshot):
    predictions = repository.list_predictions(monitor["user_id"], monitor["_id"], limit=20)
    pending = next((row for row in predictions if row.get("status") in {"pending_actual", "ai_unavailable", "waiting_for_validation"}), None)
    if not pending or not is_snapshot_after_prediction(snapshot, pending):
        return None
    if str(snapshot.get("_id")) == str(pending.get("baseline_snapshot_id")):
        return None
    evaluated = _evaluate_prediction_record(pending, snapshot)
    repository.update_prediction(pending["_id"], {
        "status": evaluated.get("status"),
        "cycle_status": evaluated.get("cycle_status"),
        "actual_snapshot_id": evaluated.get("actual_snapshot_id"),
        "validation_snapshot_id": evaluated.get("actual_snapshot_id"),
        "validation_monitor_id": monitor.get("monitor_id") or monitor["_id"],
        "validation_prediction_id": pending.get("prediction_id") or pending.get("_id"),
        "actual_average_price": evaluated.get("actual_average_price"),
        "validation_snapshot_at": evaluated.get("validation_snapshot_at"),
        "observed_average_price": evaluated.get("observed_average_price"),
        "validation_observed_price": evaluated.get("validation_observed_price"),
        "baseline_absolute_error": evaluated.get("baseline_absolute_error"),
        "baseline_percentage_error": evaluated.get("baseline_percentage_error"),
        "benchmark_absolute_error": evaluated.get("benchmark_absolute_error"),
        "benchmark_error_percent": evaluated.get("benchmark_error_percent"),
        "ai_absolute_error": evaluated.get("ai_absolute_error"),
        "ai_percentage_error": evaluated.get("ai_percentage_error"),
        "ai_error_percent": evaluated.get("ai_error_percent"),
        "baseline_mae": evaluated.get("baseline_mae"),
        "baseline_mape": evaluated.get("baseline_mape"),
        "ai_mae": evaluated.get("ai_mae"),
        "ai_mape": evaluated.get("ai_mape"),
        "absolute_error": evaluated.get("absolute_error"),
        "percentage_error": evaluated.get("percentage_error"),
        "evaluated_at": evaluated.get("evaluated_at"),
        "validated_at": evaluated.get("validated_at"),
    })
    return pending.get("_id")


def refresh_monitor(monitor_id, owner_id=None, trigger="manual", force=False, scheduled_date=None, now=None):
    """Shared manual/scheduled refresh path with safe status and daily idempotency."""
    now_utc = _normalized_utc_datetime(now) or utcnow()
    trigger = "scheduled" if trigger == "scheduled" else "manual"
    monitor = repository.get_watchlist_item(monitor_id, owner_id)
    if not monitor or monitor.get("status") != "active":
        return {"status": "skipped", "error_code": "monitor_ineligible", "snapshot_id": None}
    owner_id = monitor.get("user_id")
    run_date = scheduled_date or monitor_scheduled_date(now_utc)
    if trigger == "scheduled":
        if not force and not app.config.get("MONITOR_DAILY_REFRESH_ENABLED", False):
            return {"status": "skipped", "error_code": "scheduled_refresh_disabled", "snapshot_id": None}
        if not force and not monitor.get("auto_refresh_enabled", False):
            return {"status": "skipped", "error_code": "automatic_refresh_paused", "snapshot_id": None}
        due_at = _normalized_utc_datetime(monitor.get("next_refresh_at"))
        if not force and (not due_at or due_at > now_utc):
            return {"status": "skipped", "error_code": "not_due", "snapshot_id": None}
        if not repository.claim_scheduled_monitor_refresh(monitor["_id"], owner_id, run_date):
            safe_call(lambda: repository.log_event(owner_id, "scheduled_refresh_duplicate_skipped", None, "Scheduled monitor", {"monitor_id": str(monitor["_id"]), "trigger": trigger, "status": "skipped", "scheduled_date": run_date}), None)
            return {"status": "skipped", "error_code": "duplicate_scheduled_refresh", "snapshot_id": None}
        safe_call(lambda: repository.log_event(owner_id, "scheduled_refresh_started", None, "Scheduled monitor", {"monitor_id": str(monitor["_id"]), "trigger": trigger, "status": "started", "scheduled_date": run_date}), None)

    try:
        if app.config.get("TESTING") and not app.config.get("MONITOR_TEST_PROVIDER_CALLS"):
            records, source_statuses = _offline_monitor_records(monitor)
        else:
            records, source_statuses = _retrieve_monitor_provider_records(monitor)
    except Exception as exc:
        records = []
        source_statuses = {"monitor": {"status": "failed", "error_code": f"provider_{exc.__class__.__name__.lower()}"}}
    requested_count = len(source_statuses)
    succeeded_count = sum(row.get("status") == "success" for row in source_statuses.values())
    payload = _snapshot_payload_from_items(monitor, records)
    snapshot_id = None
    status = "failed"
    error_code = None
    if payload:
        status = "partial" if succeeded_count < requested_count else "success"
        payload.update({
            "refresh_trigger": trigger,
            "scheduled_date": run_date if trigger == "scheduled" else None,
            "source_statuses": source_statuses,
            "data_quality": "partial" if status == "partial" else payload.get("data_quality"),
        })
        snapshot_id = repository.save_price_snapshot(monitor["_id"], owner_id, payload)
        snapshot = repository.latest_price_snapshot(monitor["_id"])
        if snapshot:
            _evaluate_pending_monitor_forecast(monitor, snapshot)
            def alert_audit(event_type, details):
                repository.log_event(owner_id, event_type, None, "Price alert", details)
            # Alert delivery is deliberately isolated from refresh success.  The
            # evaluator only reads the snapshot already collected above.
            safe_call(lambda: evaluate_price_alert(
                repository, monitor, snapshot,
                mail_service_factory=PasswordResetMailService,
                monitor_url=price_alert_monitor_url(monitor["_id"]),
                alerts_enabled=app.config.get("PRICE_ALERTS_ENABLED", True),
                now=now_utc,
                audit=alert_audit,
            ), None)
    else:
        error_code = "no_valid_comparable_data" if succeeded_count else "all_sources_unavailable"

    updates = {
        "last_refreshed_at": now_utc,
        "last_refresh_status": status,
        "last_refresh_error_code": error_code,
        "last_refresh_trigger": trigger,
    }
    if trigger == "scheduled":
        updates["last_scheduled_refresh_date"] = run_date
        updates["next_refresh_at"] = next_monitor_refresh_at(now_utc)
    repository.update_watchlist_item(monitor["_id"], updates, user_id=owner_id)
    if trigger == "scheduled":
        repository.complete_scheduled_monitor_refresh(monitor["_id"], run_date, status)
        safe_call(lambda: repository.log_event(owner_id, f"scheduled_refresh_{status}", None, "Scheduled monitor", {"monitor_id": str(monitor["_id"]), "trigger": trigger, "status": status, "scheduled_date": run_date, "error_code": error_code}), None)
    return {"status": status, "error_code": error_code, "snapshot_id": snapshot_id, "source_statuses": source_statuses}


def due_monitor_candidates(now=None, monitor_id=None, force=False):
    now_utc = _normalized_utc_datetime(now) or utcnow()
    monitors = repository.list_watchlist_items(user_id=None, include_archived=True, limit=0)
    candidates = []
    for monitor in monitors:
        if monitor_id is not None and str(monitor.get("_id")) != str(monitor_id):
            continue
        if monitor.get("status") != "active":
            continue
        if not force and not monitor.get("auto_refresh_enabled", False):
            continue
        due_at = _normalized_utc_datetime(monitor.get("next_refresh_at"))
        if not force and (not due_at or due_at > now_utc):
            continue
        candidates.append(monitor)
    candidates.sort(key=lambda row: _normalized_utc_datetime(row.get("next_refresh_at")) or datetime.max.replace(tzinfo=timezone.utc))
    return candidates


@app.post("/watchlist/<watchlist_id>/refresh")
@login_required
def refresh_watchlist_snapshot(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    result = refresh_monitor(watchlist_id, session["user_id"], trigger="manual")
    if result["status"] == "failed":
        flash("No valid comparable records found for this snapshot.", "error")
    elif result["status"] == "skipped":
        flash("This Monitor is not eligible for refresh.", "warning")
    else:
        flash("Snapshot refreshed." if result["status"] == "success" else "Snapshot refreshed with partial source coverage.", "success")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/auto-refresh/enable")
@login_required
def enable_watchlist_auto_refresh(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    enabled, current, reason = repository.enable_watchlist_auto_refresh(
        watchlist_id,
        session["user_id"],
        next_monitor_refresh_at(),
        app.config.get("MONITOR_REFRESH_TIMEZONE", "Asia/Singapore"),
    )
    if not enabled:
        if reason == "limit":
            current_label = (current or {}).get("product_label") or (current or {}).get("keyword")
            detail = f" Current automatic Monitor: {current_label}." if current_label else ""
            flash(f"Pause the current automatic monitor before enabling another one.{detail}", "warning")
        else:
            flash("Only active Monitors can enable daily refresh.", "warning")
        return _watchlist_redirect(watchlist_id)
    repository.log_event(session["user_id"], "daily_refresh_enabled", session.get("active_role") or session.get("role"), session.get("username"), {"monitor_id": watchlist_id, "owner_id": session["user_id"], "trigger": "manual", "status": "enabled"})
    flash("Daily refresh enabled.", "success")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/auto-refresh/pause")
@login_required
def pause_watchlist_auto_refresh(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    if item.get("status") != "active":
        flash("Only active Monitors can pause daily refresh.", "warning")
        return _watchlist_redirect(watchlist_id)
    repository.pause_watchlist_auto_refresh(watchlist_id, session["user_id"])
    repository.log_event(session["user_id"], "daily_refresh_paused", session.get("active_role") or session.get("role"), session.get("username"), {"monitor_id": watchlist_id, "owner_id": session["user_id"], "trigger": "manual", "status": "paused"})
    flash("Daily refresh paused.", "success")
    return _watchlist_redirect(watchlist_id)


def _owned_alert_monitor(watchlist_id):
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    return item


@app.post("/watchlist/<watchlist_id>/price-alert/settings")
@login_required
def configure_watchlist_price_alert(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    item = _owned_alert_monitor(watchlist_id)
    direction = str(request.form.get("alert_direction") or "").strip().lower()
    try:
        threshold = Decimal(str(request.form.get("alert_threshold_percent") or ""))
    except Exception:
        threshold = None
    if direction not in VALID_DIRECTIONS:
        flash("Choose drop, increase, or either for the price alert direction.", "error")
        return _watchlist_redirect(watchlist_id)
    if threshold is None or threshold < Decimal("1") or threshold > Decimal("50"):
        flash("Price alert threshold must be between 1% and 50%.", "error")
        return _watchlist_redirect(watchlist_id)
    now = utcnow()
    repository.update_watchlist_item(watchlist_id, {
        "alert_direction": direction,
        "alert_threshold_percent": float(threshold.quantize(Decimal("0.01"))),
        "alert_email_enabled": request.form.get("alert_email_enabled") == "on",
        "alert_cooldown_hours": int(app.config.get("PRICE_ALERT_COOLDOWN_HOURS", 24)),
        "alert_rule_updated_at": now,
        "alert_rule_version": int(item.get("alert_rule_version") or 0) + 1,
    }, user_id=session["user_id"])
    repository.log_event(session["user_id"], "price_alert_configured", session.get("role"), session.get("username"), {
        "monitor_id": watchlist_id, "direction": direction, "threshold_percent": float(threshold),
        "email_enabled": request.form.get("alert_email_enabled") == "on",
    })
    flash("Price alert settings saved.", "success")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/price-alert/enable")
@login_required
def enable_watchlist_price_alert(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    item = _owned_alert_monitor(watchlist_id)
    if item.get("status") != "active":
        flash("Archived or deleted Monitors cannot enable price alerts.", "warning")
        return _watchlist_redirect(watchlist_id)
    threshold = item.get("alert_threshold_percent")
    if threshold is None:
        try:
            threshold = float(app.config.get("PRICE_ALERT_DEFAULT_THRESHOLD_PERCENT", 5))
        except (TypeError, ValueError):
            threshold = 5.0
    repository.update_watchlist_item(watchlist_id, {
        "alert_enabled": True,
        "alert_threshold_percent": min(50.0, max(1.0, float(threshold))),
        "alert_rule_updated_at": utcnow(),
        "alert_rule_version": int(item.get("alert_rule_version") or 0) + 1,
    }, user_id=session["user_id"])
    repository.log_event(session["user_id"], "price_alert_enabled", session.get("role"), session.get("username"), {"monitor_id": watchlist_id})
    flash("Price alert enabled. It will evaluate future Monitor snapshots.", "success")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/price-alert/pause")
@login_required
def pause_watchlist_price_alert(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    _owned_alert_monitor(watchlist_id)
    repository.update_watchlist_item(watchlist_id, {"alert_enabled": False, "alert_rule_updated_at": utcnow()}, user_id=session["user_id"])
    repository.log_event(session["user_id"], "price_alert_paused", session.get("role"), session.get("username"), {"monitor_id": watchlist_id})
    flash("Price alert paused. Historical snapshots will not be resent when you resume.", "success")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/price-alert/test")
@login_required
def test_watchlist_price_alert_email(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    item = _owned_alert_monitor(watchlist_id)
    user = repository.get_user_by_id(session["user_id"])
    try:
        delivery = PasswordResetMailService().send_price_alert_test(
            user.get("email"), item.get("product_label") or item.get("keyword") or "Watchlist Monitor",
            price_alert_monitor_url(watchlist_id),
        )
        repository.log_event(session["user_id"], "test_alert_email_sent", session.get("role"), session.get("username"), {"monitor_id": watchlist_id, "delivery_mode": delivery, "notification_type": "test"})
        flash("Test price alert email sent. No marketplace threshold was triggered.", "success")
    except Exception as exc:
        repository.log_event(session["user_id"], "test_alert_email_failed", session.get("role"), session.get("username"), {"monitor_id": watchlist_id, "error_code": normalize_mail_error(exc), "notification_type": "test"})
        flash("The test email could not be delivered. Check the configured mail service and try again.", "error")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/archive")
@login_required
def archive_watchlist_item(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    repository.archive_watchlist_item(watchlist_id, session["user_id"])
    flash("Monitor moved to Archived.", "success")
    return redirect(url_for("watchlist"))


@app.post("/watchlist/<watchlist_id>/restore")
@login_required
def restore_watchlist_item(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item or item.get("status") != "archived":
        abort(404)
    repository.restore_watchlist_item(watchlist_id, session["user_id"])
    flash("Monitor restored to Active.", "success")
    return redirect(url_for("watchlist", item_id=watchlist_id))


@app.post("/watchlist/<watchlist_id>/delete")
@login_required
def delete_watchlist_item(watchlist_id):
    blocked = ensure_feature_access("watchlist", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    repository.delete_watchlist_item(watchlist_id, session["user_id"])
    flash("Watchlist item deleted.", "success")
    return redirect(url_for("watchlist"))


@app.route("/watchlist")
@login_required
def watchlist():
    if not can_access(current_user(), "watchlist"):
        return locked_feature_response("watchlist", title="Watchlist")
    include_archived = request.args.get("archived") == "1"
    debug_charts = request.args.get("debug_charts") == "1"
    debug_validation = request.args.get("debug_validation") == "1"
    selected_id = request.args.get("item_id")
    all_items = _watchlist_items_for_user(session["user_id"], include_archived=True)
    active_items = [item for item in all_items if item.get("status") not in {"archived", "deleted"}]
    archived_items = [item for item in all_items if item.get("status") == "archived"]
    automatic_items = [item for item in active_items if item.get("auto_refresh_enabled")]
    items = archived_items if include_archived else active_items
    selected_item = _watchlist_selected_item(items, selected_id=selected_id)
    if selected_id and not any(str(item.get("_id")) == str(selected_id) for item in items):
        flash("That watchlist item could not be found. Showing the most recent item instead.", "info")
    snapshots = repository.list_price_snapshots(selected_item["_id"], limit=100) if selected_item else []
    snapshots = _snapshots_with_forecast_context(snapshots)
    summary = selected_item.get("latest_snapshot") if selected_item and selected_item.get("latest_snapshot") else None
    predictions = repository.list_predictions(session["user_id"], selected_item["_id"], limit=100) if selected_item else []
    alert_events = repository.list_price_alert_events(
        session["user_id"], selected_item.get("monitor_id") or selected_item["_id"], limit=50,
    ) if selected_item else []
    forecast_cycles = [_forecast_cycle_view(row, snapshots) for row in predictions]
    current_forecast_cycle = forecast_cycles[0] if forecast_cycles else None
    selected_forecast_id = str(request.args.get("forecast_id") or "")
    requested_forecast_cycle = next((row for row in forecast_cycles if selected_forecast_id in {str(row.get("forecast_id")), str(row.get("_id"))}), None)
    selected_forecast_cycle = requested_forecast_cycle or current_forecast_cycle
    trend_chart = _watchlist_trend_chart_data(snapshots)
    trend_svg = _watchlist_trend_svg_data(trend_chart)
    trend_summary = _watchlist_trend_summary(trend_chart)
    actual_snapshot_used = _prediction_actual_snapshot(selected_forecast_cycle, snapshots)
    prediction_has_actual = actual_snapshot_used is not None
    if current_forecast_cycle and current_forecast_cycle.get("visible_status") == "validated":
        forecast_workflow_status = "Forecast validated"
    elif current_forecast_cycle:
        forecast_workflow_status = "Waiting for validation snapshot"
    else:
        forecast_workflow_status = "No forecast generated"
    snapshots_total = len(snapshots)
    snapshots_after_prediction_count = len([row for row in snapshots if selected_forecast_cycle and is_snapshot_after_prediction(row, selected_forecast_cycle)])
    pending_prediction = next((row for row in predictions if row.get("status") in {"pending_actual", "ai_unavailable"}), None)
    chart_filename = _prediction_chart_filename(selected_forecast_cycle, actual_snapshot_used) if selected_forecast_cycle and prediction_has_actual else None
    prediction_chart = _forecast_cycle_chart_data(selected_forecast_cycle)
    reason_validation_not_available = prediction_chart.get("reason")
    env_loaded = Path(".env").exists()
    gemini_api_key_present = bool(os.getenv("GEMINI_API_KEY"))
    gemini_model = os.getenv("GEMINI_MODEL", "gemini-2.5-flash-lite")
    gemini_call_reason = None
    if selected_forecast_cycle:
        gemini_call_reason = selected_forecast_cycle.get("ai_failure_reason") or selected_forecast_cycle.get("ai_reason") or selected_forecast_cycle.get("ai_provider")
        if selected_forecast_cycle.get("ai_predicted_average_price") is not None:
            gemini_call_reason = "api_call_success"
        elif selected_forecast_cycle.get("ai_status") == "ai_unavailable" or selected_forecast_cycle.get("status") == "ai_unavailable":
            gemini_call_reason = _normalize_gemini_status_reason(gemini_call_reason)
        else:
            gemini_call_reason = "record_has_no_gemini_output"
    plotted_snapshot_count = len([value for value in trend_chart.get("average", []) if value is not None]) if trend_chart else 0
    validation_error = None
    if selected_forecast_cycle and prediction_has_actual:
        validation_error = {
            "benchmarkError": selected_forecast_cycle.get("benchmark_absolute_error"),
            "benchmarkPctError": selected_forecast_cycle.get("benchmark_error_percent"),
            "aiError": selected_forecast_cycle.get("ai_absolute_error"),
            "aiPctError": selected_forecast_cycle.get("ai_error_percent"),
            "bestMethod": "Benchmark" if selected_forecast_cycle.get("ai_absolute_error") is None or (selected_forecast_cycle.get("benchmark_absolute_error") is not None and selected_forecast_cycle.get("benchmark_absolute_error") <= (selected_forecast_cycle.get("ai_absolute_error") or selected_forecast_cycle.get("benchmark_absolute_error"))) else "AI-assisted forecast",
        }
    return render_template(
        "watchlist.html",
        items=items,
        snapshots=snapshots,
        selected_item=selected_item,
        summary=summary,
        include_archived=include_archived,
        active_count=len(active_items),
        archived_count=len(archived_items),
        automatic_count=len(automatic_items),
        selected_id=selected_id,
        predictions=predictions,
        forecast_cycles=forecast_cycles,
        selected_forecast_cycle=selected_forecast_cycle,
        current_forecast_cycle=current_forecast_cycle,
        selected_forecast_id=str((selected_forecast_cycle or {}).get("forecast_id") or ""),
        chart_filename=chart_filename,
        trend_chart=make_json_safe(trend_chart),
        trend_svg=make_json_safe(trend_svg),
        trend_summary=make_json_safe(trend_summary),
        prediction_chart=make_json_safe(prediction_chart),
        prediction_error=make_json_safe(validation_error),
        prediction_has_actual=prediction_has_actual,
        forecast_workflow_status=forecast_workflow_status,
        actual_snapshot_used=actual_snapshot_used,
        selected_forecast_cycle_safe=make_json_safe(selected_forecast_cycle),
        debug_charts=debug_charts,
        debug_validation=debug_validation,
        chart_js_loaded=True,
        pending_prediction=pending_prediction,
        alert_events=alert_events,
        snapshots_total=snapshots_total,
        plotted_snapshot_count=plotted_snapshot_count,
        snapshots_after_prediction_count=snapshots_after_prediction_count,
        validation_status=selected_forecast_cycle.get("status") if selected_forecast_cycle else None,
        last_evaluation_error=validation_error,
        reason_validation_not_available=reason_validation_not_available,
        gemini_status={
            "env_loaded": env_loaded,
            "api_key_present": gemini_api_key_present,
            "model": gemini_model,
            "call_reason": gemini_call_reason,
        },
    )


@app.route("/watchlist/export/snapshots.csv")
@login_required
def export_watchlist_snapshots_csv():
    if not can_access(current_user(), "watchlist"):
        return export_forbidden_response("Watchlist exports require Premium membership.")
    monitor_id = request.args.get("monitor_id", "").strip()
    if not monitor_id:
        abort(404)
    monitor = repository.get_watchlist_item(monitor_id, session["user_id"])
    if not monitor:
        abort(404)
    monitor = _monitor_with_category_metadata(monitor)
    snapshots = _monitor_snapshot_history(monitor, session["user_id"])
    rows = [
        ["Monitor Name", "Monitor ID", "Snapshot ID", "Collected At (SGT)", "Average Price", "Lowest Price", "Highest Price", "Listing Count", "Platform Coverage", "Source", "Forecast Context", "Currency", "Category"],
    ]
    for snapshot in _snapshots_with_forecast_context(snapshots):
        rows.append([
            monitor.get("product_label") or monitor.get("keyword") or "",
            str(monitor.get("monitor_id") or monitor.get("_id")),
            str(snapshot.get("_id") or ""),
            singapore_time_string(snapshot.get("collected_at") or snapshot.get("created_at")),
            _display_price(snapshot.get("average_price")),
            _display_price(snapshot.get("lowest_price") if snapshot.get("lowest_price") is not None else snapshot.get("low_price")),
            _display_price(snapshot.get("highest_price") if snapshot.get("highest_price") is not None else snapshot.get("high_price")),
            snapshot.get("listing_count") if snapshot.get("listing_count") is not None else snapshot.get("record_count", ""),
            snapshot.get("platform_scope") or monitor.get("platform_scope") or "",
            snapshot.get("source_label") or monitor.get("source_label") or "",
            snapshot["forecast_context"]["label"],
            snapshot.get("currency") or monitor.get("currency") or "USD",
            monitor.get("category_display") or "General products",
        ])
    return csv_response(rows, "watchlist-snapshots.csv")


@app.route("/watchlist/export/validation.csv")
@login_required
def export_watchlist_validation_csv():
    if not can_access(current_user(), "prediction_validation"):
        return export_forbidden_response("Prediction validation exports require Professional membership.")
    monitor_id = request.args.get("monitor_id", "").strip()
    if not monitor_id:
        abort(404)
    monitor = repository.get_watchlist_item(monitor_id, session["user_id"])
    if not monitor:
        abort(404)
    monitor = _monitor_with_category_metadata(monitor)
    snapshots = _monitor_snapshot_history(monitor, session["user_id"])
    predictions = repository.list_predictions(session["user_id"], monitor["_id"], limit=100)
    requested_forecast_id = request.args.get("forecast_id", "").strip()
    if requested_forecast_id:
        selected = next((row for row in predictions if requested_forecast_id in {
            str(row.get("forecast_id") or ""), str(row.get("prediction_id") or ""), str(row.get("_id") or ""),
        }), None)
        if not selected:
            abort(404)
        predictions = [selected]
    rows = [
        ["Monitor Name", "Monitor ID", "Forecast Cycle ID", "Baseline Snapshot ID", "Validation Snapshot ID", "Forecast Created At (SGT)", "Validated At (SGT)", "Benchmark Prediction", "AI-assisted Forecast", "Validation Observed Price", "Benchmark Absolute Error", "AI Absolute Error", "Benchmark Error Percent", "AI Error Percent", "Forecast Direction", "Confidence", "Provider Used", "Fallback Used", "Validation Result", "Benchmark Method", "Baseline Snapshot At (SGT)", "Validation Snapshot At (SGT)", "Baseline Snapshot Observed Price", "Category"],
    ]
    for prediction in predictions:
        valid, validation_snapshot = _valid_prediction_validation(prediction, monitor, snapshots)
        if not valid:
            continue
        cycle = _forecast_cycle_view(prediction, snapshots)
        ai_value = cycle.get("ai_predicted_price")
        fallback_used = ai_value is None
        rows.append([
            monitor.get("product_label") or monitor.get("keyword") or "",
            str(monitor.get("monitor_id") or monitor.get("_id")),
            str(cycle.get("forecast_id") or ""),
            str(cycle.get("baseline_snapshot_id") or ""),
            str(cycle.get("validation_snapshot_id") or cycle.get("actual_snapshot_id") or ""),
            singapore_time_string(cycle.get("generated_at")),
            singapore_time_string(cycle.get("validated_at_display") or cycle.get("validation_snapshot_at")),
            _display_price(cycle.get("benchmark_predicted_price")),
            _display_price(ai_value),
            _display_price(cycle.get("validation_observed_price")),
            _display_price(cycle.get("benchmark_absolute_error")),
            _display_price(cycle.get("ai_absolute_error")) if not fallback_used else "",
            _display_percentage(cycle.get("benchmark_error_percent")),
            _display_percentage(cycle.get("ai_error_percent")) if not fallback_used else "",
            prediction.get("ai_predicted_direction") or "Not available",
            prediction.get("ai_confidence_level") or "Not available",
            prediction.get("ai_provider") if not fallback_used else "Fallback",
            "Yes" if fallback_used else "No",
            "Valid independent validation",
            cycle.get("benchmark_method_label"),
            singapore_time_string(cycle.get("baseline_snapshot_at")),
            singapore_time_string(cycle.get("validation_snapshot_at")),
            _display_price(cycle.get("baseline_snapshot_observed_price")),
            monitor.get("category_display") or "General products",
        ])
    return csv_response(rows, "watchlist-validation.csv")


@app.route("/watchlist/report")
@login_required
def watchlist_report():
    monitor_id = request.args.get("monitor_id", "")
    if not monitor_id or not is_valid_object_id(monitor_id):
        abort(404)
    monitor = repository.get_watchlist_item(monitor_id, session["user_id"])
    if not monitor:
        abort(404)
    monitor = _monitor_with_category_metadata(monitor)
    snapshots = _monitor_snapshot_history(monitor, session["user_id"])
    predictions = repository.list_predictions(session["user_id"], monitor_id, limit=50)
    latest_snapshot = snapshots[-1] if snapshots else None
    requested_forecast_id = request.args.get("forecast_id", "").strip()
    selected_prediction = next((row for row in predictions if requested_forecast_id in {
        str(row.get("forecast_id") or ""), str(row.get("prediction_id") or ""), str(row.get("_id") or ""),
    }), None) if requested_forecast_id else (predictions[0] if predictions else None)
    if requested_forecast_id and not selected_prediction:
        abort(404)
    selected_forecast_cycle = _forecast_cycle_view(selected_prediction, snapshots)
    validation_is_valid, validation_snapshot = _valid_prediction_validation(selected_prediction, monitor, snapshots) if selected_prediction else (False, None)
    trend_chart = _watchlist_trend_chart_data(snapshots)
    return render_template("watchlist_monitor_report.html", monitor=monitor, snapshots=snapshots,
                           latest_snapshot=latest_snapshot, selected_forecast_cycle=selected_forecast_cycle,
                           forecast_cycle_scope="Selected Forecast Cycle" if requested_forecast_id else "Latest Forecast Cycle",
                           validation_is_valid=validation_is_valid, validation_snapshot=validation_snapshot,
                           trend_svg=_watchlist_trend_svg_data(trend_chart),
                           forecast_chart=_report_forecast_chart_data(selected_forecast_cycle),
                           generated_at=utcnow())


@app.post("/watchlist/<watchlist_id>/prediction")
@login_required
def generate_watchlist_prediction(watchlist_id):
    blocked = ensure_feature_access("prediction", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    snapshots = _valid_forecast_snapshots(_sorted_watchlist_snapshots(watchlist_id, session["user_id"]))
    if any(row.get("status") in {"pending_actual", "ai_unavailable", "waiting_for_validation"} for row in repository.list_predictions(session["user_id"], watchlist_id, limit=20)):
        flash("A forecast is already waiting for validation.", "info")
        return _watchlist_redirect(watchlist_id)
    prediction = _build_prediction_record(item, snapshots)
    if prediction.get("status") == "insufficient_history":
        flash("More snapshots are needed for reliable prediction. Refresh this Watchlist item over time.", "warning")
        return _watchlist_redirect(watchlist_id)
    prediction_id = repository.save_prediction(session["user_id"], watchlist_id, prediction)
    repository.log_event(session["user_id"], "prediction_created", session.get("active_role") or session.get("role"), session.get("username"), {"watchlist_id": watchlist_id, "prediction_id": prediction_id, "method": prediction.get("baseline_method"), "ai_provider": prediction.get("ai_provider")})
    flash("Prediction generated.", "success")
    return _watchlist_redirect(watchlist_id)


@app.post("/watchlist/<watchlist_id>/prediction/evaluate")
@login_required
def evaluate_watchlist_prediction(watchlist_id):
    blocked = ensure_feature_access("prediction_validation", redirect_endpoint="watchlist")
    if blocked:
        return blocked
    if not is_valid_object_id(watchlist_id):
        abort(404)
    item = repository.get_watchlist_item(watchlist_id, session["user_id"])
    if not item:
        abort(404)
    predictions = repository.list_predictions(session["user_id"], watchlist_id, limit=20)
    pending = next((row for row in predictions if row.get("status") in {"pending_actual", "ai_unavailable"}), None)
    if not pending:
        flash("No pending prediction found for this Watchlist item.", "info")
        return _watchlist_redirect(watchlist_id)
    snapshots = _sorted_watchlist_snapshots(watchlist_id, session["user_id"])
    future_snapshot = next((row for row in snapshots if is_snapshot_after_prediction(row, pending)), None)
    if not future_snapshot:
        flash("No later actual snapshot is available yet. Refresh this watchlist later before evaluation.", "warning")
        return _watchlist_redirect(watchlist_id)
    evaluated = _evaluate_prediction_record(pending, future_snapshot)
    update_fields = {
        "status": evaluated.get("status"),
        "cycle_status": evaluated.get("cycle_status"),
        "actual_snapshot_id": evaluated.get("actual_snapshot_id"),
        "validation_snapshot_id": evaluated.get("actual_snapshot_id"),
        "validation_monitor_id": pending.get("monitor_id") or watchlist_id,
        "validation_prediction_id": pending.get("prediction_id") or pending.get("_id"),
        "actual_average_price": evaluated.get("actual_average_price"),
        "validation_snapshot_at": evaluated.get("validation_snapshot_at"),
        "observed_average_price": evaluated.get("observed_average_price"),
        "validation_observed_price": evaluated.get("validation_observed_price"),
        "baseline_absolute_error": evaluated.get("baseline_absolute_error"),
        "baseline_percentage_error": evaluated.get("baseline_percentage_error"),
        "benchmark_absolute_error": evaluated.get("benchmark_absolute_error"),
        "benchmark_error_percent": evaluated.get("benchmark_error_percent"),
        "ai_absolute_error": evaluated.get("ai_absolute_error"),
        "ai_percentage_error": evaluated.get("ai_percentage_error"),
        "ai_error_percent": evaluated.get("ai_error_percent"),
        "baseline_mae": evaluated.get("baseline_mae"),
        "baseline_mape": evaluated.get("baseline_mape"),
        "ai_mae": evaluated.get("ai_mae"),
        "ai_mape": evaluated.get("ai_mape"),
        "absolute_error": evaluated.get("absolute_error"),
        "percentage_error": evaluated.get("percentage_error"),
        "evaluated_at": evaluated.get("evaluated_at"),
        "validated_at": evaluated.get("validated_at"),
    }
    repository.update_prediction(pending["_id"], update_fields)
    flash("Prediction evaluated against the latest snapshot.", "success")
    return _watchlist_redirect(watchlist_id)


@app.route("/audit")
@login_required
def audit():
    user = current_user()
    if not can_access(user, "source_audit"):
        return locked_feature_response("source_audit", title="Source Audit")
    return render_template("audit.html", records=build_audit_records(audit_scope_user_id(user)), locked=False)


@app.post("/audit/archive")
@role_required("administrator")
def archive_audit_logs():
    references = [value for value in request.form.getlist("record_ref") if value]
    legacy_ids = [value for value in request.form.getlist("audit_id") if value]
    references.extend(f"audit_logs:{value}" for value in legacy_ids)
    if not references:
        flash("Select at least one audit record.", "error")
        return redirect(url_for("audit"))
    known_ids = {
        "search_records": {str(row.get("_id")) for row in repository.list_searches(limit=0)},
        "evidence_records": {str(row.get("_id")) for row in repository.list_evidence(limit=0)},
        "ai_search_logs": {str(row.get("_id")) for row in repository.list_ai_logs(limit=0)},
        "audit_logs": {str(row.get("_id")) for row in repository.list_audit_logs(limit=0)},
    }
    archived = 0
    for reference in references:
        collection, separator, record_id = str(reference).partition(":")
        if not separator or record_id not in known_ids.get(collection, set()):
            continue
        if collection == "search_records":
            repository.delete_search(record_id, deleted_by=session.get("user_id"), reason="administrator_audit_cleanup")
        elif collection == "evidence_records":
            repository.delete_evidence(record_id, deleted_by=session.get("user_id"), reason="administrator_audit_cleanup")
        elif collection == "ai_search_logs":
            repository.archive_ai_log(record_id, archived_by=session.get("user_id"), reason="administrator_audit_cleanup")
        else:
            repository.archive_audit_log(record_id, archived_by=session.get("user_id"), reason="administrator_audit_cleanup")
        archived += 1
    repository.log_event(session["user_id"], "audit_log_archived", session.get("role"), session.get("username"), {"archived_count": archived, "selected_count": len(references), "mode": "administrator_audit_cleanup"})
    if archived:
        flash(f"Archived {archived} selected record(s) from active views.", "success")
    else:
        flash("No eligible records were archived.", "info")
    return redirect(url_for("audit"))


@app.route("/audit/export/source-audit.csv")
@login_required
def export_audit_csv():
    if not can_access(current_user(), "source_audit"):
        return export_forbidden_response("Source audit exports require Professional membership.")
    rows_data = _filter_audit_records(build_audit_records(audit_scope_user_id()), _audit_export_filters())
    rows = [["Audit ID", "Record Type", "Query / Product", "User", "Role", "Workspace", "Source Scope", "eBay Status", "Walmart Status", "Matched Records", "Comparable Records", "Action / Event", "Status", "Timestamp (SGT)", "Details"]]
    for row in rows_data:
        rows.append([
            row.get("audit_id", ""), row.get("type_label", ""), row.get("query", ""), row.get("user", ""),
            row.get("role", ""), row.get("workspace", ""), row.get("source_scope", ""), row.get("ebay_status", ""),
            row.get("walmart_status", ""), row.get("matched_records", ""), row.get("comparable_records", ""),
            row.get("action", ""), row.get("status", ""), singapore_time_string(row.get("time")), row.get("detail", ""),
        ])
    return csv_response(rows, _dated_export_filename("source-audit"))


@app.route("/audit/export/ai-activity-log.csv")
@login_required
def export_ai_activity_log_csv():
    if not can_access(current_user(), "source_audit"):
        return export_forbidden_response("AI activity exports require Professional membership.")
    rows = [["Event ID", "Operation", "User", "Role", "Workspace", "Query / Analysis", "AI Provider", "Model", "Input Record Count", "Scope", "Fallback Used", "Status", "Started At (SGT)", "Completed At (SGT)", "Duration (ms)", "Error Summary"]]
    for row in build_ai_activity_export_records(audit_scope_user_id()):
        rows.append([
            row.get("event_id", ""), row.get("operation", ""), row.get("user", ""), row.get("role", ""),
            row.get("workspace", ""), row.get("query", ""), row.get("provider", ""), row.get("model", ""),
            row.get("input_record_count", ""), row.get("scope", ""), row.get("fallback_used", ""), row.get("status", ""),
            singapore_time_string(row.get("started_at")), singapore_time_string(row.get("completed_at")),
            row.get("duration_ms", ""), row.get("error_summary", ""),
        ])
    return csv_response(rows, _dated_export_filename("ai-activity-log"))


@app.route("/audit/export/activity-log.csv")
@login_required
def export_activity_log_csv():
    if not can_access(current_user(), "source_audit"):
        return export_forbidden_response("Source audit exports require Professional membership.")
    rows_data = build_activity_records(audit_scope_user_id())
    rows = [["Action", "User", "Role", "Workspace", "Detail", "Time (SGT)", "Status"]]
    for row in rows_data[:200]:
        rows.append([
            str(row.get("action") or "Event").replace("_", " ").title(), row.get("user", ""), row.get("role", ""),
            str(row.get("workspace") or row.get("role") or "Precision Curator").replace("_", " ").title(),
            row.get("detail", ""), singapore_time_string(row.get("time")), row.get("status", "Recorded"),
        ])
    return csv_response(rows, _dated_export_filename("activity-log"))


@app.route("/audit/report")
@login_required
def audit_report():
    if not can_access(current_user(), "source_audit"):
        return export_forbidden_response("Source audit report requires Professional membership.")
    filters = _audit_export_filters()
    records = _filter_audit_records(build_audit_records(audit_scope_user_id()), filters)
    summary = {
        "search": sum(row.get("category") == "Search records" for row in records),
        "evidence": sum(row.get("category") == "Saved evidence" for row in records),
        "ai": sum(row.get("category") == "AI analysis" for row in records),
        "source": sum(row.get("category") == "Source events" for row in records),
        "completed": sum(row.get("status") == "Completed" for row in records),
        "partial": sum(row.get("status") == "Partial" for row in records),
        "failed": sum(row.get("status") == "Failed" for row in records),
    }
    marketplace_summary = {
        "ebay_success": sum(row.get("ebay_status") == "Success" for row in records),
        "ebay_no_results": sum(row.get("ebay_status") == "No results" for row in records),
        "ebay_failure": sum(row.get("ebay_status") not in {"Success", "No results", "Not requested"} for row in records),
        "walmart_success": sum(row.get("walmart_status") == "Success" for row in records),
        "walmart_no_results": sum(row.get("walmart_status") == "No results" for row in records),
        "walmart_correction": sum(row.get("walmart_status") == "Query correction suggested" for row in records),
        "walmart_failure": sum(row.get("walmart_status") not in {"Success", "No results", "Not requested", "Query correction suggested"} for row in records),
    }
    user = current_user() or {}
    return render_template(
        "audit_report.html", records=records, summary=summary, marketplace_summary=marketplace_summary,
        generated_at=utcnow(), report_user=user.get("display_name") or user.get("username") or session.get("username") or "System",
        workspace=(user.get("active_role") or user.get("role") or session.get("role") or "Precision Curator").replace("_", " ").title(),
        filters=filters, singapore_time_string=singapore_time_string,
    )


@app.route("/logs")
@login_required
def logs():
    user = current_user()
    active_role = user.get("active_role") or user.get("role")
    if active_role != "administrator" and (active_role != "researcher" or not can_access(user, "logs")):
        return locked_feature_response("logs", title="Activity Logs")
    return render_template("logs_complete.html", records=build_activity_records(audit_scope_user_id(user)))


@app.post("/logs/archive")
@role_required("researcher", "administrator")
def archive_activity_logs():
    blocked = ensure_feature_access("logs", redirect_endpoint="logs")
    if blocked:
        return blocked
    user = current_user()
    if (user.get("active_role") or user.get("role")) not in {"researcher", "administrator"}:
        return locked_feature_response("logs", title="Activity Logs")
    collection = request.form.get("collection", "activity_logs")
    ids = [value for value in request.form.getlist("log_id") if value]
    if not ids:
        flash("Select at least one log entry.", "error")
        return redirect(url_for("logs"))
    if collection == "audit_logs" and session.get("role") != "administrator":
        flash("You do not have permission to archive audit records.", "error")
        return redirect(url_for("logs"))
    if collection == "audit_logs":
        archived = repository.archive_many_audit_logs(ids, archived_by=session.get("user_id"), reason="log_archive_selected")
        repository.log_event(session["user_id"], "audit_log_archived", session.get("role"), session.get("username"), {"archived_count": archived, "collection": collection})
    else:
        if session.get("role") != "administrator":
            owned_ids = {str(row.get("_id")) for row in repository.list_activity_logs(user_id=session.get("user_id"), limit=0)}
            ids = [value for value in ids if str(value) in owned_ids]
        archived = repository.archive_many_activity_logs(ids, archived_by=session.get("user_id"), reason="log_archive_selected")
        repository.log_event(session["user_id"], "activity_log_archived", session.get("role"), session.get("username"), {"archived_count": archived, "collection": collection})
    flash(f"Archived {archived} selected log entry(s).", "success")
    return redirect(url_for("logs"))


@app.post("/logs/archive-old")
@role_required("administrator")
def archive_old_logs():
    cutoff = utcnow() - timedelta(days=30)
    archived = 0
    for row in repository.list_activity_logs(limit=0):
        created = row.get("created_at")
        if created and created < cutoff:
            repository.archive_activity_log(row["_id"], archived_by=session.get("user_id"), reason="archive_old_logs")
            archived += 1
    for row in repository.list_audit_logs(limit=0):
        created = row.get("created_at")
        if created and created < cutoff:
            repository.archive_audit_log(row["_id"], archived_by=session.get("user_id"), reason="archive_old_logs")
            archived += 1
    repository.log_event(session["user_id"], "audit_log_archived", session.get("role"), session.get("username"), {"archived_count": archived, "mode": "old_records"})
    flash(f"Archived {archived} old log entry(s).", "success")
    return redirect(url_for("logs"))


@app.route("/search-legacy", methods=["GET", "POST"])
@login_required
def legacy_search():
    params = request.form if request.method == "POST" else {}
    if request.method == "GET" and request.args.get("q"):
        repository.log_event(session["user_id"], "state_changing_get_blocked", session.get("role"), session.get("username"), {"endpoint": "legacy_search"})
    query = clean_search_query(params.get("q", ""))
    source = params.get("source", params.get("platform", "demo"))
    if source.startswith("demo_"):
        source = "demo"
    sort_option, item_type = params.get("sort", "default"), params.get("item_type", "all")
    platform_filter, category = params.get("platform_filter", ""), params.get("category", "")
    region, currency, condition = params.get("region", "US"), params.get("currency", ""), params.get("condition", "")
    min_price, max_price = optional_float(params.get("min_price")), optional_float(params.get("max_price"))
    items, summary, insight, error, notice, search_record_id = [], None, None, None, None, None
    if query and is_blocked_search_query(query):
        record_blocked_search_attempt(session["user_id"], query, "unsafe keyword pattern")
        error = "The search term contains unsupported characters. Please enter a product name or model."
        query = ""
    if query:
        if source not in SOURCES:
            error = "Please choose one of the available data sources."
        else:
            requested_source = source
            repository.log_event(session["user_id"], "search_started", session["role"], session["username"], {"query": query, "source": requested_source, "region": region})
            unavailable = (source == "ebay" and not (CLIENT_ID and CLIENT_SECRET)) or (source == "ai" and not (os.getenv("OPENAI_API_KEY") or os.getenv("AI_API_KEY")))
            if unavailable:
                notice = f"{SOURCES[requested_source]} is not configured; market records are shown instead."
                repository.log_event(session["user_id"], "api_call_failed", session["role"], session["username"], {"query": query, "source": requested_source, "error_type": "not_configured"})
                source = "demo"
            search_record_id = repository.create_search(session["user_id"], query, source, role=session.get("role"), data_mode="synthetic" if source == "demo" else "api")
            try:
                try:
                    raw_items = perform_search(query, source, item_type, region)
                except (AISearchError, RuntimeError, requests.RequestException) as exc:
                    if requested_source not in {"ebay", "ai"}:
                        raise
                    app.logger.warning("Price source failed; using demo fallback (%s)", exc.__class__.__name__)
                    repository.log_event(session["user_id"], "api_call_failed", session["role"], session["username"], {"query": query, "source": requested_source, "error_type": exc.__class__.__name__})
                    source = "demo"
                    repository.update_search_source(search_record_id, source, "synthetic")
                    notice = f"{SOURCES[requested_source]} could not be reached; market records are shown instead."
                    raw_items = perform_search(query, source, item_type, region)
                raw_items = normalize_price_items(raw_items)
                items = apply_result_filters(raw_items, platform_filter, category, min_price, max_price, currency, condition)
                items = annotate_comparison(items, query)
                items = sort_items(items, sort_option)
                repository.save_products(search_record_id, session["user_id"], items)
                repository.complete_search(search_record_id, len(items))
                repository.log_event(session["user_id"], "search_query", session["role"], session["username"], {"query": query, "requested_source": requested_source, "actual_source": source, "region": region, "currency": currency, "result_count": len(items)})
                summary = calculate_summary(items)
                repository.log_event(session["user_id"], "records_collected", session["role"], session["username"], {"query": query, "record_count": len(items), "actual_source": source})
                repository.log_event(session["user_id"], "comparison_generated", session["role"], session["username"], {"query": query, "record_count": len(items), "basis": "total_price"})
                insight, insight_label = (build_price_insight(items, summary), "Rule-based insight") if items else (None, None)
                repository.log_event(session["user_id"], "insight_generated", session["role"], session["username"], {"query": query, "mode": insight_label})
            except (AISearchError, RuntimeError, requests.RequestException, ValueError) as exc:
                repository.complete_search(search_record_id, 0, "failed")
                error = str(exc)
            except Exception as exc:
                app.logger.warning("Search source failed (%s)", exc.__class__.__name__)
                repository.complete_search(search_record_id, 0, "failed")
                error = f"{SOURCES.get(requested_source, 'Search')} is temporarily unavailable. Please try another source."
    demo_rows = load_demo_data()
    max_total = max((item.get("total_price") or 0 for item in items), default=0)
    return render_template("price_search.html", query=query, source=source, sources=SOURCES, items=items, summary=summary, insight=insight, insight_label=insight_label if query and items else None, error=error, notice=notice, sort_option=sort_option, item_type=item_type, search_record_id=search_record_id, recent_searches=repository.list_searches(session["user_id"], limit=5), platform_filter=platform_filter, category=category, region=region, currency=currency, condition=condition, min_price=params.get("min_price", ""), max_price=params.get("max_price", ""), platforms=sorted({row.get("platform") for row in demo_rows}), categories=sorted({row.get("category") for row in demo_rows}), max_total=max_total)


@app.route("/search-mongo", methods=["GET", "POST"])
@login_required
def mongo_search():
    params = request.form if request.method == "POST" else {}
    if request.method == "GET" and (request.args.get("q") or request.args.get("action")):
        repository.log_event(session["user_id"], "state_changing_get_blocked", session.get("role"), session.get("username"), {"endpoint": "mongo_search"})
    query = clean_search_query(params.get("q", ""))
    platform = params.get("platform", "").strip()
    # Marketplace currently exposes only a placeholder category. Keep the
    # backend parameter compatible, but do not present it as a real filter.
    category = ""
    min_price, max_price = optional_float(params.get("min_price")), optional_float(params.get("max_price"))
    sort_option = params.get("sort", "normalized_price_asc")
    include_related_variants = params.get("include_related_variants") in {"1", "true", "on", "yes"}
    action = params.get("action", "search" if request.method == "POST" else "")
    items, summary, insight, insight_label, error, notice, search_record_id = [], None, None, None, None, None, None
    facets = repository.market_facets()
    if query and is_blocked_search_query(query):
        record_blocked_search_attempt(session["user_id"], query, "unsafe keyword pattern")
        error = "The search term contains unsupported characters. Please enter a product name or model."
        query = ""
    if query and action == "search":
        repository.log_event(session["user_id"], "search_started", session["role"], session["username"], {"query": query, "source": "market_records", "platform": platform, "category": category})
        try:
            rows = _apply_search_relevance(repository.search_market_records(query, platform, category, min_price, max_price, sort_option), query)
            actual_source, data_mode = "mongodb_market_records", "mongodb"
            if not rows and repository.mode != "mongodb" and _should_show_demo_fallback():
                rows = _apply_search_relevance(repository.store_fallback_market_records(normalize_price_items(demo_search_items(query, limit=20, platform="demo_all"))), query)
                actual_source, data_mode = "memory_fallback", "synthetic"
                notice = "MongoDB is unavailable; clearly labeled in-memory test data is shown."
            items = annotate_comparison(normalize_price_items(rows), query)
            items = sort_items(items, {"normalized_price_asc": "price_asc", "normalized_price_desc": "price_desc"}.get(sort_option, sort_option))
            search_record_id = repository.create_search(session["user_id"], query, actual_source, role=session["role"], data_mode=data_mode)
            repository.attach_market_results(search_record_id, [item["_id"] for item in items if item.get("_id")], {"platform": platform, "category": category, "min_price": min_price, "max_price": max_price, "sort": sort_option})
            completion_status = "completed" if items else "no_results"
            if notice and items:
                completion_status = "partial_success"
            elif notice and not items:
                completion_status = "source_unavailable"
            repository.complete_search(search_record_id, len(items), completion_status)
            summary = calculate_summary(items)
            insight, insight_label = (build_price_insight(items, summary), "Rule-based insight") if items else (None, None)
            repository.log_event(session["user_id"], "records_collected", session["role"], session["username"], {"query": query, "record_count": len(items), "actual_source": actual_source})
            repository.log_event(session["user_id"], "comparison_generated", session["role"], session["username"], {"query": query, "record_count": len(items), "basis": "normalized_price"})
            repository.log_event(session["user_id"], "insight_generated", session["role"], session["username"], {"query": query, "mode": insight_label})
        except Exception as exc:
            app.logger.warning("MongoDB market search failed (%s)", exc.__class__.__name__)
            repository.log_event(session["user_id"], "api_call_failed", session["role"], session["username"], {"query": query, "source": "market_records", "error_type": exc.__class__.__name__})
            error = "Market records could not be queried. Please try again."
    max_total = max((item.get("total_price") or 0 for item in items), default=0)
    return render_template("mongo_search.html", query=query, source="mongo", sources={"mongo": "MongoDB market_records"}, items=items, summary=summary, insight=insight, insight_label=insight_label, error=error, notice=notice, sort_option=sort_option, item_type="all", search_record_id=search_record_id, recent_searches=repository.list_searches(session["user_id"], limit=5), platform_filter=platform, category=category, region="", currency="", condition="", min_price=params.get("min_price", ""), max_price=params.get("max_price", ""), platforms=facets["platforms"], categories=facets["categories"], max_total=max_total, mongo_search=True)


@app.route("/search", methods=["GET", "POST"])
@app.route("/search/results/<search_run_id>", methods=["GET", "POST"], endpoint="search_results")
@role_required("consumer", "retailer", "researcher")
def search(search_run_id=None):
    if request.method == "GET" and not search_run_id and (request.args.get("action") == "search" or request.args.get("refresh_live")):
        repository.log_event(session["user_id"], "state_changing_get_blocked", session.get("role"), session.get("username"), {"endpoint": "search"})
        return redirect(url_for("search"))
    params = (request.form if request.form else request.args) if request.method == "POST" else request.args
    query = clean_search_query(params.get("q", ""))
    requested_query = query
    search_scope = _normalize_search_scope(params.get("search_scope", params.get("data_source", "both")))
    platform = params.get("platform", "").strip()
    category = ""
    min_price, max_price = optional_float(params.get("min_price")), optional_float(params.get("max_price"))
    sort_option = params.get("sort", "normalized_price_asc")
    include_related_variants = params.get("match_mode", "exact") == "related" or params.get("include_related_variants") in {"1", "true", "on", "yes"}
    action = params.get("action", "search" if request.method == "POST" and not search_run_id else "")
    use_stored_walmart = params.get("use_stored_walmart") in {"1", "true", "on", "yes"}
    refresh_live = params.get("refresh_live") in {"1", "true", "on", "yes"}
    storage_filter = str(params.get("storage", "") or "").strip().upper()
    condition_filter = str(params.get("condition", "") or "").strip()
    items, summary, insight, error, notice, search_record_id = [], None, None, None, None, None
    rejected_records = []
    public_search_run_id, stale_ai_insight, ai_insight, record = None, False, None, None
    search_diagnostics = _new_search_diagnostics(query, search_scope)
    search_diagnostics["request_args"] = {
        "q": params.get("q", ""), "search_scope": params.get("search_scope"),
        "data_source": params.get("data_source"), "platform": params.get("platform"),
        "platform_filter": params.get("platform_filter"), "action": params.get("action"),
        "category": params.get("category"), "min_price": params.get("min_price"),
        "max_price": params.get("max_price"), "storage": storage_filter,
        "condition": condition_filter, "use_stored_walmart": params.get("use_stored_walmart"),
        "refresh_live": params.get("refresh_live"),
    }
    facets = repository.market_facets()
    restore_id = str(search_run_id or params.get("search_record_id", "")).strip()
    if restore_id and not (query and action == "search"):
        record = repository.get_search(restore_id, session["user_id"])
        if record:
            if record.get("status") == "running":
                return render_template("search_pending.html", query=record.get("keyword", ""), search_run_id=record.get("search_run_id") or record.get("_id")), 202
            visible_query_changed = bool(requested_query and requested_query != record.get("keyword", ""))
            query, search_scope, search_record_id = record["keyword"], _normalize_search_scope(record.get("selected_source", "both")), str(record["_id"])
            public_search_run_id = str(record.get("search_run_id") or search_record_id)
            stored_filters = record.get("filters") or {}
            platform = stored_filters.get("platform", platform)
            min_price, max_price = stored_filters.get("min_price"), stored_filters.get("max_price")
            sort_option = stored_filters.get("sort", sort_option)
            storage_filter = stored_filters.get("storage", storage_filter)
            condition_filter = stored_filters.get("condition", condition_filter)
            include_related_variants = stored_filters.get("match_mode") == "related" or bool(stored_filters.get("include_related_variants"))
            search_diagnostics = record.get("source_diagnostics") or _new_search_diagnostics(query, search_scope)
            search_diagnostics.setdefault("search_run_mode", "existing_search_run")
            search_diagnostics["page_load_mode"] = "reused_search_run"
            search_diagnostics["result_id_source"] = "persisted_search_run"
            rejected_records = [dict(row) for row in record.get("excluded_search_records") or []]
            notice = record.get("notice")
            if record.get("result_tokens"):
                items = annotate_comparison(normalize_price_items(resolve_result_tokens(record["result_tokens"])), query)
            elif record.get("market_record_ids"):
                items = annotate_comparison(normalize_price_items(repository.get_market_records(record["market_record_ids"])), query)
            else:
                items = annotate_comparison(normalize_price_items(repository.get_products(restore_id)), query)
            summary = calculate_summary(items) if items else None
            ai_insight = repository.get_ai_insight_for_search(search_record_id, session["user_id"])
            valid_insight_ids = {search_record_id, public_search_run_id}
            insight_matches_run = ai_insight and str(ai_insight.get("search_run_id")) in valid_insight_ids
            insight_matches_query = ai_insight and clean_search_query(ai_insight.get("query", "")) == query
            stored_filter_signature = hashlib.sha256(json.dumps({"selected_facets": record.get("selected_facets") or {}, "active_filters": record.get("active_filters") or {}}, sort_keys=True, default=str, separators=(",", ":")).encode("utf-8")).hexdigest()
            if insight_matches_run and insight_matches_query and ai_insight.get("filter_signature") == stored_filter_signature and not visible_query_changed:
                insight = ai_insight.get("summary")
            elif ai_insight:
                # A stored insight is only valid for the exact Search Run and query that produced it.
                stale_ai_insight, ai_insight = True, None
    if query and is_blocked_search_query(query):
        record_blocked_search_attempt(session["user_id"], query, "unsafe keyword pattern")
        error = "The search term contains unsupported characters. Please enter a product name or model."
        query = ""
    if action == "search" and not query and not error:
        error = "Enter a product keyword or model before searching."
    clear_supported_category_query = bool(query) and detect_product_category(query, query).category_key in {"smartphones", "laptops", "headphones", "cameras"} and classify_query_intent(query) == "broad"
    legacy_memory_fallback = (
        (
            request.method == "POST"
            and str(params.get("data_source", "")).strip().lower() in {"mongodb", "market_records"}
            and not params.get("search_scope")
            and repository.mode != "mongodb"
            and _should_show_demo_fallback()
        )
        or bool(record and record.get("selected_source") == "memory_fallback")
    )
    if query and action == "search" and legacy_memory_fallback:
        repository.log_event(session["user_id"], "search_started", session["role"], session["username"], {"query": query, "search_scope": "memory_fallback", "offline_fallback": True})
        parsed_query = parse_canonical_product_model(query)
        canonical_query = " ".join(str(value) for value in (parsed_query.get("brand"), parsed_query.get("product_family"), parsed_query.get("generation"), None if parsed_query.get("edition") == "Base" else parsed_query.get("edition"), parsed_query.get("storage")) if value)
        fallback_rows = repository.store_fallback_market_records(normalize_price_items(demo_search_items(query, limit=20, platform="demo_all")))
        # The legacy MongoDB selector is an offline demonstration path.  It
        # intentionally preserves its deterministic fallback catalogue even
        # when a free-text query has no exact marketplace-model match.
        items = annotate_comparison(normalize_price_items(fallback_rows), query)
        items = sort_items(items, {"normalized_price_asc": "price_asc", "normalized_price_desc": "price_desc"}.get(sort_option, sort_option))
        search_record_id = repository.create_search(
            session["user_id"], query, "memory_fallback", role=session["role"], data_mode="synthetic",
            raw_query=query, canonical_query=canonical_query or query,
            parsed_model={key: parsed_query.get(key) for key in ("brand", "product_family", "generation", "edition", "storage", "carrier", "product_type")},
            platform_scope="memory_fallback",
        )
        stored = repository.save_external_results(search_record_id, session["user_id"], items)
        for row in stored:
            row["_selection_token"] = f"external:{row['_id']}"
        items = annotate_comparison(stored, query)
        result_tokens = [item["_selection_token"] for item in items]
        repository.attach_result_tokens(search_record_id, result_tokens, {"search_scope": "memory_fallback", "platform": platform, "storage": storage_filter, "condition": condition_filter, "min_price": min_price, "max_price": max_price, "sort": sort_option, "include_related_variants": include_related_variants})
        search_diagnostics.update({"offline_fallback": True, "source_status": {"ebay": "not_requested", "walmart": "not_requested"}, "final_display_count": len(items)})
        repository.update_search_run(search_record_id, {"filters": {"platform": platform, "storage": storage_filter, "condition": condition_filter, "min_price": min_price, "max_price": max_price, "sort": sort_option, "match_mode": "related" if include_related_variants else "exact"}, "result_ids": result_tokens, "source_diagnostics": search_diagnostics, "notice": "MongoDB is unavailable; clearly labeled in-memory test data is shown."})
        repository.complete_search(search_record_id, len(items), "partial_success" if items else "no_results")
        repository.log_event(session["user_id"], "records_collected", session["role"], session["username"], {"query": query, "record_count": len(items), "search_scope": "memory_fallback", "offline_fallback": True})
        notice = "MongoDB is unavailable; clearly labeled in-memory test data is shown."
        summary = calculate_summary(items) if items else None
    elif query and action == "search":
        for key in ("selected_result_ids", "temporary_comparison_state", "active_search_run_id", "active_analysis_id", "active_ai_insight", "temporary_filters"):
            session.pop(key, None)
        search_diagnostics = _new_search_diagnostics(query, search_scope)
        search_diagnostics["force_refresh"] = bool(refresh_live)
        if search_scope not in {"ebay", "walmart", "both"}:
            error = "Choose a valid search scope."
        else:
            parsed_query = parse_canonical_product_model(query)
            canonical_query = " ".join(str(value) for value in (parsed_query.get("brand"), parsed_query.get("product_family"), parsed_query.get("generation"), None if parsed_query.get("edition") == "Base" else parsed_query.get("edition"), parsed_query.get("storage")) if value)
            search_record_id = repository.create_search(session["user_id"], query, search_scope, role=session["role"], data_mode="mixed" if search_scope == "both" else search_scope, raw_query=query, canonical_query=canonical_query or query, parsed_model={key: parsed_query.get(key) for key in ("brand", "product_family", "generation", "edition", "storage", "carrier", "product_type")}, platform_scope=search_scope)
            search_diagnostics["search_run_mode"] = "new_collection"
            search_diagnostics["page_load_mode"] = "collection_request"
            repository.log_event(session["user_id"], "search_started", session["role"], session["username"], {"query": query, "search_scope": search_scope, "platform": platform})
            collected = []
            if search_scope in {"ebay", "both"}:
                try:
                    raw_ebay_rows, ebay_meta = search_ebay_cached(query, category=category, min_price=min_price, max_price=max_price, sort=sort_option, limit=20, refresh_live=refresh_live)
                    search_diagnostics["ebay_cache_hit"] = ebay_meta["cache_status"] == "hit"
                    search_diagnostics["api_response_status"]["ebay"] = ebay_meta.get("response_status_code") or ("cache" if search_diagnostics["ebay_cache_hit"] else 200)
                    search_diagnostics["source_status"]["ebay"] = "cached" if search_diagnostics["ebay_cache_hit"] else "live"
                    search_diagnostics["raw_ebay_count"] = len(raw_ebay_rows)
                    if not raw_ebay_rows:
                        search_diagnostics["source_status"]["ebay"] = "no results"
                    search_diagnostics["ebay_status"] = search_diagnostics["source_status"]["ebay"]
                    if not search_diagnostics["first_records"]:
                        search_diagnostics["first_records"] = _diagnostic_sample_n(raw_ebay_rows)
                    normalized_ebay_rows = normalize_price_items(raw_ebay_rows)
                    for row in normalized_ebay_rows:
                        row["_include_related_variants"] = include_related_variants
                    search_diagnostics["normalized_ebay_count"] += len(normalized_ebay_rows)
                    search_diagnostics["first_normalized_ebay_records"] = _diagnostic_sample_n(normalized_ebay_rows)
                    ebay_rows, ebay_rejected = (normalized_ebay_rows, []) if clear_supported_category_query else _filter_relevant_search_items_with_rejections(normalized_ebay_rows, query)
                    rejected_records.extend(ebay_rejected)
                    classified_ebay = ebay_rows + ebay_rejected
                    search_diagnostics["product_type_eligible_ebay_count"] = sum((row.get("parsed_model") or {}).get("product_type") != "accessory" for row in classified_ebay)
                    search_diagnostics["exact_ebay_count"] = sum(row.get("match_type") == "exact_match" for row in classified_ebay)
                    search_diagnostics["related_ebay_count"] = sum(row.get("match_type") == "related_variant" for row in classified_ebay)
                    search_diagnostics["accessory_ebay_count"] = sum(row.get("match_type") == "accessory_risk" for row in classified_ebay)
                    search_diagnostics["low_confidence_ebay_count"] = sum(row.get("match_type") == "low_confidence" for row in classified_ebay)
                    ebay_rows = apply_result_filters(ebay_rows, category=category)
                    search_diagnostics["category_filtered_count"] += len(ebay_rows)
                    ebay_rows = apply_result_filters(ebay_rows, min_price=min_price, max_price=max_price)
                    search_diagnostics["price_filtered_count"] += len(ebay_rows)
                    if raw_ebay_rows and not ebay_rows:
                        search_diagnostics["source_status"]["ebay"] = "filtered out"
                    search_diagnostics["ebay_status"] = search_diagnostics["source_status"]["ebay"]
                    stored = repository.save_external_results(search_record_id, session["user_id"], ebay_rows)
                    for row in stored:
                        row["_selection_token"] = f"external:{row['_id']}"
                        collected.append(row)
                    repository.log_event(session["user_id"], "ebay_api_call", session["role"], session["username"], {"query": query, "record_count": len(stored), "status": "success", "provider": "ebay", "cache_status": ebay_meta.get("cache_status")})
                except Exception as exc:
                    app.logger.warning("eBay Browse API failed (%s)", exc.__class__.__name__)
                    search_diagnostics["source_status"]["ebay"] = "unavailable"
                    search_diagnostics["api_response_status"]["ebay"] = "error"
                    repository.log_event(session["user_id"], "api_call_failed", session["role"], session["username"], {"query": query, "source": "eBay API", "error_type": exc.__class__.__name__})
                    notice = "eBay is temporarily unavailable. " + ("Walmart results are still shown." if collected else "Try Walmart or retry shortly.")
            if search_scope in {"walmart", "both"}:
                walmart_rows = []
                stored_walmart_rows = []
                try:
                    walmart_request_diagnostics = {}
                    serpapi_meta = None
                    if _serpapi_walmart_enabled():
                        raw_walmart_rows, serpapi_meta = search_serpapi_cached(query, "walmart", category=category, min_price=min_price, max_price=max_price, sort=sort_option, limit=12, refresh_live=refresh_live)
                        walmart_parser_diagnostics = serpapi_meta.get("parser_diagnostics") or {}
                        provider_raw_walmart_count = int(walmart_parser_diagnostics.get("raw_result_count", len(raw_walmart_rows)) or 0)
                        search_diagnostics["walmart_cache_hit"] = serpapi_meta["cache_status"] == "hit"
                        search_diagnostics["walmart_provider_response_mode"] = "cached_provider_response" if search_diagnostics["walmart_cache_hit"] else "new_collection"
                        search_diagnostics["api_response_status"]["walmart"] = serpapi_meta.get("response_status_code") or ("cache" if search_diagnostics["walmart_cache_hit"] else 200)
                        search_diagnostics["serpapi_status"] = "cache_hit" if search_diagnostics["walmart_cache_hit"] else "live"
                        walmart_request_diagnostics = {
                            "provider": "serpapi",
                            "app_query": query,
                            "normalized_query": normalize_serpapi_query(query),
                            "serpapi_params": serpapi_meta.get("serpapi_params"),
                            "cache_key": serpapi_meta.get("cache_key"),
                            "cache_status": serpapi_meta.get("cache_status"),
                            "record_field": walmart_parser_diagnostics.get("record_field") or "organic_results",
                            "page_type": "product_results" if provider_raw_walmart_count else "no_results",
                            "raw_record_count": provider_raw_walmart_count,
                            "parsed_record_count": int(walmart_parser_diagnostics.get("parsed_result_count", len(raw_walmart_rows)) or 0),
                            "normalized_record_count": int(walmart_parser_diagnostics.get("normalized_result_count", len(raw_walmart_rows)) or 0),
                            "excluded_record_count": int(walmart_parser_diagnostics.get("excluded_result_count", 0) or 0),
                            "final_matched_count": len(raw_walmart_rows),
                            "rejection_reasons": walmart_parser_diagnostics.get("rejection_counts") or {},
                        }
                        search_diagnostics.setdefault("serpapi", {})["walmart"] = serpapi_meta
                        search_diagnostics["source_status"]["walmart"] = "cached" if serpapi_meta["cache_status"] == "hit" else "live"
                        provider_suggestion = str(serpapi_meta.get("spelling_suggestion") or "").strip()
                        if provider_suggestion and clean_search_query(provider_suggestion).casefold() != query.casefold():
                            search_diagnostics["provider_spelling_suggestion"] = provider_suggestion
                            search_diagnostics["corrected_walmart_raw_count"] = provider_raw_walmart_count
                            # Corrected-query rows belong to a new user-accepted
                            # Search Run and are never merged into this run.
                            raw_walmart_rows = []
                            search_diagnostics["source_status"]["walmart"] = "query_correction_suggested"
                    elif _legacy_walmart_scraper_enabled():
                        raw_walmart_rows = search_walmart_items(query, limit=12, diagnostics=walmart_request_diagnostics)
                        provider_raw_walmart_count = len(raw_walmart_rows)
                    else:
                        raw_walmart_rows = []
                        provider_raw_walmart_count = 0
                        walmart_request_diagnostics = {
                            "provider": "serpapi",
                            "app_query": query,
                            "normalized_query": normalize_serpapi_query(query),
                            "serpapi_params": None,
                            "page_type": "unavailable",
                            "error": "authentication_error",
                            "safe_message": "Marketplace provider credentials are not configured.",
                        }
                        search_diagnostics["serpapi_status"] = "unavailable"
                        search_diagnostics["serpapi_error"] = walmart_request_diagnostics["error"]
                        search_diagnostics["walmart_safe_message"] = walmart_request_diagnostics["safe_message"]
                    search_diagnostics["walmart_request"] = walmart_request_diagnostics
                    app.logger.info(
                        "Walmart route diagnostics query=%r normalized_query=%r search_scope=%s data_source=%r platform_filter=%r request_url=%r status=%s final_url=%r preview=%r page_type=%s raw_count=%s normalized_count=%s final_matched_count=%s rejection_reasons=%s",
                        query,
                        walmart_request_diagnostics.get("normalized_query"),
                        search_scope,
                        params.get("data_source"),
                        platform,
                        walmart_request_diagnostics.get("request_url"),
                        walmart_request_diagnostics.get("response_status_code"),
                        walmart_request_diagnostics.get("response_final_url"),
                        walmart_request_diagnostics.get("response_text_preview"),
                        walmart_request_diagnostics.get("page_type"),
                        walmart_request_diagnostics.get("raw_record_count"),
                        walmart_request_diagnostics.get("normalized_record_count"),
                        walmart_request_diagnostics.get("final_matched_count"),
                        walmart_request_diagnostics.get("rejection_reasons"),
                    )
                    search_diagnostics["raw_walmart_count"] = provider_raw_walmart_count
                    search_diagnostics["parsed_walmart_count"] = int(walmart_request_diagnostics.get("parsed_record_count", len(raw_walmart_rows)) or 0)
                    search_diagnostics["walmart_parser_rejection_counts"] = dict(walmart_request_diagnostics.get("rejection_reasons") or {})
                    search_diagnostics["walmart_response_field"] = walmart_request_diagnostics.get("record_field")
                    search_diagnostics["walmart_response_type"] = walmart_request_diagnostics.get("page_type") or ("product_results" if provider_raw_walmart_count else "no_results")
                    search_diagnostics["raw_walmart_status"] = search_diagnostics["walmart_response_type"]
                    if not _serpapi_walmart_enabled() and _legacy_walmart_scraper_enabled():
                        search_diagnostics["source_status"]["walmart"] = "matched" if raw_walmart_rows else "no results"
                    elif not _serpapi_walmart_enabled():
                        search_diagnostics["source_status"]["walmart"] = "unavailable"
                    elif search_diagnostics["source_status"].get("walmart") == "query_correction_suggested":
                        pass
                    elif not provider_raw_walmart_count:
                        search_diagnostics["source_status"]["walmart"] = "no results"
                    search_diagnostics["walmart_status"] = search_diagnostics["source_status"]["walmart"]
                    search_diagnostics["first_raw_walmart_records"] = _diagnostic_sample_n(raw_walmart_rows)
                    normalized_walmart_rows = normalize_price_items(raw_walmart_rows)
                    for row in normalized_walmart_rows:
                        row["_include_related_variants"] = include_related_variants
                    search_diagnostics["normalized_walmart_count"] += len(normalized_walmart_rows)
                    search_diagnostics["walmart_valid_price_count"] = sum(
                        1 for row in normalized_walmart_rows if row.get("analytics_eligible")
                    )
                    search_diagnostics["first_normalized_walmart_records"] = _diagnostic_sample_n(normalized_walmart_rows)
                    walmart_rows, walmart_rejected = (normalized_walmart_rows, []) if clear_supported_category_query else _filter_relevant_search_items_with_rejections(normalized_walmart_rows, query)
                    walmart_rejected = [_prepare_walmart_rejection(row) for row in walmart_rejected]
                    rejected_records.extend(walmart_rejected)
                    classified_walmart = walmart_rows + walmart_rejected
                    search_diagnostics["product_type_eligible_walmart_count"] = sum((row.get("parsed_model") or {}).get("product_type") != "accessory" for row in classified_walmart)
                    search_diagnostics["exact_walmart_count"] = sum(row.get("match_type") == "exact_match" for row in classified_walmart)
                    search_diagnostics["related_walmart_count"] = sum(row.get("match_type") == "related_variant" for row in classified_walmart)
                    search_diagnostics["accessory_walmart_count"] = sum(row.get("match_type") == "accessory_risk" for row in classified_walmart)
                    search_diagnostics["low_confidence_walmart_count"] = sum(row.get("match_type") == "low_confidence" for row in classified_walmart)
                    walmart_rows = apply_result_filters(walmart_rows, category=category)
                    search_diagnostics["category_filtered_count"] += len(walmart_rows)
                    walmart_rows = apply_result_filters(walmart_rows, min_price=min_price, max_price=max_price)
                    search_diagnostics["price_filtered_count"] += len(walmart_rows)
                    if provider_raw_walmart_count and not walmart_rows and search_diagnostics["source_status"].get("walmart") != "query_correction_suggested":
                        search_diagnostics["source_status"]["walmart"] = "filtered out"
                    search_diagnostics["walmart_status"] = search_diagnostics["source_status"]["walmart"]
                    if not walmart_rows and not provider_raw_walmart_count and search_diagnostics["source_status"].get("walmart") != "query_correction_suggested":
                        search_diagnostics["walmart_response_type"] = "no_results"
                        search_diagnostics["raw_walmart_status"] = "no_results"
                except Exception as exc:
                    app.logger.warning("Walmart source failed (%s)", exc.__class__.__name__)
                    if isinstance(exc, SerpApiError):
                        provider_error_code = getattr(exc, "code", str(exc))
                        search_diagnostics["source_status"]["walmart"] = "unavailable"
                        search_diagnostics["walmart_status"] = "unavailable"
                        search_diagnostics["walmart_response_type"] = "serpapi_error"
                        search_diagnostics["raw_walmart_status"] = "serpapi_error"
                        search_diagnostics["serpapi_status"] = "error"
                        search_diagnostics["serpapi_error"] = provider_error_code
                        search_diagnostics["walmart_safe_message"] = getattr(
                            exc,
                            "safe_message",
                            "Marketplace provider request failed.",
                        )
                    elif isinstance(exc, CollectorError) and "bot check" in str(exc).lower():
                        search_diagnostics["source_status"]["walmart"] = "unavailable_bot_check"
                        search_diagnostics["walmart_status"] = "unavailable_bot_check"
                        search_diagnostics["walmart_response_type"] = "bot_check"
                        search_diagnostics["raw_walmart_status"] = "bot_check"
                        search_diagnostics["walmart_status"] = "unavailable_bot_check"
                    else:
                        search_diagnostics["source_status"]["walmart"] = "unavailable"
                        search_diagnostics["walmart_status"] = "unavailable"
                        search_diagnostics["walmart_response_type"] = "unavailable"
                        search_diagnostics["raw_walmart_status"] = "unavailable"
                    search_diagnostics["walmart_error"] = (
                        getattr(exc, "code", str(exc))
                        if isinstance(exc, SerpApiError)
                        else str(exc)
                    )
                    search_diagnostics["api_response_status"]["walmart"] = "error"
                    repository.log_event(session["user_id"], "api_call_failed", session["role"], session["username"], {"query": query, "source": "Walmart", "error_type": exc.__class__.__name__})
                stored_walmart_count = _has_stored_walmart_evidence(query)
                search_diagnostics["stored_walmart_count"] = stored_walmart_count
                if use_stored_walmart and search_diagnostics.get("walmart_status") == "unavailable_bot_check" and stored_walmart_count:
                    stored_walmart_rows = _stored_walmart_evidence(query)
                    stored_walmart_rows = annotate_comparison(normalize_price_items(stored_walmart_rows), query)
                    stored_walmart_rows = apply_result_filters(stored_walmart_rows, category=category)
                    stored_walmart_rows = apply_result_filters(stored_walmart_rows, min_price=min_price, max_price=max_price)
                    for row in stored_walmart_rows:
                        row["_include_related_variants"] = include_related_variants
                    search_diagnostics["stored_walmart_loaded"] = len(stored_walmart_rows)
                    search_diagnostics["stored_walmart_source_type"] = "stored_walmart_evidence"
                    search_diagnostics["source_status"]["walmart"] = "matched_stored_evidence"
                    search_diagnostics["walmart_status"] = "matched_stored_evidence"
                    search_diagnostics["walmart_response_type"] = "stored_evidence"
                    search_diagnostics["raw_walmart_status"] = "stored_evidence"
                    for row in stored_walmart_rows:
                        row["source_type"] = "stored_walmart_evidence"
                        row["source_label"] = "Stored Walmart evidence"
                    walmart_rows = stored_walmart_rows
                # Walmart availability and result suitability are presented by
                # the canonical source-status message below. Do not compose a
                # second, contradictory notice from an empty retained set.
                stored = repository.save_external_results(search_record_id, session["user_id"], walmart_rows)
                for row in stored:
                    row["_selection_token"] = f"external:{row['_id']}"
                    collected.append(row)
                repository.log_event(session["user_id"], "walmart_source_call", session["role"], session["username"], {"query": query, "record_count": len(stored), "status": search_diagnostics["walmart_response_type"], "provider": "serpapi" if serpapi_meta else "walmart", "cache_status": serpapi_meta.get("cache_status") if serpapi_meta else None})
            items = annotate_comparison(collected, query)
            before_deduplication = len(items)
            items = _deduplicate_search_items(items, rejected_records)
            search_diagnostics["duplicate_count"] = before_deduplication - len(items)
            items = sort_items(items, {"normalized_price_asc": "price_asc", "normalized_price_desc": "price_desc"}.get(sort_option, sort_option))
            items = sorted(items, key=_match_rank)
            _record_rejections(search_diagnostics, rejected_records)
            if items:
                search_diagnostics["first_records"] = _diagnostic_sample_n(items)
            search_diagnostics["final_ebay_count"] = sum(
                1 for row in items
                if canonical_marketplace_platform(row.get("platform"), row.get("source_type")) == "eBay"
            )
            search_diagnostics["final_walmart_count"] = sum(
                1 for row in items
                if canonical_marketplace_platform(row.get("platform"), row.get("source_type")) == "Walmart"
            )
            search_diagnostics["comparable_walmart_count"] = search_diagnostics["final_walmart_count"]
            app.logger.info(
                "Search routing diagnostics query=%r search_scope=%s platform=%r data_source=%r platform_filter=%r final_ebay=%s final_walmart=%s walmart_status=%s",
                query,
                search_scope,
                platform,
                params.get("data_source"),
                params.get("platform_filter"),
                search_diagnostics["final_ebay_count"],
                search_diagnostics["final_walmart_count"],
                search_diagnostics.get("walmart_status"),
            )
            tokens = [item["_selection_token"] for item in items]
            search_diagnostics["retained_result_ids"] = list(tokens)
            search_diagnostics["result_id_source"] = "newly_persisted"
            repository.attach_result_tokens(search_record_id, tokens, {"search_scope": search_scope, "platform": platform, "category": category, "min_price": min_price, "max_price": max_price, "sort": sort_option, "include_related_variants": include_related_variants})
            requested_statuses = [search_diagnostics["source_status"].get(source) for source in ("ebay", "walmart") if search_scope in {source, "both"}]
            has_source_failure = any(str(status).startswith("unavailable") for status in requested_statuses)
            completion_status = "partial_success" if items and has_source_failure else ("completed" if items else ("source_unavailable" if has_source_failure else "no_results"))
            repository.complete_search(search_record_id, len(items), completion_status)
            session["last_search_record_id"] = str(search_record_id)
            session["active_search_run_id"] = str(search_record_id)
            summary = calculate_summary(items)
            insight = build_price_insight(items, summary) if items else None
            repository.log_event(session["user_id"], "records_collected", session["role"], session["username"], {"query": query, "record_count": len(items), "search_scope": search_scope})
            if not items:
                if search_scope == "ebay":
                    notice = notice or ("eBay source is currently unavailable." if search_diagnostics["ebay_status"] == "unavailable" else "No matching listings found on eBay.")
                elif search_scope == "walmart":
                    notice = notice or ("Walmart source is currently unavailable." if str(search_diagnostics["walmart_status"]).startswith("unavailable") else "No matching listings found on Walmart.")
                else:
                    ebay_unavailable = str(search_diagnostics.get("ebay_status") or "").startswith("unavailable")
                    walmart_unavailable = str(search_diagnostics.get("walmart_status") or "").startswith("unavailable")
                    if ebay_unavailable and walmart_unavailable:
                        notice = "No matching records found because eBay and Walmart are currently unavailable. Try again later."
                    elif walmart_unavailable:
                        notice = "No matching records found. Walmart source is currently unavailable, and eBay returned no matching listings."
                    elif ebay_unavailable:
                        notice = "No matching records found. eBay is currently unavailable, and Walmart returned no matching listings."
                    elif search_diagnostics["raw_ebay_count"] == 0 and search_diagnostics["raw_walmart_count"] == 0:
                        notice = "No matching records found from eBay or Walmart. Try a broader product keyword."
                    elif search_diagnostics["raw_ebay_count"] > 0 and search_diagnostics["raw_walmart_count"] == 0:
                        notice = "Records were collected from eBay but none matched this product. Walmart returned no matching listings."
                    elif search_diagnostics["raw_ebay_count"] == 0 and search_diagnostics["raw_walmart_count"] > 0:
                        notice = "Records were collected from Walmart but none matched this product. eBay returned no matching listings."
                    else:
                        notice = "Some source records were returned but excluded by product-match filters."
    # Category normalization is deliberately local to the persisted Search Run:
    # applying facets never triggers a marketplace or AI request.
    items = [enrich_listing_category(item, query) for item in items]
    if clear_supported_category_query:
        broad_category_key = detect_product_category(query, query).category_key
        for item in items:
            if item.get("category_key") == broad_category_key and item.get("product_role") == "complete_product" and item.get("total_price") is not None and item.get("price_type") not in {"installment", "range", "missing"}:
                item.update(analytics_eligible=True, exclusion_reason=None, match_type="broad_category_match")
    source_statuses = (record.get("source_statuses") or {}) if record and restore_id else _canonical_source_statuses(query, search_scope, search_diagnostics, items)
    source_messages = _source_status_messages(source_statuses) if search_record_id else []
    query_intent_status = classify_query_intent(query)
    query_category = detect_product_category(query, query).category_key
    requested_category_key = params.get("selected_category_key") if "selected_category_key" in params else None
    stored_category_key = record.get("selected_category_key") if record and restore_id and requested_category_key is None else None
    auto_category_key = query_category if query_category in {"smartphones", "laptops", "headphones", "cameras"} else None
    selected_category_key = requested_category_key or stored_category_key or auto_category_key
    scope = category_scope(items, selected_category_key, query_intent_status=query_intent_status)
    if legacy_memory_fallback:
        # This backwards-compatible synthetic catalogue intentionally contains
        # unrelated demo departments. It is not a marketplace result set and
        # retains its established General comparison behavior.
        scope = {
            "detected_category_keys": ["generic"] if items else [],
            "category_counts": {"generic": len(items)} if items else {},
            "selected_category_key": None,
            "category_mode": "single_category",
            "comparison_enabled": True,
        }
    result_categories = {
        str(item.get("_selection_token") or item.get("_id")): {
            "category_key": item.get("category_key"),
            "category_display": item.get("category_display"),
            "category_source": item.get("category_source"),
            "category_confidence": item.get("category_confidence"),
            "product_type": (item.get("attributes") or {}).get("product_type"),
            "product_type_display": (item.get("attributes") or {}).get("product_type_display"),
            "product_role": item.get("product_role"),
            "classification_confidence": item.get("classification_confidence"),
            "classification_source": item.get("classification_source"),
        }
        for item in items if item.get("_selection_token") or item.get("_id")
    }
    selected_category_key = scope["selected_category_key"]
    category_mode = scope["category_mode"]
    comparison_enabled = scope["comparison_enabled"]
    if category_mode == "multi_category":
        query_intent_status = "ambiguous" if query_intent_status != "broad" else "broad"
    category_key = selected_category_key or (scope["detected_category_keys"][0] if len(scope["detected_category_keys"]) == 1 else "generic")
    category_profile = CATEGORY_PROFILES.get(category_key, CATEGORY_PROFILES["generic"])
    all_facet_keys = {facet for profile in CATEGORY_PROFILES.values() for facet in profile["facets"]}
    requested_facets = {key: [value for value in params.getlist(key) if value] for key in all_facet_keys if hasattr(params, "getlist") and params.getlist(key)}
    stored_selected_facets = (record.get("selected_facets") or {}) if restore_id and record else {}
    has_filter_request = any(key in params for key in all_facet_keys | {"platform", "condition", "min_price", "max_price", "selected_category_key", "selected_product_type"})
    selected_facets = requested_facets if has_filter_request else stored_selected_facets
    selected_facets = {key: values if isinstance(values, list) else [values] for key, values in selected_facets.items() if key in category_profile["facets"]}
    selected_facets = {key: [canonical_facet_value(key, value) for value in values] for key, values in selected_facets.items()}
    category_options = [
        {"key": key, "display_name": CATEGORY_PROFILES.get(key, CATEGORY_PROFILES["generic"])["display_name"], "count": scope["category_counts"][key]}
        for key in scope["detected_category_keys"]
    ]
    role_excluded_items = [item for item in items if item.get("product_role") != "complete_product"]
    scoped_items = [item for item in items if item.get("product_role") == "complete_product" and (not selected_category_key or item.get("category_key") == selected_category_key)]
    requested_product_type = params.get("selected_product_type") if "selected_product_type" in params else None
    stored_product_type = record.get("selected_product_type") if record and restore_id and requested_product_type is None else None
    selected_product_type = requested_product_type or stored_product_type
    type_scope = product_type_scope(scoped_items, selected_product_type) if category_key == "generic" and comparison_enabled and not legacy_memory_fallback else {"detected_product_types": [], "product_type_counts": {}, "selected_product_type": None, "multi_product_type": False, "facet_coverage": {"eligible_record_count": len(scoped_items), "populated_record_count": 0, "coverage_ratio": 0.0, "distinct_value_count": 0}}
    selected_product_type = type_scope["selected_product_type"]
    if type_scope["multi_product_type"]:
        comparison_enabled = False
        category_mode = "multi_product_type"
    elif category_key == "generic" and not selected_product_type and len(type_scope["detected_product_types"]) == 1:
        selected_product_type = type_scope["detected_product_types"][0]
    if selected_product_type:
        scoped_items = [item for item in scoped_items if (item.get("attributes") or {}).get("product_type") == selected_product_type]
    scoped_tokens = {str(item.get("_selection_token") or item.get("_id")) for item in scoped_items}
    scope_excluded_items = [
        _prepare_scope_exclusion(item, selected_category_key, selected_product_type)
        for item in items
        if str(item.get("_selection_token") or item.get("_id")) not in scoped_tokens
    ]
    parsed_scope_model = parse_canonical_product_model(query) if query else {}
    exact_product_query = bool(parsed_scope_model.get("generation")) or bool(re.search(r"\b(?:wh|wf)-?\d{3,}|\beos\s+[a-z0-9]+", query, re.I))
    if legacy_memory_fallback and scoped_items:
        # The deterministic in-memory catalogue backs legacy/offline tests and
        # demonstrations rather than a real marketplace query. Keep that
        # explicitly labelled path usable without weakening ambiguity gates for
        # persisted or live source results.
        search_mode = "general_coherent"
        comparison_enabled = True
    elif category_mode in {"multi_category", "multi_product_type"}:
        search_mode = "multi_product_type"
    elif category_key in {"smartphones", "laptops", "headphones", "cameras"}:
        search_mode = "exact_product" if exact_product_query else "broad_single_category"
        if search_mode == "broad_single_category" and auto_category_key == selected_category_key and len(scoped_items) < 3:
            comparison_enabled = False
    elif category_key == "food_and_grocery" and selected_category_key and scoped_items:
        search_mode = "general_coherent"
        comparison_enabled = True
    elif selected_product_type and len(scoped_items) >= 3:
        search_mode = "general_coherent"
        comparison_enabled = True
    else:
        search_mode = "ambiguous_low_confidence"
        comparison_enabled = False
    product_type_options = [
        {"key": key, "display_name": PRODUCT_TYPE_LABELS.get(key, key.replace("_", " ").title()), "count": type_scope["product_type_counts"][key]}
        for key in type_scope["detected_product_types"]
    ]
    available_facets = generate_available_facets(scoped_items, category_key if comparison_enabled else "generic", selected_facets=selected_facets)
    available_facets.pop("product_type", None)
    if not comparison_enabled:
        available_facets.pop("brand", None)
        available_facets.pop("model", None)
    exact_model_scope = not include_related_variants and bool(parse_canonical_product_model(query).get("generation"))
    if exact_model_scope:
        available_facets.pop("model", None)
    selected_facets = {key: [value for value in values if value in {entry.get("filter_value", entry["value"]) for entry in available_facets.get(key, {}).get("values", [])}] for key, values in selected_facets.items()}
    selected_facets = {key: values for key, values in selected_facets.items() if values}
    technical_facets_limited = category_key in {"smartphones", "laptops", "headphones", "cameras"} and any(
        key not in available_facets for key in category_profile["facets"] if key not in {"condition", "brand"}
    )
    available_filter_platforms = sorted({item.get("platform") for item in scoped_items if item.get("platform")})
    items = apply_category_filters(scoped_items, selected_facets, min_price=min_price, max_price=max_price, platform=platform or None, condition=condition_filter or None)
    available_storage_variants = [entry["value"] for entry in available_facets.get("storage", {}).get("values", [])]
    available_conditions = [entry["value"] for entry in available_facets.get("condition", {}).get("values", [])]
    walmart_business_exclusions = [
        _excluded_walmart_snapshot(row)
        for row in rejected_records
        if canonical_marketplace_platform(row.get("platform"), row.get("source_type")) == "Walmart"
    ]
    excluded_items = scope_excluded_items + [
        row for row in items
        if not row.get("analytics_eligible")
        and str(row.get("_selection_token") or row.get("_id")) in scoped_tokens
    ] + walmart_business_exclusions
    items = [row for row in items if row.get("analytics_eligible")]
    filtered_result_ids = [str(row.get("_selection_token") or row.get("_id")) for row in items if row.get("_selection_token") or row.get("_id")]
    search_diagnostics["comparable_result_ids"] = filtered_result_ids
    search_diagnostics["filtered_result_ids"] = filtered_result_ids
    result_platforms = sorted({row.get("platform") for row in items if row.get("platform")})
    stored_walmart_visible = any(row.get("source_type") == "stored_walmart_evidence" for row in items)
    live_walmart_visible = any(row.get("source_type") == "live_walmart" for row in items)
    active_platform = platform if platform in {"", "All results", "eBay", "Walmart"} else ""
    if search_scope == "ebay":
        result_view = "eBay"
    elif search_scope == "walmart":
        result_view = "Walmart"
    else:
        result_view = active_platform or "All results"
    if search_scope != "both" and result_view in {"eBay", "Walmart"}:
        items = [
            row for row in items
            if canonical_marketplace_platform(row.get("platform"), row.get("source_type")) == result_view
        ]
    visible_result_ids = [str(row.get("_selection_token") or row.get("_id")) for row in items if row.get("_selection_token") or row.get("_id")]
    search_diagnostics["visible_result_ids"] = visible_result_ids
    if query and not items and not error:
        has_active_result_filters = bool(selected_facets or platform or condition_filter or min_price is not None or max_price is not None)
        if has_active_result_filters:
            notice = "No comparable listings match the current filters. Reset the filters or broaden the price range."
        elif source_messages and search_scope == "walmart":
            notice = None
        elif source_messages and search_scope == "both":
            ebay_unavailable = str(search_diagnostics.get("ebay_status") or "").startswith("unavailable")
            notice = "eBay is temporarily unavailable, and no comparable listings were available from the other requested source." if ebay_unavailable else None
        else:
            notice = empty_marketplace_search_notice(search_scope, search_diagnostics)
    search_diagnostics["final_display_count"] = len(items)
    search_diagnostics["displayed_walmart_count"] = sum(
        1 for row in items
        if canonical_marketplace_platform(row.get("platform"), row.get("source_type")) == "Walmart"
    )
    if source_statuses.get("walmart"):
        source_statuses["walmart"]["displayed_count"] = search_diagnostics["displayed_walmart_count"]
    summary = calculate_summary(items) if items and comparison_enabled else None
    current_filter_signature = hashlib.sha256(json.dumps({"selected_facets": selected_facets, "active_filters": {"selected_category_key": selected_category_key, "selected_product_type": selected_product_type, "platform": platform, "condition": condition_filter, "min_price": min_price, "max_price": max_price, "sort": sort_option}}, sort_keys=True, default=str, separators=(",", ":")).encode("utf-8")).hexdigest()
    if ai_insight and not stale_ai_insight and ai_insight.get("filter_signature") == current_filter_signature and str(ai_insight.get("search_run_id")) in {str(search_record_id), str(public_search_run_id or "")}:
        insight = ai_insight.get("summary")
        insight, insight_replaced = ground_market_summary(query, insight, items)
        if insight_replaced:
            ai_insight = {**ai_insight, "summary_source": "Rule-based fallback", "generation_mode": "rule_based_fallback", "fallback_reason": "unsupported_condition_claim"}
    max_total = max((item.get("total_price") or 0 for item in items), default=0)
    if search_scope == "both":
        platforms = ["All results", "eBay", "Walmart"]
    else:
        platforms = []
    result_platform_counts = Counter(row.get("platform") or "Unknown" for row in items)
    if search_scope == "both" and stored_walmart_visible:
        result_view = params.get("result_view", "All comparable records")
    elif search_scope == "both" and search_diagnostics.get("walmart_status") == "unavailable_bot_check":
        result_view = params.get("result_view", "All results")
    search_diagnostics["source_mode_note"] = (
        "Comparison includes live eBay records and stored Walmart evidence."
        if stored_walmart_visible else
        ("Only eBay has live matched records. Cross-platform comparison is limited."
         if search_scope == "both" and search_diagnostics.get("walmart_status") in {"unavailable_bot_check", "unavailable", "no results"} and search_diagnostics.get("final_ebay_count", 0) > 0
         else None)
    )
    walmart_unavailable = search_diagnostics.get("walmart_status") == "unavailable_bot_check"
    walmart_chip_label = "Walmart unavailable (bot check detected)" if walmart_unavailable else "Walmart"
    walmart_pipeline_summary = {
        "provider_status": (source_statuses.get("walmart") or {}).get("status"),
        "raw_result_count": search_diagnostics.get("raw_walmart_count", 0),
        "parsed_result_count": search_diagnostics.get("parsed_walmart_count", 0),
        "normalized_result_count": search_diagnostics.get("normalized_walmart_count", 0),
        "valid_price_result_count": search_diagnostics.get("walmart_valid_price_count", 0),
        "model_matched_result_count": search_diagnostics.get("exact_walmart_count", 0),
        "excluded_result_count": search_diagnostics.get("excluded_walmart_count", 0),
        "excluded_count_by_reason": search_diagnostics.get("walmart_rejected_count_by_reason") or {},
        "comparable_result_count": search_diagnostics.get("comparable_walmart_count", 0),
        "displayed_result_count": search_diagnostics.get("displayed_walmart_count", 0),
    }
    search_diagnostics["walmart_pipeline_summary"] = walmart_pipeline_summary
    runtime_pipeline_summary = {
        "internal_search_run_id": str(public_search_run_id or search_record_id or ""),
        "search_run_mode": search_diagnostics.get("search_run_mode"),
        "page_load_mode": search_diagnostics.get("page_load_mode"),
        "provider_response_mode": search_diagnostics.get("walmart_provider_response_mode"),
        "result_id_source": search_diagnostics.get("result_id_source"),
        "retained_result_count": len(search_diagnostics.get("retained_result_ids") or []),
        "comparable_result_count": len(search_diagnostics.get("comparable_result_ids") or []),
        "visible_result_count": len(search_diagnostics.get("visible_result_ids") or []),
    }
    search_diagnostics["runtime_pipeline_summary"] = runtime_pipeline_summary
    if _search_diagnostics_enabled():
        app.logger.info("walmart_pipeline_summary=%s", walmart_pipeline_summary)
        app.logger.info("runtime_pipeline_summary=%s", runtime_pipeline_summary)
        app.logger.info("search_diagnostics=%s", {k: search_diagnostics.get(k) for k in (
            "query", "normalized_query", "search_scope", "force_refresh",
            "ebay_cache_hit", "walmart_cache_hit",
            "raw_ebay_count", "raw_walmart_count",
            "parsed_walmart_count", "normalized_ebay_count", "normalized_walmart_count",
            "excluded_walmart_count", "comparable_walmart_count", "displayed_walmart_count",
            "exact_ebay_count", "exact_walmart_count",
            "final_ebay_count", "final_walmart_count",
            "ebay_status", "walmart_status", "serpapi_status", "serpapi_error",
            "first_normalized_ebay_records", "first_normalized_walmart_records",
            "first_rejected_records", "rejected_count_by_reason",
            "source_status", "walmart_error",
        )})
    if request.method == "POST" and search_record_id:
        repository.update_search_run(search_record_id, {
            "filters": {"platform": platform, "storage": storage_filter, "condition": condition_filter, "min_price": min_price, "max_price": max_price, "sort": sort_option, "match_mode": "related" if include_related_variants else "exact"},
            "detected_category_key": category_key,
            "detected_category_display": category_profile["display_name"],
            "query_intent_status": query_intent_status,
            "detected_category_keys": scope["detected_category_keys"],
            "category_counts": scope["category_counts"],
            "result_categories": result_categories,
            "selected_category_key": selected_category_key,
            "detected_product_types": type_scope["detected_product_types"],
            "product_type_counts": type_scope["product_type_counts"],
            "product_type_facet_coverage": type_scope["facet_coverage"],
            "selected_product_type": selected_product_type,
            "category_mode": category_mode,
            "comparison_enabled": comparison_enabled,
            "search_mode": search_mode,
            "source_statuses": source_statuses,
            "excluded_search_records": walmart_business_exclusions,
            "submitted_query": query,
            "effective_query": query,
            "completed_at": utcnow(),
            "available_facets": available_facets,
            "selected_facets": selected_facets,
            "active_filters": {"selected_category_key": selected_category_key, "selected_product_type": selected_product_type, "platform": platform, "condition": condition_filter, "min_price": min_price, "max_price": max_price, "sort": sort_option},
            "result_ids": list(repository.get_search(search_record_id, session["user_id"]).get("result_tokens") or []),
            "retained_result_ids": list(search_diagnostics.get("retained_result_ids") or []),
            "comparable_result_ids": list(search_diagnostics.get("comparable_result_ids") or []),
            "filtered_result_ids": list(search_diagnostics.get("filtered_result_ids") or []),
            "visible_result_ids": list(search_diagnostics.get("visible_result_ids") or []),
            "source_diagnostics": search_diagnostics,
            "notice": notice,
        })
        current_run = repository.get_search(search_record_id, session["user_id"])
        return redirect(url_for("search_results", search_run_id=current_run.get("search_run_id") or search_record_id))
    canonical_model = parse_canonical_product_model(query) if query else {}
    spelling_suggestion = _search_spelling_suggestion(query, items, source_statuses) if search_record_id else None
    return render_template("source_search_roles.html", query=query, search_scope=search_scope, result_view=result_view, platform_filter=result_view, category=category, min_price=params.get("min_price", min_price or ""), max_price=params.get("max_price", max_price or ""), sort_option=sort_option, platforms=platforms, categories=facets["categories"], items=items, excluded_items=excluded_items, summary=summary, insight=insight, ai_insight=ai_insight, stale_ai_insight=stale_ai_insight, error=error, notice=notice, search_record_id=search_record_id, search_run_id=public_search_run_id or search_record_id, max_total=max_total, show_platform_filter=(search_scope == "both"), result_platforms=result_platforms, result_platform_counts=result_platform_counts, include_related_variants=include_related_variants, query_is_product_like=_is_product_like_query(query), search_diagnostics=search_diagnostics, show_search_diagnostics=_search_ui_diagnostics_enabled(), walmart_unavailable=walmart_unavailable, walmart_chip_label=walmart_chip_label, source_mode_note=search_diagnostics["source_mode_note"], stored_walmart_visible=stored_walmart_visible, live_walmart_visible=live_walmart_visible, stored_walmart_count=search_diagnostics.get("stored_walmart_count", 0), comparison_max=COMPARISON_MAX_RECORDS, available_storage_variants=available_storage_variants, available_conditions=available_conditions, storage_filter=storage_filter, condition_filter=condition_filter, canonical_model=canonical_model, category_profile=category_profile, category_key=category_key, category_mode=category_mode, search_mode=search_mode, comparison_enabled=comparison_enabled, category_options=category_options, selected_category_key=selected_category_key, product_type_options=product_type_options, selected_product_type=selected_product_type, technical_facets_limited=technical_facets_limited, spelling_suggestion=spelling_suggestion, source_messages=source_messages, query_intent_status=query_intent_status, available_facets=available_facets, selected_facets=selected_facets, available_filter_platforms=available_filter_platforms, general_comparison_mode=(comparison_enabled and category_key in {"generic", "food_and_grocery"}))


@app.post("/api/search")
@login_required
def api_search():
    payload = request.get_json(silent=True) or request.form or request.args
    query, source = clean_search_query(payload.get("q", "")), payload.get("source", "demo")
    if source.startswith("demo_"):
        source = "demo"
    if not query:
        return jsonify({"error": "Missing q parameter"}), 400
    if is_blocked_search_query(query):
        record_blocked_search_attempt(session["user_id"], query, "unsafe keyword pattern")
        return jsonify({"error": "The search term contains unsupported characters. Please enter a product name or model."}), 400
    if source not in SOURCES:
        return jsonify({"error": "Invalid source"}), 400
    requested_source = source
    unavailable = (source == "ebay" and not (CLIENT_ID and CLIENT_SECRET)) or (source == "ai" and not (os.getenv("OPENAI_API_KEY") or os.getenv("AI_API_KEY")))
    if unavailable:
        source = "demo"
    search_id = repository.create_search(session["user_id"], query, source, role=session.get("role"), data_mode="synthetic" if source == "demo" else "api")
    try:
        try:
            items = perform_search(query, source, payload.get("item_type", "all"), payload.get("region", "US"))
        except (AISearchError, RuntimeError, requests.RequestException):
            if requested_source not in {"ebay", "ai"}:
                raise
            source, unavailable = "demo", True
            repository.update_search_source(search_id, source, "synthetic")
            items = perform_search(query, source, payload.get("item_type", "all"), payload.get("region", "US"))
        items = normalize_price_items(items)
        if source == "demo" and not _should_show_demo_fallback():
            items = []
        repository.save_products(search_id, session["user_id"], items)
        completion_status = "completed" if items else "no_results"
        if unavailable and items:
            completion_status = "partial_success"
        elif unavailable and not items:
            completion_status = "source_unavailable"
        repository.complete_search(search_id, len(items), completion_status)
        repository.log_event(session["user_id"], "search_query", session["role"], session["username"], {"query": query, "requested_source": requested_source, "actual_source": source, "result_count": len(items)})
    except Exception as exc:
        repository.complete_search(search_id, 0, "failed")
        return jsonify({"error": str(exc)}), 503
    return jsonify({"search_record_id": search_id, "items": items, "requested_source": requested_source, "actual_source": source, "data_mode": "synthetic" if source == "demo" else "api", "notice": "Requested source is unavailable; market records are returned instead." if unavailable else None})


@app.route("/api/admin/serpapi/account")
@role_required("administrator")
def api_serpapi_account():
    try:
        return jsonify(serpapi_account())
    except SerpApiError as exc:
        return jsonify({"error": str(exc)}), 503


@app.route("/debug/walmart-records")
@role_required("consumer", "retailer", "researcher")
def debug_walmart_records():
    if not (app.debug or _search_diagnostics_enabled()):
        abort(404)
    limit = max(1, min(optional_int(request.args.get("limit"), 100), 500))
    if repository.db is not None:
        cursor = repository.db.market_records.find({"platform": "Walmart"}).sort("created_at", -1).limit(limit)
        rows = [repository._public(row) for row in cursor]
    else:
        rows = [repository._public(row) for row in repository._memory["market_records"] if row.get("platform") == "Walmart"]
        rows.sort(key=lambda row: row.get("created_at") or row.get("collected_at") or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
        rows = rows[:limit]
    payload = []
    for row in rows:
        source_text = str(row.get("source_type") or row.get("source") or row.get("data_source") or "")
        payload.append({
            "platform": row.get("platform"),
            "source": source_text,
            "product_url": row.get("source_url") or row.get("item_url") or row.get("url"),
            "is_demo": bool(row.get("demo_mode")) or "demo" in source_text.lower(),
            "is_sample": "sample" in source_text.lower() or str(row.get("confidence_level") or "").lower() == "sample",
            "collected_at": row.get("collected_at") or row.get("created_at") or row.get("timestamp"),
            "query": row.get("query"),
            "title": row.get("title") or row.get("product_name"),
        })
    return jsonify({"count": len(payload), "records": payload})


@app.post("/api/ai-discover")
@login_required
def ai_discover():
    if not can_access(current_user(), "ai_summary"):
        return jsonify({"error": locked_message("ai_summary"), "summary_source": None, "record_count": 0}), 403
    payload = request.get_json(silent=True) or request.form
    search_id = str(payload.get("search_record_id", "")).strip()
    record = repository.get_search(search_id, session["user_id"]) if search_id else None
    if record and record.get("comparison_enabled") is False:
        return jsonify({"error": "Choose one product category or product type before generating a price insight.", "summary_source": None, "record_count": 0}), 400
    if record and record.get("result_tokens"):
        items = resolve_result_tokens(record["result_tokens"])
    elif record and record.get("market_record_ids"):
        items = repository.get_market_records(record["market_record_ids"])
    else:
        items = repository.get_products(search_id) if record else []
    items = annotate_comparison(normalize_price_items(items), record.get("keyword", "") if record else "")
    if record:
        query_for_filters = record.get("keyword", "")
        items = [enrich_listing_category(item, query_for_filters) for item in items]
        items = [item for item in items if item.get("product_role") == "complete_product"]
        if record.get("selected_category_key"):
            items = [item for item in items if item.get("category_key") == record["selected_category_key"]]
        if record.get("selected_product_type"):
            items = [item for item in items if (item.get("attributes") or {}).get("product_type") == record["selected_product_type"]]
        active_filters = record.get("active_filters") or {}
        items = apply_category_filters(items, record.get("selected_facets") or {}, min_price=active_filters.get("min_price"), max_price=active_filters.get("max_price"), platform=active_filters.get("platform") or None, condition=active_filters.get("condition") or None)
    items = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    if not record or not items:
        return jsonify({"error": "Please search sources first before using AI Discover.", "summary_source": None, "record_count": 0}), 400
    query = record.get("keyword", "current market") if record else "current market"
    summary, mode, fallback_reason = summarize_market(query, items)
    model = os.getenv("GEMINI_MODEL", "gemini-2.5-flash-lite")
    summary_source = "Gemini API" if mode == "gemini_api" else "Rule-based fallback"
    repository.log_ai_search(session["user_id"], query, model, "Gemini market summary grounded in current normalized records.", summary)
    repository.log_ai_activity(session["user_id"], query, model, summary_source, len(items), fallback_reason)
    repository.log_event(session["user_id"], "ai_discover", session["role"], session["username"], {"query": query, "provider": "Gemini", "model": model, "generation_mode": mode, "summary_source": summary_source, "fallback_reason": fallback_reason, "record_count": len(items)})
    generated_at = utcnow()
    included_result_ids = [str(item.get("_selection_token") or item.get("_id")) for item in items]
    source_data_version = hashlib.sha256(json.dumps([
        {"id": result_id, "price": item.get("total_price"), "updated": str(item.get("updated_at") or item.get("collected_at") or item.get("created_at") or "")}
        for result_id, item in zip(included_result_ids, items)
    ], sort_keys=True, default=str).encode("utf-8")).hexdigest()
    insight_document = {
        "search_run_id": str(record.get("_id")), "query": query,
        "selected_category_key": record.get("selected_category_key"),
        "selected_product_type": record.get("selected_product_type"),
        "result_ids": included_result_ids,
        "result_set_signature": hashlib.sha256(json.dumps(sorted(included_result_ids)).encode("utf-8")).hexdigest(),
        "source_data_version": source_data_version,
        "filter_signature": hashlib.sha256(json.dumps({"selected_facets": record.get("selected_facets") or {}, "active_filters": record.get("active_filters") or {}}, sort_keys=True, default=str, separators=(",", ":")).encode("utf-8")).hexdigest(),
        "summary": summary, "summary_source": summary_source, "generation_mode": mode,
        "model": model, "generated_at": generated_at, "fallback_reason": fallback_reason,
    }
    ai_insight_id = repository.save_ai_insight(session["user_id"], insight_document)
    return jsonify({"ai_insight_id": ai_insight_id, "summary": summary, "summary_source": summary_source, "generation_mode": mode, "model": model,
                    "search_run_id": str(record.get("_id")), "user_id": str(session["user_id"]), "query": query,
                    "included_result_ids": included_result_ids,
                    "comparable_result_count": len(items), "record_count": len(items),
                    "platforms": sorted({item.get("platform") for item in items if item.get("platform")} ),
                    "generated_at": display_sgt_datetime(generated_at), "fallback_reason": fallback_reason})


@app.post("/search/results/<search_run_id>/ai-insight")
@login_required
def create_search_ai_insight(search_run_id):
    record = repository.get_search(search_run_id, session["user_id"])
    if not record:
        abort(404)
    # Reuse the API implementation so both entry points create the same bound,
    # persisted AI Insight object. The browser-facing action follows PRG.
    response = ai_discover()
    status_code = response[1] if isinstance(response, tuple) else response.status_code
    if status_code >= 400:
        payload = response[0].get_json() if isinstance(response, tuple) else response.get_json()
        flash((payload or {}).get("error") or "AI price insight could not be generated.", "error")
    return redirect(url_for("search_results", search_run_id=search_run_id))


@app.route("/analytics")
@login_required
def analytics_compatibility():
    if not can_access(current_user(), "analytics_dashboard"):
        return locked_feature_response("analytics_dashboard", title="Analytics")
    analyses = repository.list_analysis_records(session["user_id"], limit=30)
    return render_template("analytics_index.html", analyses=analyses)


def _normalized_analysis_filters(filters):
    """Return a stable, JSON-safe representation of the filters used for analysis."""
    if not isinstance(filters, dict):
        return {}

    def normalize(value):
        if isinstance(value, dict):
            return {str(key).strip().lower(): normalize(value[key]) for key in sorted(value, key=lambda key: str(key).lower())}
        if isinstance(value, (list, tuple, set)):
            return [normalize(item) for item in value]
        if isinstance(value, str):
            return value.strip()
        return value

    return {str(key).strip().lower(): normalize(filters[key]) for key in sorted(filters, key=lambda key: str(key).lower()) if filters[key] not in (None, "")}


def _analysis_record_payload(scope, record, items, included_tokens, comparison_set_id=None, analysis_filters=None):
    annotated = annotate_comparison(normalize_price_items(items), record.get("keyword", "") if record else "")
    eligible = [item for item in annotated if item.get("analytics_eligible") and item.get("total_price") is not None]
    eligible_tokens = {item.get("_selection_token") for item in eligible if item.get("_selection_token")}
    included_ids = [token for token in included_tokens if token in eligible_tokens]
    excluded_ids = [token for token in included_tokens if token not in eligible_tokens]
    excluded_by_token = {item.get("_selection_token"): item.get("exclusion_reason") or "Not eligible for analytics" for item in annotated}
    metrics = comparison_metrics(annotated)
    source_version_payload = [
        {
            "id": str(item.get("_selection_token") or item.get("_id") or ""),
            "price": item.get("total_price"),
            "source_price": item.get("price"),
            "source_type": item.get("source_type"),
            "updated": str(item.get("updated_at") or item.get("collected_at") or item.get("created_at") or ""),
        }
        for item in sorted(annotated, key=lambda row: str(row.get("_selection_token") or row.get("_id") or ""))
    ]
    source_data_version = hashlib.sha256(json.dumps(source_version_payload, sort_keys=True, default=str).encode("utf-8")).hexdigest()
    filters = _normalized_analysis_filters(record.get("filters") if analysis_filters is None else analysis_filters)
    signature_payload = {
        "user_id": str(session.get("user_id") or ""),
        "workspace": scope,
        "comparison_set_id": str(comparison_set_id or ""),
        "selected_result_ids": sorted(str(token) for token in included_tokens),
        "filters": filters,
        "source_data_version": source_data_version,
        "analysis_version": "comparison-v2",
    }
    analysis_signature = hashlib.sha256(json.dumps(signature_payload, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
    generated_at = utcnow()
    return {
        "search_record_id": record.get("_id") if record else None,
        "comparison_set_id": comparison_set_id,
        "scope": scope,
        "query": record.get("keyword", "") if record else "Selected comparison",
        "included_result_ids": included_ids,
        "excluded_result_ids": excluded_ids,
        "preview_result_ids": list(included_tokens),
        "excluded_reasons": {token: excluded_by_token.get(token, "Not eligible for analytics") for token in excluded_ids},
        "metrics_snapshot": {key: value for key, value in metrics.items() if key not in {"excluded", "comparable"}},
        "filters": filters, "source_data_version": source_data_version, "analysis_version": "comparison-v2",
        "analysis_schema_version": CURRENT_ANALYSIS_SCHEMA_VERSION, "normalized_filters": filters,
        "generated_at": generated_at, "updated_at": generated_at,
        "analysis_signature": analysis_signature, "saved_status": "saved", "status": "active",
    }


def _validate_analytics_scope(analysis_record, payload):
    """Reject a stale or internally inconsistent frozen analysis before display/export."""
    expected = len(analysis_record.get("included_result_ids") or [])
    checks = {
        "summary": payload.get("summary", {}).get("total_records", 0),
        "records": len(payload.get("records") or []),
        "platform": sum(payload.get("platform_distribution", {}).get("counts") or []),
        "condition": sum(payload.get("condition_distribution", {}).get("counts") or []),
        "histogram": sum(payload.get("price_distribution", {}).get("counts") or []),
        "price_curve": len(payload.get("comparable_price_curve", {}).get("record_ids") or []),
    }
    if not analysis_record.get("analysis_schema_version"):
        preview_ids = analysis_record.get("preview_result_ids") or []
        if preview_ids:
            checks["legacy_preview"] = len(preview_ids)
    if expected and all(value == expected for value in checks.values()):
        return
    raise AnalysisScopeMismatch(analysis_record, checks)


def _frozen_analysis_items(analysis_record):
    included_ids = [str(value) for value in analysis_record.get("included_result_ids") or []]
    if not included_ids or len(set(included_ids)) != len(included_ids):
        return None
    resolved = annotate_comparison(normalize_price_items(resolve_result_tokens(included_ids)), analysis_record.get("query", ""))
    by_token = {str(item.get("_selection_token") or item.get("_id")): item for item in resolved}
    if any(token not in by_token for token in included_ids):
        return None
    return [by_token[token] for token in included_ids]


def _analytics_context_from_frozen_analysis(analysis_record):
    items = _frozen_analysis_items(analysis_record)
    if items is None:
        raise AnalysisScopeMismatch(analysis_record, {"records": 0})
    scope = analysis_record.get("scope") or "all_valid_results"
    payload = build_analytics_payload(items, query=analysis_record.get("query", ""), source="frozen analysis", analysis_scope=scope, mode="selected" if scope == "selected_comparison" else "full")
    _validate_analytics_scope(analysis_record, payload)
    source_record = repository.get_search(analysis_record.get("search_record_id"), session["user_id"]) if analysis_record.get("search_record_id") else None
    return source_record or {"_id": analysis_record.get("search_record_id") or "", "keyword": analysis_record.get("query", "")}, payload, items


@app.post("/analytics/from-search")
@login_required
def create_search_analysis():
    if not can_access(current_user(), "analytics_dashboard"):
        return locked_feature_response("analytics_dashboard", title="Analytics")
    search_id = request.form.get("search_record_id", "")
    record = repository.get_search(search_id, session["user_id"])
    if not record:
        abort(404)
    if record.get("comparison_enabled") is False:
        flash("Choose one product category or product type to enable reliable price comparison and analysis.", "error")
        return redirect(url_for("search_results", search_run_id=search_id))
    tokens = list(record.get("result_tokens") or [])
    items = [enrich_listing_category(item, record.get("keyword", "")) for item in normalize_price_items(resolve_result_tokens(tokens))]
    items = [item for item in items if item.get("product_role") == "complete_product"]
    if record.get("selected_category_key"):
        items = [item for item in items if item.get("category_key") == record["selected_category_key"]]
    if record.get("selected_product_type"):
        items = [item for item in items if (item.get("attributes") or {}).get("product_type") == record["selected_product_type"]]
    active_filters = record.get("active_filters") or {}
    items = apply_category_filters(items, record.get("selected_facets") or {}, min_price=active_filters.get("min_price"), max_price=active_filters.get("max_price"), platform=active_filters.get("platform") or None, condition=active_filters.get("condition") or None)
    tokens = [item.get("_selection_token") for item in items if item.get("_selection_token")]
    analysis_filters = {**(record.get("selected_facets") or {}), **{key: value for key, value in active_filters.items() if value not in (None, "")}}
    payload = _analysis_record_payload("all_valid_results", record, items, tokens, analysis_filters=analysis_filters)
    if not payload["included_result_ids"]:
        flash("No analytics-eligible full-price records are available.", "error")
        return redirect(url_for("search", search_record_id=search_id))
    existing = repository.find_analysis_by_signature(session["user_id"], payload["analysis_signature"])
    analysis_id = existing["_id"] if existing else repository.create_analysis_record(session["user_id"], payload)
    session["active_analysis_id"] = analysis_id
    return redirect(url_for("analytics", search_record_id=analysis_id))


@app.post("/analytics/from-comparison/<comparison_id>")
@login_required
def create_comparison_analysis(comparison_id):
    if not can_access(current_user(), "analytics_dashboard"):
        return locked_feature_response("analytics_dashboard", title="Analytics")
    group = repository.get_comparison_group(comparison_id, session["user_id"])
    if not group:
        abort(404)
    record = repository.get_search(group.get("search_record_id"), session["user_id"])
    tokens = list(group.get("selected_record_ids") or [])
    items = _comparison_group_records(group)
    payload = _analysis_record_payload("selected_comparison", record or {"_id": group.get("search_record_id"), "keyword": group.get("query", "")}, items, tokens, comparison_set_id=comparison_id)
    if not payload["included_result_ids"]:
        flash("No comparable full-price records are available for analysis.", "error")
        return redirect(url_for("compare_group", comparison_id=comparison_id))
    existing = repository.find_analysis_by_signature(session["user_id"], payload["analysis_signature"])
    if existing and request.form.get("refresh") != "1":
        analysis_id = existing["_id"]
    else:
        if existing:
            payload["analysis_signature"] = hashlib.sha256(f"{payload['analysis_signature']}:{utcnow().isoformat()}".encode("utf-8")).hexdigest()
            payload["refreshed_from_analysis_id"] = existing["_id"]
        analysis_id = repository.create_analysis_record(session["user_id"], payload)
    repository.update_comparison_group(comparison_id, {"last_analysis_id": analysis_id, "last_analysis_signature": payload["analysis_signature"]})
    session["active_analysis_id"] = analysis_id
    return redirect(url_for("analytics", search_record_id=analysis_id))


@app.route("/analytics/<search_record_id>")
@login_required
def analytics(search_record_id):
    if not can_access(current_user(), "analytics_dashboard"):
        return locked_feature_response("analytics_dashboard", title="Analytics")
    if not is_valid_object_id(search_record_id):
        abort(404)
    analysis_record = repository.get_analysis_record(search_record_id, session["user_id"])
    if analysis_record:
        analysis_record = repository.touch_analysis_record(search_record_id) or analysis_record
        source_record, analytics_payload, items = _analytics_context_from_frozen_analysis(analysis_record)
        scope = analysis_record.get("scope") or "all_valid_results"
        display_query = analysis_record.get("query") or analytics_display_query(source_record)
        summary = analytics_payload["summary"]
        return render_template(
            "analytics_dashboard.html", summary=summary, analytics_payload=analytics_payload,
            preview_items=analytics_payload["records"][:20], record=source_record,
            display_query=display_query, platform_distribution=dict(zip(analytics_payload["platform_distribution"]["labels"], analytics_payload["platform_distribution"]["counts"])),
            category_distribution={}, accessories_excluded=len(analysis_record.get("excluded_result_ids") or []), data_source_label="Stored marketplace evidence",
            analysis_scope=scope, record_count=summary["total_records"], single_platform_message=analytics_payload["singlePlatformMessage"],
            analytics_label="Comparison analytics" if scope == "selected_comparison" else "Search analytics", platform_chart_title=analytics_payload["platformChartTitle"],
            comparison_id=analysis_record.get("comparison_set_id"), analytics_heading="Comparison Analytics — Selected records" if scope == "selected_comparison" else "Search Analytics — All valid results",
            analysis_record=analysis_record,
        )
    user_scope = None if session.get("role") == "researcher" else session["user_id"]
    record, items, source_kind = _analytics_records_for_search(search_record_id, user_scope)
    if not record:
        abort(404)
    display_query = analytics_display_query(record)
    mode = request.args.get("mode", "full")
    selected_tokens = request.args.getlist("result_token")
    if mode == "selected" and selected_tokens:
        allowed = set(record.get("result_tokens", []))
        items = annotate_comparison(normalize_price_items(resolve_result_tokens([token for token in selected_tokens if token in allowed])), display_query)
    elif mode == "saved":
        items = annotate_comparison(normalize_price_items(safe_call(lambda: repository.list_evidence(record.get("user_id"), limit=200), [])), display_query)
        items = [item for item in items if str(item.get("search_record_id")) == str(record.get("_id"))]
    else:
        items = annotate_comparison(items, display_query)
    if request.args.get("include_demo") != "1":
        items = [item for item in items if _normalize_export_source_type(item) != "demo_sample"]
    analysis_items = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    analytics_payload = build_analytics_payload(
        items,
        query=display_query,
        source=record.get("selected_source", ""),
        analysis_scope="selected records" if mode == "selected" else ("saved evidence records" if mode == "saved" else "all current search results"),
        mode=mode,
    )

    summary = analytics_payload["summary"]
    return render_template(
        "analytics_dashboard.html",
        summary=summary,
        analytics_payload=analytics_payload,
        preview_items=analytics_payload["records"][:20],
        record=record,
        display_query=display_query,
        platform_distribution=dict(zip(analytics_payload["platform_distribution"]["labels"], analytics_payload["platform_distribution"]["counts"])),
        category_distribution={},
        accessories_excluded=len(items) - len(analysis_items),
        data_source_label=source_kind,
        analysis_scope=analytics_payload["analysisScope"],
        record_count=summary["total_records"],
        single_platform_message=analytics_payload["singlePlatformMessage"],
        analytics_label=analytics_payload["analyticsLabel"],
        platform_chart_title=analytics_payload["platformChartTitle"],
        comparison_id=None,
    )


@app.post("/analytics/<analysis_id>/delete")
@login_required
def delete_analysis(analysis_id):
    if not can_access(current_user(), "analytics_dashboard") or not is_valid_object_id(analysis_id):
        abort(404)
    analysis = repository.get_analysis_record(analysis_id, session["user_id"])
    if not analysis:
        abort(404)
    repository.delete_analysis_record(analysis_id, session["user_id"])
    repository.log_event(session["user_id"], "analysis_deleted", session.get("role"), session.get("username"), {"analysis_id": str(analysis_id)})
    flash("Analysis deleted", "success")
    return redirect(url_for("analytics_compatibility"))


@app.route("/analytics/comparison/<comparison_id>")
@login_required
def analytics_comparison(comparison_id):
    if not can_access(current_user(), "analytics_dashboard"):
        return locked_feature_response("analytics_dashboard", title="Analytics")
    group = repository.get_comparison_group(comparison_id, session["user_id"])
    if not group:
        flash("Selected comparison records were not found. Please select records again.", "error")
        return redirect(url_for("search"))
    search_id = group.get("search_record_id")
    record = repository.get_search(search_id, session["user_id"]) if search_id else None
    display_query = group.get("query") or (analytics_display_query(record) if record else "Selected comparison")
    items = annotate_comparison(_comparison_group_records(group), display_query)
    if request.args.get("include_demo") != "1":
        items = [item for item in items if _normalize_export_source_type(item) != "demo_sample"]
    analysis_items = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    _comparison_debug("analytics_comparison", comparison_id=comparison_id, selected_record_ids=group.get("selected_record_ids", []), record_count=len(analysis_items))
    analytics_payload = build_analytics_payload(
        items,
        query=display_query,
        source="selected comparison records",
        analysis_scope="selected comparison records",
        mode="selected",
    )
    summary = analytics_payload["summary"]
    return render_template(
        "analytics_dashboard.html",
        summary=summary,
        analytics_payload=analytics_payload,
        preview_items=analytics_payload["records"][:20],
        record=record or {"keyword": display_query, "_id": search_id or ""},
        display_query=display_query,
        platform_distribution=dict(zip(analytics_payload["platform_distribution"]["labels"], analytics_payload["platform_distribution"]["counts"])),
        category_distribution={},
        accessories_excluded=len(items) - len(analysis_items),
        data_source_label="selected comparison records",
        analysis_scope=analytics_payload["analysisScope"],
        record_count=summary["total_records"],
        single_platform_message=analytics_payload["singlePlatformMessage"],
        analytics_label=analytics_payload["analyticsLabel"],
        platform_chart_title=analytics_payload["platformChartTitle"],
        comparison_id=comparison_id,
    )


def _analytics_export_context_for_comparison(comparison_id):
    group = repository.get_comparison_group(comparison_id, session["user_id"])
    if not group:
        abort(404)
    search_id = group.get("search_record_id")
    record = repository.get_search(search_id, session["user_id"]) if search_id else {"_id": "", "keyword": group.get("query", ""), "selected_source": "selected comparison records"}
    if not record:
        record = {"_id": search_id or "", "keyword": group.get("query", ""), "selected_source": "selected comparison records"}
    display_query = group.get("query") or analytics_display_query(record)
    items = annotate_comparison(_comparison_group_records(group), display_query)
    analysis_items = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    analytics_payload = build_analytics_payload(
        analysis_items,
        query=display_query,
        source="selected comparison records",
        analysis_scope="selected comparison records",
        mode="selected",
    )
    return record, analytics_payload, items


def _analytics_export_context(search_record_id):
    if not is_valid_object_id(search_record_id):
        abort(404)
    analysis_record = repository.get_analysis_record(search_record_id, session["user_id"])
    if analysis_record:
        return _analytics_context_from_frozen_analysis(analysis_record)
    user_scope = None if session.get("role") == "researcher" else session["user_id"]
    record, items, _source_kind = _analytics_records_for_search(search_record_id, user_scope)
    if not record:
        abort(404)
    display_query = analytics_display_query(record)
    mode = request.args.get("mode", "full")
    selected_tokens = request.args.getlist("result_token")
    if mode == "selected" and selected_tokens:
        allowed = set(record.get("result_tokens", []))
        items = annotate_comparison(normalize_price_items(resolve_result_tokens([token for token in selected_tokens if token in allowed])), display_query)
    elif mode == "saved":
        items = annotate_comparison(normalize_price_items(safe_call(lambda: repository.list_evidence(record.get("user_id"), limit=200), [])), display_query)
        items = [item for item in items if str(item.get("search_record_id")) == str(record.get("_id"))]
    else:
        items = annotate_comparison(items, display_query)
    analysis_items = [item for item in items if item.get("analytics_eligible") and item.get("total_price") is not None]
    analytics_payload = build_analytics_payload(
        analysis_items,
        query=record.get("keyword", ""),
        source=record.get("selected_source", ""),
        analysis_scope="selected records" if mode == "selected" else ("saved evidence records" if mode == "saved" else "all current search results"),
        mode=mode,
    )
    return record, analytics_payload, items


def _report_dashboard_filters(items, submitted_filters):
    """Apply only validated dashboard filters to server-owned analysis rows."""
    if not isinstance(submitted_filters, dict):
        return list(items or []), {}
    rows = list(items or [])
    field_values = {
        "platform": lambda row: str(row.get("platform") or "Unknown"),
        "category": lambda row: str(row.get("category_display") or row.get("category") or "Not available"),
        "condition": lambda row: str(row.get("condition_display") or "Not specified by source"),
    }
    active = {}
    for key, getter in field_values.items():
        value = str(submitted_filters.get(key) or "").strip()
        if not value or value.lower() == "all":
            continue
        allowed = {getter(row) for row in rows}
        if value not in allowed:
            abort(400, description=f"Invalid report filter: {key}")
        rows = [row for row in rows if getter(row) == value]
        active[key] = value
    return rows, active


def _report_persisted_filters(search_record_id, record):
    search_filters = {
        key: value
        for key, value in (record.get("filters") or {}).items()
        if value not in (None, "", [], {})
    }
    search_filters.update({
        key: value
        for key, value in (record.get("selected_facets") or {}).items()
        if value not in (None, "", [], {})
    })
    search_filters.update({
        key: value
        for key, value in (record.get("active_filters") or {}).items()
        if value not in (None, "", [], {})
    })
    analysis_record = repository.get_analysis_record(search_record_id, session["user_id"])
    if analysis_record:
        analysis_filters = analysis_record.get("normalized_filters") or analysis_record.get("filters") or {}
        search_filters.update({
            key: value
            for key, value in analysis_filters.items()
            if value not in (None, "", [], {})
        })
    return _normalized_analysis_filters(search_filters)


def _report_filter_text(persisted_filters, dashboard_filters):
    labels = {
        "storage": "Storage",
        "platform": "Platform",
        "condition": "Condition",
        "category": "Category",
        "selected_category_key": "Category",
        "selected_product_type": "Product type",
        "min_price": "Minimum price",
        "max_price": "Maximum price",
    }
    parts = []
    for key, value in (persisted_filters or {}).items():
        if key == "sort" or value in (None, "", [], {}) or str(value).lower() == "all":
            continue
        display_value = ", ".join(str(item) for item in value) if isinstance(value, list) else str(value)
        parts.append(f"{labels.get(key, key.replace('_', ' ').title())}: {display_value}")
    for key, value in (dashboard_filters or {}).items():
        parts.append(f"Analytics {labels.get(key, key.title())}: {value}")
    return "; ".join(dict.fromkeys(parts)) or "No additional filters"


@app.route("/analytics/<search_record_id>/export/results.csv")
@login_required
def analytics_export_results_csv(search_record_id):
    if not can_access(current_user(), "watchlist"):
        return export_forbidden_response("Export is available on Premium and Professional plans.")
    record, analytics_payload, _items = _analytics_export_context(search_record_id)
    rows = [["Analysis ID", "Query", "Platform", "Product Title", "Observed Price", "Currency", "Normalized Price", "Condition", "Category", "Seller", "Collected At (SGT)", "Record Source", "Source URL", "Analysis Included", "Exclusion Reason"]]
    for item in analytics_payload["records"]:
        rows.append([
            search_record_id,
            record.get("keyword", ""),
            item.get("platform", ""),
            item.get("title", ""),
            _display_price(item.get("price")),
            item.get("currency") or "USD",
            _display_price(item.get("normalized_price")),
            item.get("condition_display") or "Not available",
            item.get("category_display") or "Not available",
            item.get("seller", ""),
            item.get("collected_at", ""),
            item.get("record_source") or "Stored marketplace evidence",
            item.get("source_url") or "",
            "Yes",
            "",
        ])
    return csv_response(rows, f"analytics-results-{search_record_id}.csv")


def _analytics_workbook_response(analysis_id, record, payload):
    summary = payload["summary"]
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Query", record.get("keyword", "")])
    summary_sheet.append(["Total Records", summary.get("total_records", 0)])
    summary_sheet.append(["Platforms", summary.get("platform_count", 0)])
    summary_sheet.append(["Average Price", summary.get("average_price")])
    summary_sheet.append(["Median Price", summary.get("median_price")])
    summary_sheet.append(["Minimum Price", summary.get("minimum_price")])
    summary_sheet.append(["Maximum Price", summary.get("maximum_price")])
    summary_sheet.append(["Price Range", summary.get("price_range")])
    summary_sheet.append(["Analysis Scope", f"{summary.get('total_records', 0)} comparable listings"])
    summary_sheet.append(["Applied Filters", "No additional filters"])
    records_sheet = workbook.create_sheet("Records")
    records_sheet.append(["Platform", "Product Title", "Observed Price", "Currency", "Normalized Price", "Condition", "Category", "Seller", "Collected At (SGT)", "Record Source", "Source URL"])
    for item in payload["records"]:
        records_sheet.append([item.get("platform"), item.get("title"), item.get("price"), item.get("currency"), item.get("normalized_price"), item.get("condition_display") or "Not available", item.get("category_display") or "Not available", item.get("seller") or "", item.get("collected_at"), item.get("record_source"), item.get("source_url")])
    platform_sheet = workbook.create_sheet("Platform Breakdown")
    platform_sheet.append(["Platform", "Count", "Percentage", "Average", "Median", "Minimum", "Maximum"])
    comparison = payload["platform_price_comparison"]
    counts = dict(zip(payload["platform_distribution"]["labels"], payload["platform_distribution"]["counts"]))
    percentages = dict(zip(payload["platform_distribution"]["labels"], payload["platform_distribution"]["percentages"]))
    for label, average, median, minimum, maximum in zip(comparison["labels"], comparison["average"], comparison["median"], comparison["minimum"], comparison["maximum"]):
        platform_sheet.append([label, counts.get(label, 0), percentages.get(label, 0), average, median, minimum, maximum])
    condition_sheet = workbook.create_sheet("Condition Breakdown")
    condition_sheet.append(["Condition", "Count", "Percentage"])
    for row in zip(payload["condition_distribution"]["labels"], payload["condition_distribution"]["counts"], payload["condition_distribution"]["percentages"]):
        condition_sheet.append(list(row))
    distribution_sheet = workbook.create_sheet("Price Distribution")
    distribution_sheet.append(["Lower Bound", "Upper Bound", "Label", "Count"])
    for label, count in zip(payload["price_distribution"]["bins"], payload["price_distribution"]["counts"]):
        distribution_sheet.append(["", "", label, count])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(buffer.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="analytics-workbook-{analysis_id}.xlsx"'})


@app.route("/analytics/<search_record_id>/export/workbook.xlsx")
@login_required
def analytics_export_workbook(search_record_id):
    if not can_access(current_user(), "watchlist"):
        return export_forbidden_response("Export is available on Premium and Professional plans.")
    record, analytics_payload, _items = _analytics_export_context(search_record_id)
    return _analytics_workbook_response(search_record_id, record, analytics_payload)


@app.route("/analytics/comparison/<comparison_id>/export/workbook.xlsx")
@login_required
def analytics_comparison_export_workbook(comparison_id):
    if not can_access(current_user(), "watchlist"):
        return export_forbidden_response("Export is available on Premium and Professional plans.")
    record, analytics_payload, _items = _analytics_export_context_for_comparison(comparison_id)
    return _analytics_workbook_response(comparison_id, record, analytics_payload)


@app.route("/analytics/<search_record_id>/export/chart-data.json")
@login_required
def analytics_export_chart_data_json(search_record_id):
    if not can_access(current_user(), "export_report"):
        return export_forbidden_response("Full report and chart-data export require Professional.")
    _record, analytics_payload, _items = _analytics_export_context(search_record_id)
    return Response(json.dumps(make_json_safe(analytics_payload), ensure_ascii=False, indent=2), mimetype="application/json; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="analytics-chart-data-{search_record_id}.json"'})


@app.route("/analytics/<search_record_id>/report")
@login_required
def analytics_report(search_record_id):
    if not can_access(current_user(), "export_report"):
        return export_forbidden_response("Full report and chart-data export require Professional.")
    record, analytics_payload, _items = _analytics_export_context(search_record_id)
    return render_template("analytics_report.html", record=record, analytics_payload=analytics_payload, items=analytics_payload["records"][:20], analysis_id=search_record_id)


@app.post("/analytics/<search_record_id>/export/report")
@login_required
def analytics_export_report_post(search_record_id):
    if not can_access(current_user(), "export_report"):
        return export_forbidden_response("Full report export requires Professional.")
    if not is_valid_object_id(search_record_id):
        abort(404)
    record, payload, source_items = _analytics_export_context(search_record_id)

    def parse_json_field(name, fallback):
        try:
            return json.loads(request.form.get(name, "") or "")
        except (TypeError, ValueError, json.JSONDecodeError):
            return fallback

    submitted_metadata = parse_json_field("metadata", {})
    charts = parse_json_field("charts", {})
    filtered_items, dashboard_filters = _report_dashboard_filters(
        source_items, submitted_metadata.get("filters")
    )
    if dashboard_filters:
        payload = build_analytics_payload(
            filtered_items,
            query=analytics_display_query(record),
            source="frozen analysis",
            analysis_scope="filtered frozen analysis",
            mode="full",
        )
    summary = payload["summary"]
    persisted_filters = _report_persisted_filters(search_record_id, record)
    storage_values = summary.get("storage_variants") or []
    persisted_storage = persisted_filters.get("storage")
    if not storage_values and persisted_storage:
        storage_values = persisted_storage if isinstance(persisted_storage, list) else [persisted_storage]
        storage_values = [str(value) for value in storage_values if str(value).strip()]
    if len(storage_values) == 1:
        for report_row in payload.get("records") or []:
            if not report_row.get("storage"):
                report_row["storage"] = storage_values[0]
    storage_scope = ", ".join(storage_values) if storage_values else "Not consistently provided"
    condition_values = payload.get("condition_distribution", {}).get("labels") or []
    condition_scope = ", ".join(condition_values) if condition_values else "Not consistently provided"
    query = analytics_display_query(record)
    title_scope = query
    if len(storage_values) == 1 and storage_values[0].lower() not in query.lower():
        title_scope = f"{query} · {storage_values[0]}"
    analysis_record = repository.get_analysis_record(search_record_id, session["user_id"])
    excluded_count = len((analysis_record or {}).get("excluded_result_ids") or [])
    collected_count = len(record.get("result_tokens") or record.get("market_record_ids") or [])
    metadata = {
        "generated_at": singapore_time_string(utcnow()),
        "query": query,
        "title_scope": title_scope,
        "scope": f"{summary.get('total_records', 0)} comparable listings",
        "included": summary.get("total_records", 0),
        "excluded": excluded_count,
        "collected": collected_count or summary.get("total_records", 0) + excluded_count,
        "platforms": ", ".join(payload["platform_distribution"].get("labels") or []) or "Not available",
        "filters": _report_filter_text(persisted_filters, dashboard_filters),
        "storage_scope": storage_scope,
        "condition_scope": condition_scope,
        "collected_at": display_sgt_datetime(record.get("created_at")),
        "analysis_id": search_record_id,
        "mixed_variants": bool(summary.get("mixed_variants")),
        "mixed_conditions": bool(summary.get("mixed_conditions")),
    }
    return render_template(
        "export_report.html",
        title="Precision Curator Analytics Report",
        record=record,
        metadata=metadata,
        summary=summary,
        records=payload["records"][:200],
        charts=charts,
    )


@app.route("/demo")
@app.route("/test-data")
@demo_tools_required
def demo():
    items = repository.list_testing_records(limit=50)
    prices = [float(item["price"]) for item in items if item.get("price") is not None]
    summary = {"total_results": len(items), "platform_count": len({item.get("platform") for item in items}), "category_count": len({item.get("category") for item in items}), "average_price": round(sum(prices) / len(prices), 2) if prices else None}
    return render_template("test_data.html", items=items[:20], summary=summary, analysis_image="data_analysis_demo.png", visualization_image="visualization_demo.png")


@app.post("/test-data/seed-watchlist-snapshots")
@demo_tools_required
def seed_watchlist_snapshots_from_test_data():
    user = current_user()
    if not user:
        abort(403)
    existing = [
        row
        for row in repository.list_watchlist_items(user_id=user["_id"], include_archived=False, limit=20)
        if (row.get("source_label") or row.get("data_source_label")) == "demo"
    ]
    if existing:
        item = existing[0]
    else:
        item_id, _created = repository.create_watchlist_item(
            user["_id"],
            {
                "keyword": "demo watchlist item",
                "product_label": "Demo watchlist item",
                "tracking_mode": "search_scope",
                "platform_scope": "demo_all",
                "source_label": "demo",
                "data_source_label": "demo",
            },
        )
        item = repository.get_watchlist_item(item_id, user["_id"])
    created = seed_demo_snapshots([item["_id"]], user_id=user["_id"])
    flash(f"Seeded {len(created)} demo snapshot(s) for presentation.", "success")
    return redirect(url_for("watchlist", item_id=item["_id"]))


@app.post("/test-data/seed-watchlist-validation")
@demo_tools_required
def seed_watchlist_validation_from_test_data():
    user = current_user()
    if not user:
        abort(403)
    result = seed_demo_prediction_validation(user["_id"])
    if not result:
        flash("Could not seed demo prediction validation.", "error")
        return redirect(url_for("demo"))
    flash("Seeded demo prediction validation.", "success")
    return redirect(url_for("watchlist", item_id=result["watchlist_id"]))


@app.post("/test-data/delete")
@demo_tools_required
def delete_test_data():
    ids = [value for value in request.form.getlist("record_id") if value]
    if not ids:
        flash("Select at least one test record.", "error")
        return redirect(url_for("demo"))
    deleted = repository.delete_testing_records(ids, deleted_by=session.get("user_id"), reason="test_data_deleted")
    repository.log_event(session["user_id"], "testing_records_deleted", session.get("role"), session.get("username"), {"deleted_count": deleted, "action": "delete_selected"})
    flash(f"Deleted {deleted} test record(s).", "success")
    return redirect(url_for("demo"))


@app.post("/test-data/reset")
@demo_tools_required
def reset_test_data():
    repository.restore_demo_test_records()
    repository.log_event(session["user_id"], "testing_records_reset", session.get("role"), session.get("username"), {"action": "reset_demo_data"})
    flash("Demo data reset.", "success")
    return redirect(url_for("demo"))


@app.errorhandler(403)
def forbidden(_error):
    return render_template("error.html", code=403, message="This dashboard belongs to a different role."), 403


@app.errorhandler(404)
def not_found(_error):
    return render_template("error.html", code=404, message="That record or page could not be found."), 404


@app.errorhandler(AnalysisScopeMismatch)
def analysis_scope_mismatch(error):
    analysis = error.analysis_record or {}
    return render_template("analysis_recovery.html", analysis=analysis), 409


@app.errorhandler(500)
def internal_error(_error):
    app.logger.exception("Unhandled application error")
    return render_template("error.html", code=500, message="Something went wrong. Please return to the workspace and try again."), 500
