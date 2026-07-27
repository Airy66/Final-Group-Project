import csv
import io
import json
from datetime import datetime, timezone

from openpyxl import load_workbook

import precision_app
from services.database import MongoRepository


def _login(client, name, role="consumer"):
    email = f"{''.join(character.lower() for character in name if character.isalnum())}@example.test"
    password = "OfflineTestPassword1!"
    user = precision_app.repository.get_user_by_email(email)
    if not user:
        user = precision_app.repository.create_user(name, email, precision_app.generate_password_hash(password, method="pbkdf2:sha256:1000"), role, roles=[role], active_role=role)
    client.post("/login", data={"email": email, "password": password})
    return user


def _fixture(legacy_preview=False):
    precision_app.app.config.update(TESTING=True, SECRET_KEY="analytics-export-tests")
    precision_app.repository = MongoRepository(uri="", database_name="test_analytics_exports")
    client = precision_app.app.test_client()
    user = _login(client, "Analyst")
    precision_app.repository.update_user(user["_id"], {"membership_tier": "professional", "plan": "professional"})
    search_id = precision_app.repository.create_search(user["_id"], "iPhone 12", "mock", role="consumer", data_mode="mock")
    products = []
    for index in range(24):
        products.append({"title": f"iPhone 12 {64 if index < 12 else 128}GB #{index}", "platform": "eBay" if index < 15 else "Walmart", "price": 300 + index, "currency": "USD", "condition": "Used" if index % 2 else "New", "condition_display": "Used" if index % 2 else "New", "category_display": "Smartphones", "seller": "" if index == 0 else f"Seller {index}", "source_type": "live_ebay" if index < 15 else "serpapi_walmart", "source_url": f"https://example.test/{index}", "collected_at": datetime(2026, 7, 1, tzinfo=timezone.utc)})
    stored = precision_app.repository.save_external_results(search_id, user["_id"], products)
    precision_app.repository.attach_result_tokens(search_id, [f"external:{row['_id']}" for row in stored])
    precision_app.repository.complete_search(search_id, len(products))
    search = precision_app.repository.get_search(search_id, user["_id"])
    included = search["result_tokens"][:23]
    analysis_id = precision_app.repository.create_analysis_record(user["_id"], {"search_record_id": search_id, "scope": "all_valid_results", "query": "iPhone 12", "included_result_ids": included, "preview_result_ids": search["result_tokens"] if legacy_preview else included, "excluded_result_ids": [search["result_tokens"][23]], "analysis_signature": "scope-test", **({} if legacy_preview else {"analysis_schema_version": precision_app.CURRENT_ANALYSIS_SCHEMA_VERSION})})
    return client, analysis_id


def test_frozen_analysis_scope_drives_page_csv_and_workbook():
    client, analysis_id = _fixture()
    page = client.get(f"/analytics/{analysis_id}")
    assert page.status_code == 200
    body = page.data.decode("utf-8")
    assert "24 filtered records" not in body
    assert "23 comparable listings analyzed" in body
    assert ">More</summary>" not in body
    assert "Delete analysis" not in body
    assert '<option value="all">All listings</option>' not in body
    assert "Lowest 8" in body and "Highest 8" in body
    assert "right: 112" in body and "overflow: 'truncate'" in body
    assert "normalizedSource.includes('walmart') ? 'Walmart'" in body

    response = client.get(f"/analytics/{analysis_id}/export/results.csv")
    rows = list(csv.reader(io.StringIO(response.data.decode("utf-8"))))
    assert rows[0][0] == "Analysis ID" and len(rows) == 24
    assert all(row[8] == "Smartphones" for row in rows[1:])
    assert all(row[11] in {"eBay Browse API", "Walmart via SerpAPI"} for row in rows[1:])

    workbook_response = client.get(f"/analytics/{analysis_id}/export/workbook.xlsx")
    book = load_workbook(io.BytesIO(workbook_response.data), data_only=True)
    assert book.sheetnames == ["Summary", "Records", "Platform Breakdown", "Condition Breakdown", "Price Distribution"]
    assert book["Summary"]["B3"].value == 23
    assert sum(row[1] for row in book["Platform Breakdown"].iter_rows(min_row=2, values_only=True)) == 23
    assert sum(row[1] for row in book["Condition Breakdown"].iter_rows(min_row=2, values_only=True)) == 23


def test_print_report_describes_mixed_storage_and_frozen_scope():
    client, analysis_id = _fixture()
    report = client.post(f"/analytics/{analysis_id}/export/report", data={})
    text = report.data.decode("utf-8")
    assert report.status_code == 200
    assert "Analysis scope" in text and "23 comparable listings" in text
    assert "Mixed storage limitation" in text
    assert "A lowest-price platform is not determined for this mixed-storage result set." in text
    assert "24 collected · 23 analyzed" in text
    assert "frozen analysis scope" not in text
    assert "{'" not in text


def test_condition_mix_does_not_hide_lowest_price_platform():
    payload = precision_app.build_analytics_payload([
        {
            "title": "iPhone 15 Used",
            "platform": "eBay",
            "normalized_price": 349,
            "price": 349,
            "analytics_eligible": True,
            "condition_display": "Used",
            "attributes": {"storage": "128GB"},
        },
        {
            "title": "iPhone 15 Refurbished",
            "platform": "Walmart",
            "normalized_price": 389,
            "price": 389,
            "analytics_eligible": True,
            "condition_display": "Refurbished",
            "attributes": {"storage": "128GB"},
        },
    ])
    assert payload["summary"]["storage_variants"] == ["128GB"]
    assert payload["summary"]["mixed_storage"] is False
    assert payload["summary"]["mixed_conditions"] is True
    assert payload["summary"]["mixed_variants"] is False
    assert payload["summary"]["best_platform"] == "eBay"


def test_print_report_filters_are_validated_and_recalculated_server_side():
    client, analysis_id = _fixture()
    report = client.post(
        f"/analytics/{analysis_id}/export/report",
        data={
            "metadata": json.dumps({"filters": {"platform": "Walmart", "category": "All", "condition": "All"}}),
            "summary": json.dumps({"totalRecords": 999, "averagePrice": 1}),
            "records": json.dumps([{"title": "Browser forged record", "platform": "Fake"}]),
            "charts": json.dumps({}),
        },
    )
    text = report.data.decode("utf-8")
    assert report.status_code == 200
    assert "8 comparable listings" in text
    assert "Analytics Platform: Walmart" in text
    assert "iPhone 12 · 128GB" in text
    assert "Lowest-price platform: Walmart" in text
    assert "Browser forged record" not in text
    assert "999 comparable listings" not in text
    assert ">eBay<" not in text

    invalid = client.post(
        f"/analytics/{analysis_id}/export/report",
        data={"metadata": json.dumps({"filters": {"platform": "Injected platform"}})},
    )
    assert invalid.status_code == 400


def test_legacy_analysis_report_filters_fall_back_to_persisted_search_facets():
    client, analysis_id = _fixture()
    analysis = precision_app.repository.get_analysis_record(analysis_id)
    search = precision_app.repository.get_search(analysis["search_record_id"], analysis["user_id"])
    precision_app.repository.update_search_run(search["_id"], {
        "filters": {"storage": ""},
        "selected_facets": {"storage": ["128GB"]},
    })
    precision_app.repository.update_analysis_record(analysis_id, {
        "filters": {},
        "normalized_filters": {},
    })
    with client.session_transaction() as browser_session:
        user_id = browser_session["user_id"]
    with precision_app.app.test_request_context():
        precision_app.session["user_id"] = user_id
        persisted = precision_app._report_persisted_filters(
            analysis_id,
            precision_app.repository.get_search(search["_id"], user_id),
        )
    assert persisted["storage"] == ["128GB"]


def test_legacy_scope_mismatch_has_styled_delete_and_blocks_exports():
    client, analysis_id = _fixture(legacy_preview=True)
    blocked = client.get(f"/analytics/{analysis_id}")
    body = blocked.data.decode("utf-8")
    assert blocked.status_code == 409
    assert "Analysis no longer available" in body and "Delete analysis" in body
    assert "Conflict" not in body
    assert client.get(f"/analytics/{analysis_id}/export/results.csv").status_code == 409

    deleted = client.post(f"/analytics/{analysis_id}/delete")
    assert deleted.status_code == 302
    assert precision_app.repository.get_analysis_record(analysis_id) is None
    raw = next(row for row in precision_app.repository._memory["analytics_reports"] if str(row["_id"]) == str(analysis_id))
    assert raw["status"] == "deleted"
    assert client.get(f"/analytics/{analysis_id}").status_code == 404


def test_unrecoverable_legacy_analysis_has_clear_delete_message():
    client, analysis_id = _fixture(legacy_preview=True)
    analysis = precision_app.repository.get_analysis_record(analysis_id)
    precision_app.repository.update_analysis_record(analysis_id, {"included_result_ids": ["external:missing-record"]})
    response = client.get(f"/analytics/{analysis_id}")
    assert response.status_code == 409
    assert "not compatible with the current analytics version" in response.data.decode("utf-8")
    assert analysis


def test_owner_can_soft_delete_without_removing_source_and_other_user_cannot():
    client, analysis_id = _fixture()
    analysis = precision_app.repository.get_analysis_record(analysis_id)
    search_id = analysis["search_record_id"]
    history = client.get("/analytics")
    assert "Delete" in history.data.decode("utf-8")

    client.post("/logout")
    _login(client, "Other")
    other = precision_app.repository.get_user_by_display_name("Other")
    precision_app.repository.update_user(other["_id"], {"membership_tier": "professional", "plan": "professional"})
    assert client.post(f"/analytics/{analysis_id}/delete").status_code == 404

    client.post("/logout")
    _login(client, "Analyst")
    precision_app.repository.update_user(precision_app.repository.get_user_by_display_name("Analyst")["_id"], {"membership_tier": "professional", "plan": "professional"})
    assert client.post(f"/analytics/{analysis_id}/delete").status_code == 302
    assert precision_app.repository.get_search(search_id) is not None
    assert precision_app.repository.find_analysis_by_signature(analysis["user_id"], analysis["analysis_signature"]) is None
